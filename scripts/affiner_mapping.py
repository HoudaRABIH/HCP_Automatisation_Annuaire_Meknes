from pathlib import Path
from collections import Counter
from difflib import SequenceMatcher
from openpyxl import load_workbook, Workbook
from docx import Document
import re
import unicodedata
import warnings

warnings.filterwarnings("ignore")

# ============================================================
# CHEMINS
# ============================================================

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_MODELES = Path("modeles")
DOSSIER_SORTIE = Path("sorties")

FICHIER_MAPPING = DOSSIER_SORTIE / "mapping_revu.xlsx"
FICHIER_CATALOGUE = DOSSIER_SORTIE / "catalogue_tables_reelles.xlsx"
FICHIER_SORTIE = DOSSIER_SORTIE / "mapping_affine.xlsx"

# ============================================================
# ROUTAGE DES THEMES
# ============================================================

ROUTAGE_BRUT = {
    "DIVISION ADMINISTRATIVE": ["1.Division Administrative_AS 2025.xlsx"],
    "POPULATION": ["2. population_AS 2025.xlsx"],
    "SANTE": ["12. Santé_AS 2025.xlsx"],
    "EDUCATION": ["2. population_AS 2025.xlsx"],
    "ENSEIGNEMENT PRESCOLAIRE": ["11. Enseignement_AS 2025.xlsx"],
    "ENSEIGNEMENT PRIMAIRE PUBLIC": ["11. Enseignement_AS 2025.xlsx"],
    "ENSEIGNEMENT PRIMAIRE PRIVE": ["11. Enseignement_AS 2025.xlsx"],
    "ENSEIGNEMENT SECONDAIRE COLLEGIAL PUBLIC": ["11. Enseignement_AS 2025.xlsx"],
    "ENSEIGNEMENT SECONDAIRE COLLEGIAL PRIVE": ["11. Enseignement_AS 2025.xlsx"],
    "ENSEIGNEMENT SECONDAIRE QUALIFIANT PUBLIC": ["11. Enseignement_AS 2025.xlsx"],
    "ENSEIGNEMENT SECONDAIRE QUALIFIANT PRIVE": ["11. Enseignement_AS 2025.xlsx"],
    "ENSEIGNEMENT POSTE SECONDAIRE": ["11. Enseignement_AS 2025.xlsx"],
    "ENSEIGNEMENT SUPERIEUR PUBLIC": ["11. Enseignement_AS 2025.xlsx"],
    "FORMATION PROFESSIONNELLE": ["11. Enseignement_AS 2025.xlsx"],
    "ACTIVITES CULTURELLES ET LOISIRS": ["16. Activités culturelles_Loisirs et Entraide_AS 2025.xlsx"],
    "EMPLOI ET SALAIRE": ["14-Emploi et salaire_AS 2025.xlsx"],
    "JUSTICE": ["15.Justice et Sauvegarde de l'Enfance _AS2025.xlsx"],
    "AGRICULTURE": ["3. Agriculture_AS 20252.xlsx"],
    "ELEVAGE": ["3. Agriculture_AS 20252.xlsx"],
    "FORET": ["3. Agriculture_AS 20252.xlsx"],
    "INDUSTRIE": ["6. Industrie et artisanat_AS 2025.xlsx"],
    "EAU POTABLE": ["5.Energie et Eau_AS 2025.xlsx"],
    "ENERGIE": ["5.Energie et Eau_AS 2025.xlsx"],
    "TOURISME": ["10. Tourisme_AS 2025.xlsx"],
    "CONDITION D HABITAT DES MENAGES": ["2. population_AS 2025.xlsx"],
    "CONSTRUCTION": [
        "7. Foncier_AS 2025.xlsx",
        "2. population_AS 2025.xlsx",
        "6. Industrie et artisanat_AS 2025.xlsx"
    ],
    "FONCIER": ["7. Foncier_AS 2025.xlsx"],
    "TRANSPORT ROUTIER": ["8. Transport_AS 2025.xlsx"],
    "PRIX": ["17.Indice des prix_AS  2025.xlsx"],
    "ENVIRONNEMENT": ["23. Environnement_AS 2025.xlsx"],
}

MOTS_VIDES = {
    "de", "du", "des", "la", "le", "les", "un", "une",
    "et", "en", "par", "selon", "au", "aux", "a", "dans",
    "pour", "sur", "avec", "sans", "nombre", "effectif",
    "effectifs", "total", "totale", "totaux", "annee",
    "annees", "source", "tableau", "tableaux"
}

MOTIF_NOMBRE = re.compile(r"[-+]?\d+(?:[\s\u00A0]\d{3})*(?:[.,]\d+)?")
MOTIF_ANNEE = re.compile(r"\b(?:19|20)\d{2}\b")

# ============================================================
# NORMALISATION
# ============================================================

def sans_accents(texte):
    texte = str(texte or "").replace("œ", "oe").replace("Œ", "OE")
    return "".join(
        c for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )

def normaliser(texte):
    texte = sans_accents(texte).lower().replace("\n", " ").replace("’", "'")
    texte = re.sub(r"^\s*\d{1,2}\s*[-–—]\s*\d{1,3}\s*", "", texte)
    texte = re.sub(r"\(\s*\d+\s*\)", " ", texte)
    texte = re.sub(r"[^a-z0-9%]+", " ", texte)
    texte = re.sub(r"\s+", " ", texte)
    return texte.strip()

def normaliser_theme(texte):
    return normaliser(texte).upper()

ROUTAGE = {
    normaliser_theme(cle): valeur
    for cle, valeur in ROUTAGE_BRUT.items()
}

def mots_significatifs(texte):
    mots = set()
    for mot in normaliser(texte).split():
        if len(mot) >= 3 and mot not in MOTS_VIDES:
            mots.add(mot)
    return mots

def score_titre(a, b):
    na = normaliser(a)
    nb = normaliser(b)
    if not na or not nb:
        return 0.0

    seq = SequenceMatcher(None, na, nb).ratio() * 100
    ma = mots_significatifs(a)
    mb = mots_significatifs(b)

    chev = 0.0
    if ma and mb:
        chev = len(ma & mb) / min(len(ma), len(mb)) * 100

    inclusion = 100.0 if min(len(na), len(nb)) >= 12 and (na in nb or nb in na) else 0.0
    return min(100.0, 0.5 * seq + 0.4 * chev + 0.1 * inclusion)

# ============================================================
# EMPREINTES
# ============================================================

def normaliser_nombre(v):
    try:
        n = float(v)
    except Exception:
        return None

    if n.is_integer() and 1900 <= n <= 2035:
        return None
    if n.is_integer() and abs(n) <= 2:
        return None

    if n.is_integer():
        return str(int(n))
    return f"{n:.4f}".rstrip("0").rstrip(".")

def extraire_nombres(valeur):
    resultats = []

    if valeur is None or isinstance(valeur, bool):
        return resultats

    if isinstance(valeur, (int, float)):
        n = normaliser_nombre(valeur)
        return [n] if n is not None else []

    texte = str(valeur).replace("\u00A0", " ")

    for m in MOTIF_NOMBRE.finditer(texte):
        brut = m.group(0).replace(" ", "").replace(",", ".")
        try:
            n = normaliser_nombre(float(brut))
        except Exception:
            n = None
        if n is not None:
            resultats.append(n)

    return resultats

def empreinte(valeurs):
    nombres = Counter()
    mots = set()
    annees = set()

    for valeur in valeurs:
        if valeur in (None, ""):
            continue

        for n in extraire_nombres(valeur):
            nombres[n] += 1

        texte = str(valeur)

        for a in MOTIF_ANNEE.findall(texte):
            annees.add(a)

        mots |= mots_significatifs(texte)

    return nombres, mots, annees

def score_nombres(word, excel):
    if not word:
        return 0.0

    communs = word & excel
    total = sum(word.values())

    return 100.0 * sum(communs.values()) / total if total else 0.0

def score_mots(word, excel):
    if not word:
        return 0.0
    return 100.0 * len(word & excel) / len(word)

def score_annees(word, excel):
    if not word:
        return 0.0
    return 100.0 * len(word & excel) / len(word)

# ============================================================
# CHARGEMENT DU MAPPING REVU
# ============================================================

wb_mapping = load_workbook(FICHIER_MAPPING, data_only=True)
ws_mapping = wb_mapping["Mapping_revu"]

hmap = {c.value: c.column for c in ws_mapping[1]}

def vm(ligne, nom):
    return ws_mapping.cell(ligne, hmap[nom]).value

lignes_mapping = []
for r in range(2, ws_mapping.max_row + 1):
    lignes_mapping.append({
        "table_word_titre": vm(r, "Table Word titre"),
        "table_word_donnees": vm(r, "Table Word données"),
        "theme": vm(r, "Thème Word") or "",
        "numero_word": vm(r, "N° Word") or "",
        "titre_word": vm(r, "Titre Word") or "",
        "decision_actuelle": vm(r, "Décision") or "",
        "fichier_actuel": vm(r, "Fichier attendu") or "",
        "numero_excel_actuel": vm(r, "Nouveau N° Excel") or "",
        "score_actuel": vm(r, "Nouveau score titre") or 0,
    })

wb_mapping.close()

# ============================================================
# CHARGEMENT DU CATALOGUE
# ============================================================

wb_cat = load_workbook(FICHIER_CATALOGUE, data_only=True)
ws_cat = wb_cat["Tables_logiques"]

hcat = {c.value: c.column for c in ws_cat[1]}

def vc(ligne, nom):
    return ws_cat.cell(ligne, hcat[nom]).value

tables_excel = []
for r in range(2, ws_cat.max_row + 1):
    tables_excel.append({
        "id": vc(r, "ID"),
        "fichier": vc(r, "Fichier"),
        "numero": vc(r, "Numéro complet"),
        "titre": vc(r, "Titre principal") or "",
        "feuilles": vc(r, "Feuille(s)") or "",
        "plages": vc(r, "Plage(s)") or "",
        "meknes": vc(r, "Meknès détecté"),
        "fes_meknes": vc(r, "Fès-Meknès détecté"),
        "annees_meta": vc(r, "Années") or "",
    })

wb_cat.close()

# ============================================================
# OUVERTURE WORD
# ============================================================

fichiers_word = [
    f for f in DOSSIER_MODELES.glob("*.docx")
    if not f.name.startswith("~$")
]

if not fichiers_word:
    raise FileNotFoundError("Aucun fichier Word dans modeles.")

document = Document(fichiers_word[0])

def valeurs_table_word(numero_table):
    table = document.tables[int(numero_table) - 1]
    valeurs = []
    vus = set()

    for ligne in table.rows:
        for cellule in ligne.cells:
            identifiant = id(cellule._tc)
            if identifiant in vus:
                continue
            vus.add(identifiant)

            if cellule.text.strip():
                valeurs.append(cellule.text.strip())

    return valeurs

# ============================================================
# CACHE DES SOURCES EXCEL
# ============================================================

classeurs = {}
empreintes_excel = {}

def ouvrir_classeur(nom):
    if nom not in classeurs:
        classeurs[nom] = load_workbook(
            DOSSIER_SOURCE / nom,
            data_only=True,
            read_only=False
        )
    return classeurs[nom]

def valeurs_table_excel(table):
    valeurs = []
    wb = ouvrir_classeur(table["fichier"])

    for morceau in str(table["plages"]).split("|"):
        morceau = morceau.strip()
        if not morceau or "!" not in morceau:
            continue

        feuille_nom, lignes = morceau.rsplit("!", 1)
        if ":" not in lignes or feuille_nom not in wb.sheetnames:
            continue

        debut, fin = lignes.split(":", 1)
        try:
            debut = int(debut)
            fin = int(fin)
        except ValueError:
            continue

        ws = wb[feuille_nom]
        max_col = min(ws.max_column, 80)

        for ligne in ws.iter_rows(
            min_row=debut,
            max_row=fin,
            min_col=1,
            max_col=max_col
        ):
            for cellule in ligne:
                if cellule.value not in (None, ""):
                    valeurs.append(cellule.value)

    return valeurs

def empreinte_excel(table):
    if table["id"] not in empreintes_excel:
        empreintes_excel[table["id"]] = empreinte(
            valeurs_table_excel(table)
        )
    return empreintes_excel[table["id"]]

# ============================================================
# AFFINAGE
# ============================================================

rapport = Workbook()
ws_res = rapport.active
ws_res.title = "Mapping_affine"

ws_res.append([
    "Table Word titre",
    "Table Word données",
    "Thème Word",
    "N° Word",
    "Titre Word",
    "Décision précédente",
    "Fichier(s) testé(s)",
    "N° Excel proposé",
    "Titre Excel proposé",
    "Score final",
    "Score titre",
    "Score nombres",
    "Score texte",
    "Score années",
    "Meknès",
    "Feuille(s)",
    "Plage(s)",
    "Décision affinée"
])

ws_top3 = rapport.create_sheet("Top3_affine")
ws_top3.append([
    "N° Word", "Titre Word", "Thème", "Rang",
    "Score final", "Score titre", "Score nombres",
    "Score texte", "Score années", "Fichier Excel",
    "N° Excel", "Titre Excel", "Feuille(s)", "Plage(s)"
])

compteurs = Counter()

print("=" * 78)
print("AFFINAGE DU MAPPING WORD <-> EXCEL")
print("=" * 78)

for i, ligne in enumerate(lignes_mapping, start=1):
    theme_norm = normaliser_theme(ligne["theme"])
    fichiers = ROUTAGE.get(theme_norm)

    if not fichiers:
        fichiers = sorted({t["fichier"] for t in tables_excel})

    candidats = [
        t for t in tables_excel
        if t["fichier"] in fichiers
    ]

    valeurs_word = valeurs_table_word(
        ligne["table_word_donnees"]
    )
    n_word, m_word, a_word = empreinte(valeurs_word)

    scores = []

    for table in candidats:
        n_excel, m_excel, a_excel = empreinte_excel(table)

        st = score_titre(
            ligne["titre_word"],
            table["titre"]
        )

        sn = score_nombres(
            n_word,
            n_excel
        )

        sm = score_mots(
            m_word,
            m_excel
        )

        sa = score_annees(
            a_word,
            a_excel
        )

        nb_nombres_word = sum(n_word.values())

        if nb_nombres_word >= 3:
            sf = (
                0.35 * st
                + 0.35 * sn
                + 0.20 * sm
                + 0.10 * sa
            )
        else:
            sf = (
                0.55 * st
                + 0.30 * sm
                + 0.15 * sa
            )

        if (
            "meknes" in normaliser(
                ligne["titre_word"] + " " + ligne["theme"]
            )
            and table["meknes"] == "OUI"
        ):
            sf = min(100.0, sf + 3.0)

        scores.append((
            sf, st, sn, sm, sa, table
        ))

    scores.sort(
        key=lambda x: x[0],
        reverse=True
    )

    top = scores[:3]

    if not top:
        continue

    sf, st, sn, sm, sa, meilleur = top[0]

    deuxieme = top[1][0] if len(top) > 1 else 0.0
    marge = sf - deuxieme

    if sf >= 82 and marge >= 6:
        decision = "VALIDE AUTO"
    elif sf >= 70 and marge >= 3:
        decision = "TRES PROBABLE"
    elif sf >= 58:
        decision = "A VERIFIER"
    else:
        decision = "SOURCE / TABLE A CONFIRMER"

    compteurs[decision] += 1

    ws_res.append([
        ligne["table_word_titre"],
        ligne["table_word_donnees"],
        ligne["theme"],
        ligne["numero_word"],
        ligne["titre_word"],
        ligne["decision_actuelle"],
        " | ".join(fichiers),
        meilleur["numero"],
        meilleur["titre"],
        round(sf, 1),
        round(st, 1),
        round(sn, 1),
        round(sm, 1),
        round(sa, 1),
        meilleur["meknes"],
        meilleur["feuilles"],
        meilleur["plages"],
        decision
    ])

    for rang, element in enumerate(top, start=1):
        sf2, st2, sn2, sm2, sa2, table2 = element

        ws_top3.append([
            ligne["numero_word"],
            ligne["titre_word"],
            ligne["theme"],
            rang,
            round(sf2, 1),
            round(st2, 1),
            round(sn2, 1),
            round(sm2, 1),
            round(sa2, 1),
            table2["fichier"],
            table2["numero"],
            table2["titre"],
            table2["feuilles"],
            table2["plages"]
        ])

    if i == 1 or i % 20 == 0 or i == len(lignes_mapping):
        print(
            f"[{i}/{len(lignes_mapping)}] "
            f"{ligne['numero_word']} "
            f"-> {meilleur['numero']} "
            f"({round(sf,1)}%)"
        )

# ============================================================
# RESUME
# ============================================================

ws_resume = rapport.create_sheet("Resume")
ws_resume.append(["Indicateur", "Valeur"])
ws_resume.append(["Tableaux Word analysés", len(lignes_mapping)])

for categorie in [
    "VALIDE AUTO",
    "TRES PROBABLE",
    "A VERIFIER",
    "SOURCE / TABLE A CONFIRMER"
]:
    ws_resume.append([
        categorie,
        compteurs[categorie]
    ])

for feuille in [ws_res, ws_top3]:
    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = feuille.dimensions

largeurs = {
    "A": 16, "B": 18, "C": 32, "D": 12, "E": 75,
    "F": 22, "G": 55, "H": 15, "I": 75, "J": 14,
    "K": 14, "L": 15, "M": 15, "N": 15, "O": 12,
    "P": 35, "Q": 60, "R": 26
}

for col, largeur in largeurs.items():
    ws_res.column_dimensions[col].width = largeur

ws_top3.column_dimensions["A"].width = 12
ws_top3.column_dimensions["B"].width = 75
ws_top3.column_dimensions["C"].width = 32
ws_top3.column_dimensions["J"].width = 45
ws_top3.column_dimensions["K"].width = 15
ws_top3.column_dimensions["L"].width = 75
ws_top3.column_dimensions["M"].width = 35
ws_top3.column_dimensions["N"].width = 60

ws_resume.column_dimensions["A"].width = 40
ws_resume.column_dimensions["B"].width = 20

rapport.save(FICHIER_SORTIE)

for wb in classeurs.values():
    wb.close()

print()
print("Résultat :")
for categorie in [
    "VALIDE AUTO",
    "TRES PROBABLE",
    "A VERIFIER",
    "SOURCE / TABLE A CONFIRMER"
]:
    print(
        f"{categorie} : "
        f"{compteurs[categorie]}"
    )

print()
print(
    f"Fichier créé : "
    f"{FICHIER_SORTIE}"
)
print("=" * 78)
print("AFFINAGE TERMINE")
print("=" * 78)
