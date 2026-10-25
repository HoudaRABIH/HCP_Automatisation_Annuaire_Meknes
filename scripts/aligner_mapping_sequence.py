from pathlib import Path
from openpyxl import load_workbook, Workbook
from difflib import SequenceMatcher
from collections import defaultdict, Counter
import re
import unicodedata


DOSSIER_SORTIE = Path("sorties")

FICHIER_MAPPING = DOSSIER_SORTIE / "mapping_affine.xlsx"
FICHIER_CATALOGUE = DOSSIER_SORTIE / "catalogue_tables_reelles.xlsx"
FICHIER_SORTIE = DOSSIER_SORTIE / "mapping_sequence.xlsx"


# ============================================================
# ROUTAGE WORD -> FICHIER EXCEL
# ============================================================

# Important :
# Les thèmes qui utilisent le même fichier Excel seront alignés ENSEMBLE
# dans l'ordre du Word et dans l'ordre des tableaux Excel.
#
# Cela évite, par exemple, que plusieurs tableaux scolaires portant tous
# le titre "Etablissements selon le milieu" soient associés au même tableau
# Excel : l'ordre global est utilisé pour distinguer primaire / collège /
# qualifiant, etc.

ROUTAGE_BRUT = {
    "DIVISION ADMINISTRATIVE": [
        "1.Division Administrative_AS 2025.xlsx"
    ],

    # Le fichier Population contient plusieurs familles reprises dans le Word :
    # population, certains indicateurs d'éducation du RGPH, habitat des ménages.
    "POPULATION": [
        "2. population_AS 2025.xlsx"
    ],
    "EDUCATION": [
        "2. population_AS 2025.xlsx"
    ],
    "CONDITION D’HABITAT DES MENAGES": [
        "2. population_AS 2025.xlsx"
    ],

    "SANTE": [
        "12. Santé_AS 2025.xlsx"
    ],

    "ENSEIGNEMENT PRESCOLAIRE": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT PRIMAIRE PUBLIC": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT PRIMAIRE PRIVE": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT SECONDAIRE COLLEGIAL PUBLIC": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT SECONDAIRE COLLEGIAL PRIVE": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT SECONDAIRE QUALIFIANT PUBLIC": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT SECONDAIRE QUALIFIANT PRIVE": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT POSTE - SECONDAIRE": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT SUPERIEUR PUBLIC": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "FORMATION PROFESSIONNELLE": [
        "11. Enseignement_AS 2025.xlsx"
    ],

    "ACTIVITES CULTURELLES ET LOISIRS": [
        "16. Activités culturelles_Loisirs et Entraide_AS 2025.xlsx"
    ],

    "EMPLOI ET SALAIRE": [
        "14-Emploi et salaire_AS 2025.xlsx"
    ],

    "JUSTICE": [
        "15.Justice et Sauvegarde de l'Enfance _AS2025.xlsx"
    ],

    # Ces trois thèmes appartiennent au même classeur et doivent rester
    # monotones dans l'ordre Agriculture -> Elevage -> Forêt.
    "AGRICULTURE": [
        "3. Agriculture_AS 20252.xlsx"
    ],
    "ELEVAGE": [
        "3. Agriculture_AS 20252.xlsx"
    ],
    "FORET": [
        "3. Agriculture_AS 20252.xlsx"
    ],

    "INDUSTRIE": [
        "6. Industrie et artisanat_AS 2025.xlsx"
    ],

    "EAU POTABLE": [
        "5.Energie et Eau_AS 2025.xlsx"
    ],
    "ENERGIE": [
        "5.Energie et Eau_AS 2025.xlsx"
    ],

    "TOURISME": [
        "10. Tourisme_AS 2025.xlsx"
    ],

    # Construction reste volontairement multi-source car le catalogue
    # actuel ne permet pas encore d'affirmer le fichier exact.
    "CONSTRUCTION": [
        "7. Foncier_AS 2025.xlsx",
        "6. Industrie et artisanat_AS 2025.xlsx",
        "2. population_AS 2025.xlsx"
    ],

    "FONCIER": [
        "7. Foncier_AS 2025.xlsx"
    ],

    "TRANSPORT ROUTIER": [
        "8. Transport_AS 2025.xlsx"
    ],

    "PRIX": [
        "17.Indice des prix_AS  2025.xlsx"
    ],

    "ENVIRONNEMENT": [
        "23. Environnement_AS 2025.xlsx"
    ],
}


MOTS_VIDES = {
    "de", "du", "des", "la", "le", "les", "un", "une",
    "et", "en", "par", "selon", "au", "aux", "a", "à",
    "dans", "d", "l", "pour", "sur", "avec", "sans",
    "nombre", "effectif", "effectifs", "total", "totale",
    "totaux", "annee", "annees", "année", "années",
    "source", "tableau", "tableaux"
}


# ============================================================
# NORMALISATION
# ============================================================

def sans_accents(texte):
    texte = str(texte or "").replace("œ", "oe").replace("Œ", "OE")

    return "".join(
        c
        for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )


def normaliser(texte):
    texte = sans_accents(texte).lower()
    texte = texte.replace("\n", " ")
    texte = texte.replace("’", "'")

    # Retirer le numéro de tableau au début.
    texte = re.sub(
        r"^\s*\d{1,2}\s*[-–—]\s*\d{1,3}\s*",
        "",
        texte
    )

    # Retirer notes (1), (2)...
    texte = re.sub(
        r"\(\s*\d+\s*\)",
        " ",
        texte
    )

    texte = re.sub(
        r"[^a-z0-9%]+",
        " ",
        texte
    )

    texte = re.sub(
        r"\s+",
        " ",
        texte
    )

    return texte.strip()


def normaliser_theme(texte):
    return normaliser(texte).upper()


ROUTAGE = {
    normaliser_theme(cle): valeur
    for cle, valeur in ROUTAGE_BRUT.items()
}


def mots_significatifs(texte):
    resultat = set()

    for mot in normaliser(texte).split():
        if len(mot) < 3:
            continue

        if mot in MOTS_VIDES:
            continue

        resultat.add(mot)

    return resultat


def score_semantique(titre_word, titre_excel):
    """
    Score basé UNIQUEMENT sur le sens du titre.
    On ne compare PAS les valeurs numériques :
    le Word peut contenir l'ancienne année et Excel la nouvelle année.
    """

    a = normaliser(titre_word)
    b = normaliser(titre_excel)

    if not a or not b:
        return 0.0

    sequence = (
        SequenceMatcher(
            None,
            a,
            b
        ).ratio()
        * 100
    )

    mots_a = mots_significatifs(
        titre_word
    )

    mots_b = mots_significatifs(
        titre_excel
    )

    if mots_a and mots_b:
        jaccard = (
            len(mots_a & mots_b)
            / len(mots_a | mots_b)
            * 100
        )

        couverture_word = (
            len(mots_a & mots_b)
            / len(mots_a)
            * 100
        )
    else:
        jaccard = 0.0
        couverture_word = 0.0

    inclusion = 0.0

    if min(len(a), len(b)) >= 10:
        if a in b or b in a:
            inclusion = 100.0

    score = (
        0.42 * sequence
        + 0.28 * couverture_word
        + 0.20 * jaccard
        + 0.10 * inclusion
    )

    return min(
        100.0,
        score
    )


# ============================================================
# CHARGEMENT WORD DEPUIS mapping_affine
# ============================================================

if not FICHIER_MAPPING.exists():
    raise FileNotFoundError(
        f"Introuvable : {FICHIER_MAPPING}"
    )

wb_mapping = load_workbook(
    FICHIER_MAPPING,
    data_only=True
)

ws_mapping = wb_mapping[
    "Mapping_affine"
]

hmap = {
    cellule.value: cellule.column
    for cellule in ws_mapping[1]
}


def vm(ligne, nom):
    return ws_mapping.cell(
        row=ligne,
        column=hmap[nom]
    ).value


word_rows = []

for ligne in range(
    2,
    ws_mapping.max_row + 1
):
    word_rows.append({
        "ordre": ligne - 1,
        "table_word_titre": vm(
            ligne,
            "Table Word titre"
        ),
        "table_word_donnees": vm(
            ligne,
            "Table Word données"
        ),
        "theme": vm(
            ligne,
            "Thème Word"
        ) or "",
        "numero_word": vm(
            ligne,
            "N° Word"
        ) or "",
        "titre_word": vm(
            ligne,
            "Titre Word"
        ) or "",
        "decision_precedente": vm(
            ligne,
            "Décision affinée"
        ) or "",
    })


wb_mapping.close()


# ============================================================
# CHARGEMENT CATALOGUE EXCEL
# ============================================================

if not FICHIER_CATALOGUE.exists():
    raise FileNotFoundError(
        f"Introuvable : {FICHIER_CATALOGUE}"
    )

wb_catalogue = load_workbook(
    FICHIER_CATALOGUE,
    data_only=True
)

ws_catalogue = wb_catalogue[
    "Tables_logiques"
]

hcat = {
    cellule.value: cellule.column
    for cellule in ws_catalogue[1]
}


def vc(ligne, nom):
    return ws_catalogue.cell(
        row=ligne,
        column=hcat[nom]
    ).value


excel_rows = []

for ligne in range(
    2,
    ws_catalogue.max_row + 1
):
    excel_rows.append({
        "id": vc(
            ligne,
            "ID"
        ),
        "ordre_global": ligne - 1,
        "fichier": vc(
            ligne,
            "Fichier"
        ),
        "chapitre": vc(
            ligne,
            "Chapitre source"
        ),
        "numero_source": vc(
            ligne,
            "N° source"
        ),
        "numero_excel": vc(
            ligne,
            "Numéro complet"
        ),
        "titre_excel": vc(
            ligne,
            "Titre principal"
        ) or "",
        "feuilles": vc(
            ligne,
            "Feuille(s)"
        ) or "",
        "plages": vc(
            ligne,
            "Plage(s)"
        ) or "",
        "meknes": vc(
            ligne,
            "Meknès détecté"
        ),
        "fes_meknes": vc(
            ligne,
            "Fès-Meknès détecté"
        ),
        "annees": vc(
            ligne,
            "Années"
        ) or "",
    })


wb_catalogue.close()


# ============================================================
# ALIGNEMENT MONOTONE
# ============================================================

def aligner_sequence(words, excels):
    """
    Alignement dynamique :
    - conserve l'ordre Word
    - conserve l'ordre Excel
    - autorise à sauter des tableaux Excel
    - autorise à laisser un tableau Word non apparié si aucun candidat
      n'est suffisamment crédible

    Cela impose implicitement une correspondance 1 -> 1.
    """

    n = len(words)
    m = len(excels)

    # dp[i][j] = meilleur score pour les i premiers Word
    #            et j premiers Excel
    dp = [
        [-10**9] * (m + 1)
        for _ in range(n + 1)
    ]

    action = [
        [None] * (m + 1)
        for _ in range(n + 1)
    ]

    dp[0][0] = 0.0

    # Sauter un tableau Excel coûte très peu :
    # le Word est un sous-ensemble du gros classeur Excel.
    for j in range(
        1,
        m + 1
    ):
        dp[0][j] = (
            dp[0][j - 1]
            - 0.7
        )
        action[0][j] = (
            "SKIP_EXCEL",
            0,
            j - 1
        )

    # Laisser un tableau Word sans correspondance est possible,
    # mais coûte davantage.
    for i in range(
        1,
        n + 1
    ):
        dp[i][0] = (
            dp[i - 1][0]
            - 18.0
        )
        action[i][0] = (
            "SKIP_WORD",
            i - 1,
            0
        )

    for i in range(
        1,
        n + 1
    ):
        for j in range(
            1,
            m + 1
        ):
            similarite = (
                score_semantique(
                    words[i - 1][
                        "titre_word"
                    ],
                    excels[j - 1][
                        "titre_excel"
                    ]
                )
            )

            # Convertir la similarité en "gain".
            # < 30 : association très improbable.
            if similarite < 30:
                gain_match = -35.0

            else:
                gain_match = (
                    similarite
                    - 42.0
                )

            options = [
                (
                    dp[i - 1][j - 1]
                    + gain_match,
                    "MATCH",
                    i - 1,
                    j - 1
                ),
                (
                    dp[i][j - 1]
                    - 0.7,
                    "SKIP_EXCEL",
                    i,
                    j - 1
                ),
                (
                    dp[i - 1][j]
                    - 18.0,
                    "SKIP_WORD",
                    i - 1,
                    j
                ),
            ]

            meilleur = max(
                options,
                key=lambda x: x[0]
            )

            dp[i][j] = meilleur[0]

            action[i][j] = (
                meilleur[1],
                meilleur[2],
                meilleur[3]
            )

    # Reconstruction
    associations = {}

    i = n
    j = m

    while i > 0 or j > 0:
        info = action[i][j]

        if info is None:
            break

        type_action = info[0]

        if type_action == "MATCH":
            word_index = i - 1
            excel_index = j - 1

            associations[
                word_index
            ] = excel_index

            i -= 1
            j -= 1

        elif type_action == "SKIP_EXCEL":
            j -= 1

        elif type_action == "SKIP_WORD":
            associations[
                i - 1
            ] = None

            i -= 1

    return associations


# ============================================================
# GROUPER PAR FICHIER ROUTE
# ============================================================

# Pour chaque ligne Word, déterminer la route.
groupes_mono = defaultdict(list)
groupes_multi = []

for word in word_rows:
    theme_normalise = (
        normaliser_theme(
            word["theme"]
        )
    )

    fichiers = ROUTAGE.get(
        theme_normalise
    )

    if not fichiers:
        # Thème non reconnu : on le gardera en recherche globale.
        groupes_multi.append(
            (
                word,
                sorted(
                    {
                        e["fichier"]
                        for e in excel_rows
                    }
                )
            )
        )

    elif len(fichiers) == 1:
        groupes_mono[
            fichiers[0]
        ].append(
            word
        )

    else:
        groupes_multi.append(
            (
                word,
                fichiers
            )
        )


# ============================================================
# CALCUL DES ASSOCIATIONS
# ============================================================

resultats = []


print("=" * 80)
print("ALIGNEMENT SEMANTIQUE + ORDRE WORD / EXCEL")
print("=" * 80)


# ------------------------------------------------------------
# GROUPES MONO-FICHIER :
# ALIGNEMENT GLOBAL DANS L'ORDRE
# ------------------------------------------------------------

for fichier, words in groupes_mono.items():
    words = sorted(
        words,
        key=lambda x: x["ordre"]
    )

    excels = [
        e
        for e in excel_rows
        if e["fichier"] == fichier
    ]

    excels.sort(
        key=lambda x: (
            x["chapitre"]
            if x["chapitre"] is not None
            else 9999,
            x["numero_source"]
            if x["numero_source"] is not None
            else 9999,
            x["ordre_global"]
        )
    )

    print(
        f"{fichier} : "
        f"{len(words)} tableaux Word / "
        f"{len(excels)} tableaux Excel"
    )

    associations = (
        aligner_sequence(
            words,
            excels
        )
    )

    for index_word, word in enumerate(
        words
    ):
        index_excel = associations.get(
            index_word
        )

        if index_excel is None:
            resultats.append({
                "word": word,
                "fichier": fichier,
                "excel": None,
                "score": 0.0,
                "marge": 0.0,
                "mode": "SEQUENCE"
            })

            continue

        excel = excels[
            index_excel
        ]

        score = score_semantique(
            word["titre_word"],
            excel["titre_excel"]
        )

        # Calculer une marge indicative par rapport au meilleur candidat
        # sémantique alternatif dans le même fichier.
        alternatives = []

        for candidat in excels:
            if candidat["id"] == excel["id"]:
                continue

            alternatives.append(
                score_semantique(
                    word["titre_word"],
                    candidat[
                        "titre_excel"
                    ]
                )
            )

        deuxieme = (
            max(alternatives)
            if alternatives
            else 0.0
        )

        marge = (
            score
            - deuxieme
        )

        resultats.append({
            "word": word,
            "fichier": fichier,
            "excel": excel,
            "score": score,
            "marge": marge,
            "mode": "SEQUENCE"
        })


# ------------------------------------------------------------
# ROUTES MULTI-FICHIERS :
# RECHERCHE SEMANTIQUE SIMPLE
# ------------------------------------------------------------

for word, fichiers in groupes_multi:
    candidats = [
        e
        for e in excel_rows
        if e["fichier"] in fichiers
    ]

    scores = [
        (
            score_semantique(
                word["titre_word"],
                e["titre_excel"]
            ),
            e
        )
        for e in candidats
    ]

    scores.sort(
        key=lambda x: x[0],
        reverse=True
    )

    if scores:
        meilleur_score = (
            scores[0][0]
        )
        meilleur_excel = (
            scores[0][1]
        )

        deuxieme = (
            scores[1][0]
            if len(scores) > 1
            else 0.0
        )
    else:
        meilleur_score = 0.0
        meilleur_excel = None
        deuxieme = 0.0

    resultats.append({
        "word": word,
        "fichier": " | ".join(fichiers),
        "excel": meilleur_excel,
        "score": meilleur_score,
        "marge": (
            meilleur_score
            - deuxieme
        ),
        "mode": "MULTI-SOURCE"
    })


resultats.sort(
    key=lambda x: x["word"][
        "ordre"
    ]
)


# ============================================================
# DECISION
# ============================================================

compteurs = Counter()

for resultat in resultats:
    score = resultat["score"]
    excel = resultat["excel"]
    mode = resultat["mode"]

    if excel is None:
        decision = (
            "SOURCE / TABLE A CONFIRMER"
        )

    elif mode == "SEQUENCE":
        # La séquence elle-même est un signal important :
        # un score 70 avec ordre cohérent peut être plus fiable
        # qu'un score indépendant légèrement supérieur.
        if score >= 88:
            decision = "VALIDE AUTO"

        elif score >= 74:
            decision = "TRES PROBABLE"

        elif score >= 58:
            decision = "A VERIFIER"

        else:
            decision = (
                "SOURCE / TABLE A CONFIRMER"
            )

    else:
        # En multi-source on reste plus prudent.
        if (
            score >= 90
            and resultat["marge"] >= 8
        ):
            decision = "TRES PROBABLE"

        elif score >= 68:
            decision = "A VERIFIER"

        else:
            decision = (
                "SOURCE / TABLE A CONFIRMER"
            )

    resultat[
        "decision"
    ] = decision

    compteurs[
        decision
    ] += 1


# ============================================================
# RAPPORT
# ============================================================

rapport = Workbook()

ws = rapport.active
ws.title = "Mapping_sequence"

ws.append([
    "Table Word titre",
    "Table Word données",
    "Thème Word",
    "N° Word",
    "Titre Word",
    "Décision précédente",
    "Mode de mapping",
    "Fichier Excel",
    "N° Excel",
    "Titre Excel",
    "Score sémantique",
    "Marge indicative",
    "Feuille(s)",
    "Plage(s)",
    "Meknès",
    "Fès-Meknès",
    "Années",
    "Décision"
])


for resultat in resultats:
    word = resultat["word"]
    excel = resultat["excel"]

    ws.append([
        word["table_word_titre"],
        word["table_word_donnees"],
        word["theme"],
        word["numero_word"],
        word["titre_word"],
        word["decision_precedente"],
        resultat["mode"],
        (
            excel["fichier"]
            if excel
            else resultat["fichier"]
        ),
        (
            excel["numero_excel"]
            if excel
            else ""
        ),
        (
            excel["titre_excel"]
            if excel
            else ""
        ),
        round(
            resultat["score"],
            1
        ),
        round(
            resultat["marge"],
            1
        ),
        (
            excel["feuilles"]
            if excel
            else ""
        ),
        (
            excel["plages"]
            if excel
            else ""
        ),
        (
            excel["meknes"]
            if excel
            else ""
        ),
        (
            excel["fes_meknes"]
            if excel
            else ""
        ),
        (
            excel["annees"]
            if excel
            else ""
        ),
        resultat["decision"]
    ])


# ------------------------------------------------------------
# A VERIFIER
# ------------------------------------------------------------

ws_verif = rapport.create_sheet(
    "A_verifier"
)

ws_verif.append([
    "N° Word",
    "Thème",
    "Titre Word",
    "Fichier Excel proposé",
    "N° Excel proposé",
    "Titre Excel proposé",
    "Score",
    "Décision"
])


for resultat in resultats:
    if resultat["decision"] in {
        "A VERIFIER",
        "SOURCE / TABLE A CONFIRMER"
    }:
        word = resultat["word"]
        excel = resultat["excel"]

        ws_verif.append([
            word["numero_word"],
            word["theme"],
            word["titre_word"],
            (
                excel["fichier"]
                if excel
                else resultat["fichier"]
            ),
            (
                excel["numero_excel"]
                if excel
                else ""
            ),
            (
                excel["titre_excel"]
                if excel
                else ""
            ),
            round(
                resultat["score"],
                1
            ),
            resultat["decision"]
        ])


# ------------------------------------------------------------
# RESUME
# ------------------------------------------------------------

ws_resume = rapport.create_sheet(
    "Resume"
)

ws_resume.append([
    "Indicateur",
    "Valeur"
])

ws_resume.append([
    "Tableaux Word analysés",
    len(resultats)
])

for categorie in [
    "VALIDE AUTO",
    "TRES PROBABLE",
    "A VERIFIER",
    "SOURCE / TABLE A CONFIRMER"
]:
    ws_resume.append([
        categorie,
        compteurs[
            categorie
        ]
    ])


# ============================================================
# FORMAT
# ============================================================

for feuille in [
    ws,
    ws_verif
]:
    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = (
        feuille.dimensions
    )


largeurs = {
    "A": 16,
    "B": 18,
    "C": 32,
    "D": 12,
    "E": 75,
    "F": 24,
    "G": 18,
    "H": 48,
    "I": 14,
    "J": 75,
    "K": 18,
    "L": 18,
    "M": 35,
    "N": 60,
    "O": 12,
    "P": 16,
    "Q": 24,
    "R": 28
}

for colonne, largeur in (
    largeurs.items()
):
    ws.column_dimensions[
        colonne
    ].width = largeur


ws_verif.column_dimensions[
    "A"
].width = 12

ws_verif.column_dimensions[
    "B"
].width = 32

ws_verif.column_dimensions[
    "C"
].width = 75

ws_verif.column_dimensions[
    "D"
].width = 48

ws_verif.column_dimensions[
    "E"
].width = 14

ws_verif.column_dimensions[
    "F"
].width = 75

ws_verif.column_dimensions[
    "G"
].width = 14

ws_verif.column_dimensions[
    "H"
].width = 28

ws_resume.column_dimensions[
    "A"
].width = 42

ws_resume.column_dimensions[
    "B"
].width = 20


rapport.save(
    FICHIER_SORTIE
)


print()
print("Résultat :")

for categorie in [
    "VALIDE AUTO",
    "TRES PROBABLE",
    "A VERIFIER",
    "SOURCE / TABLE A CONFIRMER"
]:
    print(
        f"{categorie} : "
        f"{compteurs[categorie]}"
    )

print()
print(
    f"Fichier créé : "
    f"{FICHIER_SORTIE}"
)

print("=" * 80)
print("ALIGNEMENT TERMINE")
print("=" * 80)
