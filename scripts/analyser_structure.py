from pathlib import Path
from openpyxl import load_workbook
from openpyxl import Workbook
import warnings

# Masquer uniquement les avertissements liés aux zones d'impression Excel
warnings.filterwarnings(
    "ignore",
    message="Print area cannot be set*"
)

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_SORTIE = Path("sorties")
FICHIER_SORTIE = DOSSIER_SORTIE / "inventaire_structure.xlsx"

DOSSIER_SORTIE.mkdir(exist_ok=True)

fichiers_excel = sorted(DOSSIER_SOURCE.glob("*.xlsx"))

# Création du fichier Excel de résultat
rapport = Workbook()
feuille_rapport = rapport.active
feuille_rapport.title = "Inventaire"

# En-têtes
feuille_rapport.append([
    "Fichier",
    "Feuille",
    "Nombre de lignes",
    "Nombre de colonnes",
    "Cellules fusionnées",
    "Premier contenu trouvé"
])

print("=" * 70)
print("ANALYSE DE LA STRUCTURE DES FICHIERS HCP")
print("=" * 70)

nombre_feuilles = 0

for numero_fichier, chemin in enumerate(fichiers_excel, start=1):

    print(
        f"[{numero_fichier}/{len(fichiers_excel)}] "
        f"{chemin.name}"
    )

    try:
        classeur = load_workbook(
            chemin,
            read_only=False,
            data_only=True
        )

        for feuille in classeur.worksheets:

            nombre_feuilles += 1

            # Recherche du premier contenu non vide
            premier_contenu = ""

            for ligne in feuille.iter_rows(
                min_row=1,
                max_row=min(feuille.max_row, 20),
                min_col=1,
                max_col=min(feuille.max_column, 15)
            ):
                for cellule in ligne:
                    if cellule.value not in (None, ""):
                        premier_contenu = str(cellule.value).strip()
                        break

                if premier_contenu:
                    break

            feuille_rapport.append([
                chemin.name,
                feuille.title,
                feuille.max_row,
                feuille.max_column,
                len(feuille.merged_cells.ranges),
                premier_contenu
            ])

        classeur.close()

    except Exception as erreur:

        print(f"   ERREUR : {erreur}")

print()
print(f"Nombre total de feuilles analysées : {nombre_feuilles}")

rapport.save(FICHIER_SORTIE)

print(f"Rapport créé : {FICHIER_SORTIE}")
print("=" * 70)
print("ANALYSE TERMINEE")
print("=" * 70)