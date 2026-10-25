from pathlib import Path
from docx import Document
from openpyxl import Workbook


DOSSIER_MODELES = Path("modeles")
DOSSIER_SORTIE = Path("sorties")

FICHIER_SORTIE = DOSSIER_SORTIE / "inventaire_word.xlsx"

DOSSIER_SORTIE.mkdir(exist_ok=True)


# --------------------------------------------------
# RECHERCHE DU DOCUMENT WORD
# --------------------------------------------------

fichiers_word = [
    fichier
    for fichier in DOSSIER_MODELES.glob("*.docx")
    if not fichier.name.startswith("~$")
]


if not fichiers_word:
    raise FileNotFoundError(
        "Aucun fichier .docx trouvé dans le dossier modeles."
    )


if len(fichiers_word) > 1:
    print("ATTENTION : plusieurs fichiers Word trouvés.")

    for fichier in fichiers_word:
        print(f" - {fichier.name}")


fichier_word = fichiers_word[0]


print("=" * 72)
print("ANALYSE DU MODELE WORD HCP")
print("=" * 72)

print()
print(f"Document analysé : {fichier_word.name}")


# --------------------------------------------------
# OUVERTURE DU DOCUMENT
# --------------------------------------------------

document = Document(fichier_word)


# --------------------------------------------------
# CREATION DU RAPPORT EXCEL
# --------------------------------------------------

rapport = Workbook()


# ==================================================
# FEUILLE RESUME
# ==================================================

resume = rapport.active
resume.title = "Resume"

resume.append([
    "Indicateur",
    "Valeur"
])

resume.append([
    "Document Word",
    fichier_word.name
])

resume.append([
    "Nombre de paragraphes",
    len(document.paragraphs)
])

resume.append([
    "Nombre de tableaux",
    len(document.tables)
])

resume.append([
    "Nombre de sections",
    len(document.sections)
])


# ==================================================
# FEUILLE PARAGRAPHES
# ==================================================

paragraphes = rapport.create_sheet("Paragraphes")

paragraphes.append([
    "N°",
    "Style",
    "Texte"
])


nombre_paragraphes_non_vides = 0


for numero, paragraphe in enumerate(
    document.paragraphs,
    start=1
):

    texte = paragraphe.text.strip()

    if not texte:
        continue

    nombre_paragraphes_non_vides += 1

    style = ""

    if paragraphe.style is not None:
        style = paragraphe.style.name

    paragraphes.append([
        numero,
        style,
        texte
    ])


resume.append([
    "Paragraphes non vides",
    nombre_paragraphes_non_vides
])


# ==================================================
# FEUILLE TABLEAUX
# ==================================================

tableaux = rapport.create_sheet("Tableaux")

tableaux.append([
    "N° tableau",
    "Nombre de lignes",
    "Nombre de colonnes",
    "Style Word",
    "Premier contenu",
    "Nombre de cellules non vides"
])


# ==================================================
# FEUILLE CELLULES
# ==================================================

cellules = rapport.create_sheet("Cellules")

cellules.append([
    "N° tableau",
    "Ligne",
    "Colonne",
    "Contenu"
])


nombre_cellules_non_vides_total = 0


for numero_tableau, tableau in enumerate(
    document.tables,
    start=1
):

    nombre_lignes = len(tableau.rows)

    if nombre_lignes > 0:
        nombre_colonnes = len(tableau.rows[0].cells)
    else:
        nombre_colonnes = 0

    style_tableau = ""

    if tableau.style is not None:
        style_tableau = tableau.style.name

    premier_contenu = ""

    nombre_cellules_non_vides = 0


    for numero_ligne, ligne in enumerate(
        tableau.rows,
        start=1
    ):

        for numero_colonne, cellule in enumerate(
            ligne.cells,
            start=1
        ):

            texte = cellule.text.strip()

            if texte:

                nombre_cellules_non_vides += 1
                nombre_cellules_non_vides_total += 1

                if not premier_contenu:
                    premier_contenu = texte

                cellules.append([
                    numero_tableau,
                    numero_ligne,
                    numero_colonne,
                    texte
                ])


    tableaux.append([
        numero_tableau,
        nombre_lignes,
        nombre_colonnes,
        style_tableau,
        premier_contenu,
        nombre_cellules_non_vides
    ])


resume.append([
    "Cellules Word non vides",
    nombre_cellules_non_vides_total
])


# ==================================================
# FEUILLE SECTIONS
# ==================================================

sections = rapport.create_sheet("Sections")

sections.append([
    "N° section",
    "Largeur page",
    "Hauteur page",
    "Marge gauche",
    "Marge droite",
    "Marge haute",
    "Marge basse"
])


for numero, section in enumerate(
    document.sections,
    start=1
):

    sections.append([
        numero,
        section.page_width.cm,
        section.page_height.cm,
        section.left_margin.cm,
        section.right_margin.cm,
        section.top_margin.cm,
        section.bottom_margin.cm
    ])


# --------------------------------------------------
# MISE EN FORME DU RAPPORT
# --------------------------------------------------

resume.column_dimensions["A"].width = 35
resume.column_dimensions["B"].width = 40

paragraphes.column_dimensions["A"].width = 10
paragraphes.column_dimensions["B"].width = 30
paragraphes.column_dimensions["C"].width = 120

tableaux.column_dimensions["A"].width = 15
tableaux.column_dimensions["B"].width = 20
tableaux.column_dimensions["C"].width = 20
tableaux.column_dimensions["D"].width = 30
tableaux.column_dimensions["E"].width = 100
tableaux.column_dimensions["F"].width = 25

cellules.column_dimensions["A"].width = 15
cellules.column_dimensions["B"].width = 12
cellules.column_dimensions["C"].width = 12
cellules.column_dimensions["D"].width = 120

sections.column_dimensions["A"].width = 15
sections.column_dimensions["B"].width = 20
sections.column_dimensions["C"].width = 20
sections.column_dimensions["D"].width = 20
sections.column_dimensions["E"].width = 20
sections.column_dimensions["F"].width = 20
sections.column_dimensions["G"].width = 20


paragraphes.freeze_panes = "A2"
tableaux.freeze_panes = "A2"
cellules.freeze_panes = "A2"


paragraphes.auto_filter.ref = paragraphes.dimensions
tableaux.auto_filter.ref = tableaux.dimensions
cellules.auto_filter.ref = cellules.dimensions


# --------------------------------------------------
# ENREGISTREMENT
# --------------------------------------------------

rapport.save(FICHIER_SORTIE)


print()
print(f"Paragraphes : {len(document.paragraphs)}")
print(f"Tableaux Word : {len(document.tables)}")
print(f"Sections : {len(document.sections)}")

print()
print(f"Rapport créé : {FICHIER_SORTIE}")

print("=" * 72)
print("ANALYSE WORD TERMINEE")
print("=" * 72)