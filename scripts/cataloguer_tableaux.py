from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
import re
import warnings


warnings.filterwarnings("ignore")


DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_SORTIE = Path("sorties")

FICHIER_SORTIE = DOSSIER_SORTIE / "catalogue_tableaux.xlsx"

DOSSIER_SORTIE.mkdir(exist_ok=True)


def normaliser(texte):
    if texte is None:
        return ""

    return str(texte).strip()


def est_feuille_ignoree(nom):
    nom = nom.strip().lower()

    if nom.startswith("pg"):
        return True

    if "sommaire" in nom:
        return True

    if (
        "graph" in nom
        or "grraph" in nom
        or "graphique" in nom
    ):
        return True

    return False


def cellules_reellement_utilisees(feuille):
    """
    Utilise uniquement les cellules réellement présentes dans
    le fichier Excel.

    Ceci évite les faux max_row = 1048576 ou
    max_column = 16384 provoqués par du formatage Excel.
    """

    cellules = []

    for cellule in feuille._cells.values():

        if cellule.value not in (None, ""):
            cellules.append(cellule)

    return cellules


def trouver_titre(cellules):
    candidats = []

    for cellule in cellules:

        if cellule.row > 15:
            continue

        texte = normaliser(cellule.value)

        if len(texte) < 8:
            continue

        # Les titres HCP contiennent souvent un numéro de tableau
        if re.search(r"\d+\s*[-–]\s*\d+", texte):
            candidats.append(
                (cellule.row, texte)
            )

    if candidats:
        candidats.sort(key=lambda x: x[0])
        return candidats[0][1]

    # Deuxième stratégie :
    # premier texte suffisamment long en haut de la feuille
    candidats = []

    for cellule in cellules:

        if cellule.row > 12:
            continue

        texte = normaliser(cellule.value)

        if len(texte) >= 20:
            candidats.append(
                (cellule.row, texte)
            )

    if candidats:
        candidats.sort(key=lambda x: x[0])
        return candidats[0][1]

    return ""


def trouver_annees(cellules):
    annees = set()

    motif = re.compile(
        r"\b(?:19|20)\d{2}"
        r"(?:\s*[-/]\s*(?:19|20)?\d{2})?\b"
    )

    for cellule in cellules:

        # Les années sont généralement dans le haut
        if cellule.row > 25:
            continue

        texte = normaliser(cellule.value)

        for resultat in motif.findall(texte):
            annees.add(resultat)

    return ", ".join(sorted(annees))


def trouver_unite(cellules):
    mots = [
        "en milliers",
        "en millions",
        "en millier",
        "pourcentage",
        "en %",
        "(%)",
        "unité",
        "unites",
        "unités",
        "dh",
        "dirham"
    ]

    for cellule in cellules:

        if cellule.row > 25:
            continue

        texte = normaliser(cellule.value)

        texte_min = texte.lower()

        if any(mot in texte_min for mot in mots):
            return texte

    return ""


def trouver_source(cellules):
    for cellule in cellules:

        texte = normaliser(cellule.value)

        texte_min = texte.lower()

        if (
            texte_min.startswith("source")
            or "source :" in texte_min
            or "المصدر" in texte
        ):
            return texte

    return ""


def contient_meknes(cellules):
    for cellule in cellules:

        texte = normaliser(cellule.value)

        if re.search(
            r"\bmekn[eè]s\b",
            texte,
            re.IGNORECASE
        ):
            return True

        if "مكناس" in texte:
            return True

    return False


def contient_region_fes_meknes(cellules):
    variantes = [
        "fès - meknès",
        "fès-meknès",
        "fès -meknès",
        "fès- meknès",
        "fes - meknes",
        "fes-meknes",
        "فاس ــ مكناس",
        "فاس - مكناس"
    ]

    for cellule in cellules:

        texte = normaliser(cellule.value).lower()

        if any(v.lower() in texte for v in variantes):
            return True

    return False


def detecter_niveau_geographique(cellules):

    texte_global = " ".join(
        normaliser(cellule.value).lower()
        for cellule in cellules
        if cellule.row <= 20
    )

    niveaux = []

    if "région" in texte_global or "region" in texte_global:
        niveaux.append("Région")

    if "province" in texte_global:
        niveaux.append("Province")

    if "préfecture" in texte_global or "prefecture" in texte_global:
        niveaux.append("Préfecture")

    if "commune" in texte_global:
        niveaux.append("Commune")

    if not niveaux:
        return "Non détecté"

    return " / ".join(niveaux)


fichiers_excel = sorted(
    fichier
    for fichier in DOSSIER_SOURCE.glob("*.xlsx")
    if not fichier.name.startswith("~$")
)


rapport = Workbook()

catalogue = rapport.active
catalogue.title = "Catalogue"

catalogue.append([
    "N°",
    "Fichier",
    "Feuille",
    "Titre probable",
    "Année(s)",
    "Unité",
    "Meknès détecté",
    "Fès-Meknès détecté",
    "Niveau géographique",
    "Source",
    "Dernière ligne réelle",
    "Dernière colonne réelle",
    "Lettre dernière colonne",
    "Cellules avec données",
    "Cellules fusionnées"
])


resume = rapport.create_sheet("Resume")

resume.append([
    "Indicateur",
    "Valeur"
])


anomalies = rapport.create_sheet("Anomalies")

anomalies.append([
    "Fichier",
    "Feuille",
    "Problème détecté"
])


print("=" * 72)
print("CATALOGAGE INTELLIGENT DES TABLEAUX HCP")
print("=" * 72)


numero_tableau = 0
nombre_meknes = 0
nombre_region = 0
nombre_anomalies = 0


for numero_fichier, fichier in enumerate(
    fichiers_excel,
    start=1
):

    print(
        f"[{numero_fichier}/{len(fichiers_excel)}] "
        f"{fichier.name}"
    )

    try:

        classeur = load_workbook(
            fichier,
            data_only=False,
            read_only=False
        )

        for feuille in classeur.worksheets:

            if est_feuille_ignoree(feuille.title):
                continue

            numero_tableau += 1

            cellules = cellules_reellement_utilisees(
                feuille
            )

            if cellules:

                derniere_ligne = max(
                    cellule.row
                    for cellule in cellules
                )

                derniere_colonne = max(
                    cellule.column
                    for cellule in cellules
                )

            else:

                derniere_ligne = 0
                derniere_colonne = 0

            meknes = contient_meknes(cellules)

            fes_meknes = contient_region_fes_meknes(
                cellules
            )

            if meknes:
                nombre_meknes += 1

            if fes_meknes:
                nombre_region += 1

            titre = trouver_titre(cellules)
            annees = trouver_annees(cellules)
            unite = trouver_unite(cellules)
            source = trouver_source(cellules)

            niveau = detecter_niveau_geographique(
                cellules
            )

            if derniere_colonne > 100:

                anomalies.append([
                    fichier.name,
                    feuille.title,
                    (
                        "Nombre réel de colonnes "
                        f"élevé : {derniere_colonne}"
                    )
                ])

                nombre_anomalies += 1

            if derniere_ligne > 5000:

                anomalies.append([
                    fichier.name,
                    feuille.title,
                    (
                        "Nombre réel de lignes "
                        f"élevé : {derniere_ligne}"
                    )
                ])

                nombre_anomalies += 1

            catalogue.append([
                numero_tableau,
                fichier.name,
                feuille.title,
                titre,
                annees,
                unite,
                "OUI" if meknes else "NON",
                "OUI" if fes_meknes else "NON",
                niveau,
                source,
                derniere_ligne,
                derniere_colonne,
                (
                    get_column_letter(derniere_colonne)
                    if derniere_colonne > 0
                    else ""
                ),
                len(cellules),
                len(feuille.merged_cells.ranges)
            ])

        classeur.close()

    except Exception as erreur:

        anomalies.append([
            fichier.name,
            "",
            str(erreur)
        ])

        nombre_anomalies += 1

        print(
            f"   ERREUR : {erreur}"
        )


resume.append([
    "Fichiers Excel analysés",
    len(fichiers_excel)
])

resume.append([
    "Tableaux analysés",
    numero_tableau
])

resume.append([
    "Tableaux contenant Meknès",
    nombre_meknes
])

resume.append([
    "Tableaux contenant Fès-Meknès",
    nombre_region
])

resume.append([
    "Anomalies détectées",
    nombre_anomalies
])


# Largeur des colonnes du catalogue
largeurs = {
    "A": 8,
    "B": 45,
    "C": 25,
    "D": 65,
    "E": 20,
    "F": 35,
    "G": 18,
    "H": 22,
    "I": 28,
    "J": 65,
    "K": 20,
    "L": 22,
    "M": 22,
    "N": 22,
    "O": 20
}

for colonne, largeur in largeurs.items():
    catalogue.column_dimensions[
        colonne
    ].width = largeur


catalogue.freeze_panes = "A2"
catalogue.auto_filter.ref = catalogue.dimensions

resume.column_dimensions["A"].width = 35
resume.column_dimensions["B"].width = 20

anomalies.column_dimensions["A"].width = 45
anomalies.column_dimensions["B"].width = 25
anomalies.column_dimensions["C"].width = 60


rapport.save(FICHIER_SORTIE)


print()
print(f"Tableaux catalogués : {numero_tableau}")
print(f"Tableaux avec Meknès : {nombre_meknes}")
print(
    f"Tableaux avec Fès-Meknès : "
    f"{nombre_region}"
)
print(f"Anomalies : {nombre_anomalies}")

print()
print(
    f"Fichier créé : {FICHIER_SORTIE}"
)

print("=" * 72)
print("CATALOGAGE TERMINE")
print("=" * 72)