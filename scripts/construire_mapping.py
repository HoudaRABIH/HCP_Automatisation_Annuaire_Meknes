from pathlib import Path
from collections import Counter, defaultdict
from difflib import SequenceMatcher
from openpyxl import load_workbook, Workbook
from docx import Document
import re
import unicodedata
import warnings

warnings.filterwarnings("ignore")

# ============================================================
# CHEMINS
# ============================================================

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_MODELES = Path("modeles")
DOSSIER_SORTIE = Path("sorties")

FICHIER_CATALOGUE = DOSSIER_SORTIE / "catalogue_tables_reelles.xlsx"
FICHIER_SORTIE = DOSSIER_SORTIE / "mapping_candidats.xlsx"

DOSSIER_SORTIE.mkdir(exist_ok=True)

# ============================================================
# CONSTANTES
# ============================================================

MOTIF_TITRE_WORD = re.compile(
    r"^\s*(\d{1,2})\s*[-–—]\s*(\d{1,3})(?!\d)"
)

MOTIF_NOMBRE = re.compile(
    r"[-+]?\d+(?:[\s\u00A0]\d{3})*(?:[.,]\d+)?"
)

MOTS_VIDES = {
    "de", "du", "des", "la", "le", "les", "un", "une",
    "et", "en", "par", "selon", "au", "aux", "a", "à",
    "dans", "d", "l", "pour", "sur", "avec", "sans",
    "plus", "moins", "total", "totale", "totaux",
    "annee", "annees", "année", "années",
    "nombre", "effectif", "effectifs",
    "source", "tableau", "tableaux",
}

# ============================================================
# NORMALISATION
# ============================================================

def sans_accents(texte):
    texte = str(texte or "").replace("œ", "oe").replace("Œ", "OE")
    return "".join(
        caractere
        for caractere in unicodedata.normalize("NFD", texte)
        if unicodedata.category(caractere) != "Mn"
    )


def normaliser(texte):
    texte = sans_accents(texte).lower()
    texte = texte.replace("\n", " ")
    texte = texte.replace("’", "'")
    texte = re.sub(r"\s+", " ", texte)
    return texte.strip()


def normaliser_titre(texte):
    texte = normaliser(texte)

    # Retirer le numéro Word/Excel au début : 2-5, 12 - 6, etc.
    texte = re.sub(
        r"^\s*\d{1,2}\s*[-–—]\s*\d{1,3}\s*",
        "",
        texte
    )

    # Retirer les notes (1), (2), etc.
    texte = re.sub(r"\(\s*\d+\s*\)", " ", texte)

    texte = re.sub(r"[^a-z0-9%]+", " ", texte)
    texte = re.sub(r"\s+", " ", texte)

    return texte.strip()


def mots_significatifs(texte):
    mots = set()

    for mot in normaliser_titre(texte).split():
        if len(mot) < 2:
            continue
        if mot in MOTS_VIDES:
            continue
        mots.add(mot)

    return mots


def similarite_texte(texte1, texte2):
    a = normaliser_titre(texte1)
    b = normaliser_titre(texte2)

    if not a or not b:
        return 0.0

    sequence = SequenceMatcher(None, a, b).ratio() * 100

    mots_a = mots_significatifs(texte1)
    mots_b = mots_significatifs(texte2)

    if mots_a and mots_b:
        intersection = len(mots_a & mots_b)
        denominateur = min(len(mots_a), len(mots_b))
        chevauchement = (intersection / denominateur) * 100
    else:
        chevauchement = 0.0

    inclusion = 0.0
    if min(len(a), len(b)) >= 12:
        if a in b or b in a:
            inclusion = 100.0

    score = (
        0.50 * sequence
        + 0.40 * chevauchement
        + 0.10 * inclusion
    )

    return min(100.0, score)


# ============================================================
# NOMBRES / EMPREINTES
# ============================================================

def normaliser_nombre(valeur):
    try:
        nombre = float(valeur)
    except Exception:
        return None

    # Exclure les années : elles sont trop communes pour identifier un tableau
    if nombre.is_integer() and 1900 <= nombre <= 2035:
        return None

    # Exclure quelques petits entiers très fréquents
    if nombre.is_integer() and abs(nombre) <= 2:
        return None

    if nombre.is_integer():
        return str(int(nombre))

    return f"{nombre:.4f}".rstrip("0").rstrip(".")


def extraire_nombres_valeur(valeur):
    resultat = []

    if valeur is None:
        return resultat

    if isinstance(valeur, bool):
        return resultat

    if isinstance(valeur, (int, float)):
        nombre = normaliser_nombre(valeur)
        if nombre is not None:
            resultat.append(nombre)
        return resultat

    texte = str(valeur).replace("\u00A0", " ")

    for correspondance in MOTIF_NOMBRE.finditer(texte):
        brut = correspondance.group(0)
        brut = brut.replace(" ", "")
        brut = brut.replace(",", ".")

        try:
            nombre = normaliser_nombre(float(brut))
        except Exception:
            nombre = None

        if nombre is not None:
            resultat.append(nombre)

    return resultat


def empreinte_valeurs(valeurs):
    nombres = Counter()
    mots = set()

    for valeur in valeurs:
        if valeur in (None, ""):
            continue

        for nombre in extraire_nombres_valeur(valeur):
            nombres[nombre] += 1

        texte = str(valeur)

        for mot in mots_significatifs(texte):
            mots.add(mot)

    return nombres, mots


def score_nombres(nombres_word, nombres_excel):
    if not nombres_word:
        return 0.0

    intersection = nombres_word & nombres_excel

    nb_commun = sum(intersection.values())
    nb_word = sum(nombres_word.values())

    if nb_word == 0:
        return 0.0

    return (nb_commun / nb_word) * 100


def score_mots(mots_word, mots_excel):
    if not mots_word:
        return 0.0

    communs = mots_word & mots_excel

    return (len(communs) / len(mots_word)) * 100


# ============================================================
# THEMES / FICHIERS
# ============================================================

def nom_theme_fichier(nom):
    texte = re.sub(r"\.xlsx$", "", nom, flags=re.IGNORECASE)
    texte = re.sub(r"^\s*\d+\s*[.\-]?\s*", "", texte)
    texte = re.sub(
        r"_?\s*AS_?\s*20\d{2,3}.*$",
        "",
        texte,
        flags=re.IGNORECASE
    )
    texte = texte.replace("_", " ")
    return texte.strip()


def score_theme_fichier(theme, nom_fichier):
    return similarite_texte(
        theme,
        nom_theme_fichier(nom_fichier)
    )


# ============================================================
# LECTURE DU CATALOGUE EXCEL
# ============================================================

if not FICHIER_CATALOGUE.exists():
    raise FileNotFoundError(
        f"Fichier introuvable : {FICHIER_CATALOGUE}"
    )

catalogue_wb = load_workbook(
    FICHIER_CATALOGUE,
    data_only=True,
    read_only=True
)

ws = catalogue_wb["Tables_logiques"]

entetes = {
    cellule.value: index
    for index, cellule in enumerate(ws[1], start=1)
}

tables_excel = []

for ligne in ws.iter_rows(min_row=2, values_only=True):
    def valeur(colonne):
        return ligne[entetes[colonne] - 1]

    tables_excel.append({
        "id": valeur("ID"),
        "fichier": valeur("Fichier"),
        "numero": valeur("Numéro complet"),
        "titre": valeur("Titre principal") or "",
        "feuilles": valeur("Feuille(s)") or "",
        "plages": valeur("Plage(s)") or "",
        "segments": valeur("Nombre segments"),
        "meknes": valeur("Meknès détecté"),
        "fes_meknes": valeur("Fès-Meknès détecté"),
        "annees": valeur("Années") or "",
    })

catalogue_wb.close()

fichiers_sources = sorted({
    table["fichier"]
    for table in tables_excel
})


# ============================================================
# OUVERTURE DES EXCEL SOURCES + CACHE DES EMPREINTES
# ============================================================

classeurs_sources = {}
empreintes_excel = {}


def ouvrir_source(nom_fichier):
    if nom_fichier not in classeurs_sources:
        chemin = DOSSIER_SOURCE / nom_fichier

        classeurs_sources[nom_fichier] = load_workbook(
            chemin,
            data_only=True,
            read_only=False
        )

    return classeurs_sources[nom_fichier]


def lire_valeurs_table_excel(table):
    valeurs = []

    classeur = ouvrir_source(
        table["fichier"]
    )

    plages = [
        morceau.strip()
        for morceau in str(table["plages"]).split("|")
        if morceau.strip()
    ]

    for plage in plages:
        if "!" not in plage:
            continue

        nom_feuille, lignes = plage.rsplit("!", 1)

        if ":" not in lignes:
            continue

        debut_txt, fin_txt = lignes.split(":", 1)

        try:
            debut = int(debut_txt)
            fin = int(fin_txt)
        except ValueError:
            continue

        if nom_feuille not in classeur.sheetnames:
            continue

        feuille = classeur[nom_feuille]

        # Les tableaux HCP sont beaucoup moins larges que 80 colonnes.
        # Cette limite évite les feuilles artificiellement formatées
        # jusqu'à XFD.
        max_col = min(feuille.max_column, 80)

        for ligne in feuille.iter_rows(
            min_row=debut,
            max_row=fin,
            min_col=1,
            max_col=max_col
        ):
            for cellule in ligne:
                if cellule.value not in (None, ""):
                    valeurs.append(cellule.value)

    return valeurs


print("=" * 78)
print("PREPARATION DES EMPREINTES DES TABLEAUX EXCEL")
print("=" * 78)

for index, table in enumerate(tables_excel, start=1):
    if index == 1 or index % 50 == 0 or index == len(tables_excel):
        print(
            f"[{index}/{len(tables_excel)}] "
            f"{table['numero']} - {table['fichier']}"
        )

    valeurs = lire_valeurs_table_excel(table)

    nombres, mots = empreinte_valeurs(valeurs)

    empreintes_excel[table["id"]] = {
        "nombres": nombres,
        "mots": mots
    }


# ============================================================
# LECTURE DU WORD
# ============================================================

fichiers_word = [
    fichier
    for fichier in DOSSIER_MODELES.glob("*.docx")
    if not fichier.name.startswith("~$")
]

if not fichiers_word:
    raise FileNotFoundError(
        "Aucun fichier .docx trouvé dans le dossier modeles."
    )

if len(fichiers_word) > 1:
    print()
    print("ATTENTION : plusieurs fichiers Word trouvés.")
    print("Le premier sera utilisé :")
    for fichier in fichiers_word:
        print(f" - {fichier.name}")

fichier_word = fichiers_word[0]

document = Document(fichier_word)


def textes_uniques_table_word(table):
    textes = []
    deja_vus = set()

    for ligne in table.rows:
        for cellule in ligne.cells:
            identifiant = id(cellule._tc)

            if identifiant in deja_vus:
                continue

            deja_vus.add(identifiant)

            texte = cellule.text.strip()

            if texte:
                textes.append(texte)

    return textes


def premier_texte_table_word(table):
    textes = textes_uniques_table_word(table)

    if textes:
        return textes[0]

    return ""


def trouver_prochain_tableau_donnees(index_titre):
    for index in range(
        index_titre + 1,
        len(document.tables)
    ):
        table = document.tables[index]

        # Les tableaux de titres/chapitres ont généralement 1 ligne.
        # Le premier tableau de plusieurs lignes qui suit est considéré
        # comme le tableau de données.
        if len(table.rows) >= 2:
            return index

    return None


tables_word = []
theme_actuel = ""

for index, table in enumerate(document.tables):
    texte = premier_texte_table_word(table)

    if len(table.rows) == 1 and len(table.columns) <= 2:
        resultat = MOTIF_TITRE_WORD.match(texte)

        if resultat:
            index_donnees = trouver_prochain_tableau_donnees(index)

            if index_donnees is None:
                continue

            table_donnees = document.tables[index_donnees]
            valeurs_word = textes_uniques_table_word(table_donnees)

            nombres_word, mots_word = empreinte_valeurs(
                valeurs_word
            )

            tables_word.append({
                "table_titre": index + 1,
                "table_donnees": index_donnees + 1,
                "theme": theme_actuel,
                "numero": (
                    f"{resultat.group(1)}-"
                    f"{resultat.group(2)}"
                ),
                "titre": texte,
                "nb_lignes": len(table_donnees.rows),
                "nb_colonnes": len(table_donnees.columns),
                "nombres": nombres_word,
                "mots": mots_word
            })

        else:
            texte_min = normaliser(texte)

            if (
                texte
                and "liste des tableaux" not in texte_min
                and len(texte) <= 120
            ):
                theme_actuel = texte


# ============================================================
# CALCUL DU MAPPING
# ============================================================

print()
print("=" * 78)
print("CONSTRUCTION DU MAPPING WORD <-> EXCEL")
print("=" * 78)

meilleurs = []
top5_lignes = []


for position, word in enumerate(tables_word, start=1):
    print(
        f"[{position}/{len(tables_word)}] "
        f"{word['numero']} {word['titre'][:55]}"
    )

    scores_fichiers = sorted(
        [
            (
                score_theme_fichier(
                    word["theme"],
                    fichier
                ),
                fichier
            )
            for fichier in fichiers_sources
        ],
        reverse=True
    )

    meilleur_score_theme = (
        scores_fichiers[0][0]
        if scores_fichiers
        else 0
    )

    # Si le thème Word correspond clairement à un fichier,
    # on limite la recherche aux meilleurs fichiers.
    if meilleur_score_theme >= 75:
        fichiers_autorises = {
            scores_fichiers[0][1]
        }
    elif meilleur_score_theme >= 50:
        fichiers_autorises = {
            fichier
            for _, fichier in scores_fichiers[:2]
        }
    else:
        fichiers_autorises = {
            fichier
            for _, fichier in scores_fichiers[:3]
        }

    candidats = []

    for excel in tables_excel:
        if excel["fichier"] not in fichiers_autorises:
            continue

        empreinte = empreintes_excel[
            excel["id"]
        ]

        score_titre = similarite_texte(
            word["titre"],
            excel["titre"]
        )

        score_num = score_nombres(
            word["nombres"],
            empreinte["nombres"]
        )

        score_texte_donnees = score_mots(
            word["mots"],
            empreinte["mots"]
        )

        score_theme = score_theme_fichier(
            word["theme"],
            excel["fichier"]
        )

        # Si le tableau Word contient assez de nombres,
        # les valeurs deviennent le signal principal.
        if sum(word["nombres"].values()) >= 3:
            score_final = (
                0.45 * score_num
                + 0.30 * score_titre
                + 0.15 * score_texte_donnees
                + 0.10 * score_theme
            )
        else:
            score_final = (
                0.55 * score_titre
                + 0.30 * score_texte_donnees
                + 0.15 * score_theme
            )

        # Petit bonus lorsque le Word parle explicitement de Meknès
        # et que le bloc Excel contient également Meknès.
        texte_word_global = (
            word["titre"] + " " + word["theme"]
        )

        if (
            "meknes" in normaliser(texte_word_global)
            and excel["meknes"] == "OUI"
        ):
            score_final = min(
                100.0,
                score_final + 3.0
            )

        candidats.append({
            "score": score_final,
            "score_titre": score_titre,
            "score_nombres": score_num,
            "score_texte": score_texte_donnees,
            "score_theme": score_theme,
            "excel": excel
        })

    candidats.sort(
        key=lambda element: element["score"],
        reverse=True
    )

    top = candidats[:5]

    if not top:
        continue

    meilleur = top[0]

    deuxieme_score = (
        top[1]["score"]
        if len(top) >= 2
        else 0
    )

    marge = (
        meilleur["score"]
        - deuxieme_score
    )

    if (
        meilleur["score"] >= 80
        and marge >= 8
    ):
        statut = "TRES PROBABLE"

    elif (
        meilleur["score"] >= 68
        and marge >= 4
    ):
        statut = "PROBABLE"

    elif meilleur["score"] >= 52:
        statut = "A VERIFIER"

    else:
        statut = "FAIBLE"

    meilleurs.append({
        "word": word,
        "candidat": meilleur,
        "marge": marge,
        "statut": statut
    })

    for rang, candidat in enumerate(
        top,
        start=1
    ):
        excel = candidat["excel"]

        top5_lignes.append([
            word["table_titre"],
            word["table_donnees"],
            word["theme"],
            word["numero"],
            word["titre"],
            rang,
            round(candidat["score"], 1),
            round(candidat["score_titre"], 1),
            round(candidat["score_nombres"], 1),
            round(candidat["score_texte"], 1),
            round(candidat["score_theme"], 1),
            excel["fichier"],
            excel["numero"],
            excel["titre"],
            excel["feuilles"],
            excel["plages"],
            excel["meknes"],
            excel["fes_meknes"],
            excel["annees"]
        ])


# ============================================================
# RAPPORT EXCEL
# ============================================================

rapport = Workbook()

ws_meilleurs = rapport.active
ws_meilleurs.title = "Meilleurs"

ws_meilleurs.append([
    "Table Word titre",
    "Table Word données",
    "Thème Word",
    "N° Word",
    "Titre Word",
    "Score final",
    "Marge sur 2e",
    "Statut",
    "Fichier Excel",
    "N° Excel",
    "Titre Excel",
    "Feuille(s)",
    "Plage(s)",
    "Meknès",
    "Fès-Meknès",
    "Années",
    "Score titre",
    "Score nombres",
    "Score texte données",
    "Score thème"
])

for resultat in meilleurs:
    word = resultat["word"]
    candidat = resultat["candidat"]
    excel = candidat["excel"]

    ws_meilleurs.append([
        word["table_titre"],
        word["table_donnees"],
        word["theme"],
        word["numero"],
        word["titre"],
        round(candidat["score"], 1),
        round(resultat["marge"], 1),
        resultat["statut"],
        excel["fichier"],
        excel["numero"],
        excel["titre"],
        excel["feuilles"],
        excel["plages"],
        excel["meknes"],
        excel["fes_meknes"],
        excel["annees"],
        round(candidat["score_titre"], 1),
        round(candidat["score_nombres"], 1),
        round(candidat["score_texte"], 1),
        round(candidat["score_theme"], 1)
    ])


ws_top5 = rapport.create_sheet("Top5")

ws_top5.append([
    "Table Word titre",
    "Table Word données",
    "Thème Word",
    "N° Word",
    "Titre Word",
    "Rang",
    "Score final",
    "Score titre",
    "Score nombres",
    "Score texte données",
    "Score thème",
    "Fichier Excel",
    "N° Excel",
    "Titre Excel",
    "Feuille(s)",
    "Plage(s)",
    "Meknès",
    "Fès-Meknès",
    "Années"
])

for ligne in top5_lignes:
    ws_top5.append(ligne)


ws_a_verifier = rapport.create_sheet("A_verifier")

ws_a_verifier.append([
    "Table Word titre",
    "Table Word données",
    "Thème Word",
    "N° Word",
    "Titre Word",
    "Score",
    "Marge",
    "Statut",
    "Fichier proposé",
    "N° Excel proposé",
    "Titre Excel proposé"
])

for resultat in meilleurs:
    if resultat["statut"] in {
        "A VERIFIER",
        "FAIBLE"
    }:
        word = resultat["word"]
        candidat = resultat["candidat"]
        excel = candidat["excel"]

        ws_a_verifier.append([
            word["table_titre"],
            word["table_donnees"],
            word["theme"],
            word["numero"],
            word["titre"],
            round(candidat["score"], 1),
            round(resultat["marge"], 1),
            resultat["statut"],
            excel["fichier"],
            excel["numero"],
            excel["titre"]
        ])


# Réutilisations : un même tableau Excel peut alimenter plusieurs
# tableaux Word. Ce n'est pas forcément une erreur, mais il faut le voir.
reutilisations = defaultdict(list)

for resultat in meilleurs:
    excel = resultat["candidat"]["excel"]
    cle = (
        excel["fichier"],
        excel["numero"]
    )
    reutilisations[cle].append(
        resultat["word"]["numero"]
    )


ws_reutilisations = rapport.create_sheet(
    "Reutilisations"
)

ws_reutilisations.append([
    "Fichier Excel",
    "N° Excel",
    "Nombre de tableaux Word",
    "Tableaux Word concernés"
])

for cle, numeros_word in reutilisations.items():
    if len(numeros_word) > 1:
        ws_reutilisations.append([
            cle[0],
            cle[1],
            len(numeros_word),
            ", ".join(numeros_word)
        ])


ws_resume = rapport.create_sheet("Resume")

ws_resume.append([
    "Indicateur",
    "Valeur"
])

ws_resume.append([
    "Document Word",
    fichier_word.name
])

ws_resume.append([
    "Tables Word numérotées détectées",
    len(tables_word)
])

ws_resume.append([
    "Tables Excel logiques disponibles",
    len(tables_excel)
])

for statut in [
    "TRES PROBABLE",
    "PROBABLE",
    "A VERIFIER",
    "FAIBLE"
]:
    ws_resume.append([
        statut,
        sum(
            1
            for resultat in meilleurs
            if resultat["statut"] == statut
        )
    ])


# ============================================================
# MISE EN FORME SIMPLE
# ============================================================

for feuille in [
    ws_meilleurs,
    ws_top5,
    ws_a_verifier,
    ws_reutilisations
]:
    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = feuille.dimensions

largeurs_meilleurs = {
    "A": 16, "B": 18, "C": 28, "D": 12,
    "E": 70, "F": 14, "G": 14, "H": 18,
    "I": 45, "J": 14, "K": 70, "L": 35,
    "M": 55, "N": 12, "O": 16, "P": 25,
    "Q": 14, "R": 16, "S": 20, "T": 14
}

for colonne, largeur in largeurs_meilleurs.items():
    ws_meilleurs.column_dimensions[colonne].width = largeur

ws_top5.column_dimensions["E"].width = 70
ws_top5.column_dimensions["L"].width = 45
ws_top5.column_dimensions["N"].width = 70
ws_top5.column_dimensions["O"].width = 35
ws_top5.column_dimensions["P"].width = 55

ws_a_verifier.column_dimensions["E"].width = 70
ws_a_verifier.column_dimensions["I"].width = 45
ws_a_verifier.column_dimensions["K"].width = 70

ws_reutilisations.column_dimensions["A"].width = 45
ws_reutilisations.column_dimensions["B"].width = 15
ws_reutilisations.column_dimensions["C"].width = 22
ws_reutilisations.column_dimensions["D"].width = 60

ws_resume.column_dimensions["A"].width = 40
ws_resume.column_dimensions["B"].width = 25


rapport.save(FICHIER_SORTIE)


# ============================================================
# FERMETURE DES CLASSEURS SOURCES
# ============================================================

for classeur in classeurs_sources.values():
    classeur.close()


# ============================================================
# RESULTAT
# ============================================================

print()
print("=" * 78)
print("MAPPING TERMINE")
print("=" * 78)

print(
    f"Tables Word détectées : "
    f"{len(tables_word)}"
)

print(
    f"Tables Excel disponibles : "
    f"{len(tables_excel)}"
)

for statut in [
    "TRES PROBABLE",
    "PROBABLE",
    "A VERIFIER",
    "FAIBLE"
]:
    nombre = sum(
        1
        for resultat in meilleurs
        if resultat["statut"] == statut
    )

    print(
        f"{statut} : {nombre}"
    )

print()
print(
    f"Fichier créé : "
    f"{FICHIER_SORTIE}"
)

print("=" * 78)
