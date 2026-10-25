from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


DOSSIER_SORTIE = Path("sorties")

FICHIER_MAPPING = DOSSIER_SORTIE / "mapping_sequence.xlsx"
FICHIER_CATALOGUE = DOSSIER_SORTIE / "catalogue_tables_reelles.xlsx"
FICHIER_SORTIE = DOSSIER_SORTIE / "mapping_final_recettes.xlsx"


# ============================================================
# REGLES DETERMINISTES
# ============================================================
#
# DIRECT      : 1 tableau Word <- 1 tableau Excel
# COMPOSITE   : 1 tableau Word <- plusieurs tableaux Excel
# CONSERVER   : la source détaillée n'est pas présente dans les 23
#               classeurs fournis. Le tableau du modèle reste inchangé.
#
# On évite volontairement tout "fuzzy matching" ici :
# ce fichier devient le contrat de mapping définitif du projet.
# ============================================================


REGLES = {}


def direct(word, source):
    REGLES[word] = {
        "type": "DIRECT",
        "sources": [source],
        "raison": "Correspondance déterministe validée par le thème et la structure."
    }


def composite(word, *sources):
    REGLES[word] = {
        "type": "COMPOSITE",
        "sources": list(sources),
        "raison": "Le tableau Word regroupe plusieurs tableaux Excel sources."
    }


def conserver(word, raison):
    REGLES[word] = {
        "type": "CONSERVER_MODELE",
        "sources": [],
        "raison": raison
    }


# ------------------------------------------------------------
# 1. DIVISION ADMINISTRATIVE / POPULATION
# ------------------------------------------------------------

direct("1-1", "1-1")
direct("1-2", "1-2")

# Le Word regroupe Ensemble + Urbain + Rural.
composite("1-3", "2-5", "2-6", "2-7")

direct("1-4", "2-14")

for code in [
    "1-5", "1-6", "1-7", "1-8", "1-9",
    "1-10", "1-11", "1-12", "1-13", "1-14"
]:
    conserver(
        code,
        "Le Word contient des données RGPH détaillées au niveau communal/"
        "préfectoral qui ne figurent pas dans les 23 classeurs Annuaire fournis. "
        "Conserver le tableau existant évite toute substitution erronée."
    )


# ------------------------------------------------------------
# 2. SANTE
# ------------------------------------------------------------

direct("2-1", "12-1")
composite("2-2", "12-2", "12-3")          # hôpitaux + lits
direct("2-3", "12-10")                    # cliniques privées
direct("2-4", "12-5")
direct("2-5", "12-6")
direct("2-6", "12-7")
direct("2-7", "12-8")
composite("2-8", "12-9", "12-11")         # chirurgiens + cabinets
direct("2-9", "12-14")
direct("2-10", "12-15")
direct("2-11", "12-16")
direct("2-12", "12-17")
direct("2-13", "12-18")
direct("2-14", "12-19")
direct("2-15", "12-20")
direct("2-16", "12-21")
direct("2-17", "12-22")
direct("2-18", "12-23")
direct("2-20", "12-25")
direct("2-21", "12-26")
direct("2-22", "12-27")
composite("2-23", "12-12", "12-13")       # pharmacies + laboratoires

conserver(
    "2-24",
    "La source détaillée du taux de prévalence du handicap n'est pas "
    "présente dans les 23 classeurs fournis."
)


# ------------------------------------------------------------
# 3. EDUCATION RGPH / ENSEIGNEMENT
# ------------------------------------------------------------

for code in ["3-2", "3-3", "3-4"]:
    conserver(
        code,
        "Indicateur RGPH détaillé absent des 23 classeurs Annuaire fournis."
    )

# Préscolaire
direct("3-5", "11-2")
direct("3-6", "11-3")

# Primaire public / privé
direct("3-7", "11-5")
direct("3-8", "11-6")
direct("3-9", "11-10")
direct("3-10", "11-12")
direct("3-11", "11-13")
direct("3-12", "11-15")

# Collégial public / privé
direct("3-13", "11-19")
direct("3-14", "11-20")
direct("3-15", "11-23")
direct("3-16", "11-26")

# Qualifiant public / privé
direct("3-17", "11-30")
direct("3-18", "11-31")
direct("3-19", "11-33")
direct("3-20", "11-36")

# Post-secondaire
direct("3-21", "11-43")

# Supérieur
direct("4-1", "11-47")
direct("4-2", "11-49")
direct("4-4", "11-52")
direct("4-5", "11-54")

# Formation professionnelle
direct("5-1", "11-68")


# ------------------------------------------------------------
# 4. ACTIVITES CULTURELLES ET LOISIRS
# ------------------------------------------------------------

direct("6-1", "16-1")
direct("6-2", "16-2")     # Word "Jardins d'enfants" / source "Crèches"
direct("6-3", "16-3")
direct("6-4", "16-4")
direct("6-5", "16-5")
direct("6-6", "16-8")
direct("6-7", "16-11")
direct("6-8", "16-13")
direct("6-9", "16-15")
direct("6-10", "16-16")
direct("6-11", "16-17")


# ------------------------------------------------------------
# 5. EMPLOI ET SALAIRE
# ------------------------------------------------------------

conserver(
    "7-2",
    "Le tableau SMIG/SMAG n'est pas présent dans le classeur Emploi fourni."
)

conserver(
    "7-3",
    "Le Word contient une ventilation Nation/Région/Préfecture du personnel "
    "de l'Etat qui n'existe pas dans le classeur Emploi fourni."
)


# ------------------------------------------------------------
# 6. JUSTICE
# ------------------------------------------------------------

# Word 8-1 synthétise plusieurs catégories d'auxiliaires.
composite("8-1", "15-2", "15-3", "15-4", "15-5", "15-6", "15-7")

direct("8-2", "15-10")
direct("8-3", "15-12")
direct("8-4", "15-14")

# "Chambres d'appel du TPI Meknès" : pas de table locale équivalente
# suffisamment explicite dans les sources fournies.
conserver(
    "8-5",
    "La table locale détaillée des chambres d'appel du TPI Meknès n'est pas "
    "identifiée comme table source autonome dans le classeur Justice."
)

direct("8-6", "15-15")
direct("8-7", "15-16")
direct("8-8", "15-21")
direct("8-9", "15-22")
direct("8-10", "15-23")
direct("8-11", "15-33")


# ------------------------------------------------------------
# 7. AGRICULTURE / ELEVAGE / FORET
# ------------------------------------------------------------

direct("9-1", "3-3")
composite("9-2", "3-4", "3-5")      # superficies + productions céréales
direct("9-3", "3-6")
direct("9-4", "3-7")
direct("9-5", "3-8")
direct("9-6", "3-10")

direct("9-7", "3-18")
direct("9-8", "3-19")
direct("9-9", "3-20")

direct("9-10", "3-23")
direct("9-11", "3-24")


# ------------------------------------------------------------
# 8. INDUSTRIE / EAU / ENERGIE
# ------------------------------------------------------------

direct("10-1", "6-18")

direct("11-1", "5-20")
direct("11-2", "5-7")


# ------------------------------------------------------------
# 9. TOURISME
# ------------------------------------------------------------

direct("12-1", "10-1")
direct("12-2", "10-2")
direct("12-3", "10-3")
direct("12-4", "10-4")
direct("12-6", "10-8")


# ------------------------------------------------------------
# 10. HABITAT / CONSTRUCTION / FONCIER
# ------------------------------------------------------------

for code in ["13-1", "13-4", "13-6", "13-7", "13-8"]:
    conserver(
        code,
        "Indicateur RGPH Habitat détaillé absent des 23 classeurs fournis."
    )

for code in ["13-9", "13-10"]:
    conserver(
        code,
        "La source Construction/Autorisations de construire n'est pas "
        "présente dans les 23 classeurs fournis."
    )

direct("13-11", "7-2")
direct("13-12", "7-4")
direct("13-13", "7-5")
direct("13-14", "7-6")
direct("13-15", "7-7")
direct("13-16", "7-9")


# ------------------------------------------------------------
# 11. TRANSPORT
# ------------------------------------------------------------

direct("14-1", "8-8")
direct("14-2", "8-9")


# ------------------------------------------------------------
# 12. PRIX
# ------------------------------------------------------------

# Correction importante : Meknès = source 17-11, et non Agadir 17-3.
direct("15-1", "17-11")
direct("15-2", "17-23")
direct("15-3", "17-25")
direct("15-4", "17-26")


# ------------------------------------------------------------
# 13. ENVIRONNEMENT
# ------------------------------------------------------------

direct("16-1", "23-3")
direct("16-2", "23-4")
direct("16-3", "23-5")
direct("16-4", "23-6")
direct("16-5", "23-7")
direct("16-6", "23-8")


# ============================================================
# CHARGEMENT DU MAPPING WORD
# ============================================================

if not FICHIER_MAPPING.exists():
    raise FileNotFoundError(
        f"Introuvable : {FICHIER_MAPPING}"
    )

wb_mapping = load_workbook(
    FICHIER_MAPPING,
    data_only=True
)

ws_mapping = wb_mapping["Mapping_sequence"]

entetes_mapping = {
    cellule.value: cellule.column
    for cellule in ws_mapping[1]
}


def vm(ligne, nom):
    return ws_mapping.cell(
        row=ligne,
        column=entetes_mapping[nom]
    ).value


lignes_word = []

for ligne in range(
    2,
    ws_mapping.max_row + 1
):
    lignes_word.append({
        "table_word_titre": vm(ligne, "Table Word titre"),
        "table_word_donnees": vm(ligne, "Table Word données"),
        "theme": vm(ligne, "Thème Word") or "",
        "numero_word": str(vm(ligne, "N° Word") or "").strip(),
        "titre_word": vm(ligne, "Titre Word") or "",
    })

wb_mapping.close()


# ============================================================
# CHARGEMENT DU CATALOGUE EXCEL
# ============================================================

if not FICHIER_CATALOGUE.exists():
    raise FileNotFoundError(
        f"Introuvable : {FICHIER_CATALOGUE}"
    )

wb_cat = load_workbook(
    FICHIER_CATALOGUE,
    data_only=True
)

ws_cat = wb_cat["Tables_logiques"]

entetes_cat = {
    cellule.value: cellule.column
    for cellule in ws_cat[1]
}


def vc(ligne, nom):
    return ws_cat.cell(
        row=ligne,
        column=entetes_cat[nom]
    ).value


catalogue = {}

for ligne in range(
    2,
    ws_cat.max_row + 1
):
    numero = str(
        vc(ligne, "Numéro complet") or ""
    ).strip()

    catalogue[numero] = {
        "fichier": vc(ligne, "Fichier") or "",
        "titre": vc(ligne, "Titre principal") or "",
        "feuilles": vc(ligne, "Feuille(s)") or "",
        "plages": vc(ligne, "Plage(s)") or "",
        "meknes": vc(ligne, "Meknès détecté") or "",
        "fes_meknes": vc(ligne, "Fès-Meknès détecté") or "",
        "annees": vc(ligne, "Années") or "",
    }

wb_cat.close()


# ============================================================
# VERIFICATION : 100 % DES TABLEAUX WORD DOIVENT AVOIR UNE REGLE
# ============================================================

codes_word = [
    ligne["numero_word"]
    for ligne in lignes_word
]

codes_uniques = set(codes_word)

manquants = sorted(
    code
    for code in codes_uniques
    if code not in REGLES
)

if manquants:
    raise RuntimeError(
        "REGLES MANQUANTES POUR : "
        + ", ".join(manquants)
    )


# Vérifier que toutes les tables Excel citées existent.
sources_absentes = []

for code, regle in REGLES.items():
    for source in regle["sources"]:
        if source not in catalogue:
            sources_absentes.append(
                f"{code} -> {source}"
            )

if sources_absentes:
    raise RuntimeError(
        "TABLES EXCEL CITEES MAIS ABSENTES DU CATALOGUE : "
        + " | ".join(sources_absentes)
    )


# ============================================================
# CREATION DU RAPPORT FINAL
# ============================================================

rapport = Workbook()

ws = rapport.active
ws.title = "Mapping_final"

ws.append([
    "Table Word titre",
    "Table Word données",
    "Thème Word",
    "N° Word",
    "Titre Word",
    "Type de recette",
    "Sources Excel",
    "Fichier(s) source",
    "Titre(s) source",
    "Feuille(s)",
    "Plage(s)",
    "Année(s) source",
    "Meknès détecté",
    "Décision",
    "Raison / règle"
])


nb_direct = 0
nb_composite = 0
nb_conserver = 0


for ligne in lignes_word:
    code = ligne["numero_word"]
    regle = REGLES[code]

    sources = regle["sources"]

    if regle["type"] == "DIRECT":
        nb_direct += 1
        decision = "MISE A JOUR AUTOMATIQUE"

    elif regle["type"] == "COMPOSITE":
        nb_composite += 1
        decision = "MISE A JOUR PAR RECETTE"

    else:
        nb_conserver += 1
        decision = "CONSERVER LE MODELE"

    fichiers = []
    titres = []
    feuilles = []
    plages = []
    annees = []
    meknes = []

    for source in sources:
        info = catalogue[source]

        fichiers.append(
            info["fichier"]
        )

        titres.append(
            f"{source} — {info['titre']}"
        )

        feuilles.append(
            f"{source}: {info['feuilles']}"
        )

        plages.append(
            f"{source}: {info['plages']}"
        )

        if info["annees"]:
            annees.append(
                str(info["annees"])
            )

        meknes.append(
            f"{source}: {info['meknes']}"
        )

    ws.append([
        ligne["table_word_titre"],
        ligne["table_word_donnees"],
        ligne["theme"],
        code,
        ligne["titre_word"],
        regle["type"],
        " + ".join(sources),
        " | ".join(dict.fromkeys(fichiers)),
        " | ".join(titres),
        " | ".join(feuilles),
        " | ".join(plages),
        " | ".join(dict.fromkeys(annees)),
        " | ".join(meknes),
        decision,
        regle["raison"]
    ])


# ============================================================
# RESUME
# ============================================================

resume = rapport.create_sheet("Resume")

resume.append([
    "Indicateur",
    "Valeur"
])

resume.append([
    "Tableaux Word analysés",
    len(lignes_word)
])

resume.append([
    "Codes Word uniques",
    len(codes_uniques)
])

resume.append([
    "DIRECT",
    nb_direct
])

resume.append([
    "COMPOSITE",
    nb_composite
])

resume.append([
    "CONSERVER_MODELE",
    nb_conserver
])

resume.append([
    "Tables sans règle",
    0
])

resume.append([
    "Sources Excel citées mais absentes",
    0
])


# ============================================================
# SOURCES A FOURNIR PLUS TARD
# ============================================================

manquantes = rapport.create_sheet("Sources_manquantes")

manquantes.append([
    "N° Word",
    "Thème",
    "Titre Word",
    "Motif"
])

for ligne in lignes_word:
    regle = REGLES[
        ligne["numero_word"]
    ]

    if regle["type"] == "CONSERVER_MODELE":
        manquantes.append([
            ligne["numero_word"],
            ligne["theme"],
            ligne["titre_word"],
            regle["raison"]
        ])


# ============================================================
# MISE EN FORME
# ============================================================

bleu = "1F4E78"
bleu_clair = "D9EAF7"
vert = "E2F0D9"
orange = "FCE4D6"
gris = "E7E6E6"
blanc = "FFFFFF"

header_fill = PatternFill(
    "solid",
    fgColor=bleu
)

thin = Side(
    style="thin",
    color="BFBFBF"
)

for cellule in ws[1]:
    cellule.fill = header_fill
    cellule.font = Font(
        bold=True,
        color=blanc
    )
    cellule.alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

for ligne in range(
    2,
    ws.max_row + 1
):
    type_recette = ws.cell(
        ligne,
        6
    ).value

    if type_recette == "DIRECT":
        remplissage = PatternFill(
            "solid",
            fgColor=vert
        )

    elif type_recette == "COMPOSITE":
        remplissage = PatternFill(
            "solid",
            fgColor=bleu_clair
        )

    else:
        remplissage = PatternFill(
            "solid",
            fgColor=orange
        )

    for colonne in range(
        1,
        ws.max_column + 1
    ):
        cellule = ws.cell(
            ligne,
            colonne
        )

        cellule.border = Border(
            bottom=thin
        )

        cellule.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

    ws.cell(
        ligne,
        6
    ).fill = remplissage

    ws.cell(
        ligne,
        14
    ).fill = remplissage


largeurs = {
    "A": 16,
    "B": 18,
    "C": 30,
    "D": 12,
    "E": 65,
    "F": 22,
    "G": 30,
    "H": 45,
    "I": 70,
    "J": 45,
    "K": 65,
    "L": 24,
    "M": 25,
    "N": 28,
    "O": 85
}

for colonne, largeur in largeurs.items():
    ws.column_dimensions[
        colonne
    ].width = largeur


ws.freeze_panes = "A2"
ws.auto_filter.ref = ws.dimensions


for cellule in resume[1]:
    cellule.fill = header_fill
    cellule.font = Font(
        bold=True,
        color=blanc
    )

resume.column_dimensions[
    "A"
].width = 42

resume.column_dimensions[
    "B"
].width = 25


for cellule in manquantes[1]:
    cellule.fill = header_fill
    cellule.font = Font(
        bold=True,
        color=blanc
    )

manquantes.column_dimensions[
    "A"
].width = 12

manquantes.column_dimensions[
    "B"
].width = 32

manquantes.column_dimensions[
    "C"
].width = 70

manquantes.column_dimensions[
    "D"
].width = 100

manquantes.freeze_panes = "A2"


rapport.save(
    FICHIER_SORTIE
)


# ============================================================
# RESULTAT
# ============================================================

print("=" * 78)
print("MAPPING FINAL PAR RECETTES")
print("=" * 78)

print(
    f"Tableaux Word : {len(lignes_word)}"
)

print(
    f"DIRECT : {nb_direct}"
)

print(
    f"COMPOSITE : {nb_composite}"
)

print(
    f"CONSERVER_MODELE : {nb_conserver}"
)

print(
    "Tables sans règle : 0"
)

print(
    "Sources citées absentes : 0"
)

print()
print(
    f"Fichier créé : {FICHIER_SORTIE}"
)

print("=" * 78)
print("MAPPING FINAL TERMINE")
print("=" * 78)
