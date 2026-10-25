from pathlib import Path
from difflib import SequenceMatcher
from openpyxl import load_workbook, Workbook
import re
import unicodedata


DOSSIER_SORTIE = Path("sorties")

FICHIER_MAPPING = DOSSIER_SORTIE / "mapping_candidats.xlsx"
FICHIER_CATALOGUE = DOSSIER_SORTIE / "catalogue_tables_reelles.xlsx"
FICHIER_SORTIE = DOSSIER_SORTIE / "mapping_revu.xlsx"


# ============================================================
# ROUTAGE DES THEMES WORD VERS LES FICHIERS EXCEL ATTENDUS
# ============================================================

ROUTAGE = {
    "DIVISION ADMINISTRATIVE": [
        "1.Division Administrative_AS 2025.xlsx"
    ],
    "POPULATION": [
        "2. population_AS 2025.xlsx"
    ],
    "SANTE": [
        "12. Santé_AS 2025.xlsx"
    ],
    "EDUCATION": [
        "2. population_AS 2025.xlsx"
    ],
    "ENSEINEMENT PRESCOLAIRE": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT PRIMAIRE PUBLIC": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT PRIMAIRE PRIVE": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT SECONDAIRE COLLEGIAL PUBLIC": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT SECONDAIRE COLLEGIAL PRIVE": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT SECONDAIRE QUALIFIANT PUBLIC": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT SECONDAIRE QUALIFIANT PRIVE": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT POSTE - SECONDAIRE": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ENSEIGNEMENT SUPERIEUR PUBLIC": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "FORMATION PROFESSIONNELLE": [
        "11. Enseignement_AS 2025.xlsx"
    ],
    "ACTIVITES CULTURELLES ET LOISIRS": [
        "16. Activités culturelles_Loisirs et Entraide_AS 2025.xlsx"
    ],
    "EMPLOI ET SALAIRE": [
        "14-Emploi et salaire_AS 2025.xlsx"
    ],
    "JUSTICE": [
        "15.Justice et Sauvegarde de l'Enfance _AS2025.xlsx"
    ],
    "AGRICULTURE": [
        "3. Agriculture_AS 20252.xlsx"
    ],
    "ELEVAGE": [
        "3. Agriculture_AS 20252.xlsx"
    ],
    "FORET": [
        "3. Agriculture_AS 20252.xlsx"
    ],
    "INDUSTRIE": [
        "6. Industrie et artisanat_AS 2025.xlsx"
    ],
    "EAU POTABLE": [
        "5.Energie et Eau_AS 2025.xlsx"
    ],
    "ENERGIE": [
        "5.Energie et Eau_AS 2025.xlsx"
    ],
    "TOURISME": [
        "10. Tourisme_AS 2025.xlsx"
    ],
    "CONDITION D’HABITAT DES MENAGES": [
        "2. population_AS 2025.xlsx"
    ],
    "CONSTRUCTION": [
        "7. Foncier_AS 2025.xlsx"
    ],
    "FONCIER": [
        "7. Foncier_AS 2025.xlsx"
    ],
    "TRANSPORT ROUTIER": [
        "8. Transport_AS 2025.xlsx"
    ],
    "PRIX": [
        "17.Indice des prix_AS  2025.xlsx"
    ],
    "ENVIRONNEMENT": [
        "23. Environnement_AS 2025.xlsx"
    ],
}


MOTS_VIDES = {
    "de", "du", "des", "la", "le", "les",
    "un", "une", "et", "en", "par", "selon",
    "au", "aux", "a", "à", "dans", "pour",
    "sur", "avec", "sans", "nombre", "effectif",
    "effectifs", "total", "totale", "totaux",
    "annee", "annees", "année", "années",
}


def sans_accents(texte):
    texte = str(texte or "")
    texte = texte.replace("œ", "oe").replace("Œ", "OE")

    return "".join(
        caractere
        for caractere in unicodedata.normalize("NFD", texte)
        if unicodedata.category(caractere) != "Mn"
    )


def normaliser(texte):
    texte = sans_accents(texte).lower()
    texte = texte.replace("\n", " ")
    texte = texte.replace("’", "'")

    # Retirer le numéro de tableau au début
    texte = re.sub(
        r"^\s*\d{1,2}\s*[-–—]\s*\d{1,3}\s*",
        "",
        texte
    )

    # Retirer les notes (1), (2), ...
    texte = re.sub(
        r"\(\s*\d+\s*\)",
        " ",
        texte
    )

    texte = re.sub(
        r"[^a-z0-9%]+",
        " ",
        texte
    )

    texte = re.sub(
        r"\s+",
        " ",
        texte
    )

    return texte.strip()


def normaliser_theme(texte):
    texte = normaliser(texte)
    return texte.upper()


def mots_significatifs(texte):
    resultat = set()

    for mot in normaliser(texte).split():
        if len(mot) < 3:
            continue

        if mot in MOTS_VIDES:
            continue

        resultat.add(mot)

    return resultat


def score_titre(titre_word, titre_excel):
    a = normaliser(titre_word)
    b = normaliser(titre_excel)

    if not a or not b:
        return 0.0

    sequence = (
        SequenceMatcher(
            None,
            a,
            b
        ).ratio()
        * 100
    )

    mots_a = mots_significatifs(
        titre_word
    )

    mots_b = mots_significatifs(
        titre_excel
    )

    if mots_a and mots_b:
        chevauchement = (
            len(mots_a & mots_b)
            / min(
                len(mots_a),
                len(mots_b)
            )
            * 100
        )
    else:
        chevauchement = 0.0

    inclusion = 0.0

    if min(len(a), len(b)) >= 12:
        if a in b or b in a:
            inclusion = 100.0

    score = (
        0.50 * sequence
        + 0.40 * chevauchement
        + 0.10 * inclusion
    )

    return min(
        100.0,
        score
    )


# ============================================================
# CHARGEMENT DU MAPPING ACTUEL
# ============================================================

if not FICHIER_MAPPING.exists():
    raise FileNotFoundError(
        f"Introuvable : {FICHIER_MAPPING}"
    )

mapping_wb = load_workbook(
    FICHIER_MAPPING,
    data_only=True
)

mapping_ws = mapping_wb["Meilleurs"]

entetes_mapping = {
    cellule.value: cellule.column
    for cellule in mapping_ws[1]
}


def valeur_mapping(ligne, nom):
    return mapping_ws.cell(
        row=ligne,
        column=entetes_mapping[nom]
    ).value


# ============================================================
# CHARGEMENT DU CATALOGUE DES TABLES EXCEL
# ============================================================

if not FICHIER_CATALOGUE.exists():
    raise FileNotFoundError(
        f"Introuvable : {FICHIER_CATALOGUE}"
    )

catalogue_wb = load_workbook(
    FICHIER_CATALOGUE,
    data_only=True
)

catalogue_ws = catalogue_wb[
    "Tables_logiques"
]

entetes_catalogue = {
    cellule.value: cellule.column
    for cellule in catalogue_ws[1]
}


def valeur_catalogue(ligne, nom):
    return catalogue_ws.cell(
        row=ligne,
        column=entetes_catalogue[nom]
    ).value


tables_excel = []

for ligne in range(
    2,
    catalogue_ws.max_row + 1
):
    tables_excel.append({
        "id": valeur_catalogue(
            ligne,
            "ID"
        ),
        "fichier": valeur_catalogue(
            ligne,
            "Fichier"
        ),
        "numero": valeur_catalogue(
            ligne,
            "Numéro complet"
        ),
        "titre": valeur_catalogue(
            ligne,
            "Titre principal"
        ) or "",
        "feuilles": valeur_catalogue(
            ligne,
            "Feuille(s)"
        ) or "",
        "plages": valeur_catalogue(
            ligne,
            "Plage(s)"
        ) or "",
        "meknes": valeur_catalogue(
            ligne,
            "Meknès détecté"
        ),
        "fes_meknes": valeur_catalogue(
            ligne,
            "Fès-Meknès détecté"
        ),
        "annees": valeur_catalogue(
            ligne,
            "Années"
        ) or "",
    })


# ============================================================
# CREATION DU RAPPORT REVISE
# ============================================================

rapport = Workbook()

ws_revu = rapport.active
ws_revu.title = "Mapping_revu"

ws_revu.append([
    "Table Word titre",
    "Table Word données",
    "Thème Word",
    "N° Word",
    "Titre Word",
    "Ancien statut",
    "Ancien score",
    "Ancien fichier Excel",
    "Ancien N° Excel",
    "Fichier attendu",
    "Nouveau N° Excel",
    "Nouveau titre Excel",
    "Nouveau score titre",
    "Feuille(s)",
    "Plage(s)",
    "Meknès",
    "Fès-Meknès",
    "Années",
    "Décision"
])


ws_top3 = rapport.create_sheet(
    "Top3_routage"
)

ws_top3.append([
    "N° Word",
    "Titre Word",
    "Thème",
    "Rang",
    "Score titre",
    "Fichier Excel",
    "N° Excel",
    "Titre Excel",
    "Feuille(s)",
    "Plage(s)"
])


ws_a_confirmer = rapport.create_sheet(
    "Sources_a_confirmer"
)

ws_a_confirmer.append([
    "N° Word",
    "Thème",
    "Titre Word",
    "Meilleur fichier testé",
    "Meilleur N° Excel",
    "Meilleur titre Excel",
    "Score",
    "Observation"
])


compteurs = {
    "MATCH FORT": 0,
    "MATCH BON": 0,
    "A VERIFIER": 0,
    "SOURCE A CONFIRMER": 0,
}


print("=" * 78)
print("REVISION DU MAPPING WORD <-> EXCEL")
print("=" * 78)


for ligne in range(
    2,
    mapping_ws.max_row + 1
):

    numero_word = valeur_mapping(
        ligne,
        "N° Word"
    )

    titre_word = valeur_mapping(
        ligne,
        "Titre Word"
    ) or ""

    theme_word = valeur_mapping(
        ligne,
        "Thème Word"
    ) or ""

    theme_normalise = (
        normaliser_theme(
            theme_word
        )
    )

    # Les retours à la ligne ne doivent pas changer le routage
    theme_normalise = re.sub(
        r"\s+",
        " ",
        theme_normalise
    ).strip()

    fichiers_attendus = ROUTAGE.get(
        theme_normalise
    )

    if fichiers_attendus:
        candidats = [
            table
            for table in tables_excel
            if table["fichier"]
            in fichiers_attendus
        ]
    else:
        # Sécurité : si un nouveau thème apparaît,
        # ne pas perdre le tableau.
        candidats = tables_excel

    scores = []

    for table in candidats:

        score = score_titre(
            titre_word,
            table["titre"]
        )

        scores.append(
            (
                score,
                table
            )
        )

    scores.sort(
        key=lambda element: element[0],
        reverse=True
    )

    top3 = scores[:3]

    if not top3:
        continue

    score_meilleur = top3[0][0]
    meilleur = top3[0][1]

    # --------------------------------------------------------
    # DECISION
    # --------------------------------------------------------

    if score_meilleur >= 85:
        decision = "MATCH FORT"

    elif score_meilleur >= 72:
        decision = "MATCH BON"

    elif score_meilleur >= 58:
        decision = "A VERIFIER"

    else:
        decision = "SOURCE A CONFIRMER"

    compteurs[decision] += 1

    fichier_attendu = (
        " | ".join(fichiers_attendus)
        if fichiers_attendus
        else "ROUTAGE NON DEFINI"
    )

    ws_revu.append([
        valeur_mapping(
            ligne,
            "Table Word titre"
        ),
        valeur_mapping(
            ligne,
            "Table Word données"
        ),
        theme_word,
        numero_word,
        titre_word,
        valeur_mapping(
            ligne,
            "Statut"
        ),
        valeur_mapping(
            ligne,
            "Score final"
        ),
        valeur_mapping(
            ligne,
            "Fichier Excel"
        ),
        valeur_mapping(
            ligne,
            "N° Excel"
        ),
        fichier_attendu,
        meilleur["numero"],
        meilleur["titre"],
        round(
            score_meilleur,
            1
        ),
        meilleur["feuilles"],
        meilleur["plages"],
        meilleur["meknes"],
        meilleur["fes_meknes"],
        meilleur["annees"],
        decision
    ])

    for rang, (
        score,
        candidat
    ) in enumerate(
        top3,
        start=1
    ):

        ws_top3.append([
            numero_word,
            titre_word,
            theme_word,
            rang,
            round(
                score,
                1
            ),
            candidat["fichier"],
            candidat["numero"],
            candidat["titre"],
            candidat["feuilles"],
            candidat["plages"]
        ])

    if decision == "SOURCE A CONFIRMER":

        ws_a_confirmer.append([
            numero_word,
            theme_word,
            titre_word,
            meilleur["fichier"],
            meilleur["numero"],
            meilleur["titre"],
            round(
                score_meilleur,
                1
            ),
            (
                "Aucun titre suffisamment proche "
                "dans le fichier source attendu. "
                "Il faut vérifier si la source Excel "
                "correspondante est présente dans donnees_brutes."
            )
        ])


# ============================================================
# RESUME
# ============================================================

ws_resume = rapport.create_sheet(
    "Resume"
)

ws_resume.append([
    "Indicateur",
    "Valeur"
])

ws_resume.append([
    "Tableaux Word analysés",
    mapping_ws.max_row - 1
])

for cle in [
    "MATCH FORT",
    "MATCH BON",
    "A VERIFIER",
    "SOURCE A CONFIRMER"
]:

    ws_resume.append([
        cle,
        compteurs[cle]
    ])


# ============================================================
# MISE EN FORME
# ============================================================

for feuille in [
    ws_revu,
    ws_top3,
    ws_a_confirmer
]:
    feuille.freeze_panes = "A2"
    feuille.auto_filter.ref = (
        feuille.dimensions
    )


largeurs = {
    "A": 16,
    "B": 18,
    "C": 32,
    "D": 12,
    "E": 75,
    "F": 18,
    "G": 15,
    "H": 45,
    "I": 15,
    "J": 48,
    "K": 15,
    "L": 75,
    "M": 18,
    "N": 35,
    "O": 60,
    "P": 12,
    "Q": 16,
    "R": 25,
    "S": 22,
}

for colonne, largeur in (
    largeurs.items()
):
    ws_revu.column_dimensions[
        colonne
    ].width = largeur


ws_top3.column_dimensions["A"].width = 12
ws_top3.column_dimensions["B"].width = 75
ws_top3.column_dimensions["C"].width = 32
ws_top3.column_dimensions["D"].width = 10
ws_top3.column_dimensions["E"].width = 14
ws_top3.column_dimensions["F"].width = 45
ws_top3.column_dimensions["G"].width = 14
ws_top3.column_dimensions["H"].width = 75
ws_top3.column_dimensions["I"].width = 35
ws_top3.column_dimensions["J"].width = 60

ws_a_confirmer.column_dimensions["A"].width = 12
ws_a_confirmer.column_dimensions["B"].width = 32
ws_a_confirmer.column_dimensions["C"].width = 75
ws_a_confirmer.column_dimensions["D"].width = 45
ws_a_confirmer.column_dimensions["E"].width = 15
ws_a_confirmer.column_dimensions["F"].width = 75
ws_a_confirmer.column_dimensions["G"].width = 12
ws_a_confirmer.column_dimensions["H"].width = 90

ws_resume.column_dimensions["A"].width = 35
ws_resume.column_dimensions["B"].width = 18


rapport.save(
    FICHIER_SORTIE
)


mapping_wb.close()
catalogue_wb.close()


print()
print("Résultat :")

for cle in [
    "MATCH FORT",
    "MATCH BON",
    "A VERIFIER",
    "SOURCE A CONFIRMER"
]:

    print(
        f"{cle} : "
        f"{compteurs[cle]}"
    )


print()
print(
    f"Fichier créé : "
    f"{FICHIER_SORTIE}"
)

print("=" * 78)
print("REVISION TERMINEE")
print("=" * 78)
