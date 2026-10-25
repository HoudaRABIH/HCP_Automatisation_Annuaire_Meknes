from pathlib import Path
from docx import Document


DOSSIER_SORTIE = Path("sorties")

FICHIER_ENTREE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_OK.docx"
)

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_FINAL.docx"
)


def remplacer_dans_runs(paragraphe, ancien, nouveau):
    nb = 0

    for run in paragraphe.runs:
        if ancien in run.text:
            run.text = run.text.replace(
                ancien,
                nouveau
            )
            nb += 1

    return nb


def remplacer_cellule_exacte(cellule, ancien, nouveau):
    if (cellule.text or "").strip() != ancien:
        return 0

    for paragraphe in cellule.paragraphs:
        if ancien in (paragraphe.text or ""):
            return remplacer_dans_runs(
                paragraphe,
                ancien,
                nouveau
            )

    return 0


if not FICHIER_ENTREE.exists():
    raise FileNotFoundError(
        f"Document introuvable : {FICHIER_ENTREE}"
    )


document = Document(
    FICHIER_ENTREE
)

nb = 0
lignes_corrigees = 0


# ============================================================
# CORRECTION ULTRA CIBLEE DE LA NOTE DU TABLEAU 9-11
# ============================================================
#
# On ne modifie une ligne que si elle contient explicitement
# "Inventaire National Forestier" ou "الجرد الوطني".
#
# Dans cette même ligne :
#   (3) -> (1)
#
# Ainsi, aucune autre note (3) du document ne peut être touchée.
# ============================================================

for tableau in document.tables:
    for ligne in tableau.rows:
        texte_ligne = " | ".join(
            cellule.text or ""
            for cellule in ligne.cells
        )

        if (
            "Inventaire National Forestier" not in texte_ligne
            and "الجرد الوطني" not in texte_ligne
        ):
            continue

        lignes_corrigees += 1

        cellules_vues = set()

        for cellule in ligne.cells:
            identifiant = id(
                cellule._tc
            )

            if identifiant in cellules_vues:
                continue

            cellules_vues.add(
                identifiant
            )

            # Cas de la cellule contenant uniquement "(3)"
            nb += remplacer_cellule_exacte(
                cellule,
                "(3)",
                "(1)"
            )

            # Cas des phrases FR / AR qui contiennent "(3)"
            for paragraphe in cellule.paragraphs:
                texte = paragraphe.text or ""

                if (
                    "Inventaire National Forestier" in texte
                    or "الجرد الوطني" in texte
                ):
                    nb += remplacer_dans_runs(
                        paragraphe,
                        "(3)",
                        "(1)"
                    )


# Sécurité supplémentaire :
# si la note est exceptionnellement hors tableau dans Word.
for paragraphe in document.paragraphs:
    texte = paragraphe.text or ""

    if (
        "Inventaire National Forestier" in texte
        or "الجرد الوطني" in texte
    ):
        nb += remplacer_dans_runs(
            paragraphe,
            "(3)",
            "(1)"
        )


document.save(
    FICHIER_SORTIE
)


print("=" * 82)
print("NOTE 9-11 CORRIGEE")
print("=" * 82)
print(f"Ligne(s) de note identifiée(s) : {lignes_corrigees}")
print(f"Remplacement(s) effectué(s)   : {nb}")
print()
print("Correction : (3) -> (1)")
print("Uniquement pour la note de l'Inventaire National Forestier.")
print()
print(f"Document créé : {FICHIER_SORTIE}")
print("=" * 82)
