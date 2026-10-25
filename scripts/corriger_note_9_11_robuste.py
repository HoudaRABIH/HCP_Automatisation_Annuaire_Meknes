from pathlib import Path
from docx import Document


DOSSIER_SORTIE = Path("sorties")

CANDIDATS_ENTREE = [
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_FINAL.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_OK.docx",
]

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_VALIDE.docx"
)


def remplacer_paragraphe_preserver_style(paragraphe, ancien, nouveau):
    """Remplace un texte même s'il est découpé entre plusieurs runs."""
    texte = paragraphe.text or ""

    if ancien not in texte:
        return 0

    # D'abord essayer run par run.
    for run in paragraphe.runs:
        if ancien in run.text:
            run.text = run.text.replace(ancien, nouveau)
            return 1

    # Sinon Word a probablement séparé "(3)" sur plusieurs runs.
    nouveau_texte = texte.replace(ancien, nouveau)

    if paragraphe.runs:
        paragraphe.runs[0].text = nouveau_texte
        for run in paragraphe.runs[1:]:
            run.text = ""
    else:
        paragraphe.add_run(nouveau_texte)

    return 1


def remplacer_cellule_exacte(cellule, ancien, nouveau):
    if (cellule.text or "").strip() != ancien:
        return 0

    if cellule.paragraphs:
        p = cellule.paragraphs[0]

        if p.runs:
            p.runs[0].text = nouveau
            for run in p.runs[1:]:
                run.text = ""
        else:
            p.add_run(nouveau)

        for p2 in cellule.paragraphs[1:]:
            for run in p2.runs:
                run.text = ""

    else:
        cellule.text = nouveau

    return 1


fichier_entree = None

for candidat in CANDIDATS_ENTREE:
    if candidat.exists():
        fichier_entree = candidat
        break

if fichier_entree is None:
    raise FileNotFoundError(
        "Aucun fichier Agriculture FINAL/OK trouvé dans sorties."
    )


document = Document(fichier_entree)

print("=" * 82)
print("RECHERCHE DE LA NOTE 9-11")
print("=" * 82)
print(f"Document : {fichier_entree}")
print()


# ============================================================
# 1) TROUVER LE TABLEAU QUI CONTIENT LA PHRASE DE NOTE
# ============================================================

indices_note = []

for idx, table in enumerate(document.tables):
    contenu = "\n".join(
        cell.text or ""
        for row in table.rows
        for cell in row.cells
    )

    if (
        "Inventaire National Forestier" in contenu
        or "الجرد الوطني" in contenu
    ):
        indices_note.append(idx)


if not indices_note:
    raise RuntimeError(
        "La phrase 'Inventaire National Forestier' n'a pas été trouvée "
        "dans les tableaux du document."
    )


print(
    "Tableau(x) contenant la note : "
    + ", ".join(str(i + 1) for i in indices_note)
)


# ============================================================
# 2) CORRIGER LA PHRASE + LE NUMERO ISOLE DANS LA ZONE
# ============================================================

nb = 0
tables_zone = set()

for idx_note in indices_note:
    # Le numéro (3) peut être dans le même tableau ou juste avant/après.
    for idx in range(
        max(0, idx_note - 2),
        min(len(document.tables), idx_note + 3)
    ):
        tables_zone.add(idx)


for idx in sorted(tables_zone):
    table = document.tables[idx]
    cellules_vues = set()

    for row in table.rows:
        for cell in row.cells:
            identifiant = id(cell._tc)

            if identifiant in cellules_vues:
                continue

            cellules_vues.add(identifiant)

            contenu = cell.text or ""

            # Numéro isolé
            nb += remplacer_cellule_exacte(
                cell,
                "(3)",
                "(1)"
            )

            # Phrase française/arabe de la note
            if (
                "Inventaire National Forestier" in contenu
                or "الجرد الوطني" in contenu
            ):
                for p in cell.paragraphs:
                    nb += remplacer_paragraphe_preserver_style(
                        p,
                        "(3)",
                        "(1)"
                    )


# ============================================================
# 3) VERIFICATION DE SECURITE :
#    EN-TETE 9-11 DOIT RESTER EN 2024(1)
# ============================================================

for table in document.tables:
    contenu = "\n".join(
        cell.text or ""
        for row in table.rows
        for cell in row.cells
    )

    if (
        "Essences résineuses naturelles" in contenu
        and "Matorral" in contenu
    ):
        # Première cellule = "Année 2024(1)"
        cell_fr = table.cell(0, 0)
        cell_ar = table.cell(0, len(table.columns) - 1)

        if "2024" in (cell_fr.text or ""):
            for p in cell_fr.paragraphs:
                remplacer_paragraphe_preserver_style(
                    p,
                    "(3)",
                    "(1)"
                )

        if "2024" in (cell_ar.text or ""):
            for p in cell_ar.paragraphs:
                remplacer_paragraphe_preserver_style(
                    p,
                    "(3)",
                    "(1)"
                )

        break


document.save(FICHIER_SORTIE)


print()
print(f"Remplacements effectués : {nb}")
print()
print("Résultat attendu :")
print("  Année 2024(1)")
print("  (1) Statistiques de l'Inventaire National Forestier.")
print("  note arabe également en (1)")
print()
print(f"Document créé : {FICHIER_SORTIE}")
print("=" * 82)
