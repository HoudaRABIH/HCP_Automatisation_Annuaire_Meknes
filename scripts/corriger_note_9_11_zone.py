from pathlib import Path
from docx import Document


DOSSIER_SORTIE = Path("sorties")

FICHIER_ENTREE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_FINAL.docx"
)

# Si FINAL n'existe pas encore, on repart de Agriculture_OK.
FICHIER_SECOURS = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_OK.docx"
)

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_VALIDE.docx"
)


def remplacer_texte_cellule_preserver_style(cellule, ancien, nouveau):
    nb = 0

    for paragraphe in cellule.paragraphs:
        texte_total = paragraphe.text or ""

        if ancien not in texte_total:
            continue

        # Cas normal : (3) est dans un run existant.
        trouve_run = False

        for run in paragraphe.runs:
            if ancien in run.text:
                run.text = run.text.replace(
                    ancien,
                    nouveau
                )
                nb += 1
                trouve_run = True

        # Cas où Word coupe "(3)" entre plusieurs runs.
        if not trouve_run:
            nouveau_total = texte_total.replace(
                ancien,
                nouveau
            )

            if paragraphe.runs:
                paragraphe.runs[0].text = nouveau_total

                for run in paragraphe.runs[1:]:
                    run.text = ""
            else:
                paragraphe.add_run(
                    nouveau_total
                )

            nb += 1

    return nb


def texte_table(table):
    return " ".join(
        cellule.text or ""
        for ligne in table.rows
        for cellule in ligne.cells
    )


fichier_entree = (
    FICHIER_ENTREE
    if FICHIER_ENTREE.exists()
    else FICHIER_SECOURS
)

if not fichier_entree.exists():
    raise FileNotFoundError(
        "Aucun document Agriculture FINAL/OK trouvé."
    )


document = Document(
    fichier_entree
)


# ============================================================
# LOCALISER LE TABLEAU DE DONNEES 9-11
# ============================================================

index_9_11 = None

for i, table in enumerate(document.tables):
    contenu = texte_table(table)

    # Le tableau de données 9-11 contient simultanément ces marqueurs.
    if (
        "Essences résineuses naturelles" in contenu
        and "Matorral" in contenu
        and "Préfecture : Meknès" in contenu
    ):
        index_9_11 = i
        break


if index_9_11 is None:
    raise RuntimeError(
        "Impossible de localiser le tableau de données 9-11."
    )


# ============================================================
# CORRIGER UNIQUEMENT LA ZONE IMMEDIATEMENT APRES 9-11
# ============================================================
#
# On inspecte le tableau 9-11 lui-même + les 3 tableaux suivants.
# C'est la zone où se trouvent la source et la note du tableau 9-11.
# On ne touche à rien plus loin dans le document.
# ============================================================

nb = 0

debut = index_9_11
fin = min(
    len(document.tables),
    index_9_11 + 4
)

for idx in range(debut, fin):
    table = document.tables[idx]

    cellules_vues = set()

    for ligne in table.rows:
        for cellule in ligne.cells:
            identifiant = id(
                cellule._tc
            )

            if identifiant in cellules_vues:
                continue

            cellules_vues.add(
                identifiant
            )

            contenu = cellule.text or ""

            # On corrige seulement les cellules de note :
            # - cellule exactement "(3)"
            # - texte contenant Inventaire National Forestier
            # - texte arabe contenant الجرد الوطني
            est_note_9_11 = (
                contenu.strip() == "(3)"
                or "Inventaire National Forestier" in contenu
                or "الجرد الوطني" in contenu
            )

            if not est_note_9_11:
                continue

            nb += remplacer_texte_cellule_preserver_style(
                cellule,
                "(3)",
                "(1)"
            )


# ============================================================
# SAUVEGARDE
# ============================================================

document.save(
    FICHIER_SORTIE
)

print("=" * 82)
print("CORRECTION FINALE NOTE 9-11")
print("=" * 82)
print(f"Document d'entrée : {fichier_entree}")
print(f"Objet Word 9-11   : {index_9_11 + 1}")
print(f"Remplacements     : {nb}")
print()
print("Résultat attendu :")
print("  Année 2024(1)")
print("  (1) Statistiques de l'Inventaire National Forestier")
print("  équivalent arabe également en (1)")
print()
print(f"Document créé : {FICHIER_SORTIE}")
print("=" * 82)
