from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


DOSSIER_SORTIE = Path("sorties")

# On prend le fichier Santé le plus avancé disponible.
CANDIDATS_ENTREE = [
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_Sante_VALIDE.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_Sante_OK.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_Sante.docx",
]

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_Sante_SOURCES_OK.docx"
)


def set_bidi(paragraphe, actif=True):
    """
    Active/désactive le sens RTL au niveau paragraphe sans recréer le texte.
    """
    pPr = paragraphe._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))

    if actif:
        if bidi is None:
            bidi = OxmlElement("w:bidi")
            pPr.append(bidi)
        bidi.set(qn("w:val"), "1")
    else:
        if bidi is not None:
            bidi.getparent().remove(bidi)


def est_source_arabe(texte):
    texte = (texte or "").strip()
    return (
        "المصدر" in texte
        or "النشرة الإحصائية السنوية للمغرب" in texte
    )


def est_source_francaise(texte):
    texte = (texte or "").strip()

    # On cible seulement les lignes de source, pas les notes explicatives.
    return (
        texte.lower().startswith("source")
        or "Annuaire Statistique du Maroc" in texte
    )


def corriger_paragraphe(paragraphe):
    texte = paragraphe.text or ""

    if est_source_arabe(texte):
        paragraphe.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        set_bidi(paragraphe, True)
        return "AR"

    if est_source_francaise(texte):
        paragraphe.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_bidi(paragraphe, False)
        return "FR"

    return None


fichier_entree = None

for candidat in CANDIDATS_ENTREE:
    if candidat.exists():
        fichier_entree = candidat
        break

if fichier_entree is None:
    raise FileNotFoundError(
        "Aucun document Santé VALIDE/OK/base trouvé dans sorties."
    )


print("=" * 88)
print("CORRECTION DE POSITION DES SOURCES FR / AR")
print("=" * 88)
print(f"Document d'entrée : {fichier_entree}")
print()

document = Document(fichier_entree)

nb_ar = 0
nb_fr = 0
cellules_vues = set()


# ------------------------------------------------------------
# Tableaux : c'est là que se trouvent pratiquement toutes les sources.
# Même si FR et AR sont dans une même grande cellule fusionnée,
# l'alignement du paragraphe place visuellement :
#   FR à gauche
#   AR à droite
# ------------------------------------------------------------

for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            identifiant = id(cell._tc)

            if identifiant in cellules_vues:
                continue

            cellules_vues.add(identifiant)

            for p in cell.paragraphs:
                resultat = corriger_paragraphe(p)

                if resultat == "AR":
                    nb_ar += 1
                elif resultat == "FR":
                    nb_fr += 1

            # Sous-tableaux éventuels
            for sous_table in cell.tables:
                for sous_row in sous_table.rows:
                    for sous_cell in sous_row.cells:
                        for p in sous_cell.paragraphs:
                            resultat = corriger_paragraphe(p)

                            if resultat == "AR":
                                nb_ar += 1
                            elif resultat == "FR":
                                nb_fr += 1


# Paragraphes hors tableaux, au cas où certaines sources seraient
# stockées dans une zone de texte/structure particulière visible par
# python-docx.
for p in document.paragraphs:
    resultat = corriger_paragraphe(p)

    if resultat == "AR":
        nb_ar += 1
    elif resultat == "FR":
        nb_fr += 1


document.save(FICHIER_SORTIE)


print(f"Sources françaises alignées à gauche : {nb_fr}")
print(f"Sources arabes alignées à droite     : {nb_ar}")
print()
print("Aucun texte de source n'a été recréé.")
print("Aucune donnée numérique n'a été modifiée.")
print()
print(f"Document créé : {FICHIER_SORTIE}")
print("=" * 88)
