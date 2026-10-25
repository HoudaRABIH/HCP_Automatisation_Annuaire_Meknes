from pathlib import Path
from openpyxl import load_workbook
from docx import Document
from docx.table import _Cell
import unicodedata
import re


DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_SORTIE = Path("sorties")

FICHIER_ENTREE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_Sante.docx"
)

FICHIER_EXCEL = (
    DOSSIER_SOURCE
    / "12. Santé_AS 2025.xlsx"
)

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_Sante_VALIDE.docx"
)


def normaliser(texte):
    texte = str(texte or "").strip().lower()

    texte = "".join(
        c
        for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )

    texte = re.sub(r"\s+", " ", texte)

    return texte


def trouver_ligne(ws, libelle):
    cible = normaliser(libelle)

    for r in range(
        1,
        ws.max_row + 1
    ):
        if normaliser(
            ws.cell(r, 1).value
        ) == cible:
            return r

    raise ValueError(
        f"{libelle!r} introuvable dans {ws.title}"
    )


def trouver_nation(ws):
    for r in range(
        ws.max_row,
        0,
        -1
    ):
        if normaliser(
            ws.cell(r, 1).value
        ) in {
            "ensemble",
            "total",
        }:
            return r

    raise ValueError(
        "Total/Ensemble national introuvable."
    )


def formater(valeur):
    if valeur in (
        None,
        "",
        "-",
        "–",
    ):
        return "-"

    try:
        nombre = float(valeur)
    except (TypeError, ValueError):
        return str(valeur).strip()

    if abs(
        nombre - round(nombre)
    ) < 1e-9:
        return (
            f"{int(round(nombre)):,}"
            .replace(",", " ")
        )

    return (
        f"{nombre:.1f}"
        .replace(".", ",")
    )


def remplacer_texte_preserver_style(cellule, nouveau_texte):
    nouveau_texte = str(
        nouveau_texte
    )

    if not cellule.paragraphs:
        cellule.text = nouveau_texte
        return

    p = cellule.paragraphs[0]

    if p.runs:
        p.runs[0].text = nouveau_texte

        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(
            nouveau_texte
        )

    for p2 in cellule.paragraphs[1:]:
        for run in p2.runs:
            run.text = ""


if not FICHIER_ENTREE.exists():
    raise FileNotFoundError(
        f"Word introuvable : {FICHIER_ENTREE}"
    )

if not FICHIER_EXCEL.exists():
    raise FileNotFoundError(
        f"Excel introuvable : {FICHIER_EXCEL}"
    )


print("=" * 86)
print("CORRECTION CIBLEE SANTE — TABLEAU 2-21")
print("=" * 86)


document = Document(
    FICHIER_ENTREE
)

wb = load_workbook(
    FICHIER_EXCEL,
    data_only=True,
    read_only=False
)

ws = wb["26"]


# ------------------------------------------------------------
# TABLEAU WORD 2-21 = OBJET 100
# ------------------------------------------------------------

table = document.tables[
    100 - 1
]

if len(table.rows) != 6:
    raise RuntimeError(
        f"2-21 : {len(table.rows)} lignes au lieu de 6."
    )


geo_excel = {
    "nation": trouver_nation(ws),
    "region": trouver_ligne(
        ws,
        "Fès - Meknès"
    ),
    "meknes": trouver_ligne(
        ws,
        "Meknès"
    ),
}

geo_word = {
    "nation": 3,
    "region": 4,
    "meknes": 5,
}


# IMPORTANT :
# python-docx expose ici 10 colonnes de grille à cause des fusions.
# Mais les lignes de données contiennent réellement 9 cellules XML :
#
#   0 libellé géographique
#   1 Total
#   2 Hygiène alimentaire
#   3 Chimie et biologie
#   4 Hématologie et transfusion
#   5 Immuno-sérologie
#   6 Parasitologie
#   7 Bactériologie
#   8 arabe
#
# On écrit donc directement dans les cellules physiques XML de la ligne,
# et non via table.cell(row, col).

for cle_geo, ligne_word in geo_word.items():

    tr = table._tbl.tr_lst[
        ligne_word
    ]

    tc_list = tr.tc_lst

    if len(tc_list) != 9:
        raise RuntimeError(
            f"2-21 ligne {ligne_word + 1}: "
            f"{len(tc_list)} cellules physiques au lieu de 9."
        )

    cellules_physiques = [
        _Cell(
            tc,
            table
        )
        for tc in tc_list
    ]

    ligne_excel = geo_excel[
        cle_geo
    ]

    # Excel B:H -> Word cellules physiques 1:7
    for idx_word, col_excel in enumerate(
        range(2, 9),
        start=1
    ):
        remplacer_texte_preserver_style(
            cellules_physiques[
                idx_word
            ],
            formater(
                ws.cell(
                    ligne_excel,
                    col_excel
                ).value
            )
        )



# ------------------------------------------------------------
# REPLACER LA SOURCE ARABE DU 2-21 DANS LA CASE DE DROITE
# ------------------------------------------------------------
#
# Le tableau de données 2-21 est l'objet 100.
# La ligne de source se trouve normalement juste après.
# On cherche dans les 2 tableaux suivants pour rester robuste.
#
# On DEPLACE le paragraphe XML arabe existant au lieu de le recréer :
# police, taille, couleur, RTL et alignement sont donc conservés.
# ------------------------------------------------------------

source_deplacee = False

for idx in range(
    100,
    min(len(document.tables), 102)
):
    table_source = document.tables[idx]

    for tr in table_source._tbl.tr_lst:
        tc_list = tr.tc_lst

        if not tc_list:
            continue

        cellule_arabe = None
        paragraphe_arabe = None

        for tc in tc_list:
            cellule = _Cell(
                tc,
                table_source
            )

            for p in cellule.paragraphs:
                contenu = p.text or ""

                if (
                    "المصدر" in contenu
                    or "النشرة الإحصائية" in contenu
                ):
                    cellule_arabe = cellule
                    paragraphe_arabe = p
                    break

            if paragraphe_arabe is not None:
                break

        if paragraphe_arabe is None:
            continue

        cellule_droite = _Cell(
            tc_list[-1],
            table_source
        )

        # Si le texte arabe n'est pas déjà dans la cellule de droite,
        # déplacer le paragraphe complet vers celle-ci.
        if cellule_arabe._tc is not cellule_droite._tc:
            p_xml = paragraphe_arabe._p
            p_xml.getparent().remove(
                p_xml
            )

            cellule_droite._tc.append(
                p_xml
            )

            source_deplacee = True

        break

    if source_deplacee:
        break


wb.close()

document.save(
    FICHIER_SORTIE
)


print()
print("Valeurs attendues après correction :")
print()
print("Nation :")
print(
    "16 310 297 | 309 | 10 467 102 | "
    "4 718 111 | 816 832 | 46 214 | 261 729"
)
print()
print("Région :")
print(
    "2 542 445 | 54 | 1 649 278 | "
    "664 621 | 182 166 | 2 649 | 43 677"
)
print()
print("Meknès :")
print(
    "258 713 | - | 153 309 | "
    "72 567 | 29 106 | 432 | 3 299"
)
print()
print(
    "Source arabe 2-21 replacée à droite : "
    + ("OUI" if source_deplacee else "DEJA A DROITE / NON TROUVEE")
)
print()
print(
    f"Document créé : {FICHIER_SORTIE}"
)
print("=" * 86)
