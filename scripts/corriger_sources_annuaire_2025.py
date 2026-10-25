from pathlib import Path
from docx import Document


DOSSIER_SORTIE = Path("sorties")

# Prend le Word cumulé le plus avancé disponible.
CANDIDATS = [
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_Transport.docx",
]

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Sources2025.docx"
)

FR_ANCIEN = "Annuaire Statistique du Maroc 2023"
FR_NOUVEAU = "Annuaire Statistique du Maroc 2025"

AR_ANCIEN = "النشرة الإحصائية السنوية للمغرب 2023"
AR_NOUVEAU = "النشرة الإحصائية السنوية للمغرب 2025"


def remplacer_dans_paragraphe(paragraphe):
    """
    Remplace UNIQUEMENT les deux références exactes de l'Annuaire 2023.
    Toutes les autres sources (RGPH, ministères, HCP, etc.) restent intactes.
    """
    texte_complet = paragraphe.text or ""

    if FR_ANCIEN not in texte_complet and AR_ANCIEN not in texte_complet:
        return 0

    nb = 0

    # Cas normal : la phrase se trouve dans un même run.
    for run in paragraphe.runs:
        avant = run.text

        if FR_ANCIEN in run.text:
            run.text = run.text.replace(
                FR_ANCIEN,
                FR_NOUVEAU
            )

        if AR_ANCIEN in run.text:
            run.text = run.text.replace(
                AR_ANCIEN,
                AR_NOUVEAU
            )

        if run.text != avant:
            nb += 1

    # Si Word a découpé la phrase sur plusieurs runs,
    # on remplace le texte complet en conservant le style du premier run.
    nouveau_complet = (
        texte_complet
        .replace(FR_ANCIEN, FR_NOUVEAU)
        .replace(AR_ANCIEN, AR_NOUVEAU)
    )

    if (
        nouveau_complet != texte_complet
        and not any(
            FR_NOUVEAU in run.text or AR_NOUVEAU in run.text
            for run in paragraphe.runs
        )
    ):
        if paragraphe.runs:
            paragraphe.runs[0].text = nouveau_complet
            for run in paragraphe.runs[1:]:
                run.text = ""
        else:
            paragraphe.add_run(nouveau_complet)

        nb += 1

    return nb


def parcourir_tableau(tableau):
    nb = 0
    cellules_vues = set()

    for ligne in tableau.rows:
        for cellule in ligne.cells:
            identifiant = id(cellule._tc)

            if identifiant in cellules_vues:
                continue

            cellules_vues.add(identifiant)

            for paragraphe in cellule.paragraphs:
                nb += remplacer_dans_paragraphe(paragraphe)

            # Au cas où il y a un tableau imbriqué.
            for sous_tableau in cellule.tables:
                nb += parcourir_tableau(sous_tableau)

    return nb


fichier_entree = None

for candidat in CANDIDATS:
    if candidat.exists():
        fichier_entree = candidat
        break

if fichier_entree is None:
    raise FileNotFoundError(
        "Aucun document cumulé trouvé dans le dossier sorties."
    )


document = Document(fichier_entree)

nb_remplacements = 0

# Paragraphes hors tableaux.
for paragraphe in document.paragraphs:
    nb_remplacements += remplacer_dans_paragraphe(paragraphe)

# Tableaux.
for tableau in document.tables:
    nb_remplacements += parcourir_tableau(tableau)

# Headers / footers, uniquement si la phrase exacte de l'annuaire s'y trouve.
for section in document.sections:
    for paragraphe in section.header.paragraphs:
        nb_remplacements += remplacer_dans_paragraphe(paragraphe)

    for tableau in section.header.tables:
        nb_remplacements += parcourir_tableau(tableau)

    for paragraphe in section.footer.paragraphs:
        nb_remplacements += remplacer_dans_paragraphe(paragraphe)

    for tableau in section.footer.tables:
        nb_remplacements += parcourir_tableau(tableau)


document.save(FICHIER_SORTIE)

print("=" * 78)
print("CORRECTION CIBLEE DES SOURCES")
print("=" * 78)
print(f"Document d'entrée : {fichier_entree}")
print(f"Document créé     : {FICHIER_SORTIE}")
print()
print(f"Occurrences modifiées : {nb_remplacements}")
print()
print("Modification autorisée uniquement :")
print("  Annuaire Statistique du Maroc 2023 -> 2025")
print("  النشرة الإحصائية السنوية للمغرب 2023 -> 2025")
print()
print("Toutes les autres sources (RGPH, HCP, ministères, etc.)")
print("ont été laissées strictement intactes.")
print("=" * 78)
