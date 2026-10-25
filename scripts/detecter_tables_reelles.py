from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
import re
import warnings


warnings.filterwarnings("ignore")


DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_SORTIE = Path("sorties")

FICHIER_SORTIE = (
    DOSSIER_SORTIE / "catalogue_tables_reelles.xlsx"
)

DOSSIER_SORTIE.mkdir(exist_ok=True)


# ============================================================
# OUTILS
# ============================================================

def nettoyer(valeur):
    if valeur is None:
        return ""

    return str(valeur).strip()


def chapitre_du_fichier(nom_fichier):
    """
    Exemple :
    11. Enseignement_AS 2025.xlsx -> 11
    14-Emploi et salaire...       -> 14
    """

    resultat = re.match(
        r"^\s*(\d+)",
        nom_fichier
    )

    if resultat:
        return int(resultat.group(1))

    return None


def est_feuille_ignoree(nom):
    nom = nom.strip().lower()

    return (
        nom.startswith("pg")
        or "sommaire" in nom
        or "graph" in nom
        or "grraph" in nom
        or "graphique" in nom
    )


# Numéro de tableau :
# 11 - 2
# 12-15
# 3 – 4
MOTIF_TABLEAU = re.compile(
    r"^\s*(\d{1,2})\s*[-–—]\s*(\d{1,3})(?!\d)"
)


MOTIF_ANNEE = re.compile(
    r"\b(?:19|20)\d{2}\b"
)


def numero_tableau_valide(
    texte,
    chapitre_attendu
):

    resultat = MOTIF_TABLEAU.search(
        texte
    )

    if not resultat:
        return None

    chapitre = int(
        resultat.group(1)
    )

    numero = int(
        resultat.group(2)
    )

    # REGLE CRITIQUE :
    # Le numéro doit correspondre au chapitre
    # indiqué dans le nom du fichier.
    if chapitre != chapitre_attendu:
        return None

    return chapitre, numero


def contient_francais(texte):
    return bool(
        re.search(
            r"[A-Za-zÀ-ÿ]",
            texte
        )
    )


def choisir_titre(textes):

    francais = [
        texte
        for texte in textes
        if contient_francais(texte)
    ]

    if francais:
        return max(
            francais,
            key=len
        )

    return max(
        textes,
        key=len
    )


def contient_meknes(cellules):

    for cellule in cellules:

        texte = nettoyer(
            cellule.value
        )

        if re.search(
            r"\bmekn[eè]s\b",
            texte,
            re.IGNORECASE
        ):
            return True

        if "مكناس" in texte:
            return True

    return False


def contient_fes_meknes(cellules):

    variantes = [
        "fès-meknès",
        "fès - meknès",
        "fes-meknes",
        "fes - meknes",
        "فاس - مكناس",
        "فاس ــ مكناس"
    ]

    for cellule in cellules:

        texte = nettoyer(
            cellule.value
        ).lower()

        for variante in variantes:

            if variante.lower() in texte:
                return True

    return False


def extraire_annees(cellules):

    annees = set()

    for cellule in cellules:

        texte = nettoyer(
            cellule.value
        )

        for resultat in MOTIF_ANNEE.finditer(
            texte
        ):
            annee = int(
                resultat.group(0)
            )

            # Limite raisonnable pour éviter que certaines
            # valeurs statistiques soient prises pour des années.
            if 1900 <= annee <= 2030:
                annees.add(
                    str(annee)
                )

    return sorted(annees)


def choisir_titre_logique(segments):

    titres = [
        segment["titre"]
        for segment in segments
    ]

    # Priorité aux titres sans "suite"
    sans_suite = [
        titre
        for titre in titres
        if "suite" not in titre.lower()
    ]

    if sans_suite:
        return max(
            sans_suite,
            key=len
        )

    return max(
        titres,
        key=len
    )


# ============================================================
# FICHIERS
# ============================================================

fichiers_excel = sorted(
    fichier
    for fichier in DOSSIER_SOURCE.glob("*.xlsx")
    if not fichier.name.startswith("~$")
)


segments_detectes = []
anomalies_detectees = []


print("=" * 76)
print("DETECTION CORRIGEE DES TABLEAUX STATISTIQUES HCP")
print("=" * 76)


# ============================================================
# ANALYSE
# ============================================================

for numero_fichier, fichier in enumerate(
    fichiers_excel,
    start=1
):

    chapitre_attendu = chapitre_du_fichier(
        fichier.name
    )

    print(
        f"[{numero_fichier}/{len(fichiers_excel)}] "
        f"{fichier.name}"
    )

    if chapitre_attendu is None:

        anomalies_detectees.append([
            fichier.name,
            "",
            "Impossible de déterminer le chapitre"
        ])

        continue

    try:

        classeur = load_workbook(
            fichier,
            data_only=False,
            read_only=False
        )

        for feuille in classeur.worksheets:

            if est_feuille_ignoree(
                feuille.title
            ):
                continue

            cellules = [
                cellule
                for cellule
                in feuille._cells.values()
                if cellule.value not in (None, "")
            ]

            if not cellules:
                continue

            derniere_ligne = max(
                cellule.row
                for cellule in cellules
            )

            # --------------------------------------------
            # Chercher UNIQUEMENT les vrais titres
            # --------------------------------------------

            titres_par_ligne = {}

            for cellule in cellules:

                texte = nettoyer(
                    cellule.value
                )

                numero = numero_tableau_valide(
                    texte,
                    chapitre_attendu
                )

                if numero is None:
                    continue

                titres_par_ligne.setdefault(
                    cellule.row,
                    []
                ).append(
                    texte
                )

            lignes_titres = sorted(
                titres_par_ligne
            )

            if not lignes_titres:

                anomalies_detectees.append([
                    fichier.name,
                    feuille.title,
                    "Aucun titre de tableau valide détecté"
                ])

                continue

            # --------------------------------------------
            # Créer les segments
            # --------------------------------------------

            for index, ligne_debut in enumerate(
                lignes_titres
            ):

                textes_titre = (
                    titres_par_ligne[
                        ligne_debut
                    ]
                )

                titre = choisir_titre(
                    textes_titre
                )

                chapitre, numero = (
                    numero_tableau_valide(
                        titre,
                        chapitre_attendu
                    )
                )

                numero_complet = (
                    f"{chapitre}-{numero}"
                )

                if index + 1 < len(
                    lignes_titres
                ):

                    ligne_fin = (
                        lignes_titres[
                            index + 1
                        ] - 1
                    )

                else:

                    ligne_fin = (
                        derniere_ligne
                    )

                cellules_segment = [
                    cellule
                    for cellule in cellules
                    if (
                        ligne_debut
                        <= cellule.row
                        <= ligne_fin
                    )
                ]

                if not cellules_segment:
                    continue

                derniere_colonne = max(
                    cellule.column
                    for cellule
                    in cellules_segment
                )

                segments_detectes.append({
                    "fichier": fichier.name,
                    "feuille": feuille.title,
                    "chapitre": chapitre,
                    "numero": numero,
                    "numero_complet": numero_complet,
                    "titre": titre,
                    "ligne_debut": ligne_debut,
                    "ligne_fin": ligne_fin,
                    "derniere_colonne": (
                        get_column_letter(
                            derniere_colonne
                        )
                    ),
                    "meknes": contient_meknes(
                        cellules_segment
                    ),
                    "fes_meknes": (
                        contient_fes_meknes(
                            cellules_segment
                        )
                    ),
                    "annees": extraire_annees(
                        cellules_segment
                    ),
                    "nb_cellules": len(
                        cellules_segment
                    )
                })

        classeur.close()

    except Exception as erreur:

        anomalies_detectees.append([
            fichier.name,
            "",
            str(erreur)
        ])

        print(
            f"   ERREUR : {erreur}"
        )


# ============================================================
# CONSOLIDATION EN TABLEAUX LOGIQUES
# ============================================================

groupes = {}

for segment in segments_detectes:

    cle = (
        segment["fichier"],
        segment["numero_complet"]
    )

    groupes.setdefault(
        cle,
        []
    ).append(
        segment
    )


tables_logiques = []


for cle, segments in groupes.items():

    fichier = cle[0]
    numero_complet = cle[1]

    premier = segments[0]

    feuilles = sorted(
        set(
            segment["feuille"]
            for segment in segments
        )
    )

    annees = sorted(
        set(
            annee
            for segment in segments
            for annee in segment["annees"]
        )
    )

    plages = []

    for segment in segments:

        plages.append(
            (
                f"{segment['feuille']}!"
                f"{segment['ligne_debut']}:"
                f"{segment['ligne_fin']}"
            )
        )

    tables_logiques.append({
        "fichier": fichier,
        "chapitre_source": (
            premier["chapitre"]
        ),
        "numero_source": (
            premier["numero"]
        ),
        "numero_complet": (
            numero_complet
        ),
        "titre": choisir_titre_logique(
            segments
        ),
        "feuilles": " | ".join(
            feuilles
        ),
        "plages": " | ".join(
            plages
        ),
        "nb_segments": len(
            segments
        ),
        "meknes": any(
            segment["meknes"]
            for segment in segments
        ),
        "fes_meknes": any(
            segment["fes_meknes"]
            for segment in segments
        ),
        "annees": ", ".join(
            annees
        ),
        "nb_cellules": sum(
            segment["nb_cellules"]
            for segment in segments
        )
    })


tables_logiques.sort(
    key=lambda t: (
        t["chapitre_source"],
        t["numero_source"]
    )
)


# ============================================================
# RAPPORT EXCEL
# ============================================================

rapport = Workbook()


# ------------------------------------------------------------
# TABLES LOGIQUES
# ------------------------------------------------------------

ws_tables = rapport.active
ws_tables.title = "Tables_logiques"

ws_tables.append([
    "ID",
    "Fichier",
    "Chapitre source",
    "N° source",
    "Numéro complet",
    "Titre principal",
    "Feuille(s)",
    "Plage(s)",
    "Nombre segments",
    "Meknès détecté",
    "Fès-Meknès détecté",
    "Années",
    "Nombre cellules"
])


for identifiant, table in enumerate(
    tables_logiques,
    start=1
):

    ws_tables.append([
        identifiant,
        table["fichier"],
        table["chapitre_source"],
        table["numero_source"],
        table["numero_complet"],
        table["titre"],
        table["feuilles"],
        table["plages"],
        table["nb_segments"],
        (
            "OUI"
            if table["meknes"]
            else "NON"
        ),
        (
            "OUI"
            if table["fes_meknes"]
            else "NON"
        ),
        table["annees"],
        table["nb_cellules"]
    ])


# ------------------------------------------------------------
# SEGMENTS
# ------------------------------------------------------------

ws_segments = rapport.create_sheet(
    "Segments"
)

ws_segments.append([
    "ID",
    "Fichier",
    "Feuille",
    "Numéro complet",
    "Titre",
    "Ligne début",
    "Ligne fin",
    "Dernière colonne",
    "Meknès",
    "Fès-Meknès",
    "Années",
    "Nombre cellules"
])


for identifiant, segment in enumerate(
    segments_detectes,
    start=1
):

    ws_segments.append([
        identifiant,
        segment["fichier"],
        segment["feuille"],
        segment["numero_complet"],
        segment["titre"],
        segment["ligne_debut"],
        segment["ligne_fin"],
        segment["derniere_colonne"],
        (
            "OUI"
            if segment["meknes"]
            else "NON"
        ),
        (
            "OUI"
            if segment["fes_meknes"]
            else "NON"
        ),
        ", ".join(
            segment["annees"]
        ),
        segment["nb_cellules"]
    ])


# ------------------------------------------------------------
# RESUME
# ------------------------------------------------------------

ws_resume = rapport.create_sheet(
    "Resume"
)

ws_resume.append([
    "Indicateur",
    "Valeur"
])

ws_resume.append([
    "Fichiers analysés",
    len(fichiers_excel)
])

ws_resume.append([
    "Segments valides",
    len(segments_detectes)
])

ws_resume.append([
    "Tables logiques",
    len(tables_logiques)
])

ws_resume.append([
    "Tables logiques contenant Meknès",
    sum(
        1
        for table in tables_logiques
        if table["meknes"]
    )
])

ws_resume.append([
    "Tables logiques contenant Fès-Meknès",
    sum(
        1
        for table in tables_logiques
        if table["fes_meknes"]
    )
])


# ------------------------------------------------------------
# ANOMALIES
# ------------------------------------------------------------

ws_anomalies = rapport.create_sheet(
    "Anomalies"
)

ws_anomalies.append([
    "Fichier",
    "Feuille",
    "Observation"
])

for anomalie in anomalies_detectees:

    ws_anomalies.append(
        anomalie
    )


# ============================================================
# LARGEURS
# ============================================================

largeurs_tables = {
    "A": 8,
    "B": 45,
    "C": 16,
    "D": 14,
    "E": 18,
    "F": 85,
    "G": 40,
    "H": 60,
    "I": 18,
    "J": 18,
    "K": 22,
    "L": 25,
    "M": 20
}

for colonne, largeur in (
    largeurs_tables.items()
):

    ws_tables.column_dimensions[
        colonne
    ].width = largeur


ws_tables.freeze_panes = "A2"
ws_tables.auto_filter.ref = (
    ws_tables.dimensions
)

ws_segments.freeze_panes = "A2"
ws_segments.auto_filter.ref = (
    ws_segments.dimensions
)

ws_resume.column_dimensions["A"].width = 45
ws_resume.column_dimensions["B"].width = 20

ws_anomalies.column_dimensions["A"].width = 45
ws_anomalies.column_dimensions["B"].width = 30
ws_anomalies.column_dimensions["C"].width = 60


# ============================================================
# SAUVEGARDE
# ============================================================

rapport.save(
    FICHIER_SORTIE
)


# ============================================================
# RESULTAT
# ============================================================

print()
print(
    f"Segments valides détectés : "
    f"{len(segments_detectes)}"
)

print(
    f"Tables logiques détectées : "
    f"{len(tables_logiques)}"
)

print(
    "Tables logiques avec Meknès : "
    f"{sum(1 for t in tables_logiques if t['meknes'])}"
)

print()
print(
    f"Fichier créé : "
    f"{FICHIER_SORTIE}"
)

print("=" * 76)
print("DETECTION CORRIGEE TERMINEE")
print("=" * 76)