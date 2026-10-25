from pathlib import Path
from openpyxl import load_workbook
from docx import Document
import re
import unicodedata

ROOT = Path(".")
BRUT = ROOT / "donnees_brutes"
SORTIES = ROOT / "sorties"

DOCX = SORTIES / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER_POPULATION_INDUSTRIE_ENERGIE_EAU.docx"
MAPPING = SORTIES / "mapping_final_recettes.xlsx"
CATALOGUE = SORTIES / "catalogue_tables_reelles.xlsx"
OUT = SORTIES / "diagnostic_alertes_qa_finales.txt"

HEALTH_OLD = ["2-11","2-12","2-14","2-15","2-16","2-17","2-20","2-21","2-22"]
OTHER_OLD = ["6-7","8-2","8-3","8-5","13-9","13-10"]

def txt(v):
    return "" if v is None else str(v).replace("\n", " ").strip()

def norm(v):
    s = txt(v).lower()
    s = (s.replace("’","'").replace("–","-").replace("—","-")
           .replace("\u202f"," ").replace("\xa0"," "))
    s = "".join(c for c in unicodedata.normalize("NFD", s)
                if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip()

def code_from_title(v):
    m = re.match(r"^\s*(\d{1,2})\s*-\s*(\d{1,2})\b", txt(v))
    if not m:
        return None
    a, b = map(int, m.groups())
    if not (1 <= a <= 16 and 1 <= b <= 30):
        return None
    return f"{a}-{b}"

def years_in(v):
    return sorted(set(re.findall(r"\b(?:19|20)\d{2}\b", str(v or ""))))

def load_mapping():
    wb = load_workbook(MAPPING, data_only=True, read_only=True)
    ws = wb["Mapping_final"]
    h = {c.value: c.column for c in ws[1]}
    out = {}
    for r in range(2, ws.max_row + 1):
        code = txt(ws.cell(r, h["N° Word"]).value)
        if code:
            out[code] = {
                "theme": txt(ws.cell(r, h["Thème Word"]).value),
                "title": txt(ws.cell(r, h["Titre Word"]).value),
                "type": txt(ws.cell(r, h["Type de recette"]).value),
                "sources": txt(ws.cell(r, h["Sources Excel"]).value),
            }
    wb.close()
    return out

def load_catalogue():
    wb = load_workbook(CATALOGUE, data_only=True, read_only=True)
    ws = wb["Tables_logiques"]
    h = {c.value: c.column for c in ws[1]}
    out = {}
    for r in range(2, ws.max_row + 1):
        num = txt(ws.cell(r, h["Numéro complet"]).value)
        if num:
            out[num] = {
                "file": txt(ws.cell(r, h["Fichier"]).value),
                "title": txt(ws.cell(r, h["Titre principal"]).value),
                "sheets": txt(ws.cell(r, h["Feuille(s)"]).value),
                "ranges": txt(ws.cell(r, h["Plage(s)"]).value),
            }
    wb.close()
    return out

def find_title_objects(doc, code):
    found = []
    for i, t in enumerate(doc.tables):
        if not t.rows:
            continue
        c = code_from_title(t.cell(0,0).text)
        if c == code:
            found.append(i)
    return found

def table_rows_preview(t, max_rows=5):
    out = []
    for r, row in enumerate(t.rows[:max_rows], start=1):
        vals = [txt(c.text) for c in row.cells]
        out.append(f"R{r}: " + " | ".join(repr(x) for x in vals))
    return out

def source_years(source_id, cat, wb_cache, report):
    if source_id not in cat:
        report.append(f"  SOURCE {source_id}: ABSENTE DU CATALOGUE")
        return

    info = cat[source_id]
    report.append(f"  SOURCE {source_id}")
    report.append(f"    Fichier : {info['file']}")
    report.append(f"    Titre catalogue : {info['title']}")
    report.append(f"    Feuille(s) : {info['sheets']}")
    report.append(f"    Plage(s) : {info['ranges']}")

    if info["file"] not in wb_cache:
        wb_cache[info["file"]] = load_workbook(
            BRUT / info["file"], data_only=True, read_only=True
        )
    wb = wb_cache[info["file"]]

    years = set()
    selected_lines = []

    for rg in [x.strip() for x in info["ranges"].split("|") if x.strip()]:
        if "!" not in rg:
            continue
        sh, bounds = rg.rsplit("!", 1)
        if sh not in wb.sheetnames or ":" not in bounds:
            continue
        try:
            r1, r2 = map(int, bounds.split(":"))
        except Exception:
            continue

        ws = wb[sh]
        for r in range(r1, min(r2, ws.max_row) + 1):
            vals = []
            for c in range(1, min(ws.max_column, 10) + 1):
                v = ws.cell(r,c).value
                if v not in (None, ""):
                    vals.append(f"{ws.cell(r,c).coordinate}={txt(v)!r}")
                    years.update(years_in(v))

            line = " | ".join(vals)
            if line and (
                any(y in line for y in ["2025","2024","2023","2022","2021","2020"])
                or "Source" in line
                or "Année" in line
            ):
                selected_lines.append(line)

    report.append(f"    Années détectées dans la plage : {sorted(years)}")
    for line in selected_lines[:12]:
        report.append(f"    {line}")

def scan_keyword_workbook(path, keyword, report):
    wb = load_workbook(path, data_only=True, read_only=True)
    nk = norm(keyword)
    matches = []
    for ws in wb.worksheets:
        for r in range(1, ws.max_row + 1):
            row_vals = [ws.cell(r,c).value for c in range(1, min(ws.max_column,10)+1)]
            row_text = " ".join(txt(v) for v in row_vals if v not in (None,""))
            if nk in norm(row_text):
                matches.append((ws.title, r, row_text))
    report.append(f"Occurrences de {keyword!r} dans {path.name}: {len(matches)}")
    for sh, r, line in matches[:30]:
        report.append(f"  {sh} R{r}: {line}")
    wb.close()

mapping = load_mapping()
cat = load_catalogue()
doc = Document(DOCX)
wb_cache = {}
report = []

def add(*x):
    report.extend(str(v) for v in x)

add("="*120)
add("DIAGNOSTIC CIBLE DES ALERTES QA FINALES")
add("="*120)
add(f"Word : {DOCX}")
add("")

# 1) Duplicate 3-2
add("#"*120)
add("1) DOUBLON 3-2 — VERIFICATION EXACTE")
add("#"*120)
objs = find_title_objects(doc, "3-2")
add(f"Objets 3-2 trouvés : {[i+1 for i in objs]}")
for i in objs:
    t = doc.tables[i]
    add(f"\nObjet {i+1} | dimensions {len(t.rows)}x{len(t.columns)}")
    add(*table_rows_preview(t, 4))
    for j in range(i+1, min(i+4, len(doc.tables))):
        nt = doc.tables[j]
        add(f"  Suivant objet {j+1} | {len(nt.rows)}x{len(nt.columns)}")
        add(*["    "+x for x in table_rows_preview(nt, 3)])

# 2) Health old years
add("")
add("#"*120)
add("2) TABLES SANTE RESTEES EN 2023 — ANNEE DE LA SOURCE EXCEL")
add("#"*120)
for code in HEALTH_OLD:
    m = mapping.get(code, {})
    add(f"\n{code} | {m.get('title','?')}")
    add(f"  Type : {m.get('type','?')}")
    add(f"  Sources mapping : {m.get('sources','')}")
    srcs = [s.strip() for s in m.get("sources","").split("+") if s.strip()]
    if not srcs:
        add("  AUCUNE SOURCE EXCEL -> modèle à conserver.")
    for src in srcs:
        source_years(src, cat, wb_cache, report)

# 3) Other old years
add("")
add("#"*120)
add("3) AUTRES ANCIENNES ANNEES SIGNALEES")
add("#"*120)
for code in OTHER_OLD:
    m = mapping.get(code, {})
    add(f"\n{code} | {m.get('title','?')}")
    add(f"  Type : {m.get('type','?')}")
    add(f"  Sources mapping : {m.get('sources','')}")
    srcs = [s.strip() for s in m.get("sources","").split("+") if s.strip()]
    if not srcs:
        add("  AUCUNE SOURCE EXCEL -> modèle à conserver.")
    for src in srcs:
        source_years(src, cat, wb_cache, report)

# 4) 13-12 semantic source scan
add("")
add("#"*120)
add("4) 13-12 — RECHERCHE SEMANTIQUE 'IMMATRIC' DANS LE FONCIER")
add("#"*120)
foncier = None
for p in BRUT.glob("*.xlsx"):
    if "foncier" in norm(p.name):
        foncier = p
        break
if foncier:
    scan_keyword_workbook(foncier, "immatric", report)
else:
    add("Fichier Foncier introuvable.")

# 5) 4-4 association check
add("")
add("#"*120)
add("5) 4-4 — VERIFICATION DE L'ASSOCIATION DU TABLEAU")
add("#"*120)
objs44 = find_title_objects(doc, "4-4")
add(f"Objets titre 4-4 : {[i+1 for i in objs44]}")
for i in objs44:
    add(f"\nTitre objet {i+1}:")
    add(*table_rows_preview(doc.tables[i], 4))
    for j in range(i+1, min(i+7, len(doc.tables))):
        add(f"  Objet suivant {j+1} | dimensions {len(doc.tables[j].rows)}x{len(doc.tables[j].columns)}")
        add(*["    "+x for x in table_rows_preview(doc.tables[j], 4)])

# 6) Statistical ellipses only
add("")
add("#"*120)
add("6) ELLIPSES STATISTIQUES REELLES (hors table des matières)")
add("#"*120)
for code in ["1-13","13-12","16-6"]:
    objs = find_title_objects(doc, code)
    add(f"\n{code} | objets titre {[i+1 for i in objs]}")
    for i in objs:
        for j in range(i+1, min(i+5, len(doc.tables))):
            t = doc.tables[j]
            hits = []
            for r, row in enumerate(t.rows, start=1):
                vals = [txt(c.text) for c in row.cells]
                if any(("..." in v or "…" in v) for v in vals):
                    hits.append((r, vals))
            if hits:
                add(f"  Objet {j+1} | {len(t.rows)}x{len(t.columns)}")
                for r, vals in hits:
                    add(f"    R{r}: " + " | ".join(repr(v) for v in vals))

# close workbooks
for wb in wb_cache.values():
    wb.close()

OUT.write_text("\n".join(report), encoding="utf-8")

print("="*100)
print("DIAGNOSTIC CIBLE QA TERMINE")
print("="*100)
print(f"Rapport : {OUT}")
print("="*100)
