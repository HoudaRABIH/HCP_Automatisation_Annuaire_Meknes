from pathlib import Path
from openpyxl import load_workbook
from docx import Document


CODES_WORD = [
    "9-1", "9-2", "9-3", "9-4", "9-5", "9-6",
    "9-7", "9-8", "9-9", "9-10", "9-11"
]

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_MODELES = Path("modeles")
DOSSIER_SORTIE = Path("sorties")

FICHIER_MAPPING = DOSSIER_SORTIE / "mapping_final_recettes.xlsx"
FICHIER_CATALOGUE = DOSSIER_SORTIE / "catalogue_tables_reelles.xlsx"
FICHIER_SORTIE = DOSSIER_SORTIE / "diagnostic_agriculture.txt"

CANDIDATS_WORD = [
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Sources2025_OK.docx",
]


def texte(valeur):
    if valeur is None:
        return ""
    return str(valeur).replace("\n", " ").strip()


def charger_mapping():
    wb = load_workbook(FICHIER_MAPPING, data_only=True, read_only=True)
    ws = wb["Mapping_final"]
    entetes = {c.value: c.column for c in ws[1]}
    resultat = {}

    for r in range(2, ws.max_row + 1):
        code = str(ws.cell(r, entetes["N° Word"]).value or "").strip()

        if code in CODES_WORD:
            resultat[code] = {
                "table_word_donnees": int(
                    ws.cell(r, entetes["Table Word données"]).value
                ),
                "theme": ws.cell(r, entetes["Thème Word"]).value or "",
                "titre_word": ws.cell(r, entetes["Titre Word"]).value or "",
                "type": ws.cell(r, entetes["Type de recette"]).value or "",
                "sources": ws.cell(r, entetes["Sources Excel"]).value or "",
            }

    wb.close()
    return resultat


def charger_catalogue():
    wb = load_workbook(FICHIER_CATALOGUE, data_only=True, read_only=True)
    ws = wb["Tables_logiques"]
    entetes = {c.value: c.column for c in ws[1]}
    resultat = {}

    for r in range(2, ws.max_row + 1):
        numero = str(
            ws.cell(r, entetes["Numéro complet"]).value or ""
        ).strip()

        resultat[numero] = {
            "fichier": ws.cell(r, entetes["Fichier"]).value or "",
            "titre": ws.cell(r, entetes["Titre principal"]).value or "",
            "plages": ws.cell(r, entetes["Plage(s)"]).value or "",
            "feuilles": ws.cell(r, entetes["Feuille(s)"]).value or "",
        }

    wb.close()
    return resultat


def lignes_word_uniques(tableau):
    sortie = []

    for no_ligne, ligne in enumerate(tableau.rows, start=1):
        vus = set()
        cellules = []

        for no_colonne, cellule in enumerate(ligne.cells, start=1):
            identifiant = id(cellule._tc)

            if identifiant in vus:
                continue

            vus.add(identifiant)
            cellules.append(
                f"C{no_colonne}={texte(cellule.text)!r}"
            )

        sortie.append(
            f"R{no_ligne}: " + " | ".join(cellules)
        )

    return sortie


# Cache : chaque fichier Excel n'est ouvert qu'UNE SEULE FOIS.
CACHE_WB = {}


def ouvrir_source(nom_fichier):
    if nom_fichier not in CACHE_WB:
        print(f"  Ouverture Excel : {nom_fichier}")
        CACHE_WB[nom_fichier] = load_workbook(
            DOSSIER_SOURCE / nom_fichier,
            data_only=True,
            read_only=True
        )

    return CACHE_WB[nom_fichier]


def extraire_source(info):
    wb = ouvrir_source(info["fichier"])
    sortie = []

    for morceau in str(info["plages"]).split("|"):
        morceau = morceau.strip()

        if not morceau or "!" not in morceau:
            continue

        nom_feuille, bornes = morceau.rsplit("!", 1)

        if ":" not in bornes:
            continue

        debut_txt, fin_txt = bornes.split(":", 1)

        try:
            debut = int(debut_txt)
            fin = int(fin_txt)
        except ValueError:
            continue

        if nom_feuille not in wb.sheetnames:
            continue

        ws = wb[nom_feuille]
        max_col = min(ws.max_column, 60)

        sortie.append(
            f"--- {nom_feuille} / lignes {debut}:{fin} ---"
        )

        for row in ws.iter_rows(
            min_row=debut,
            max_row=fin,
            min_col=1,
            max_col=max_col,
            values_only=False
        ):
            contenu = []

            for cellule in row:
                if cellule.value in (None, ""):
                    continue

                contenu.append(
                    f"{cellule.coordinate}={texte(cellule.value)!r}"
                )

            if contenu:
                sortie.append(" | ".join(contenu))

    return sortie


mapping = charger_mapping()
catalogue = charger_catalogue()

fichier_word = None

for candidat in CANDIDATS_WORD:
    if candidat.exists():
        fichier_word = candidat
        break

if fichier_word is None:
    fichiers_word = [
        f for f in DOSSIER_MODELES.glob("*.docx")
        if not f.name.startswith("~$")
    ]

    if not fichiers_word:
        raise FileNotFoundError(
            "Aucun document Word trouvé."
        )

    fichier_word = fichiers_word[0]


print("=" * 82)
print("DIAGNOSTIC AGRICULTURE - VERSION RAPIDE")
print("=" * 82)
print(f"Word : {fichier_word}")
print()

document = Document(fichier_word)

rapport = [
    "=" * 100,
    "DIAGNOSTIC DU CHAPITRE AGRICULTURE / ELEVAGE / FORET",
    "=" * 100,
    f"Document Word analysé : {fichier_word}",
]

for position, code in enumerate(CODES_WORD, start=1):
    print(f"[{position}/11] Analyse {code} ...")

    if code not in mapping:
        rapport.append("")
        rapport.append(f"ERREUR : {code} absent du mapping.")
        continue

    recette = mapping[code]

    rapport.extend([
        "",
        "#" * 100,
        f"TABLEAU WORD {code}",
        "#" * 100,
        f"Thème : {recette['theme']}",
        f"Titre Word : {texte(recette['titre_word'])}",
        f"Type : {recette['type']}",
        f"Sources : {recette['sources']}",
        f"N° objet Word : {recette['table_word_donnees']}",
    ])

    table_word = document.tables[
        recette["table_word_donnees"] - 1
    ]

    rapport.append(
        f"Dimensions Word : "
        f"{len(table_word.rows)} x {len(table_word.columns)}"
    )

    rapport.extend([
        "",
        "STRUCTURE WORD",
        "-" * 100,
    ])

    rapport.extend(lignes_word_uniques(table_word))

    sources = [
        s.strip()
        for s in str(recette["sources"]).split("+")
        if s.strip()
    ]

    for source in sources:
        rapport.extend([
            "",
            "STRUCTURE EXCEL",
            "-" * 100,
        ])

        if source not in catalogue:
            rapport.append(
                f"ERREUR : source {source} absente du catalogue."
            )
            continue

        info = catalogue[source]

        rapport.extend([
            f"Source {source}",
            f"Fichier : {info['fichier']}",
            f"Titre : {texte(info['titre'])}",
            f"Feuille(s) : {info['feuilles']}",
            f"Plage(s) : {info['plages']}",
        ])

        rapport.extend(extraire_source(info))


for wb in CACHE_WB.values():
    wb.close()


FICHIER_SORTIE.write_text(
    "\n".join(rapport),
    encoding="utf-8"
)

print()
print("=" * 82)
print("DIAGNOSTIC AGRICULTURE TERMINE")
print("=" * 82)
print(f"Fichier créé : {FICHIER_SORTIE}")
print("=" * 82)
