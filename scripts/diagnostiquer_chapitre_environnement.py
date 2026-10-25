from pathlib import Path
from openpyxl import load_workbook
from docx import Document


CODES_WORD = ["16-1", "16-2", "16-3", "16-4", "16-5", "16-6"]

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_MODELES = Path("modeles")
DOSSIER_SORTIE = Path("sorties")

FICHIER_MAPPING = DOSSIER_SORTIE / "mapping_final_recettes.xlsx"
FICHIER_CATALOGUE = DOSSIER_SORTIE / "catalogue_tables_reelles.xlsx"
FICHIER_SORTIE = DOSSIER_SORTIE / "diagnostic_environnement.txt"

# On diagnostique le Word cumulé Transport + Prix s'il existe,
# sinon on retombe sur le modèle.
FICHIER_WORD_CUMULE = DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix.docx"


def texte(valeur):
    if valeur is None:
        return ""
    return str(valeur).replace("\n", " ").strip()


def charger_mapping():
    wb = load_workbook(FICHIER_MAPPING, data_only=True)
    ws = wb["Mapping_final"]
    entetes = {c.value: c.column for c in ws[1]}

    resultat = {}

    for r in range(2, ws.max_row + 1):
        code = str(ws.cell(r, entetes["N° Word"]).value or "").strip()

        if code in CODES_WORD:
            resultat[code] = {
                "table_word_donnees": int(
                    ws.cell(r, entetes["Table Word données"]).value
                ),
                "theme": ws.cell(r, entetes["Thème Word"]).value or "",
                "titre_word": ws.cell(r, entetes["Titre Word"]).value or "",
                "type": ws.cell(r, entetes["Type de recette"]).value or "",
                "sources": ws.cell(r, entetes["Sources Excel"]).value or "",
            }

    wb.close()
    return resultat


def charger_catalogue():
    wb = load_workbook(FICHIER_CATALOGUE, data_only=True)
    ws = wb["Tables_logiques"]
    entetes = {c.value: c.column for c in ws[1]}

    resultat = {}

    for r in range(2, ws.max_row + 1):
        numero = str(
            ws.cell(r, entetes["Numéro complet"]).value or ""
        ).strip()

        resultat[numero] = {
            "fichier": ws.cell(r, entetes["Fichier"]).value or "",
            "titre": ws.cell(r, entetes["Titre principal"]).value or "",
            "plages": ws.cell(r, entetes["Plage(s)"]).value or "",
            "feuilles": ws.cell(r, entetes["Feuille(s)"]).value or "",
        }

    wb.close()
    return resultat


def lignes_word_uniques(tableau):
    sortie = []

    for no_ligne, ligne in enumerate(tableau.rows, start=1):
        vus = set()
        cellules = []

        for no_colonne, cellule in enumerate(ligne.cells, start=1):
            identifiant = id(cellule._tc)

            if identifiant in vus:
                continue

            vus.add(identifiant)
            cellules.append(
                f"C{no_colonne}={texte(cellule.text)!r}"
            )

        sortie.append(
            f"R{no_ligne}: " + " | ".join(cellules)
        )

    return sortie


def extraire_source(info):
    fichier = DOSSIER_SOURCE / info["fichier"]

    wb = load_workbook(
        fichier,
        data_only=True,
        read_only=False
    )

    sortie = []

    for morceau in str(info["plages"]).split("|"):
        morceau = morceau.strip()

        if not morceau or "!" not in morceau:
            continue

        nom_feuille, bornes = morceau.rsplit("!", 1)

        if ":" not in bornes:
            continue

        debut_txt, fin_txt = bornes.split(":", 1)

        try:
            debut = int(debut_txt)
            fin = int(fin_txt)
        except ValueError:
            continue

        if nom_feuille not in wb.sheetnames:
            continue

        ws = wb[nom_feuille]
        max_col = min(ws.max_column, 80)

        sortie.append(
            f"--- {nom_feuille} / lignes {debut}:{fin} ---"
        )

        for row in ws.iter_rows(
            min_row=debut,
            max_row=fin,
            min_col=1,
            max_col=max_col
        ):
            contenu = []

            for cellule in row:
                if cellule.value in (None, ""):
                    continue

                contenu.append(
                    f"{cellule.coordinate}={texte(cellule.value)!r}"
                )

            if contenu:
                sortie.append(" | ".join(contenu))

    wb.close()
    return sortie


mapping = charger_mapping()
catalogue = charger_catalogue()

if FICHIER_WORD_CUMULE.exists():
    fichier_word = FICHIER_WORD_CUMULE
else:
    fichiers_word = [
        f for f in DOSSIER_MODELES.glob("*.docx")
        if not f.name.startswith("~$")
    ]

    if not fichiers_word:
        raise FileNotFoundError(
            "Aucun .docx trouvé."
        )

    fichier_word = fichiers_word[0]


document = Document(fichier_word)

rapport = []

rapport.append("=" * 100)
rapport.append("DIAGNOSTIC DU CHAPITRE ENVIRONNEMENT")
rapport.append("=" * 100)
rapport.append(f"Document Word analysé : {fichier_word}")

for code in CODES_WORD:
    if code not in mapping:
        rapport.append("")
        rapport.append(f"ERREUR : {code} absent du mapping.")
        continue

    recette = mapping[code]

    rapport.append("")
    rapport.append("#" * 100)
    rapport.append(f"TABLEAU WORD {code}")
    rapport.append("#" * 100)
    rapport.append(f"Thème : {recette['theme']}")
    rapport.append(f"Titre Word : {texte(recette['titre_word'])}")
    rapport.append(f"Type : {recette['type']}")
    rapport.append(f"Sources : {recette['sources']}")
    rapport.append(f"N° objet Word : {recette['table_word_donnees']}")

    table_word = document.tables[
        recette["table_word_donnees"] - 1
    ]

    rapport.append(
        f"Dimensions Word : "
        f"{len(table_word.rows)} x {len(table_word.columns)}"
    )

    rapport.append("")
    rapport.append("STRUCTURE WORD")
    rapport.append("-" * 100)

    rapport.extend(
        lignes_word_uniques(table_word)
    )

    sources = [
        s.strip()
        for s in str(recette["sources"]).split("+")
        if s.strip()
    ]

    for source in sources:
        rapport.append("")
        rapport.append("STRUCTURE EXCEL")
        rapport.append("-" * 100)

        if source not in catalogue:
            rapport.append(
                f"ERREUR : source {source} absente du catalogue."
            )
            continue

        info = catalogue[source]

        rapport.append(f"Source {source}")
        rapport.append(f"Fichier : {info['fichier']}")
        rapport.append(f"Titre : {texte(info['titre'])}")
        rapport.append(f"Feuille(s) : {info['feuilles']}")
        rapport.append(f"Plage(s) : {info['plages']}")

        rapport.extend(
            extraire_source(info)
        )


FICHIER_SORTIE.write_text(
    "\n".join(rapport),
    encoding="utf-8"
)

print("=" * 76)
print("DIAGNOSTIC ENVIRONNEMENT TERMINE")
print("=" * 76)
print("Tableaux analysés : 16-1 à 16-6")
print(f"Fichier créé : {FICHIER_SORTIE}")
print("=" * 76)
