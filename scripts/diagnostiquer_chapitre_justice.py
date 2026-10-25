from pathlib import Path
from openpyxl import load_workbook
from zipfile import ZipFile
from lxml import etree


CODES_WORD = [
    "8-1", "8-2", "8-3", "8-4", "8-5",
    "8-6", "8-7", "8-8", "8-9", "8-10", "8-11"
]

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_MODELES = Path("modeles")
DOSSIER_SORTIE = Path("sorties")

FICHIER_MAPPING = DOSSIER_SORTIE / "mapping_final_recettes.xlsx"
FICHIER_CATALOGUE = DOSSIER_SORTIE / "catalogue_tables_reelles.xlsx"
FICHIER_SORTIE = DOSSIER_SORTIE / "diagnostic_justice.txt"

# Le fichier Education corrigé devient notre nouvelle base d'entrée.
CANDIDATS_WORD = [
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_OK.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN.docx",
]


def txt(v):
    if v is None:
        return ""
    return str(v).replace("\n", " ").strip()


def charger_mapping():
    print("1/5 Lecture du mapping...", flush=True)

    wb = load_workbook(
        FICHIER_MAPPING,
        data_only=True,
        read_only=True
    )
    ws = wb["Mapping_final"]
    h = {c.value: c.column for c in ws[1]}

    resultat = {}

    for r in range(2, ws.max_row + 1):
        code = txt(ws.cell(r, h["N° Word"]).value)

        if code not in CODES_WORD:
            continue

        resultat[code] = {
            "table_word": int(
                ws.cell(r, h["Table Word données"]).value
            ),
            "theme": txt(
                ws.cell(r, h["Thème Word"]).value
            ),
            "titre": txt(
                ws.cell(r, h["Titre Word"]).value
            ),
            "type": txt(
                ws.cell(r, h["Type de recette"]).value
            ),
            "sources": txt(
                ws.cell(r, h["Sources Excel"]).value
            ),
        }

    wb.close()
    return resultat


def charger_catalogue():
    print("2/5 Lecture du catalogue...", flush=True)

    wb = load_workbook(
        FICHIER_CATALOGUE,
        data_only=True,
        read_only=True
    )
    ws = wb["Tables_logiques"]
    h = {c.value: c.column for c in ws[1]}

    resultat = {}

    for r in range(2, ws.max_row + 1):
        numero = txt(
            ws.cell(r, h["Numéro complet"]).value
        )

        resultat[numero] = {
            "fichier": txt(
                ws.cell(r, h["Fichier"]).value
            ),
            "titre": txt(
                ws.cell(r, h["Titre principal"]).value
            ),
            "plages": txt(
                ws.cell(r, h["Plage(s)"]).value
            ),
            "feuilles": txt(
                ws.cell(r, h["Feuille(s)"]).value
            ),
        }

    wb.close()
    return resultat


def choisir_word():
    for fichier in CANDIDATS_WORD:
        if fichier.exists():
            return fichier

    # secours : prendre le fichier .docx le plus récent de sorties
    fichiers_sortie = sorted(
        [
            f for f in DOSSIER_SORTIE.glob("*.docx")
            if not f.name.startswith("~$")
        ],
        key=lambda p: p.stat().st_mtime,
        reverse=True
    )

    if fichiers_sortie:
        return fichiers_sortie[0]

    fichiers_modeles = [
        f for f in DOSSIER_MODELES.glob("*.docx")
        if not f.name.startswith("~$")
    ]

    if not fichiers_modeles:
        raise FileNotFoundError(
            "Aucun document Word trouvé."
        )

    return fichiers_modeles[0]


def lire_tables_top_level(docx_path):
    ns = {
        "w": (
            "http://schemas.openxmlformats.org/"
            "wordprocessingml/2006/main"
        )
    }

    with ZipFile(docx_path, "r") as z:
        xml = z.read("word/document.xml")

    root = etree.fromstring(xml)
    body = root.find("w:body", namespaces=ns)
    tables = body.findall("w:tbl", namespaces=ns)

    return tables, ns


def decrire_table(table, ns):
    lignes = []
    rows = table.findall("w:tr", namespaces=ns)
    max_cols = 0

    for no_ligne, row in enumerate(rows, start=1):
        cells = row.findall("w:tc", namespaces=ns)
        max_cols = max(max_cols, len(cells))

        morceaux = []

        for no_col, cell in enumerate(cells, start=1):
            textes = cell.xpath(
                ".//w:t/text()",
                namespaces=ns
            )

            contenu = " ".join(
                t.strip()
                for t in textes
                if t.strip()
            )

            morceaux.append(
                f"C{no_col}={contenu!r}"
            )

        lignes.append(
            f"R{no_ligne}: " + " | ".join(morceaux)
        )

    return len(rows), max_cols, lignes


CACHE_EXCEL = {}


def ouvrir_excel(nom_fichier):
    if nom_fichier not in CACHE_EXCEL:
        print(
            f"   Ouverture Excel : {nom_fichier}",
            flush=True
        )

        CACHE_EXCEL[nom_fichier] = load_workbook(
            DOSSIER_SOURCE / nom_fichier,
            data_only=True,
            read_only=True
        )

    return CACHE_EXCEL[nom_fichier]


def extraire_plages(info):
    wb = ouvrir_excel(info["fichier"])
    sortie = []

    plages = [
        p.strip()
        for p in info["plages"].split("|")
        if p.strip()
    ]

    for plage in plages:
        if "!" not in plage:
            continue

        nom_feuille, bornes = plage.rsplit("!", 1)

        if ":" not in bornes:
            continue

        debut_txt, fin_txt = bornes.split(":", 1)

        try:
            debut = int(debut_txt)
            fin = int(fin_txt)
        except ValueError:
            continue

        if nom_feuille not in wb.sheetnames:
            sortie.append(
                f"ERREUR: feuille {nom_feuille!r} absente"
            )
            continue

        ws = wb[nom_feuille]

        sortie.append(
            f"--- {nom_feuille} / lignes {debut}:{fin} ---"
        )

        max_col = min(ws.max_column, 35)

        for row in ws.iter_rows(
            min_row=debut,
            max_row=fin,
            min_col=1,
            max_col=max_col
        ):
            valeurs = []

            for cell in row:
                if cell.value in (None, ""):
                    continue

                valeurs.append(
                    f"{cell.coordinate}={txt(cell.value)!r}"
                )

            if valeurs:
                sortie.append(" | ".join(valeurs))

    return sortie


mapping = charger_mapping()
catalogue = charger_catalogue()
word_path = choisir_word()

print("3/5 Lecture directe du Word...", flush=True)

tables, ns = lire_tables_top_level(word_path)

print(
    f"   Word utilisé : {word_path}",
    flush=True
)
print(
    f"   Tableaux Word top-level : {len(tables)}",
    flush=True
)

print("4/5 Lecture des sources Justice...", flush=True)

rapport = [
    "=" * 110,
    "DIAGNOSTIC DU CHAPITRE JUSTICE",
    "=" * 110,
    f"Document Word analysé : {word_path}",
    f"Nombre de tableaux top-level : {len(tables)}",
]

for position, code in enumerate(CODES_WORD, start=1):
    print(
        f"   [{position}/{len(CODES_WORD)}] {code}",
        flush=True
    )

    rapport.extend([
        "",
        "#" * 110,
        f"TABLEAU WORD {code}",
        "#" * 110,
    ])

    if code not in mapping:
        rapport.append(
            "ERREUR : code absent du mapping."
        )
        continue

    m = mapping[code]

    rapport.extend([
        f"Thème : {m['theme']}",
        f"Titre Word : {m['titre']}",
        f"Type : {m['type']}",
        f"Sources Excel : {m['sources']}",
        f"N° objet Word : {m['table_word']}",
    ])

    index = m["table_word"] - 1

    if index < 0 or index >= len(tables):
        rapport.append(
            "ERREUR : objet Word hors limites."
        )
    else:
        nb_lignes, nb_cols, lignes = decrire_table(
            tables[index],
            ns
        )

        rapport.extend([
            f"Dimensions Word réelles : {nb_lignes} x {nb_cols}",
            "",
            "STRUCTURE WORD",
            "-" * 110,
        ])

        rapport.extend(lignes)

    sources = [
        s.strip()
        for s in m["sources"].split("+")
        if s.strip()
    ]

    if not sources:
        rapport.extend([
            "",
            "AUCUNE SOURCE EXCEL : CONSERVER LE MODELE.",
        ])

    for source in sources:
        rapport.extend([
            "",
            "STRUCTURE EXCEL",
            "-" * 110,
        ])

        if source not in catalogue:
            rapport.append(
                f"ERREUR : source {source} absente du catalogue."
            )
            continue

        info = catalogue[source]

        rapport.extend([
            f"Source {source}",
            f"Fichier : {info['fichier']}",
            f"Titre : {info['titre']}",
            f"Feuille(s) : {info['feuilles']}",
            f"Plage(s) : {info['plages']}",
        ])

        rapport.extend(
            extraire_plages(info)
        )


for wb in CACHE_EXCEL.values():
    wb.close()

print("5/5 Ecriture du diagnostic...", flush=True)

FICHIER_SORTIE.write_text(
    "\n".join(rapport),
    encoding="utf-8"
)

print()
print("=" * 86)
print("DIAGNOSTIC JUSTICE TERMINE")
print("=" * 86)
print(
    f"Tableaux analysés : {len(CODES_WORD)}"
)
print(
    f"Fichier créé : {FICHIER_SORTIE}"
)
print("=" * 86)
