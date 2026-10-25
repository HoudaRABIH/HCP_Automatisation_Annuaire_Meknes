from pathlib import Path
from openpyxl import load_workbook
from docx import Document
import re
import unicodedata

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_SORTIE = Path("sorties")

FICHIER_MAPPING = DOSSIER_SORTIE / "mapping_final_recettes.xlsx"
FICHIER_CATALOGUE = DOSSIER_SORTIE / "catalogue_tables_reelles.xlsx"
FICHIER_SORTIE = DOSSIER_SORTIE / "diagnostic_emploi.txt"

CANDIDATS_WORD = [
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER_POPULATION.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER.docx",
]

def txt(v):
    return "" if v is None else str(v).replace("\n", " ").strip()

def norm(v):
    s = txt(v).lower()
    s = (
        s.replace("’", "'")
         .replace("–", "-")
         .replace("—", "-")
         .replace("\u202f", " ")
         .replace("\xa0", " ")
    )
    s = "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"\s+", " ", s).strip()

def choisir_word():
    for p in CANDIDATS_WORD:
        if p.exists():
            return p

    docs = sorted(
        [p for p in DOSSIER_SORTIE.glob("*.docx") if not p.name.startswith("~$")],
        key=lambda p: p.stat().st_mtime,
        reverse=True
    )

    if not docs:
        raise FileNotFoundError("Aucun Word trouvé dans sorties.")

    return docs[0]

def charger_mapping():
    wb = load_workbook(FICHIER_MAPPING, data_only=True, read_only=True)
    ws = wb["Mapping_final"]
    h = {c.value: c.column for c in ws[1]}

    rows = []

    for r in range(2, ws.max_row + 1):
        code = txt(ws.cell(r, h["N° Word"]).value)

        if not code.startswith("7-"):
            continue

        rows.append({
            "code": code,
            "theme": txt(ws.cell(r, h["Thème Word"]).value),
            "titre": txt(ws.cell(r, h["Titre Word"]).value),
            "type": txt(ws.cell(r, h["Type de recette"]).value),
            "sources": txt(ws.cell(r, h["Sources Excel"]).value),
        })

    wb.close()

    def cle(x):
        try:
            return int(x["code"].split("-")[1])
        except Exception:
            return 999

    rows.sort(key=cle)
    return rows

def charger_catalogue():
    wb = load_workbook(FICHIER_CATALOGUE, data_only=True, read_only=True)
    ws = wb["Tables_logiques"]
    h = {c.value: c.column for c in ws[1]}

    out = {}

    for r in range(2, ws.max_row + 1):
        numero = txt(ws.cell(r, h["Numéro complet"]).value)

        out[numero] = {
            "fichier": txt(ws.cell(r, h["Fichier"]).value),
            "titre": txt(ws.cell(r, h["Titre principal"]).value),
            "feuilles": txt(ws.cell(r, h["Feuille(s)"]).value),
            "plages": txt(ws.cell(r, h["Plage(s)"]).value),
        }

    wb.close()
    return out

def trouver_table_word(doc, code):
    numero = code.split("-")[1]
    pat = re.compile(rf"^\s*7\s*-\s*{re.escape(numero)}\b", re.I)

    for i, table in enumerate(doc.tables):
        if not table.rows:
            continue

        titre = table.cell(0, 0).text or ""

        if pat.search(titre):
            for j in range(i + 1, min(i + 6, len(doc.tables))):
                cand = doc.tables[j]

                if len(cand.rows) >= 3:
                    return i + 1, cand

            return i + 1, None

    return None, None

def dump_word(table):
    if table is None:
        return ["AUCUN TABLEAU DE DONNEES WORD TROUVE."]

    out = [
        f"Dimensions Word : {len(table.rows)} x {len(table.columns)}"
    ]

    for r, row in enumerate(table.rows, start=1):
        vals = []

        for c, cell in enumerate(row.cells, start=1):
            vals.append(f"C{c}={txt(cell.text)!r}")

        out.append(f"R{r}: " + " | ".join(vals))

    return out

CACHE = {}

def ouvrir_excel(nom):
    if nom not in CACHE:
        print(f"   Ouverture Excel : {nom}", flush=True)

        CACHE[nom] = load_workbook(
            DOSSIER_SOURCE / nom,
            data_only=True,
            read_only=True
        )

    return CACHE[nom]

def dump_source(numero, info):
    wb = ouvrir_excel(info["fichier"])

    out = [
        f"Source {numero}",
        f"Fichier : {info['fichier']}",
        f"Titre catalogue : {info['titre']}",
        f"Feuille(s) catalogue : {info['feuilles']}",
        f"Plage(s) catalogue : {info['plages']}",
    ]

    dumped = False

    for plage in [p.strip() for p in info["plages"].split("|") if p.strip()]:
        if "!" not in plage:
            continue

        sheet, bounds = plage.rsplit("!", 1)

        if ":" not in bounds or sheet not in wb.sheetnames:
            continue

        try:
            r1, r2 = map(int, bounds.split(":", 1))
        except ValueError:
            continue

        ws = wb[sheet]
        out.append(f"--- {sheet} / lignes {r1}:{r2} ---")

        for row in ws.iter_rows(
            min_row=r1,
            max_row=r2,
            min_col=1,
            max_col=min(ws.max_column, 20)
        ):
            vals = [
                f"{cell.coordinate}={txt(cell.value)!r}"
                for cell in row
                if cell.value not in (None, "")
            ]

            if vals:
                out.append(" | ".join(vals))

        dumped = True

    if not dumped:
        a, b = numero.split("-")
        pat = re.compile(
            rf"^\s*{re.escape(a)}\s*-\s*{re.escape(b)}\b",
            re.I
        )

        for ws in wb.worksheets:
            found = None

            for r in range(1, ws.max_row + 1):
                for c in range(1, min(ws.max_column, 10) + 1):
                    v = ws.cell(r, c).value

                    if isinstance(v, str) and pat.match(v):
                        found = r
                        break

                if found:
                    break

            if found:
                out.append(
                    f"--- Source retrouvée dynamiquement : {ws.title} / "
                    f"lignes {found}:{min(ws.max_row, found + 100)} ---"
                )

                for row in ws.iter_rows(
                    min_row=found,
                    max_row=min(ws.max_row, found + 100),
                    min_col=1,
                    max_col=min(ws.max_column, 20)
                ):
                    vals = [
                        f"{cell.coordinate}={txt(cell.value)!r}"
                        for cell in row
                        if cell.value not in (None, "")
                    ]

                    if vals:
                        out.append(" | ".join(vals))

                dumped = True
                break

    if not dumped:
        out.append("ERREUR : source Excel introuvable dans le classeur.")

    return out

word_path = choisir_word()
mapping = charger_mapping()
catalogue = charger_catalogue()
doc = Document(word_path)

print("=" * 92)
print("DIAGNOSTIC EMPLOI")
print("=" * 92)
print(f"Word utilisé : {word_path}")
print(f"Tableaux 7-x détectés : {len(mapping)}")
print()

rapport = [
    "=" * 110,
    "DIAGNOSTIC EMPLOI",
    "=" * 110,
    f"Document Word analysé : {word_path}",
    f"Nombre de recettes 7-x : {len(mapping)}",
]

for i, m in enumerate(mapping, start=1):
    code = m["code"]
    print(f"[{i}/{len(mapping)}] {code}", flush=True)

    rapport.extend([
        "",
        "#" * 110,
        f"TABLEAU WORD {code}",
        "#" * 110,
        f"Thème : {m['theme']}",
        f"Titre : {m['titre']}",
        f"Type : {m['type']}",
        f"Sources Excel : {m['sources']}",
    ])

    obj, table = trouver_table_word(doc, code)

    if obj is None:
        rapport.append("ERREUR : titre Word introuvable.")
    else:
        rapport.append(f"Objet titre Word dynamique : {obj}")
        rapport.append("")
        rapport.append("STRUCTURE WORD")
        rapport.append("-" * 110)
        rapport.extend(dump_word(table))

    sources = [s.strip() for s in m["sources"].split("+") if s.strip()]

    if not sources:
        rapport.extend([
            "",
            "AUCUNE SOURCE EXCEL -> CONSERVER LE MODELE.",
        ])
        continue

    for source in sources:
        rapport.extend([
            "",
            "STRUCTURE EXCEL",
            "-" * 110,
        ])

        if source not in catalogue:
            rapport.append(
                f"ERREUR : source {source} absente du catalogue."
            )
            continue

        rapport.extend(
            dump_source(
                source,
                catalogue[source]
            )
        )

for wb in CACHE.values():
    wb.close()

FICHIER_SORTIE.write_text(
    "\n".join(rapport),
    encoding="utf-8"
)

print()
print("=" * 92)
print("DIAGNOSTIC EMPLOI TERMINE")
print("=" * 92)
print(f"Fichier créé : {FICHIER_SORTIE}")
print("=" * 92)
