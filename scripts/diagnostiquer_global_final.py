from pathlib import Path
from collections import Counter, defaultdict
from zipfile import ZipFile
from docx import Document
from lxml import etree
import re
import unicodedata

DOSSIER_SORTIE = Path("sorties")

CANDIDATS = [
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER_POPULATION_INDUSTRIE_ENERGIE_EAU.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER_POPULATION_INDUSTRIE.docx",
]

FICHIER_SORTIE = DOSSIER_SORTIE / "diagnostic_global_final.txt"

# Tables where old/reference years are expected and must NOT automatically be treated as errors.
ANCIENNES_ANNEES_AUTORISEES = {
    "1-7": "Comparaison 2014/2024",
    "2-13": "Source disponible en 2023",
    "2-18": "Source disponible en 2023",
    "6-3": "Source disponible en 2022",
    "6-4": "Source disponible en 2018",
    "6-5": "Source disponible en 2019",
    "7-2": "Historique du SMIG/SMAG",
    "7-3": "Personnel de l'Etat au 31-12-2020",
    "11-1": "Source ONEE Eau disponible en 2023",
    "11-2": "Série 2024*/2023R/2022R",
    "13-15": "Série 2024/2023/2022",
    "13-16": "Série 2024/2023/2022",
    "14-2": "Série 2024/2023",
    "15-3": "Tarifs historiques électricité",
    "15-4": "Tarifs historiques eau",
}

def norm(v):
    s = str(v or "").strip().lower()
    s = (
        s.replace("’", "'")
         .replace("–", "-")
         .replace("—", "-")
         .replace("\u202f", " ")
         .replace("\xa0", " ")
    )
    s = "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"\s+", " ", s).strip()

def choose_docx():
    for p in CANDIDATS:
        if p.exists():
            return p
    docs = sorted(
        [p for p in DOSSIER_SORTIE.glob("*.docx") if not p.name.startswith("~$")],
        key=lambda p: p.stat().st_mtime,
        reverse=True
    )
    if not docs:
        raise FileNotFoundError("Aucun fichier Word trouvé dans sorties.")
    return docs[0]

def canonical_code(text):
    m = re.match(r"^\s*(\d+)\s*-\s*(\d+)\b", str(text or ""))
    if not m:
        return None
    return f"{int(m.group(1))}-{int(m.group(2))}"

def table_text(table):
    return "\n".join(
        cell.text or ""
        for row in table.rows
        for cell in row.cells
    )

def year_tokens(text):
    years = re.findall(r"\b(?:19|20)\d{2}\b", text or "")
    return sorted(set(years))

def next_data_table(doc, title_idx):
    for j in range(title_idx + 1, min(title_idx + 7, len(doc.tables))):
        t = doc.tables[j]
        if len(t.rows) >= 3:
            # Avoid another title table if possible.
            first = (t.cell(0, 0).text or "").strip()
            if canonical_code(first):
                continue
            return j
    return None

def extract_all_xml_text(docx_path):
    texts = []
    with ZipFile(docx_path, "r") as zf:
        names = [
            n for n in zf.namelist()
            if n.startswith("word/") and n.endswith(".xml")
        ]
        for name in names:
            try:
                root = etree.fromstring(zf.read(name))
            except Exception:
                continue
            for node in root.iter():
                if node.text:
                    texts.append(node.text)
    return "\n".join(texts)

docx = choose_docx()
doc = Document(docx)

report = []
def add(*lines):
    report.extend(str(x) for x in lines)

add("=" * 120)
add("DIAGNOSTIC GLOBAL FINAL — ANNUAIRE HCP")
add("=" * 120)
add(f"Document analysé : {docx}")
add(f"Nombre de tableaux top-level : {len(doc.tables)}")
add("")

# ============================================================================
# 1. INDEX DES TABLEAUX WORD
# ============================================================================
title_map = {}
duplicates = defaultdict(list)

for i, t in enumerate(doc.tables):
    if not t.rows:
        continue
    first = t.cell(0, 0).text or ""
    code = canonical_code(first)
    if code:
        duplicates[code].append(i)
        if code not in title_map:
            title_map[code] = i

add("#" * 120)
add("1) INDEX DES TABLEAUX WORD")
add("#" * 120)
add(f"Codes de tableaux détectés : {len(title_map)}")
add("")

dupes = {k: v for k, v in duplicates.items() if len(v) > 1}
if dupes:
    add("DOUBLONS DE TITRES DETECTES :")
    for code, idxs in sorted(dupes.items()):
        add(f"  {code} -> objets {', '.join(str(i+1) for i in idxs)}")
else:
    add("Aucun doublon de titre détecté.")
add("")

# ============================================================================
# 2. SOURCES ANNUAIRE 2023 / 2025 — y compris XML/floating text boxes
# ============================================================================
xml_text = extract_all_xml_text(docx)
count_annuaire_fr_2023 = xml_text.count("Annuaire Statistique du Maroc 2023")
count_annuaire_ar_2023 = xml_text.count("النشرة الإحصائية السنوية للمغرب 2023")
count_annuaire_fr_2025 = xml_text.count("Annuaire Statistique du Maroc 2025")
count_annuaire_ar_2025 = xml_text.count("النشرة الإحصائية السنوية للمغرب 2025")

add("#" * 120)
add("2) SOURCES ANNUAIRE")
add("#" * 120)
add(f"Annuaire Statistique du Maroc 2023 (FR) : {count_annuaire_fr_2023}")
add(f"Annuaire ... 2023 (AR) : {count_annuaire_ar_2023}")
add(f"Annuaire Statistique du Maroc 2025 (FR) : {count_annuaire_fr_2025}")
add(f"Annuaire ... 2025 (AR) : {count_annuaire_ar_2025}")
if count_annuaire_fr_2023 or count_annuaire_ar_2023:
    add("ALERTE : au moins une ancienne source Annuaire 2023 subsiste.")
else:
    add("OK : aucune source Annuaire 2023 détectée dans les XML Word.")
add("")

# ============================================================================
# 3. ELLIPSES / DONNEES NON DISPONIBLES
# ============================================================================
add("#" * 120)
add("3) CELLULES CONTENANT ... OU …")
add("#" * 120)

ellipsis_hits = []
for i, t in enumerate(doc.tables):
    for r, row in enumerate(t.rows):
        for c, cell in enumerate(row.cells):
            text = cell.text or ""
            if "..." in text or "…" in text:
                ellipsis_hits.append((i, r, c, text.strip()))

if not ellipsis_hits:
    add("Aucune ellipse détectée.")
else:
    add(f"Nombre de cellules avec ellipse : {len(ellipsis_hits)}")
    for i, r, c, text in ellipsis_hits:
        # associate with nearest preceding title code
        code = None
        for k in range(i, -1, -1):
            if not doc.tables[k].rows:
                continue
            code = canonical_code(doc.tables[k].cell(0, 0).text or "")
            if code:
                break
        add(f"  {code or '?'} | table {i+1} R{r+1}C{c+1} = {text!r}")
add("")

# ============================================================================
# 4. ANNEES PAR TABLEAU
# ============================================================================
add("#" * 120)
add("4) ANNEES DETECTEES PAR TABLEAU")
add("#" * 120)

suspicious_old_years = []

def code_key(code):
    a, b = code.split("-")
    return (int(a), int(b))

for code in sorted(title_map, key=code_key):
    ti = title_map[code]
    di = next_data_table(doc, ti)
    title = doc.tables[ti].cell(0, 0).text.replace("\n", " ").strip()
    if di is None:
        add(f"{code}: PAS DE TABLEAU DE DONNEES TROUVE | {title}")
        continue
    data_text = table_text(doc.tables[di])
    years = year_tokens(data_text)
    add(f"{code}: {', '.join(years) if years else 'aucune année explicite'} | {title}")

    old = [y for y in years if int(y) <= 2023]
    has_2024_or_2025 = any(int(y) >= 2024 for y in years)

    if old and code not in ANCIENNES_ANNEES_AUTORISEES and not has_2024_or_2025:
        suspicious_old_years.append((code, old, title))

add("")
if suspicious_old_years:
    add("ANCIENNES ANNEES A REVOIR (hors exceptions connues) :")
    for code, years, title in suspicious_old_years:
        add(f"  {code}: {years} | {title}")
else:
    add("Aucune ancienne année isolée suspecte détectée hors exceptions connues.")
add("")

# ============================================================================
# 5. SOURCES FR/AR DANS LES TABLES
# ============================================================================
add("#" * 120)
add("5) COHERENCE DES LIGNES SOURCE FR / AR")
add("#" * 120)

source_issues = []
source_tables = 0

for i, t in enumerate(doc.tables):
    text = table_text(t)
    n = norm(text)

    has_fr = (
        "source :" in n
        or "sources :" in n
        or "source:" in n
        or "sources:" in n
    )
    has_ar = "المصدر" in text or "المصادر" in text

    if has_fr or has_ar:
        source_tables += 1
        if has_fr != has_ar:
            source_issues.append((i, has_fr, has_ar, text[:350].replace("\n", " | ")))

add(f"Tableaux contenant une ligne Source/Sources : {source_tables}")
if source_issues:
    add(f"ALERTES FR/AR : {len(source_issues)}")
    for i, fr, ar, snippet in source_issues:
        code = None
        for k in range(i, -1, -1):
            if doc.tables[k].rows:
                code = canonical_code(doc.tables[k].cell(0, 0).text or "")
                if code:
                    break
        add(f"  {code or '?'} | table {i+1} | FR={fr} AR={ar} | {snippet}")
else:
    add("OK : aucune source FR-only / AR-only détectée dans les tableaux.")
add("")

# ============================================================================
# 6. CELLULES VIDES SUSPECTES
# ============================================================================
add("#" * 120)
add("6) CELLULES VIDES SUSPECTES DANS LES ZONES NUMERIQUES")
add("#" * 120)

blank_hits = []

for code, ti in title_map.items():
    di = next_data_table(doc, ti)
    if di is None:
        continue

    t = doc.tables[di]

    for r, row in enumerate(t.rows):
        texts = [(cell.text or "").strip() for cell in row.cells]
        if not texts:
            continue

        first = norm(texts[0])

        # Ignore obvious headers/source/notes.
        if not first:
            continue
        if first.startswith(("source", "sources", "note", "dont", "cercle", "milieu")):
            continue
        if "المصدر" in " ".join(texts):
            continue

        numericish = 0
        blanks_mid = []

        for c in range(1, max(1, len(texts) - 1)):
            val = texts[c]
            if re.search(r"\d", val) or val in ("-", "...", "…"):
                numericish += 1
            elif val == "":
                blanks_mid.append(c)

        if numericish >= 1 and blanks_mid:
            blank_hits.append((code, di, r, blanks_mid, texts))

if blank_hits:
    add(f"Cellules vides potentiellement à vérifier : {len(blank_hits)}")
    for code, di, r, cols, texts in blank_hits[:200]:
        add(
            f"  {code} | table {di+1} R{r+1} | "
            f"colonnes vides {[c+1 for c in cols]} | "
            + " | ".join(repr(x) for x in texts)
        )
else:
    add("Aucune cellule vide suspecte détectée par la règle automatique.")
add("")

# ============================================================================
# 7. CONTROLES STRUCTURELS
# ============================================================================
add("#" * 120)
add("7) CONTROLES STRUCTURELS")
add("#" * 120)

small_data = []
for code, ti in title_map.items():
    di = next_data_table(doc, ti)
    if di is None:
        continue
    t = doc.tables[di]
    if len(t.rows) < 3 or len(t.columns) < 2:
        small_data.append((code, di, len(t.rows), len(t.columns)))

if small_data:
    for code, di, rows, cols in small_data:
        add(f"ALERTE {code}: table {di+1} = {rows}x{cols}")
else:
    add("OK : aucune table de données anormalement petite détectée.")
add("")

# ============================================================================
# 8. RESUME
# ============================================================================
add("#" * 120)
add("8) RESUME QA")
add("#" * 120)
add(f"Titres Word détectés : {len(title_map)}")
add(f"Doublons de titres : {len(dupes)}")
add(f"Sources Annuaire 2023 restantes FR+AR : {count_annuaire_fr_2023 + count_annuaire_ar_2023}")
add(f"Cellules avec .../… : {len(ellipsis_hits)}")
add(f"Anciennes années suspectes hors exceptions : {len(suspicious_old_years)}")
add(f"Sources FR/AR incohérentes : {len(source_issues)}")
add(f"Lignes avec cellules vides suspectes : {len(blank_hits)}")
add("")
add("IMPORTANT : ce diagnostic ne modifie pas le Word.")
add("Les alertes sont des candidats à vérification, pas des corrections automatiques.")

FICHIER_SORTIE.write_text("\n".join(report), encoding="utf-8")

print("=" * 100)
print("DIAGNOSTIC GLOBAL FINAL TERMINE")
print("=" * 100)
print(f"Word analysé : {docx}")
print(f"Titres détectés : {len(title_map)}")
print(f"Doublons : {len(dupes)}")
print(f"Annuaire 2023 restant FR+AR : {count_annuaire_fr_2023 + count_annuaire_ar_2023}")
print(f"Cellules avec ellipses : {len(ellipsis_hits)}")
print(f"Anciennes années suspectes : {len(suspicious_old_years)}")
print(f"Sources FR/AR incohérentes : {len(source_issues)}")
print(f"Cellules vides suspectes : {len(blank_hits)}")
print()
print(f"Rapport créé : {FICHIER_SORTIE}")
print("=" * 100)
