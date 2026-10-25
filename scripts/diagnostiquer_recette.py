from pathlib import Path
from openpyxl import load_workbook
from docx import Document
import re


# ============================================================
# PARAMETRES
# ============================================================

CODE_WORD = "14-2"

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_MODELES = Path("modeles")
DOSSIER_SORTIE = Path("sorties")

FICHIER_MAPPING = DOSSIER_SORTIE / "mapping_final_recettes.xlsx"
FICHIER_CATALOGUE = DOSSIER_SORTIE / "catalogue_tables_reelles.xlsx"
FICHIER_SORTIE = DOSSIER_SORTIE / f"diagnostic_{CODE_WORD}.txt"


# ============================================================
# OUTILS
# ============================================================

def texte_propre(valeur):
    if valeur is None:
        return ""
    return str(valeur).replace("\n", " ").strip()


def cellules_uniques_word(tableau):
    """Evite de répéter plusieurs fois la même cellule fusionnée."""
    lignes = []

    for no_ligne, ligne in enumerate(tableau.rows, start=1):
        cellules = []
        deja_vues = set()

        for no_colonne, cellule in enumerate(ligne.cells, start=1):
            identifiant = id(cellule._tc)

            if identifiant in deja_vues:
                continue

            deja_vues.add(identifiant)

            cellules.append(
                (
                    no_colonne,
                    texte_propre(cellule.text)
                )
            )

        lignes.append(
            (
                no_ligne,
                cellules
            )
        )

    return lignes


def charger_catalogue():
    wb = load_workbook(
        FICHIER_CATALOGUE,
        data_only=True
    )

    ws = wb["Tables_logiques"]

    entetes = {
        cellule.value: cellule.column
        for cellule in ws[1]
    }

    catalogue = {}

    for ligne in range(
        2,
        ws.max_row + 1
    ):
        numero = str(
            ws.cell(
                ligne,
                entetes["Numéro complet"]
            ).value or ""
        ).strip()

        catalogue[numero] = {
            "fichier": ws.cell(
                ligne,
                entetes["Fichier"]
            ).value,
            "titre": ws.cell(
                ligne,
                entetes["Titre principal"]
            ).value,
            "plages": ws.cell(
                ligne,
                entetes["Plage(s)"]
            ).value or "",
            "feuilles": ws.cell(
                ligne,
                entetes["Feuille(s)"]
            ).value or "",
        }

    wb.close()

    return catalogue


def charger_recette():
    wb = load_workbook(
        FICHIER_MAPPING,
        data_only=True
    )

    ws = wb["Mapping_final"]

    entetes = {
        cellule.value: cellule.column
        for cellule in ws[1]
    }

    recette = None

    for ligne in range(
        2,
        ws.max_row + 1
    ):
        numero_word = str(
            ws.cell(
                ligne,
                entetes["N° Word"]
            ).value or ""
        ).strip()

        if numero_word == CODE_WORD:
            recette = {
                "table_word_titre": ws.cell(
                    ligne,
                    entetes["Table Word titre"]
                ).value,
                "table_word_donnees": ws.cell(
                    ligne,
                    entetes["Table Word données"]
                ).value,
                "theme": ws.cell(
                    ligne,
                    entetes["Thème Word"]
                ).value,
                "titre_word": ws.cell(
                    ligne,
                    entetes["Titre Word"]
                ).value,
                "type": ws.cell(
                    ligne,
                    entetes["Type de recette"]
                ).value,
                "sources": ws.cell(
                    ligne,
                    entetes["Sources Excel"]
                ).value,
            }

            break

    wb.close()

    if recette is None:
        raise RuntimeError(
            f"Aucune recette trouvée pour {CODE_WORD}"
        )

    return recette


def extraire_bloc_excel(info_source):
    """Retourne les cellules non vides des plages d'une table logique."""
    fichier = DOSSIER_SOURCE / info_source["fichier"]

    wb = load_workbook(
        fichier,
        data_only=True,
        read_only=False
    )

    resultat = []

    plages = [
        morceau.strip()
        for morceau in str(
            info_source["plages"]
        ).split("|")
        if morceau.strip()
    ]

    for plage in plages:
        if "!" not in plage:
            continue

        nom_feuille, bornes = plage.rsplit(
            "!",
            1
        )

        if ":" not in bornes:
            continue

        debut_txt, fin_txt = bornes.split(
            ":",
            1
        )

        try:
            debut = int(
                debut_txt
            )

            fin = int(
                fin_txt
            )

        except ValueError:
            continue

        ws = wb[nom_feuille]

        # On limite à 80 colonnes pour éviter les feuilles Excel
        # artificiellement formatées jusqu'à XFD.
        max_colonne = min(
            ws.max_column,
            80
        )

        cellules = []

        for ligne in ws.iter_rows(
            min_row=debut,
            max_row=fin,
            min_col=1,
            max_col=max_colonne
        ):
            contenu_ligne = []

            for cellule in ligne:
                if cellule.value in (
                    None,
                    ""
                ):
                    continue

                contenu_ligne.append(
                    (
                        cellule.coordinate,
                        texte_propre(
                            cellule.value
                        )
                    )
                )

            if contenu_ligne:
                cellules.append(
                    contenu_ligne
                )

        resultat.append(
            {
                "feuille": nom_feuille,
                "debut": debut,
                "fin": fin,
                "lignes": cellules
            }
        )

    wb.close()

    return resultat


# ============================================================
# EXECUTION
# ============================================================

catalogue = charger_catalogue()
recette = charger_recette()

fichiers_word = [
    fichier
    for fichier in DOSSIER_MODELES.glob(
        "*.docx"
    )
    if not fichier.name.startswith(
        "~$"
    )
]

if not fichiers_word:
    raise FileNotFoundError(
        "Aucun .docx trouvé dans modeles."
    )

fichier_word = fichiers_word[0]

document = Document(
    fichier_word
)

numero_table_word = int(
    recette["table_word_donnees"]
)

table_word = document.tables[
    numero_table_word - 1
]


# ============================================================
# RAPPORT TEXTE
# ============================================================

lignes_sortie = []

lignes_sortie.append(
    "=" * 90
)

lignes_sortie.append(
    f"DIAGNOSTIC CELLULAIRE — TABLEAU WORD {CODE_WORD}"
)

lignes_sortie.append(
    "=" * 90
)

lignes_sortie.append(
    ""
)

lignes_sortie.append(
    f"Modèle Word : {fichier_word.name}"
)

lignes_sortie.append(
    f"Thème : {recette['theme']}"
)

lignes_sortie.append(
    f"Titre Word : {texte_propre(recette['titre_word'])}"
)

lignes_sortie.append(
    f"Type de recette : {recette['type']}"
)

lignes_sortie.append(
    f"Sources : {recette['sources']}"
)

lignes_sortie.append(
    f"N° objet tableau Word : {numero_table_word}"
)

lignes_sortie.append(
    f"Dimensions Word : {len(table_word.rows)} lignes x {len(table_word.columns)} colonnes"
)

lignes_sortie.append(
    ""
)

lignes_sortie.append(
    "-" * 90
)

lignes_sortie.append(
    "1) STRUCTURE DU TABLEAU WORD"
)

lignes_sortie.append(
    "-" * 90
)

for numero_ligne, cellules in cellules_uniques_word(
    table_word
):
    contenu = []

    for numero_colonne, texte in cellules:
        contenu.append(
            f"C{numero_colonne}={texte!r}"
        )

    lignes_sortie.append(
        f"R{numero_ligne}: "
        + " | ".join(
            contenu
        )
    )


lignes_sortie.append(
    ""
)

lignes_sortie.append(
    "-" * 90
)

lignes_sortie.append(
    "2) STRUCTURE EXCEL SOURCE"
)

lignes_sortie.append(
    "-" * 90
)


sources = [
    morceau.strip()
    for morceau in str(
        recette["sources"]
    ).split("+")
    if morceau.strip()
]


for source in sources:
    if source not in catalogue:
        lignes_sortie.append(
            f"ERREUR : source {source} absente du catalogue."
        )

        continue

    info = catalogue[
        source
    ]

    lignes_sortie.append(
        ""
    )

    lignes_sortie.append(
        f"SOURCE {source}"
    )

    lignes_sortie.append(
        f"Fichier : {info['fichier']}"
    )

    lignes_sortie.append(
        f"Titre : {texte_propre(info['titre'])}"
    )

    lignes_sortie.append(
        f"Feuille(s) : {info['feuilles']}"
    )

    lignes_sortie.append(
        f"Plage(s) : {info['plages']}"
    )

    blocs = extraire_bloc_excel(
        info
    )

    for bloc in blocs:
        lignes_sortie.append(
            ""
        )

        lignes_sortie.append(
            f"--- {bloc['feuille']} / lignes {bloc['debut']}:{bloc['fin']} ---"
        )

        for ligne in bloc[
            "lignes"
        ]:
            contenu = []

            for coordonnee, valeur in ligne:
                contenu.append(
                    f"{coordonnee}={valeur!r}"
                )

            lignes_sortie.append(
                " | ".join(
                    contenu
                )
            )


FICHIER_SORTIE.write_text(
    "\n".join(
        lignes_sortie
    ),
    encoding="utf-8"
)


print(
    "=" * 72
)

print(
    f"DIAGNOSTIC TERMINE POUR {CODE_WORD}"
)

print(
    "=" * 72
)

print(
    f"Tableau Word : {numero_table_word}"
)

print(
    f"Type : {recette['type']}"
)

print(
    f"Sources : {recette['sources']}"
)

print()

print(
    f"Fichier créé : {FICHIER_SORTIE}"
)

print(
    "=" * 72
)
