from pathlib import Path
from openpyxl import load_workbook
from zipfile import ZipFile
from lxml import etree


# ============================================================
# TABLEAUX SANTE
# ============================================================

CODES_WORD = [
    "2-1", "2-2", "2-3", "2-4", "2-5", "2-6",
    "2-7", "2-8", "2-9", "2-10", "2-11", "2-12",
    "2-13", "2-14", "2-15", "2-16", "2-17", "2-18",
    "2-20", "2-21", "2-22", "2-23", "2-24"
]


# ============================================================
# CHEMINS
# ============================================================

DOSSIER_SORTIE = Path("sorties")
DOSSIER_MODELES = Path("modeles")

FICHIER_MAPPING = DOSSIER_SORTIE / "mapping_final_recettes.xlsx"
FICHIER_SORTIE = DOSSIER_SORTIE / "diagnostic_sante_word_corrige.txt"

CANDIDATS_WORD = [
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_VALIDE.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_FINAL.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_OK.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture.docx",
]


# ============================================================
# OUTILS
# ============================================================

def txt(v):
    return "" if v is None else str(v).strip()


def choisir_word():
    for fichier in CANDIDATS_WORD:
        if fichier.exists():
            return fichier

    fichiers = [
        f for f in DOSSIER_MODELES.glob("*.docx")
        if not f.name.startswith("~$")
    ]

    if not fichiers:
        raise FileNotFoundError(
            "Aucun document Word trouvé."
        )

    return fichiers[0]


def charger_mapping():
    wb = load_workbook(
        FICHIER_MAPPING,
        data_only=True,
        read_only=True
    )

    ws = wb["Mapping_final"]

    h = {
        c.value: c.column
        for c in ws[1]
    }

    resultat = {}

    for r in range(2, ws.max_row + 1):
        code = txt(
            ws.cell(
                r,
                h["N° Word"]
            ).value
        )

        if code not in CODES_WORD:
            continue

        resultat[code] = {
            "table_word": int(
                ws.cell(
                    r,
                    h["Table Word données"]
                ).value
            ),
            "titre": txt(
                ws.cell(
                    r,
                    h["Titre Word"]
                ).value
            ),
            "type": txt(
                ws.cell(
                    r,
                    h["Type de recette"]
                ).value
            ),
            "sources": txt(
                ws.cell(
                    r,
                    h["Sources Excel"]
                ).value
            ),
        }

    wb.close()
    return resultat


# ============================================================
# LECTURE WORD CORRECTE
# ============================================================
#
# IMPORTANT :
# python-docx document.tables compte uniquement les tableaux
# de premier niveau du document.
#
# L'ancien diagnostic utilisait ".//w:tbl", ce qui comptait AUSSI
# les tableaux imbriqués dans des cellules et décalait les numéros.
#
# Ici on utilise uniquement :
#   /w:document/w:body/w:tbl
# ============================================================

def lire_tables_top_level(docx_path):
    ns = {
        "w": (
            "http://schemas.openxmlformats.org/"
            "wordprocessingml/2006/main"
        )
    }

    with ZipFile(docx_path, "r") as z:
        xml = z.read(
            "word/document.xml"
        )

    root = etree.fromstring(
        xml
    )

    body = root.find(
        "w:body",
        namespaces=ns
    )

    tables = body.findall(
        "w:tbl",
        namespaces=ns
    )

    return tables, ns


def decrire_table(table, ns):
    lignes = []

    rows = table.findall(
        "w:tr",
        namespaces=ns
    )

    max_cols = 0

    for no_ligne, row in enumerate(
        rows,
        start=1
    ):
        cells = row.findall(
            "w:tc",
            namespaces=ns
        )

        max_cols = max(
            max_cols,
            len(cells)
        )

        parties = []

        for no_col, cell in enumerate(
            cells,
            start=1
        ):
            textes = cell.xpath(
                ".//w:t/text()",
                namespaces=ns
            )

            contenu = " ".join(
                t.strip()
                for t in textes
                if t.strip()
            )

            parties.append(
                f"C{no_col}={contenu!r}"
            )

        lignes.append(
            f"R{no_ligne}: "
            + " | ".join(parties)
        )

    return len(rows), max_cols, lignes


# ============================================================
# EXECUTION
# ============================================================

print("=" * 84)
print("DIAGNOSTIC SANTE - WORD CORRIGE")
print("=" * 84)

word_path = choisir_word()

print(
    f"Word : {word_path}",
    flush=True
)

mapping = charger_mapping()

print(
    "Lecture des tableaux Word de premier niveau...",
    flush=True
)

tables, ns = lire_tables_top_level(
    word_path
)

print(
    f"Nombre total de tableaux Word top-level : {len(tables)}",
    flush=True
)

rapport = [
    "=" * 110,
    "DIAGNOSTIC SANTE - STRUCTURES WORD CORRIGEES",
    "=" * 110,
    f"Document : {word_path}",
    f"Nombre de tableaux top-level : {len(tables)}",
    "",
    (
        "NOTE : ce diagnostic utilise exactement la même logique de numérotation "
        "que python-docx document.tables."
    ),
]


for pos, code in enumerate(
    CODES_WORD,
    start=1
):
    print(
        f"[{pos}/{len(CODES_WORD)}] {code}",
        flush=True
    )

    rapport.extend([
        "",
        "#" * 110,
        f"TABLEAU WORD {code}",
        "#" * 110,
    ])

    if code not in mapping:
        rapport.append(
            "ERREUR : code absent du mapping."
        )
        continue

    m = mapping[code]

    no_objet = m["table_word"]
    index = no_objet - 1

    rapport.extend([
        f"Titre attendu : {m['titre']}",
        f"Type : {m['type']}",
        f"Sources Excel : {m['sources']}",
        f"N° objet Word : {no_objet}",
    ])

    if index < 0 or index >= len(tables):
        rapport.append(
            "ERREUR : objet Word hors limites."
        )
        continue

    nb_lignes, nb_cols, lignes = decrire_table(
        tables[index],
        ns
    )

    rapport.append(
        f"Dimensions réelles : {nb_lignes} x {nb_cols}"
    )

    rapport.append("")
    rapport.append("STRUCTURE WORD REELLE")
    rapport.append("-" * 110)

    rapport.extend(
        lignes
    )


FICHIER_SORTIE.write_text(
    "\n".join(rapport),
    encoding="utf-8"
)

print()
print("=" * 84)
print("TERMINE")
print("=" * 84)
print(
    f"Fichier créé : {FICHIER_SORTIE}"
)
print("=" * 84)
