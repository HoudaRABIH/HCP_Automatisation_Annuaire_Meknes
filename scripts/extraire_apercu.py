from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
import warnings

warnings.filterwarnings(
    "ignore",
    message="Unknown extension is not supported*"
)

warnings.filterwarnings(
    "ignore",
    message="Print area cannot be set*"
)

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_SORTIE = Path("sorties")

FICHIER_SORTIE = DOSSIER_SORTIE / "apercu_feuilles.xlsx"

DOSSIER_SORTIE.mkdir(exist_ok=True)

fichiers_excel = sorted(DOSSIER_SOURCE.glob("*.xlsx"))


def classifier_feuille(nom):
    nom_normalise = nom.strip().lower()

    if nom_normalise.startswith("pg"):
        return "PAGE DE GARDE"

    if "sommaire" in nom_normalise:
        return "SOMMAIRE"

    if (
        "graph" in nom_normalise
        or "grraph" in nom_normalise
        or "graphique" in nom_normalise
    ):
        return "GRAPHIQUE"

    return "TABLEAU"


rapport = Workbook()

ws_feuilles = rapport.active
ws_feuilles.title = "Feuilles"

ws_feuilles.append([
    "Fichier",
    "Feuille",
    "Type",
    "Nombre de lignes",
    "Nombre de colonnes",
    "Cellules fusionnées"
])


ws_apercu = rapport.create_sheet("Apercu")

ws_apercu.append([
    "Fichier",
    "Feuille",
    "Ligne Excel",
    "Contenu"
])


print("=" * 70)
print("EXTRACTION DES APERCUS DES TABLEAUX HCP")
print("=" * 70)

nombre_tableaux = 0


for numero, fichier in enumerate(fichiers_excel, start=1):

    print(
        f"[{numero}/{len(fichiers_excel)}] "
        f"{fichier.name}"
    )

    try:

        classeur = load_workbook(
            fichier,
            data_only=True,
            read_only=False
        )

        for feuille in classeur.worksheets:

            type_feuille = classifier_feuille(feuille.title)

            ws_feuilles.append([
                fichier.name,
                feuille.title,
                type_feuille,
                feuille.max_row,
                feuille.max_column,
                len(feuille.merged_cells.ranges)
            ])

            # On ne détaille que les vraies feuilles statistiques
            if type_feuille != "TABLEAU":
                continue

            nombre_tableaux += 1

            # Examiner les 40 premières lignes
            # et les 20 premières colonnes
            max_ligne = min(feuille.max_row, 40)
            max_colonne = min(feuille.max_column, 20)

            for numero_ligne in range(1, max_ligne + 1):

                contenu_ligne = []

                for numero_colonne in range(1, max_colonne + 1):

                    valeur = feuille.cell(
                        row=numero_ligne,
                        column=numero_colonne
                    ).value

                    if valeur not in (None, ""):

                        lettre = get_column_letter(numero_colonne)

                        texte = str(valeur).strip()

                        contenu_ligne.append(
                            f"{lettre}: {texte}"
                        )

                # Ajouter uniquement les lignes contenant quelque chose
                if contenu_ligne:

                    ws_apercu.append([
                        fichier.name,
                        feuille.title,
                        numero_ligne,
                        " | ".join(contenu_ligne)
                    ])

        classeur.close()

    except Exception as erreur:

        print(f"   ERREUR : {erreur}")


# Ajuster la largeur des colonnes
ws_feuilles.column_dimensions["A"].width = 45
ws_feuilles.column_dimensions["B"].width = 30
ws_feuilles.column_dimensions["C"].width = 20
ws_feuilles.column_dimensions["D"].width = 20
ws_feuilles.column_dimensions["E"].width = 20
ws_feuilles.column_dimensions["F"].width = 20

ws_apercu.column_dimensions["A"].width = 45
ws_apercu.column_dimensions["B"].width = 30
ws_apercu.column_dimensions["C"].width = 15
ws_apercu.column_dimensions["D"].width = 120


rapport.save(FICHIER_SORTIE)


print()
print(f"Nombre de tableaux détectés : {nombre_tableaux}")
print(f"Fichier créé : {FICHIER_SORTIE}")
print("=" * 70)
print("EXTRACTION TERMINEE")
print("=" * 70)