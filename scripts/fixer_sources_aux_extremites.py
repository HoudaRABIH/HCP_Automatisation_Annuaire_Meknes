from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
import re
import sys

SORTIES = Path("sorties")

DEFAULT_INPUTS = [
    SORTIES / "Annuaire_HCP_2025_FINAL.docx",
    SORTIES / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER_POPULATION_INDUSTRIE_ENERGIE_EAU.docx",
]

DEFAULT_OUTPUT = SORTIES / "Annuaire_HCP_2025_FINAL_SOURCES_CORNERS.docx"


def is_fr_source(text):
    t = (text or "").strip().lower()
    return (
        t.startswith("source")
        or t.startswith("sources")
    )


def is_ar_source(text):
    t = (text or "").strip()
    return (
        t.startswith("المصدر")
        or t.startswith("المصادر")
    )


def clone_rpr_from_paragraph(paragraph):
    """
    Recupere le style du premier run non vide pour conserver
    police, taille, italique et surtout la COULEUR du chapitre.
    """
    for run in paragraph.runs:
        if run.text and run.text.strip():
            rpr = run._r.find(qn("w:rPr"))
            return deepcopy(rpr) if rpr is not None else None
    return None


def set_cell_margin_zero(cell):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.find(qn("w:tcMar"))
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)

    for side in ("left", "right"):
        node = tcMar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tcMar.append(node)
        node.set(qn("w:w"), "0")
        node.set(qn("w:type"), "dxa")


def set_paragraph_source_layout(paragraph, tab_pos_twips):
    pPr = paragraph._p.get_or_add_pPr()

    # LTR paragraph so the French source starts at the true left.
    bidi = pPr.find(qn("w:bidi"))
    if bidi is not None:
        pPr.remove(bidi)

    jc = pPr.find(qn("w:jc"))
    if jc is None:
        jc = OxmlElement("w:jc")
        pPr.append(jc)
    jc.set(qn("w:val"), "left")

    # No paragraph indentation: use the full table width.
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    for a in list(ind.attrib):
        del ind.attrib[a]
    ind.set(qn("w:left"), "0")
    ind.set(qn("w:right"), "0")
    ind.set(qn("w:firstLine"), "0")

    # No extra vertical spacing.
    spacing = pPr.find(qn("w:spacing"))
    if spacing is None:
        spacing = OxmlElement("w:spacing")
        pPr.append(spacing)
    spacing.set(qn("w:before"), "0")
    spacing.set(qn("w:after"), "0")

    # One right-aligned tab at the far right edge.
    old_tabs = pPr.find(qn("w:tabs"))
    if old_tabs is not None:
        pPr.remove(old_tabs)

    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "right")
    tab.set(qn("w:pos"), str(max(1000, int(tab_pos_twips))))
    tabs.append(tab)
    pPr.append(tabs)


def add_run_with_rpr(paragraph, text, rpr=None, rtl=False):
    r = OxmlElement("w:r")

    if rpr is not None:
        rp = deepcopy(rpr)
    else:
        rp = OxmlElement("w:rPr")

    if rtl:
        rtl_node = rp.find(qn("w:rtl"))
        if rtl_node is None:
            rtl_node = OxmlElement("w:rtl")
            rp.append(rtl_node)
        rtl_node.set(qn("w:val"), "1")

        lang = rp.find(qn("w:lang"))
        if lang is None:
            lang = OxmlElement("w:lang")
            rp.append(lang)
        lang.set(qn("w:bidi"), "ar-MA")

    r.append(rp)

    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    paragraph._p.append(r)


def add_tab(paragraph):
    r = OxmlElement("w:r")
    tab = OxmlElement("w:tab")
    r.append(tab)
    paragraph._p.append(r)


def table_width_twips(table):
    """
    Uses the Word table grid. This is much more reliable than guessing
    from page size and works with the different chapter table widths.
    """
    grid = table._tbl.tblGrid
    if grid is not None:
        widths = []
        for col in grid.gridCol_lst:
            raw = col.get(qn("w:w"))
            if raw:
                try:
                    widths.append(int(raw))
                except ValueError:
                    pass
        if widths:
            return sum(widths)

    # Conservative fallback.
    return 9000


def clear_cell_content_keep_properties(cell):
    tc = cell._tc
    tcPr = tc.find(qn("w:tcPr"))

    for child in list(tc):
        if child is not tcPr:
            tc.remove(child)

    # A table cell must contain at least one paragraph.
    p = OxmlElement("w:p")
    tc.append(p)


def source_paragraphs_in_row(row):
    fr = []
    ar = []

    seen = set()
    for cell in row.cells:
        tc_id = id(cell._tc)
        if tc_id in seen:
            continue
        seen.add(tc_id)

        for p in cell.paragraphs:
            text = (p.text or "").strip()
            if is_fr_source(text):
                fr.append((cell, p, text))
            elif is_ar_source(text):
                ar.append((cell, p, text))

    return fr, ar


def fix_source_row(table, row):
    fr, ar = source_paragraphs_in_row(row)

    # Only transform genuine bilingual source rows.
    if not fr or not ar:
        return False

    fr_cell, fr_p, fr_text = fr[0]
    ar_cell, ar_p, ar_text = ar[0]

    fr_rpr = clone_rpr_from_paragraph(fr_p)
    ar_rpr = clone_rpr_from_paragraph(ar_p)

    # Merge the whole source row so FR and AR share EXACTLY the same baseline
    # and can reach the two extreme edges of the table.
    first = row.cells[0]
    last = row.cells[-1]
    merged = first if first._tc is last._tc else first.merge(last)

    clear_cell_content_keep_properties(merged)
    set_cell_margin_zero(merged)

    p = merged.paragraphs[0]
    set_paragraph_source_layout(
        p,
        table_width_twips(table) - 40
    )

    add_run_with_rpr(p, fr_text, fr_rpr, rtl=False)
    add_tab(p)
    add_run_with_rpr(p, ar_text, ar_rpr, rtl=True)

    return True


def fix_paragraph_pair(fr_p, ar_p, width_twips=9000):
    """
    Fallback for source paragraphs outside tables.
    Makes them one line: French left / Arabic right.
    """
    fr_text = (fr_p.text or "").strip()
    ar_text = (ar_p.text or "").strip()

    fr_rpr = clone_rpr_from_paragraph(fr_p)
    ar_rpr = clone_rpr_from_paragraph(ar_p)

    # Clear French paragraph but keep paragraph properties.
    for child in list(fr_p._p):
        if child.tag != qn("w:pPr"):
            fr_p._p.remove(child)

    set_paragraph_source_layout(fr_p, width_twips - 40)
    add_run_with_rpr(fr_p, fr_text, fr_rpr, rtl=False)
    add_tab(fr_p)
    add_run_with_rpr(fr_p, ar_text, ar_rpr, rtl=True)

    # Remove Arabic paragraph from document body.
    parent = ar_p._p.getparent()
    if parent is not None:
        parent.remove(ar_p._p)


def choose_input():
    if len(sys.argv) >= 2:
        return Path(sys.argv[1])

    for p in DEFAULT_INPUTS:
        if p.exists():
            return p

    candidates = sorted(
        [p for p in SORTIES.glob("*.docx") if not p.name.startswith("~$")],
        key=lambda x: x.stat().st_mtime,
        reverse=True
    )
    if not candidates:
        raise FileNotFoundError("Aucun Word trouvé dans sorties.")
    return candidates[0]


def choose_output(inp):
    if len(sys.argv) >= 3:
        return Path(sys.argv[2])

    if inp.name == "Annuaire_HCP_2025_FINAL.docx":
        return DEFAULT_OUTPUT

    return inp.with_name(inp.stem + "_SOURCES_CORNERS.docx")


input_path = choose_input()
output_path = choose_output(input_path)

if not input_path.exists():
    raise FileNotFoundError(input_path)

print("=" * 96)
print("SOURCES — FR EXTREME GAUCHE / AR EXTREME DROITE")
print("=" * 96)
print(f"Entrée : {input_path}")
print(f"Sortie : {output_path}")
print()

doc = Document(input_path)

fixed_rows = 0

# Main tables + nested tables.
def process_table(table):
    global fixed_rows

    for row in table.rows:
        if fix_source_row(table, row):
            fixed_rows += 1

        # Nested tables, if any.
        seen = set()
        for cell in row.cells:
            tc_id = id(cell._tc)
            if tc_id in seen:
                continue
            seen.add(tc_id)

            for nested in cell.tables:
                process_table(nested)


for table in doc.tables:
    process_table(table)

# Paragraphs outside tables: only pair immediately adjacent FR then AR
# (or AR then FR) source paragraphs.
paras = list(doc.paragraphs)
i = 0
fixed_body_pairs = 0

while i < len(paras) - 1:
    p1 = paras[i]
    p2 = paras[i + 1]
    t1 = (p1.text or "").strip()
    t2 = (p2.text or "").strip()

    if is_fr_source(t1) and is_ar_source(t2):
        fix_paragraph_pair(p1, p2)
        fixed_body_pairs += 1
        i += 2
        continue

    if is_ar_source(t1) and is_fr_source(t2):
        fix_paragraph_pair(p2, p1)
        fixed_body_pairs += 1
        i += 2
        continue

    i += 1

doc.save(output_path)

print(f"Lignes Source de tableaux corrigées : {fixed_rows}")
print(f"Paires Source hors tableaux corrigées : {fixed_body_pairs}")
print()
print("IMPORTANT :")
print(" - texte des sources conservé")
print(" - couleurs/polices des chapitres conservées")
print(" - données numériques non modifiées")
print(" - FR et AR placés sur la MEME ligne")
print(" - FR au bord gauche, AR au bord droit")
print()
print(f"Document créé : {output_path}")
print("=" * 96)
