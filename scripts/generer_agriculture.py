from pathlib import Path
from openpyxl import load_workbook
from docx import Document
import re
import unicodedata


# ============================================================
# CHEMINS
# ============================================================

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_SORTIE = Path("sorties")

FICHIER_AGRICULTURE = DOSSIER_SOURCE / "3. Agriculture_AS 20252.xlsx"

FICHIER_ENTREE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme.docx"
)

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture.docx"
)


# ============================================================
# OUTILS
# ============================================================

def normaliser(texte):
    texte = str(texte or "").strip().lower()
    texte = texte.replace("’", "'").replace("–", "-").replace("—", "-")

    texte = "".join(
        c
        for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )

    texte = re.sub(r"\s+", " ", texte)
    texte = re.sub(r"\s*-\s*", " - ", texte)

    return texte.strip()


def remplacer_texte_preserver_style(cellule, nouveau_texte):
    nouveau_texte = str(nouveau_texte)

    if not cellule.paragraphs:
        cellule.text = nouveau_texte
        return

    paragraphe = cellule.paragraphs[0]

    if paragraphe.runs:
        paragraphe.runs[0].text = nouveau_texte

        for run in paragraphe.runs[1:]:
            run.text = ""
    else:
        paragraphe.add_run(nouveau_texte)

    for autre in cellule.paragraphs[1:]:
        for run in autre.runs:
            run.text = ""


def formater_1_decimale(valeur, zero_tiret=False):
    if valeur in (None, "", "-", "–"):
        return "-"

    try:
        nombre = float(valeur)
    except (TypeError, ValueError):
        return str(valeur).strip()

    if zero_tiret and abs(nombre) < 1e-12:
        return "-"

    return (
        f"{nombre:,.1f}"
        .replace(",", " ")
        .replace(".", ",")
    )


def formater_entier(valeur, zero_tiret=False):
    if valeur in (None, "", "-", "–"):
        return "-"

    try:
        nombre = float(valeur)
    except (TypeError, ValueError):
        return str(valeur).strip()

    if zero_tiret and abs(nombre) < 1e-12:
        return "-"

    return f"{int(round(nombre)):,}".replace(",", " ")


def nettoyer_periode(valeur):
    """
    Exemples :
    2024-2023*   -> 2024-2023*
    2023-2022Rم -> 2023-2022R
    2022-2021   -> 2022-2021
    """
    texte = str(valeur or "")
    m = re.search(r"(\d{4}-\d{4})", texte)

    if not m:
        return texte.strip()

    resultat = m.group(1)

    if "*" in texte:
        resultat += "*"

    if re.search(r"\bR", texte[m.end():]) or "R" in texte[m.end():]:
        resultat += "R"

    return resultat


def trouver_ligne(ws, variantes):
    if isinstance(variantes, str):
        variantes = [variantes]

    cibles = {
        normaliser(v)
        for v in variantes
    }

    for r in range(1, ws.max_row + 1):
        valeur = normaliser(
            ws.cell(r, 1).value
        )

        if valeur in cibles:
            return r

    raise ValueError(
        f"Ligne introuvable dans {ws.title}: {variantes}"
    )


def trouver_dernier_total(ws):
    lignes = []

    for r in range(1, ws.max_row + 1):
        if normaliser(
            ws.cell(r, 1).value
        ) == "total":
            lignes.append(r)

    if not lignes:
        raise ValueError(
            f"Total introuvable dans {ws.title}"
        )

    return lignes[-1]


def v(ws, ligne, colonne):
    return ws.cell(
        ligne,
        colonne
    ).value


def verifier_table(document, numero_objet, lignes, colonnes):
    table = document.tables[
        numero_objet - 1
    ]

    if (
        len(table.rows) != lignes
        or len(table.columns) != colonnes
    ):
        raise RuntimeError(
            f"Objet Word {numero_objet}: "
            f"{len(table.rows)}x{len(table.columns)} "
            f"au lieu de {lignes}x{colonnes}."
        )

    return table


def changer_source_annuaire_2023_vers_2025(document, objets):
    """
    Règle validée :
    on ne touche qu'à l'année de la mention
    'Annuaire Statistique du Maroc 2023'
    (et sa version arabe).
    Les autres sources restent intactes.
    """
    marqueurs = (
        "Annuaire Statistique du Maroc",
        "النشرة الإحصائية السنوية للمغرب",
    )

    cellules_vues = set()

    for numero_objet in objets:
        index = numero_objet - 1

        for idx in range(
            max(0, index),
            min(len(document.tables), index + 3)
        ):
            table = document.tables[idx]

            for ligne in table.rows:
                for cellule in ligne.cells:
                    identifiant = id(
                        cellule._tc
                    )

                    if identifiant in cellules_vues:
                        continue

                    cellules_vues.add(
                        identifiant
                    )

                    for paragraphe in cellule.paragraphs:
                        texte = paragraphe.text or ""

                        if (
                            "2023" not in texte
                            or not any(
                                marqueur in texte
                                for marqueur in marqueurs
                            )
                        ):
                            continue

                        for run in paragraphe.runs:
                            if "2023" in run.text:
                                run.text = run.text.replace(
                                    "2023",
                                    "2025"
                                )


# ============================================================
# OUVERTURE
# ============================================================

if not FICHIER_ENTREE.exists():
    raise FileNotFoundError(
        f"Word d'entrée introuvable : {FICHIER_ENTREE}"
    )

if not FICHIER_AGRICULTURE.exists():
    raise FileNotFoundError(
        f"Excel Agriculture introuvable : {FICHIER_AGRICULTURE}"
    )


print("Ouverture du Word...", flush=True)

document = Document(
    FICHIER_ENTREE
)

print("Ouverture de l'Excel Agriculture...", flush=True)

wb = load_workbook(
    FICHIER_AGRICULTURE,
    data_only=True,
    read_only=False
)


# ============================================================
# 9-1 — CEREALES : EVOLUTION SUPERFICIE / PRODUCTION
# Excel AG,3
# Word objet 272
# ============================================================

print("[1/11] 9-1", flush=True)

ws = wb["AG,3"]
table = verifier_table(
    document,
    272,
    8,
    8
)

# Périodes
for col_word, cell_excel in (
    (1, "C6"),
    (3, "E6"),
    (5, "G6"),
):
    remplacer_texte_preserver_style(
        table.cell(2, col_word),
        nettoyer_periode(
            ws[cell_excel].value
        )
    )

geo = {
    5: trouver_dernier_total(ws),
    6: trouver_ligne(ws, "Fès - Meknès"),
    7: trouver_ligne(ws, ["Meknès", "Meknes"]),
}

# B:G = Prod/Sup x 3 campagnes
for ligne_word, ligne_excel in geo.items():
    for col_word, col_excel in zip(
        range(1, 7),
        range(2, 8)
    ):
        remplacer_texte_preserver_style(
            table.cell(ligne_word, col_word),
            formater_1_decimale(
                v(ws, ligne_excel, col_excel),
                zero_tiret=True
            )
        )


# ============================================================
# 9-2 — PRINCIPALES CEREALES
# Superficie AG,4 + Production AG,5
# Word objet 274
# ============================================================

print("[2/11] 9-2", flush=True)

ws_sup = wb["AG,4"]
ws_prod = wb["AG,5"]

table = verifier_table(
    document,
    274,
    9,
    6
)

# Campagne
remplacer_texte_preserver_style(
    table.cell(0, 0),
    str(ws_sup["A6"].value or "")
)

remplacer_texte_preserver_style(
    table.cell(0, 5),
    str(ws_sup["F6"].value or "")
)

geo_sup = {
    2: trouver_dernier_total(ws_sup),
    3: trouver_ligne(ws_sup, "Fès - Meknès"),
    4: trouver_ligne(ws_sup, ["Meknès", "Meknes"]),
}

geo_prod = {
    6: trouver_dernier_total(ws_prod),
    7: trouver_ligne(ws_prod, "Fès - Meknès"),
    8: trouver_ligne(ws_prod, ["Meknès", "Meknes"]),
}

for ligne_word, ligne_excel in geo_sup.items():
    for col_word, col_excel in zip(
        range(1, 5),
        range(2, 6)
    ):
        remplacer_texte_preserver_style(
            table.cell(ligne_word, col_word),
            formater_1_decimale(
                v(ws_sup, ligne_excel, col_excel),
                zero_tiret=True
            )
        )

for ligne_word, ligne_excel in geo_prod.items():
    for col_word, col_excel in zip(
        range(1, 5),
        range(2, 6)
    ):
        remplacer_texte_preserver_style(
            table.cell(ligne_word, col_word),
            formater_1_decimale(
                v(ws_prod, ligne_excel, col_excel),
                zero_tiret=True
            )
        )


# ============================================================
# 9-3 — LEGUMINEUSES : EVOLUTION SUPERFICIE / PRODUCTION
# Excel AG,6
# Word objet 278
# ============================================================

print("[3/11] 9-3", flush=True)

ws = wb["AG,6  "]

table = verifier_table(
    document,
    278,
    7,
    8
)

for col_word, cell_excel in (
    (1, "B6"),
    (3, "D6"),
    (5, "F6"),
):
    remplacer_texte_preserver_style(
        table.cell(1, col_word),
        nettoyer_periode(
            ws[cell_excel].value
        )
    )

geo = {
    4: trouver_dernier_total(ws),
    5: trouver_ligne(ws, "Fès - Meknès"),
    6: trouver_ligne(ws, ["Meknès", "Meknes"]),
}

for ligne_word, ligne_excel in geo.items():
    for col_word, col_excel in zip(
        range(1, 7),
        range(2, 8)
    ):
        remplacer_texte_preserver_style(
            table.cell(ligne_word, col_word),
            formater_1_decimale(
                v(ws, ligne_excel, col_excel),
                zero_tiret=True
            )
        )


# ============================================================
# 9-4 — SUPERFICIE PRINCIPALES LEGUMINEUSES
# Excel AG,7
# Word objet 280
# ============================================================

print("[4/11] 9-4", flush=True)

ws = wb["AG,7"]

table = verifier_table(
    document,
    280,
    5,
    6
)

remplacer_texte_preserver_style(
    table.cell(0, 0),
    str(ws["A6"].value or "")
)

remplacer_texte_preserver_style(
    table.cell(0, 5),
    str(ws["F6"].value or "")
)

geo = {
    2: trouver_dernier_total(ws),
    3: trouver_ligne(ws, "Fès - Meknès"),
    4: trouver_ligne(ws, ["Meknès", "Meknes"]),
}

for ligne_word, ligne_excel in geo.items():
    for col_word, col_excel in zip(
        range(1, 5),
        range(2, 6)
    ):
        remplacer_texte_preserver_style(
            table.cell(ligne_word, col_word),
            formater_1_decimale(
                v(ws, ligne_excel, col_excel),
                zero_tiret=True
            )
        )


# ============================================================
# 9-5 — PRODUCTION PRINCIPALES LEGUMINEUSES
# Excel AG,8
# Word objet 282
# ============================================================

print("[5/11] 9-5", flush=True)

ws = wb["AG,8"]

table = verifier_table(
    document,
    282,
    5,
    6
)

remplacer_texte_preserver_style(
    table.cell(0, 0),
    str(ws["A6"].value or "")
)

remplacer_texte_preserver_style(
    table.cell(0, 5),
    str(ws["F6"].value or "")
)

geo = {
    2: trouver_dernier_total(ws),
    3: trouver_ligne(ws, "Fès - Meknès"),
    4: trouver_ligne(ws, ["Meknès", "Meknes"]),
}

for ligne_word, ligne_excel in geo.items():
    for col_word, col_excel in zip(
        range(1, 5),
        range(2, 6)
    ):
        remplacer_texte_preserver_style(
            table.cell(ligne_word, col_word),
            formater_1_decimale(
                v(ws, ligne_excel, col_excel),
                zero_tiret=True
            )
        )


# ============================================================
# 9-6 — CULTURES OLEAGINEUSES
# Excel AG,10
# Word objet 285
# ============================================================

print("[6/11] 9-6", flush=True)

ws = wb["AG,10"]

table = verifier_table(
    document,
    285,
    9,
    6
)

remplacer_texte_preserver_style(
    table.cell(0, 0),
    str(ws["A7"].value or "")
)

remplacer_texte_preserver_style(
    table.cell(0, 5),
    str(ws["J7"].value or "")
)

lignes_geo = {
    2: trouver_dernier_total(ws),
    3: trouver_ligne(ws, "Fès - Meknès"),
    4: trouver_ligne(ws, ["Meknès", "Meknes"]),
}

# Word ordre : Arachide / Colza / Sésame / Tournesol
# Source superficie : C / I / E / G
cols_superficie = [3, 9, 5, 7]

for ligne_word, ligne_excel in lignes_geo.items():
    for col_word, col_excel in zip(
        range(1, 5),
        cols_superficie
    ):
        remplacer_texte_preserver_style(
            table.cell(ligne_word, col_word),
            formater_1_decimale(
                v(ws, ligne_excel, col_excel),
                zero_tiret=True
            )
        )

lignes_geo_prod = {
    6: lignes_geo[2],
    7: lignes_geo[3],
    8: lignes_geo[4],
}

# Source production : B / H / D / F
cols_production = [2, 8, 4, 6]

for ligne_word, ligne_excel in lignes_geo_prod.items():
    for col_word, col_excel in zip(
        range(1, 5),
        cols_production
    ):
        remplacer_texte_preserver_style(
            table.cell(ligne_word, col_word),
            formater_1_decimale(
                v(ws, ligne_excel, col_excel),
                zero_tiret=True
            )
        )


# ============================================================
# FONCTION COMMUNE ELEVAGE 9-7 / 9-8 / 9-9
# ============================================================

def actualiser_elevage(
    feuille,
    objet_word
):
    ws = wb[feuille]

    table = verifier_table(
        document,
        objet_word,
        5,
        5
    )

    remplacer_texte_preserver_style(
        table.cell(0, 0),
        str(ws["A6"].value or "")
    )

    remplacer_texte_preserver_style(
        table.cell(0, 4),
        str(ws["E6"].value or "")
    )

    geo = {
        2: trouver_dernier_total(ws),
        3: trouver_ligne(ws, "Fès - Meknès"),
        4: trouver_ligne(ws, ["Meknès", "Meknes"]),
    }

    for ligne_word, ligne_excel in geo.items():
        for col_word, col_excel in zip(
            range(1, 4),
            range(2, 5)
        ):
            remplacer_texte_preserver_style(
                table.cell(ligne_word, col_word),
                formater_1_decimale(
                    v(ws, ligne_excel, col_excel),
                    zero_tiret=True
                )
            )


# ============================================================
# 9-7 — CHEPTEL
# ============================================================

print("[7/11] 9-7", flush=True)
actualiser_elevage("AG,18", 288)


# ============================================================
# 9-8 — BOVINS PAR RACE
# ============================================================

print("[8/11] 9-8", flush=True)
actualiser_elevage("AG,19", 290)


# ============================================================
# 9-9 — ANIMAUX DE TRAIT
# ============================================================

print("[9/11] 9-9", flush=True)
actualiser_elevage("AG,20", 292)


# ============================================================
# 9-10 — SUPERFICIES REBOISEES
# Excel AG,23
# Word objet 295
# ============================================================

print("[10/11] 9-10", flush=True)

ws = wb["AG,23 "]

table = verifier_table(
    document,
    295,
    6,
    4
)

# Le premier bloc est verticalement fusionné dans Word.
# On garde "En hectare" et on actualise uniquement la campagne.
remplacer_texte_preserver_style(
    table.cell(0, 0),
    f"{ws['A6'].value} En hectare"
)

# On garde l'astérisque historique du modèle.
remplacer_texte_preserver_style(
    table.cell(0, 3),
    f"* {ws['D6'].value}"
)

geo = {
    3: trouver_dernier_total(ws),
    4: trouver_ligne(ws, "Fès - Meknès"),
    5: trouver_ligne(
        ws,
        [
            "Méknes/El Hajeb",
            "Meknes/El Hajeb",
            "Meknès/El Hajeb",
        ]
    ),
}

for ligne_word, ligne_excel in geo.items():
    for col_word, col_excel in zip(
        (1, 2),
        (2, 3)
    ):
        remplacer_texte_preserver_style(
            table.cell(ligne_word, col_word),
            formater_entier(
                v(ws, ligne_excel, col_excel),
                zero_tiret=True
            )
        )


# ============================================================
# 9-11 — ESSENCES FORESTIERES NATURELLES
# Excel AG,24
# Word objet 297
# ============================================================

print("[11/11] 9-11", flush=True)

ws = wb["AG,24 "]

table = verifier_table(
    document,
    297,
    5,
    8
)

# On conserve le numéro de note (3) du modèle,
# on actualise uniquement l'année.
remplacer_texte_preserver_style(
    table.cell(0, 0),
    "Année 2024(3)"
)

remplacer_texte_preserver_style(
    table.cell(0, 7),
    "سنة 2024 (3)"
)

geo = {
    2: trouver_dernier_total(ws),
    3: trouver_ligne(ws, "Fès - Meknès"),
    4: trouver_ligne(ws, ["Meknès", "Meknes"]),
}

for ligne_word, ligne_excel in geo.items():
    for col_word, col_excel in zip(
        range(1, 7),
        range(2, 8)
    ):
        remplacer_texte_preserver_style(
            table.cell(ligne_word, col_word),
            formater_entier(
                v(ws, ligne_excel, col_excel),
                zero_tiret=True
            )
        )


# ============================================================
# SOURCES : ANNUAIRE STATISTIQUE DU MAROC 2023 -> 2025
# ============================================================

changer_source_annuaire_2023_vers_2025(
    document,
    [
        272,
        274,
        278,
        280,
        282,
        285,
        288,
        290,
        292,
        295,
        297,
    ]
)


# ============================================================
# SAUVEGARDE
# ============================================================

wb.close()

print("Sauvegarde du Word...", flush=True)

document.save(
    FICHIER_SORTIE
)


print()
print("=" * 88)
print("CHAPITRE AGRICULTURE / ELEVAGE / FORET GENERE")
print("=" * 88)
print("9-1  : céréales - évolution")
print("9-2  : principales céréales")
print("9-3  : légumineuses - évolution")
print("9-4  : superficies des légumineuses")
print("9-5  : production des légumineuses")
print("9-6  : cultures oléagineuses")
print("9-7  : cheptel")
print("9-8  : bovins par race")
print("9-9  : animaux de trait")
print("9-10 : superficies reboisées")
print("9-11 : essences forestières naturelles")
print()
print(f"Document créé : {FICHIER_SORTIE}")
print()
print(
    "Sources : uniquement 'Annuaire Statistique du Maroc 2023' "
    "a été passé à 2025."
)
print(
    "Les autres sources du Word n'ont pas été modifiées."
)
print("=" * 88)
