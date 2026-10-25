from pathlib import Path
from openpyxl import load_workbook
from docx import Document
from docx.oxml.ns import qn
import re
import unicodedata
from decimal import Decimal, ROUND_HALF_UP

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_SORTIE = Path("sorties")

CANDIDATS_ENTREE = [
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER_POPULATION_INDUSTRIE.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER_POPULATION.docx",
]

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER_POPULATION_INDUSTRIE_ENERGIE_EAU.docx"
)

# ============================================================
# OUTILS
# ============================================================

def norm(v):
    s = str(v or "").strip().lower()
    s = (
        s.replace("’", "'")
         .replace("–", "-")
         .replace("—", "-")
         .replace("\u202f", " ")
         .replace("\xa0", " ")
    )
    s = "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"\s+", " ", s).strip()

def fmt_int(v):
    if v is None or v == "":
        return "-"
    if isinstance(v, str) and v.strip() in ("-", "...", "…"):
        return v.strip()

    n = Decimal(str(v).replace(",", "."))
    n = n.quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    return f"{int(n):,}".replace(",", " ")

def fmt_1(v):
    if v is None or v == "":
        return "-"
    if isinstance(v, str) and v.strip() in ("-", "...", "…"):
        return v.strip()

    n = Decimal(str(v).replace(",", "."))
    n = n.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
    s = f"{n:,.1f}"
    return s.replace(",", "X").replace(".", ",").replace("X", " ")

def fmt_exact(v):
    if v is None or v == "":
        return ""
    if isinstance(v, str):
        t = v.strip()
        if t in ("-", "...", "…"):
            return t
        try:
            n = Decimal(t.replace(",", "."))
        except Exception:
            return t
    else:
        n = Decimal(str(v))

    s = format(n.normalize(), "f")
    if "." in s:
        s = s.rstrip("0").rstrip(".")
    return s.replace(".", ",")

def set_cell(cell, value):
    value = str(value)

    if not cell.paragraphs:
        cell.text = value
        return

    p = cell.paragraphs[0]

    if p.runs:
        p.runs[0].text = value
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(value)

    for p2 in cell.paragraphs[1:]:
        for run in p2.runs:
            run.text = ""

def assert_same(label, actual, expected):
    if norm(actual) != norm(expected):
        raise RuntimeError(
            f"CONTROLE ECHEC {label}: Word={actual!r}, attendu={expected!r}"
        )

def unmerge_vertical_cells(table, row_indexes, col_indexes):
    """
    Supprime uniquement les fusions VERTICALES w:vMerge des cellules ciblées.
    Nécessaire quand l'ancien Word avait fusionné deux cellules parce que
    leurs anciennes valeurs étaient identiques, alors que les nouvelles
    données Excel sont différentes.

    Utilise le XML brut afin de ne pas réécrire le style/bordures du tableau.
    """
    for r in row_indexes:
        tr = table._tbl.tr_lst[r]

        for c in col_indexes:
            if c >= len(tr.tc_lst):
                raise RuntimeError(
                    f"Impossible de défusionner cellule brute R{r+1} C{c+1}: "
                    f"la ligne ne contient que {len(tr.tc_lst)} tc."
                )

            tc = tr.tc_lst[c]
            tcPr = tc.get_or_add_tcPr()

            # Retirer tous les w:vMerge (restart ou continue).
            for vm in list(tcPr.findall(qn("w:vMerge"))):
                tcPr.remove(vm)

def find_excel_file(words):
    candidates = []

    for p in DOSSIER_SOURCE.glob("*.xlsx"):
        if p.name.startswith("~$"):
            continue
        n = norm(p.name)
        if all(norm(w) in n for w in words):
            candidates.append(p)

    if len(candidates) != 1:
        raise RuntimeError(
            f"Excel ambigu/introuvable pour {words}: {[p.name for p in candidates]}"
        )

    return candidates[0]

def find_word_data_table(doc, code):
    a, b = code.split("-")
    pat = re.compile(
        rf"^\s*{re.escape(a)}\s*-\s*{re.escape(b)}\b",
        re.I
    )

    for i, table in enumerate(doc.tables):
        if not table.rows:
            continue

        title = table.cell(0, 0).text or ""

        if pat.search(title):
            for j in range(i + 1, min(i + 6, len(doc.tables))):
                cand = doc.tables[j]
                if len(cand.rows) >= 3:
                    return cand

    raise RuntimeError(f"Tableau Word {code} introuvable.")

def find_sheet_by_table(wb, table_code, required_terms):
    a, b = table_code.split("-")
    pat = re.compile(
        rf"^\s*{re.escape(a)}\s*-\s*{re.escape(b)}\b",
        re.I
    )
    terms = [norm(x) for x in required_terms]
    candidates = []

    for ws in wb.worksheets:
        title_found = False
        top_text = []

        for r in range(1, min(ws.max_row, 20) + 1):
            for c in range(1, min(ws.max_column, 10) + 1):
                v = ws.cell(r, c).value

                if v not in (None, ""):
                    top_text.append(str(v))

                if isinstance(v, str) and pat.match(v):
                    title_found = True

        if not title_found:
            continue

        joined = norm(" ".join(top_text))

        if not all(t in joined for t in terms):
            continue

        candidates.append(ws)

    if len(candidates) != 1:
        raise RuntimeError(
            f"Feuille {table_code} ambiguë/introuvable: {[ws.title for ws in candidates]}"
        )

    return candidates[0]

def find_row(ws, labels):
    wanted = {norm(x) for x in labels}
    found = []

    for r in range(1, ws.max_row + 1):
        if norm(ws.cell(r, 1).value) in wanted:
            found.append(r)

    if len(found) != 1:
        raise RuntimeError(
            f"Ligne {labels} ambiguë/introuvable dans {ws.title}: {found}"
        )

    return found[0]

def replace_annuaire_2023_with_2025(table):
    for row in table.rows:
        for cell in row.cells:
            text = cell.text or ""

            if (
                "Annuaire Statistique du Maroc 2023" in text
                or "النشرة الإحصائية السنوية للمغرب 2023" in text
            ):
                text = text.replace(
                    "Annuaire Statistique du Maroc 2023",
                    "Annuaire Statistique du Maroc 2025"
                )
                text = text.replace(
                    "النشرة الإحصائية السنوية للمغرب 2023",
                    "النشرة الإحصائية السنوية للمغرب 2025"
                )
                set_cell(cell, text)

# ============================================================
# FICHIERS
# ============================================================

input_docx = next(
    (p for p in CANDIDATS_ENTREE if p.exists()),
    None
)

if input_docx is None:
    raise FileNotFoundError(
        "Word Industrie d'entrée introuvable."
    )

excel_energy = find_excel_file(["energie", "eau", "2025"])
excel_prices = find_excel_file(["indice", "prix", "2025"])

print("=" * 100)
print("ENERGIE / EAU — GENERATION AVEC DOUBLE CONTROLE")
print("=" * 100)
print(f"Word d'entrée : {input_docx}")
print(f"Excel Energie/Eau : {excel_energy.name}")
print(f"Excel Prix : {excel_prices.name}")
print()

doc = Document(input_docx)
wb_energy = load_workbook(excel_energy, data_only=True, read_only=False)
wb_prices = load_workbook(excel_prices, data_only=True, read_only=False)

# ============================================================
# 11-1 — ACTIVITE ONEE EAU
# ============================================================

print("[1/4] 11-1 Activité ONEE Eau", flush=True)

t = find_word_data_table(doc, "11-1")
ws = find_sheet_by_table(
    wb_energy,
    "5-20",
    ["activite", "office national", "eau potable"]
)

# Source explicitement 2023
year_found = any(
    "2023" in str(ws.cell(r, 1).value or "")
    for r in range(1, min(ws.max_row, 20) + 1)
)

if not year_found:
    raise RuntimeError("11-1 : année 2023 non trouvée dans Excel 5-20.")

r_nat = find_row(ws, ["Ensemble"])
r_reg = find_row(ws, ["Fès - Meknès", "Fès-Meknès"])
r_mek = find_row(ws, ["Meknès"])

# Colonnes Word: abonnés, ventes abonnés, ventes régies, production
for wr, er in zip([4, 5, 6], [r_nat, r_reg, r_mek]):
    values = [
        fmt_int(ws.cell(er, 2).value),
        fmt_int(ws.cell(er, 3).value),
        fmt_int(ws.cell(er, 4).value),
        fmt_int(ws.cell(er, 5).value),
    ]

    for wc, val in zip([1, 2, 3, 4], values):
        set_cell(t.cell(wr, wc), val)

# Année FR/AR
set_cell(t.cell(0, 0), "Année 2023")
set_cell(t.cell(0, 5), "سنة 2023")

# Contrôle après écriture
expected_11_1 = {
    "Nation": [
        fmt_int(ws.cell(r_nat, 2).value),
        fmt_int(ws.cell(r_nat, 3).value),
        fmt_int(ws.cell(r_nat, 4).value),
        fmt_int(ws.cell(r_nat, 5).value),
    ],
    "Région": [
        fmt_int(ws.cell(r_reg, 2).value),
        fmt_int(ws.cell(r_reg, 3).value),
        fmt_int(ws.cell(r_reg, 4).value),
        fmt_int(ws.cell(r_reg, 5).value),
    ],
    "Meknès": [
        fmt_int(ws.cell(r_mek, 2).value),
        fmt_int(ws.cell(r_mek, 3).value),
        fmt_int(ws.cell(r_mek, 4).value),
        fmt_int(ws.cell(r_mek, 5).value),
    ],
}

for label, wr in [("Nation", 4), ("Région", 5), ("Meknès", 6)]:
    for wc, exp in zip([1, 2, 3, 4], expected_11_1[label]):
        assert_same(
            f"11-1 {label} C{wc+1}",
            t.cell(wr, wc).text,
            exp
        )

print("   -> 2023 + Nation/Région/Meknès : OK")
print(f"   -> Nation : {expected_11_1['Nation']}")
print(f"   -> Région : {expected_11_1['Région']}")
print(f"   -> Meknès : {expected_11_1['Meknès']}")

# ============================================================
# 11-2 — VENTES D'ELECTRICITE
# ============================================================

print("[2/4] 11-2 Ventes d'électricité", flush=True)

t = find_word_data_table(doc, "11-2")
ws = find_sheet_by_table(
    wb_energy,
    "5-7",
    ["ventes", "electricite"]
)

r_nat = find_row(ws, ["Total"])
r_reg = find_row(ws, ["Fès - Meknès", "Fès-Meknès"])
r_mek = find_row(ws, ["Meknès"])

# Années exactes du tableau Excel
set_cell(t.cell(0, 1), "2024*")
set_cell(t.cell(0, 2), "2023R")
set_cell(t.cell(0, 3), "2022R")

for wr, er in zip([1, 2, 3], [r_nat, r_reg, r_mek]):
    for wc, ec in zip([1, 2, 3], [2, 3, 4]):
        set_cell(
            t.cell(wr, wc),
            fmt_1(ws.cell(er, ec).value)
        )

replace_annuaire_2023_with_2025(t)

# Contrôle des chiffres
for label, wr, er in [
    ("Nation", 1, r_nat),
    ("Région", 2, r_reg),
    ("Meknès", 3, r_mek),
]:
    for wc, ec in zip([1, 2, 3], [2, 3, 4]):
        assert_same(
            f"11-2 {label} {ec}",
            t.cell(wr, wc).text,
            fmt_1(ws.cell(er, ec).value)
        )

assert_same("11-2 année 2024", t.cell(0, 1).text, "2024*")
assert_same("11-2 année 2023", t.cell(0, 2).text, "2023R")
assert_same("11-2 année 2022", t.cell(0, 3).text, "2022R")

print("   -> 2024* / 2023R / 2022R : OK")
print(
    "   -> Nation :",
    [fmt_1(ws.cell(r_nat, c).value) for c in [2, 3, 4]]
)
print(
    "   -> Région :",
    [fmt_1(ws.cell(r_reg, c).value) for c in [2, 3, 4]]
)
print(
    "   -> Meknès :",
    [fmt_1(ws.cell(r_mek, c).value) for c in [2, 3, 4]]
)

# ============================================================
# 15-3 — CORRECTION PRIX ELECTRICITE
# ============================================================

print("[3/4] 15-3 Prix réglementés électricité — revalidation", flush=True)

t = find_word_data_table(doc, "15-3")
ws = find_sheet_by_table(
    wb_prices,
    "17-25",
    ["prix reglementes", "electricite"]
)

# Contrôle dates
expected_dates = [
    "1.1.2017",
    "1.1.2016",
    "1.1.2015",
    "1.8.2014",
    "25.2.2009",
]

for wc, date in zip([1, 2, 3, 4, 5], expected_dates):
    set_cell(t.cell(0, wc), date)
    set_cell(t.cell(1, wc), date)

# IMPORTANT:
# Dans l'ancien modèle Word, certaines cellules des tranches 201-300 et
# 301-500 étaient fusionnées verticalement car leurs anciennes valeurs
# étaient identiques. Excel 2025 les distingue maintenant. On défusionne
# seulement ces cellules de données avant de les réécrire.
unmerge_vertical_cells(
    t,
    row_indexes=[5, 6, 11, 12],
    col_indexes=[1, 2, 3, 4, 5]
)

# Correspondances directes Excel -> Word.
# Les deux lignes 101-150 et 151-200 Excel ont les mêmes valeurs ;
# Word les regroupe en "101 à 200", donc on utilise la première.
mapping_15_3 = {
    3: 11,   # domestique <100
    4: 12,   # domestique 101-200
    5: 14,   # domestique 201-300
    6: 15,   # domestique 301-500
    7: 16,   # domestique >500
    9: 19,   # éclairage privé <100
    10: 20,  # éclairage privé 101-200
    11: 22,  # éclairage privé 201-300
    12: 23,  # éclairage privé 301-500
    13: 24,  # éclairage privé >500
    15: 27,  # éclairage patenté <150
    16: 28,  # éclairage patenté >150
    17: 30,  # administratif
    18: 32,  # public
    20: 35,  # force motrice <100
    21: 36,  # force motrice 101-500
    22: 37,  # force motrice >500
    25: 42,  # MT redevance puissance
    27: 44,  # MT pointes
    28: 45,  # MT pleines
    29: 46,  # MT creuses
    32: 51,  # HT/THT redevance puissance
    34: 53,  # HT/THT pointes
    35: 54,  # HT/THT pleines
    36: 55,  # HT/THT creuses
}

for wr, er in mapping_15_3.items():
    for wc, ec in zip([1, 2, 3, 4, 5], [2, 3, 4, 5, 6]):
        set_cell(
            t.cell(wr, wc),
            fmt_exact(ws.cell(er, ec).value)
        )

# Contrôle ciblé des erreurs repérées dans l'ancien Word :
assert_same("15-3 domestique 201-300 / 2015", t.cell(5, 3).text, "1,0904")
assert_same("15-3 domestique 201-300 / 2014", t.cell(5, 4).text, "1,0541")
assert_same("15-3 domestique 201-300 / 2009", t.cell(5, 5).text, "1,054")
assert_same("15-3 privé 201-300 / 2015", t.cell(11, 3).text, "1,0904")
assert_same("15-3 privé 201-300 / 2014", t.cell(11, 4).text, "1,0541")
assert_same("15-3 privé 201-300 / 2009", t.cell(11, 5).text, "1,054")

print("   -> fusions verticales anciennes séparées si nécessaire : OK")
print("   -> données revalidées depuis 17-25 : OK")
print("   -> anciennes valeurs erronées 201-300 corrigées.")

# ============================================================
# 15-4 — CORRECTION PRIX EAU MEKNES
# ============================================================

print("[4/4] 15-4 Prix réglementés eau Meknès — revalidation", flush=True)

t = find_word_data_table(doc, "15-4")
ws = find_sheet_by_table(
    wb_prices,
    "17-26",
    ["prix reglementes", "eau potable"]
)

r_mek_title = find_row(ws, ["Meknès"])

# Vérifier que les lignes Meknès suivent bien le titre :
if norm(ws.cell(r_mek_title + 1, 1).value) != norm("Consommation domestique"):
    raise RuntimeError("15-4 : structure Meknès inattendue dans Excel 17-26.")

# Dates
dates_excel = ["2014/08/01", "2013/04/04", "2006/02/06", "2003/04/25", "2000/11/01"]

for wc, date in zip([1, 2, 3, 4, 5], dates_excel):
    set_cell(t.cell(0, wc), date)

# Le modèle peut aussi avoir fusionné verticalement les lignes 6-12
# et 12-20 parce que certaines anciennes valeurs étaient identiques.
# On sépare uniquement leurs cellules numériques avant écriture.
unmerge_vertical_cells(
    t,
    row_indexes=[3, 4],
    col_indexes=[1, 2, 3, 4, 5]
)

# Excel rows relative to Meknès title:
# +2: 0-6 ; +3: 6-12 ; +4: 12-20 ; +5: 20-35 ; +6: >35 ; +7: industriel
excel_rows = [
    r_mek_title + 2,
    r_mek_title + 3,
    r_mek_title + 4,
    r_mek_title + 5,
    r_mek_title + 6,
    r_mek_title + 7,
]
word_rows = [2, 3, 4, 5, 6, 7]

for wr, er in zip(word_rows, excel_rows):
    for wc, ec in zip([1, 2, 3, 4, 5], [2, 3, 4, 5, 6]):
        set_cell(
            t.cell(wr, wc),
            fmt_exact(ws.cell(er, ec).value)
        )

# Contrôles ciblés sur la ligne qui était incomplète dans Word
expected_6_12 = ["5,12", "5,12", "3,88", "3,88", "2,84"]

for wc, exp in zip([1, 2, 3, 4, 5], expected_6_12):
    assert_same(
        f"15-4 Meknès 6-12 C{wc+1}",
        t.cell(3, wc).text,
        exp
    )

expected_12_20 = ["5,12", "", "", "", ""]

for wc, exp in zip([1, 2, 3, 4, 5], expected_12_20):
    assert_same(
        f"15-4 Meknès 12-20 C{wc+1}",
        t.cell(4, wc).text,
        exp
    )

print("   -> bloc Meknès recopié et relu depuis 17-26 : OK")
print("   -> ligne 6-12 historique corrigée.")

# ============================================================
# SAUVEGARDE
# ============================================================

wb_energy.close()
wb_prices.close()

doc.save(FICHIER_SORTIE)

print()
print("=" * 100)
print("ENERGIE / EAU TERMINE — AVEC REVALIDATION DES TABLES PRIX ENERGIE/EAU")
print("=" * 100)
print("11-1 : actualisé à 2023 (année réellement disponible dans 5-20)")
print("11-2 : actualisé à 2024* / 2023R / 2022R")
print("15-3 : revalidé et erreurs 201-300 corrigées")
print("15-4 : revalidé et valeurs historiques Meknès manquantes corrigées")
print()
print(f"Document créé : {FICHIER_SORTIE}")
print("=" * 100)
