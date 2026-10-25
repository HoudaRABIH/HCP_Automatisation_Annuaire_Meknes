from pathlib import Path
from openpyxl import load_workbook
from docx import Document
import math
import re
import unicodedata


# ============================================================
# CHEMINS
# ============================================================

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_MODELES = Path("modeles")
DOSSIER_SORTIE = Path("sorties")

FICHIER_ENV = DOSSIER_SOURCE / "23. Environnement_AS 2025.xlsx"

# On continue depuis le Word cumulé Transport + Prix.
FICHIER_ENTREE = DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix.docx"

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement.docx"
)


# ============================================================
# OUTILS
# ============================================================

def normaliser(texte):
    texte = str(texte or "").strip().lower()
    texte = texte.replace("’", "'")

    texte = "".join(
        c
        for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )

    texte = re.sub(r"\s+", " ", texte)
    return texte


def feuille(classeur, nom):
    cible = normaliser(nom)

    for ws in classeur.worksheets:
        if normaliser(ws.title) == cible:
            return ws

    raise KeyError(
        f"Feuille introuvable : {nom}"
    )


def remplacer_texte_preserver_style(cellule, nouveau_texte):
    """
    Modifie seulement le texte de la cellule afin de préserver
    bordures, trames, dimensions, alignements et style du premier run.
    """
    nouveau_texte = str(nouveau_texte)

    if not cellule.paragraphs:
        cellule.text = nouveau_texte
        return

    paragraphe = cellule.paragraphs[0]

    if paragraphe.runs:
        paragraphe.runs[0].text = nouveau_texte

        for run in paragraphe.runs[1:]:
            run.text = ""
    else:
        paragraphe.add_run(nouveau_texte)

    for autre_paragraphe in cellule.paragraphs[1:]:
        for run in autre_paragraphe.runs:
            run.text = ""


def formater_decimal(valeur, decimales=1):
    if valeur in (None, ""):
        return ""

    nombre = float(valeur)

    texte = f"{nombre:.{decimales}f}"

    return texte.replace(".", ",")


def formater_entier(valeur):
    if valeur in (None, ""):
        return ""

    return str(
        int(
            round(
                float(valeur)
            )
        )
    )


def formater_pourcentage_fraction(valeur):
    """
    Les humidités Excel sont stockées sous forme 0,61 = 61 %.
    Le Word historique les affiche comme entiers.
    """
    if valeur in (None, ""):
        return ""

    pourcentage = float(valeur) * 100

    return str(
        int(
            round(
                pourcentage
            )
        )
    )


def decimal_vers_degres_minutes(valeur):
    """
    Convertit -5.52049 -> 5°31'
    et 33.879127 -> 33°53'

    Le modèle Word historique n'affiche pas E/O/N/S,
    donc on conserve exactement cette convention visuelle.
    """
    valeur = abs(
        float(valeur)
    )

    degres = int(
        math.floor(
            valeur
        )
    )

    minutes = int(
        round(
            (valeur - degres) * 60
        )
    )

    # Cas rare où l'arrondi donne 60 minutes.
    if minutes == 60:
        degres += 1
        minutes = 0

    return f"{degres}°{minutes:02d}'"


def ouvrir_document_entree():
    if FICHIER_ENTREE.exists():
        return FICHIER_ENTREE

    fichiers_modele = [
        f
        for f in DOSSIER_MODELES.glob("*.docx")
        if not f.name.startswith("~$")
    ]

    if not fichiers_modele:
        raise FileNotFoundError(
            "Aucun modèle Word trouvé."
        )

    return fichiers_modele[0]


# ============================================================
# OUVRIR EXCEL
# ============================================================

if not FICHIER_ENV.exists():
    raise FileNotFoundError(
        f"Source Environnement introuvable : {FICHIER_ENV}"
    )

classeur = load_workbook(
    FICHIER_ENV,
    data_only=True,
    read_only=False
)

ws_16_1 = feuille(
    classeur,
    "ENV 3"
)

ws_16_2 = feuille(
    classeur,
    "ENV 4"
)

ws_16_3 = feuille(
    classeur,
    "ENV 5"
)

ws_16_4 = feuille(
    classeur,
    "ENV6"
)

ws_16_5 = feuille(
    classeur,
    "ENV 7"
)

ws_16_6 = feuille(
    classeur,
    "ENV 8"
)


# ============================================================
# OUVRIR WORD
# ============================================================

fichier_entree = ouvrir_document_entree()

document = Document(
    fichier_entree
)


# ============================================================
# 16-1 — LIEU GEOGRAPHIQUE DE LA STATION DE MEKNES
# Objet Word 401
# ============================================================

table_16_1 = document.tables[
    401 - 1
]

if (
    len(table_16_1.rows) != 4
    or len(table_16_1.columns) != 5
):
    raise RuntimeError(
        "Structure Word inattendue pour 16-1."
    )


annee_16_1 = str(
    ws_16_1["A6"].value or ""
).replace(
    "Année ",
    ""
).strip()


# Année FR / AR
remplacer_texte_preserver_style(
    table_16_1.cell(
        0,
        0
    ),
    f"Année {annee_16_1}"
)

remplacer_texte_preserver_style(
    table_16_1.cell(
        0,
        4
    ),
    f"سنة {annee_16_1}"
)


# Meknès = ligne 24
longitude = (
    ws_16_1["B24"].value
)

latitude = (
    ws_16_1["C24"].value
)

altitude = (
    ws_16_1["D24"].value
)


remplacer_texte_preserver_style(
    table_16_1.cell(
        3,
        1
    ),
    decimal_vers_degres_minutes(
        longitude
    )
)

remplacer_texte_preserver_style(
    table_16_1.cell(
        3,
        2
    ),
    decimal_vers_degres_minutes(
        latitude
    )
)

remplacer_texte_preserver_style(
    table_16_1.cell(
        3,
        3
    ),
    formater_entier(
        altitude
    )
)


# ============================================================
# FONCTION COMMUNE 16-2 A 16-5
# ============================================================

def actualiser_table_station_moyenne(
    table_word,
    feuille_excel,
    ligne_meknes,
    ligne_moyenne,
    colonnes_annees,
    ligne_annees,
    format_valeur,
):
    """
    Structure Word attendue :
      R1 : unité + 4 années + arabe
      R2 : Station
      R3 : Meknès
      R4 : Moyenne
    """

    if (
        len(table_word.rows) != 4
        or len(table_word.columns) != 6
    ):
        raise RuntimeError(
            "Structure Word 4x6 attendue."
        )

    # Années
    for colonne_word, colonne_excel in zip(
        range(
            1,
            5
        ),
        colonnes_annees
    ):
        valeur_annee = feuille_excel.cell(
            ligne_annees,
            colonne_excel
        ).value

        remplacer_texte_preserver_style(
            table_word.cell(
                0,
                colonne_word
            ),
            str(
                valeur_annee
            )
        )

    # Meknès
    for colonne_word, colonne_excel in zip(
        range(
            1,
            5
        ),
        colonnes_annees
    ):
        remplacer_texte_preserver_style(
            table_word.cell(
                2,
                colonne_word
            ),
            format_valeur(
                feuille_excel.cell(
                    ligne_meknes,
                    colonne_excel
                ).value
            )
        )

    # Moyenne
    for colonne_word, colonne_excel in zip(
        range(
            1,
            5
        ),
        colonnes_annees
    ):
        remplacer_texte_preserver_style(
            table_word.cell(
                3,
                colonne_word
            ),
            format_valeur(
                feuille_excel.cell(
                    ligne_moyenne,
                    colonne_excel
                ).value
            )
        )


# ============================================================
# 16-2 — PLUVIOMETRIE ANNUELLE
# Objet Word 403
# Excel : années C:F ; Meknès R20 ; Moyenne R33
# ============================================================

table_16_2 = document.tables[
    403 - 1
]

actualiser_table_station_moyenne(
    table_word=table_16_2,
    feuille_excel=ws_16_2,
    ligne_meknes=20,
    ligne_moyenne=33,
    colonnes_annees=[
        3,
        4,
        5,
        6
    ],
    ligne_annees=4,
    format_valeur=lambda v: formater_decimal(
        v,
        decimales=1
    )
)


# ============================================================
# 16-3 — TEMPERATURE MINIMALE
# Objet Word 405
# Excel : années B:E ; Meknès R20 ; Moyenne R33
# ============================================================

table_16_3 = document.tables[
    405 - 1
]

actualiser_table_station_moyenne(
    table_word=table_16_3,
    feuille_excel=ws_16_3,
    ligne_meknes=20,
    ligne_moyenne=33,
    colonnes_annees=[
        2,
        3,
        4,
        5
    ],
    ligne_annees=4,
    format_valeur=lambda v: formater_decimal(
        v,
        decimales=1
    )
)


# ============================================================
# 16-4 — TEMPERATURE MAXIMALE
# Objet Word 408
# Excel : années B:E ; Meknès R20 ; Moyenne R33
# ============================================================

table_16_4 = document.tables[
    408 - 1
]

actualiser_table_station_moyenne(
    table_word=table_16_4,
    feuille_excel=ws_16_4,
    ligne_meknes=20,
    ligne_moyenne=33,
    colonnes_annees=[
        2,
        3,
        4,
        5
    ],
    ligne_annees=4,
    format_valeur=lambda v: formater_decimal(
        v,
        decimales=1
    )
)


# ============================================================
# 16-5 — TEMPERATURE QUOTIDIENNE
# Objet Word 410
# Excel : années B:E ; Meknès R25 ; Moyenne R43
# ============================================================

table_16_5 = document.tables[
    410 - 1
]

actualiser_table_station_moyenne(
    table_word=table_16_5,
    feuille_excel=ws_16_5,
    ligne_meknes=25,
    ligne_moyenne=43,
    colonnes_annees=[
        2,
        3,
        4,
        5
    ],
    ligne_annees=4,
    format_valeur=lambda v: formater_decimal(
        v,
        decimales=1
    )
)


# Le fichier Excel comporte "En millimètres" par erreur pour cette
# feuille température. On conserve donc volontairement l'unité Word "En C°".


# ============================================================
# 16-6 — HUMIDITE QUOTIDIENNE
# Objet Word 412
# Excel : valeurs sous forme fractionnaire 0,61 = 61 %
# années B:E ; Meknès R25 ; Moyenne R43
# ============================================================

table_16_6 = document.tables[
    412 - 1
]

actualiser_table_station_moyenne(
    table_word=table_16_6,
    feuille_excel=ws_16_6,
    ligne_meknes=25,
    ligne_moyenne=43,
    colonnes_annees=[
        2,
        3,
        4,
        5
    ],
    ligne_annees=4,
    format_valeur=formater_pourcentage_fraction
)


# ============================================================
# SAUVEGARDE
# ============================================================

classeur.close()

DOSSIER_SORTIE.mkdir(
    exist_ok=True
)

document.save(
    FICHIER_SORTIE
)


# ============================================================
# RESULTAT CMD
# ============================================================

print("=" * 82)
print("CHAPITRE ENVIRONNEMENT GENERE")
print("=" * 82)

print()
print("16-1 Station météorologique de Meknès")
print(
    f"Année : {annee_16_1}"
)
print(
    "Longitude : "
    + decimal_vers_degres_minutes(
        longitude
    )
)
print(
    "Latitude : "
    + decimal_vers_degres_minutes(
        latitude
    )
)
print(
    "Altitude : "
    + formater_entier(
        altitude
    )
    + " m"
)

print()
print("16-2 Pluviométrie : 2023-2024 à 2020-2021")
print("16-3 Température minimale : actualisé")
print("16-4 Température maximale : actualisé")
print("16-5 Température quotidienne : actualisé")
print("16-6 Humidité : actualisé et converti en %")

print()
print(
    f"Document créé : {FICHIER_SORTIE}"
)

print(
    f"Document d'entrée : {fichier_entree}"
)

print()
print(
    "Le modèle original n'a pas été modifié."
)

print("=" * 82)
