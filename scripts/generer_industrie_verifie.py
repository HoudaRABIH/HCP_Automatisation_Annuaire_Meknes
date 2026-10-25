from pathlib import Path
from openpyxl import load_workbook
from docx import Document
import re
import unicodedata

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_SORTIE = Path("sorties")

CANDIDATS_ENTREE = [
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER_POPULATION.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER.docx",
]

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER_POPULATION_INDUSTRIE.docx"
)

def norm(v):
    s = str(v or "").strip().lower()
    s = (
        s.replace("’", "'")
         .replace("–", "-")
         .replace("—", "-")
         .replace("\u202f", " ")
         .replace("\xa0", " ")
    )
    s = "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"\s+", " ", s).strip()

def fmt(v):
    if v is None or v == "":
        return "-"
    if isinstance(v, str):
        t = v.strip()
        if t in ("-", "–", "...", "…"):
            return t
        try:
            n = float(
                t.replace("\u202f", "")
                 .replace("\xa0", "")
                 .replace(" ", "")
                 .replace(",", ".")
            )
        except ValueError:
            return t
    else:
        n = float(v)

    if abs(n - round(n)) < 1e-9:
        return f"{int(round(n)):,}".replace(",", " ")
    return str(n).replace(".", ",")

def set_cell(cell, value):
    value = str(value)

    if not cell.paragraphs:
        cell.text = value
        return

    p = cell.paragraphs[0]

    if p.runs:
        p.runs[0].text = value
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(value)

    for p2 in cell.paragraphs[1:]:
        for run in p2.runs:
            run.text = ""

def find_word_data_table(doc):
    pat = re.compile(r"^\s*10\s*-\s*1\b", re.I)

    for i, table in enumerate(doc.tables):
        if not table.rows:
            continue

        title = table.cell(0, 0).text or ""

        if pat.search(title):
            for j in range(i + 1, min(i + 6, len(doc.tables))):
                cand = doc.tables[j]

                if len(cand.rows) >= 3:
                    return cand

    raise RuntimeError("Tableau Word 10-1 introuvable.")

def find_excel_file():
    for p in DOSSIER_SOURCE.glob("*.xlsx"):
        if p.name.startswith("~$"):
            continue

        n = norm(p.name)

        if "industrie" in n and "artisanat" in n and "2025" in n:
            return p

    raise FileNotFoundError(
        "Fichier 'Industrie et artisanat_AS 2025.xlsx' introuvable."
    )

def find_source_sheet(wb):
    pat = re.compile(r"^\s*6\s*-\s*18\b", re.I)

    candidates = []

    for ws in wb.worksheets:
        found_title = False
        has_2024 = False
        has_meknes = False
        has_total = False

        for r in range(1, ws.max_row + 1):
            a = ws.cell(r, 1).value

            if isinstance(a, str) and pat.match(a):
                found_title = True

            an = norm(a)

            if "2024" in str(a):
                has_2024 = True

            if an in {"meknes", "meknes"}:
                has_meknes = True

            if an == "total":
                has_total = True

        if found_title and has_2024 and has_meknes and has_total:
            candidates.append(ws)

    if len(candidates) != 1:
        raise RuntimeError(
            f"Source 6-18 ambiguë/introuvable : {len(candidates)} feuille(s) validée(s)."
        )

    return candidates[0]

def find_row(ws, labels):
    wanted = {norm(x) for x in labels}
    found = []

    for r in range(1, ws.max_row + 1):
        if norm(ws.cell(r, 1).value) in wanted:
            found.append(r)

    if len(found) != 1:
        raise RuntimeError(
            f"Ligne {labels} ambiguë/introuvable dans {ws.title}: {found}"
        )

    return found[0]

def assert_equal(label, actual, expected):
    if norm(actual) != norm(expected):
        raise RuntimeError(
            f"CONTROLE ECHEC {label}: Word={actual!r}, attendu={expected!r}"
        )

input_docx = next(
    (p for p in CANDIDATS_ENTREE if p.exists()),
    None
)

if input_docx is None:
    raise FileNotFoundError(
        "Word Population d'entrée introuvable dans sorties."
    )

excel_path = find_excel_file()

print("=" * 96)
print("INDUSTRIE — GENERATION 10-1 AVEC DOUBLE CONTROLE")
print("=" * 96)
print(f"Word d'entrée : {input_docx}")
print(f"Excel : {excel_path.name}")
print()

doc = Document(input_docx)
wb = load_workbook(excel_path, data_only=True, read_only=False)

table = find_word_data_table(doc)
ws = find_source_sheet(wb)

print(f"Source Excel validée : feuille {ws.title}")

r_nat = find_row(ws, ["Total"])
r_mek = find_row(ws, ["Méknes", "Meknès", "Meknes"])

# Contrôle structure Excel
year_cell = None
for r in range(1, min(ws.max_row, 20) + 1):
    a = str(ws.cell(r, 1).value or "").strip()
    if "2024" in a:
        year_cell = a
        break

if year_cell is None:
    raise RuntimeError("Année 2024* introuvable dans la source 6-18.")

# Valeurs attendues depuis Excel
nation = [
    fmt(ws.cell(r_nat, 2).value),  # Mâasras
    fmt(ws.cell(r_nat, 4).value),  # Capacité
    fmt(ws.cell(r_nat, 6).value),  # Unités industrielles
]

meknes = [
    fmt(ws.cell(r_mek, 2).value),
    fmt(ws.cell(r_mek, 4).value),
    fmt(ws.cell(r_mek, 6).value),
]

print(f"Nation Excel : {nation}")
print(f"Meknès Excel : {meknes}")

# Contrôles de sécurité avec les valeurs observées dans le diagnostic
assert_equal("Nation - Mâasras", nation[0], "15 087")
assert_equal("Nation - Capacité", nation[1], "505 800")
assert_equal("Nation - Unités", nation[2], "287")

assert_equal("Meknès - Mâasras", meknes[0], "183")
assert_equal("Meknès - Capacité", meknes[1], "74 000")
assert_equal("Meknès - Unités", meknes[2], "66")

# Mise à jour année uniquement après contrôles
set_cell(table.cell(0, 0), "Année 2024*")
set_cell(table.cell(0, 4), "سنة 2024*")
set_cell(table.cell(1, 0), "Année 2024*")
set_cell(table.cell(1, 4), "سنة 2024*")

# Word order:
# C2 = Mâasras
# C3 = Capacité nominale
# C4 = unités industrielles
for wc, value in zip([1, 2, 3], nation):
    set_cell(table.cell(2, wc), value)

for wc, value in zip([1, 2, 3], meknes):
    set_cell(table.cell(3, wc), value)

# Contrôle après écriture
for wc, expected in zip([1, 2, 3], nation):
    assert_equal(
        f"Word Nation colonne {wc+1}",
        table.cell(2, wc).text,
        expected
    )

for wc, expected in zip([1, 2, 3], meknes):
    assert_equal(
        f"Word Meknès colonne {wc+1}",
        table.cell(3, wc).text,
        expected
    )

assert_equal("Année Word FR", table.cell(0, 0).text, "Année 2024*")
assert_equal("Année Word AR", table.cell(0, 4).text, "سنة 2024*")

wb.close()

doc.save(FICHIER_SORTIE)

print()
print("10-1 : OK")
print("   - titre source 6-18 vérifié")
print("   - année 2024* vérifiée")
print("   - Nation vérifiée")
print("   - Meknès vérifiée")
print("   - valeurs Word relues après écriture")
print("   - source Word non modifiée")
print()
print(f"Document créé : {FICHIER_SORTIE}")
print("=" * 96)
