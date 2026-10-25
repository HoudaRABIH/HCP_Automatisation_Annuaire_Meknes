from pathlib import Path
from openpyxl import load_workbook
from docx import Document
from docx.table import _Cell
import re
import unicodedata


# ============================================================
# CHEMINS
# ============================================================

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_SORTIE = Path("sorties")

FICHIER_JUSTICE = (
    DOSSIER_SOURCE
    / "15.Justice et Sauvegarde de l'Enfance _AS2025.xlsx"
)

CANDIDATS_ENTREE = [
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_OK.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_FINAL_CLEAN.docx",
]

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE.docx"
)


# ============================================================
# OUTILS TEXTE / FORMATAGE
# ============================================================

def normaliser(v):
    s = str(v or "").strip().lower()
    s = (
        s.replace("’", "'")
         .replace("–", "-")
         .replace("—", "-")
         .replace("‐", "-")
    )
    s = "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )
    s = re.sub(r"\s+", " ", s)
    return s.strip()


def formater(v):
    if v is None or v == "":
        return "-"

    if isinstance(v, str):
        t = v.strip()
        if t in ("-", "–", "...", "…"):
            return t
        try:
            n = float(t.replace(" ", "").replace(",", "."))
        except ValueError:
            return t
    else:
        try:
            n = float(v)
        except (TypeError, ValueError):
            return str(v)

    if abs(n - round(n)) < 1e-9:
        return f"{int(round(n)):,}".replace(",", " ")

    return f"{n:.1f}".replace(".", ",")


def remplacer_cellule(cellule, nouveau):
    nouveau = str(nouveau)

    if not cellule.paragraphs:
        cellule.text = nouveau
        return

    p = cellule.paragraphs[0]

    if p.runs:
        p.runs[0].text = nouveau
        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(nouveau)

    for p2 in cellule.paragraphs[1:]:
        for run in p2.runs:
            run.text = ""


def remplacer_annee(cellule, annee="2024"):
    ancien = cellule.text or ""

    nouveau = re.sub(
        r"\b(?:20\d{2}|2202)\b",
        str(annee),
        ancien
    )

    if nouveau != ancien:
        remplacer_cellule(cellule, nouveau)


def verifier_table(document, objet, lignes, colonnes):
    table = document.tables[objet - 1]

    if len(table.rows) != lignes or len(table.columns) != colonnes:
        raise RuntimeError(
            f"Objet Word {objet}: "
            f"{len(table.rows)}x{len(table.columns)} "
            f"au lieu de {lignes}x{colonnes}."
        )

    return table


# ============================================================
# RECHERCHE DYNAMIQUE DES TABLES EXCEL
# ============================================================

RE_NUMERO = re.compile(r"^\s*15\s*-\s*(\d+)\b", re.I)


def trouver_blocs(wb):
    """
    Repère tous les tableaux 15-x directement dans le classeur,
    sans dépendre des noms de feuilles du catalogue.
    Cela contourne les problèmes '6-7 absent', '17-17suite absent', etc.
    """
    blocs = {}

    for ws in wb.worksheets:
        ancres = []

        for r in range(1, ws.max_row + 1):
            trouve = None

            for c in range(1, min(ws.max_column, 12) + 1):
                valeur = ws.cell(r, c).value
                if not isinstance(valeur, str):
                    continue

                m = RE_NUMERO.match(valeur)
                if m:
                    trouve = int(m.group(1))
                    break

            if trouve is not None:
                ancres.append((r, trouve))

        for i, (r, numero) in enumerate(ancres):
            fin = (
                ancres[i + 1][0] - 1
                if i + 1 < len(ancres)
                else ws.max_row
            )

            texte = []

            for rr in range(r, min(fin, r + 12) + 1):
                for cc in range(1, min(ws.max_column, 12) + 1):
                    v = ws.cell(rr, cc).value
                    if v not in (None, ""):
                        texte.append(str(v))

            blocs[numero] = {
                "ws": ws,
                "debut": r,
                "fin": fin,
                "texte_entete": normaliser(" ".join(texte)),
            }

    return blocs


def bloc(blocs, numero):
    if numero not in blocs:
        raise RuntimeError(
            f"Table Excel 15-{numero} introuvable dans le classeur Justice."
        )
    return blocs[numero]


def trouver_ligne_dans_bloc(b, variantes, derniere=False):
    if isinstance(variantes, str):
        variantes = [variantes]

    cibles = {normaliser(x) for x in variantes}
    ws = b["ws"]

    plage = range(b["debut"], b["fin"] + 1)
    if derniere:
        plage = reversed(list(plage))

    for r in plage:
        val = normaliser(ws.cell(r, 1).value)
        if val in cibles:
            return r

    raise RuntimeError(
        f"Ligne {variantes!r} introuvable dans "
        f"{ws.title} / table 15-x."
    )


def trouver_ligne_total(b):
    return trouver_ligne_dans_bloc(
        b,
        ["Total", "Ensemble"],
        derniere=True
    )


def colonnes_tfm(b):
    """
    Détecte automatiquement Total / Féminin / Masculin dans les
    lignes d'en-tête du bloc.
    """
    ws = b["ws"]
    resultat = {}

    for r in range(
        b["debut"],
        min(b["fin"], b["debut"] + 15) + 1
    ):
        for c in range(1, min(ws.max_column, 15) + 1):
            t = normaliser(ws.cell(r, c).value)

            if t == "total" and "total" not in resultat:
                resultat["total"] = c
            elif t == "feminin" and "feminin" not in resultat:
                resultat["feminin"] = c
            elif t == "masculin" and "masculin" not in resultat:
                resultat["masculin"] = c

    if len(resultat) != 3:
        raise RuntimeError(
            f"Colonnes Total/Féminin/Masculin non détectées "
            f"dans {ws.title}."
        )

    return resultat


def texte_colonne_entete(b, col):
    ws = b["ws"]
    morceaux = []

    for r in range(
        b["debut"],
        min(b["fin"], b["debut"] + 15) + 1
    ):
        v = ws.cell(r, col).value
        if v not in (None, ""):
            morceaux.append(str(v))

    return normaliser(" ".join(morceaux))


def detecter_colonne(b, mots):
    if isinstance(mots, str):
        mots = [mots]

    ws = b["ws"]

    for c in range(2, min(ws.max_column, 12) + 1):
        t = texte_colonne_entete(b, c)

        if all(normaliser(m) in t for m in mots):
            return c

    raise RuntimeError(
        f"Colonne {mots!r} non détectée dans "
        f"{ws.title}."
    )


def classer_blocs_15_15_a_20(blocs):
    """
    15-15 à 15-20 sont les 6 blocs TPI :
      civil/pénal x enregistrées/jugées/en instance.
    """
    resultat = {}

    for numero in range(15, 21):
        if numero not in blocs:
            continue

        b = blocs[numero]
        t = b["texte_entete"]

        nature = None
        etat = None

        if "civil" in t or "civiles" in t:
            nature = "civil"
        elif "penal" in t or "penales" in t:
            nature = "penal"

        if "enregist" in t:
            etat = "enregistrees"
        elif "juge" in t:
            etat = "jugees"
        elif "instance" in t or "restant" in t or "restent" in t:
            etat = "instance"

        if nature and etat:
            resultat[(nature, etat)] = b

    return resultat


# ============================================================
# OUVERTURE
# ============================================================

fichier_entree = None

for candidat in CANDIDATS_ENTREE:
    if candidat.exists():
        fichier_entree = candidat
        break

if fichier_entree is None:
    docs = sorted(
        [
            f for f in DOSSIER_SORTIE.glob("*.docx")
            if not f.name.startswith("~$")
        ],
        key=lambda p: p.stat().st_mtime,
        reverse=True
    )

    if not docs:
        raise FileNotFoundError("Aucun Word d'entrée trouvé dans sorties.")

    fichier_entree = docs[0]

if not FICHIER_JUSTICE.exists():
    raise FileNotFoundError(
        f"Excel Justice introuvable : {FICHIER_JUSTICE}"
    )


print("=" * 92)
print("GENERATION DU CHAPITRE JUSTICE")
print("=" * 92)
print(f"Word d'entrée : {fichier_entree}")
print()

document = Document(fichier_entree)

wb = load_workbook(
    FICHIER_JUSTICE,
    data_only=True,
    read_only=False
)

blocs = trouver_blocs(wb)

print(
    "Tables Justice détectées : "
    + ", ".join(f"15-{n}" for n in sorted(blocs))
)


# ============================================================
# 8-1 — AUXILIAIRES DE JUSTICE
# ============================================================

print("[1/11] 8-1 Auxiliaires de justice", flush=True)

table = verifier_table(document, 241, 30, 6)

remplacer_annee(table.cell(0, 0))
remplacer_annee(table.cell(0, 5))

# Word : catégories commencent aux lignes 3,7,11,15,19,23,27
# puis Total/Féminin/Masculin sur les trois lignes suivantes.
#
# Chaque tuple :
# (bloc Excel, nom, ligne Word total, groupe spécial)
categories = [
    (2, "Interprètes", 3, "interpretes"),
    (2, "Experts judiciaires", 7, "experts"),
    (3, "Copistes", 11, None),
    (4, "Notaires", 15, None),
    (5, "Adouls", 19, None),
    (6, "Huissiers de justice", 23, None),
    (7, "Avocats titulaires", 27, None),
]

for numero, nom, ligne_word_1base, special in categories:
    b = bloc(blocs, numero)
    ws = b["ws"]

    ligne_total_excel = trouver_ligne_total(b)
    ligne_meknes = trouver_ligne_dans_bloc(
        b,
        ["Meknès", "Meknes"]
    )

    if numero == 2:
        # 15-2 contient deux groupes :
        # Interprètes B:D ; Experts E:G.
        if special == "interpretes":
            cols = {
                "total": 2,
                "feminin": 3,
                "masculin": 4,
            }
        else:
            cols = {
                "total": 5,
                "feminin": 6,
                "masculin": 7,
            }
    else:
        cols = colonnes_tfm(b)

    # indices Word 0-based :
    # colonne 2 = Nation, colonne 3 = Meknès
    lignes_word = {
        "total": ligne_word_1base - 1,
        "feminin": ligne_word_1base,
        "masculin": ligne_word_1base + 1,
    }

    for cle in ("total", "feminin", "masculin"):
        remplacer_cellule(
            table.cell(lignes_word[cle], 2),
            formater(
                ws.cell(
                    ligne_total_excel,
                    cols[cle]
                ).value
            )
        )

        remplacer_cellule(
            table.cell(lignes_word[cle], 3),
            formater(
                ws.cell(
                    ligne_meknes,
                    cols[cle]
                ).value
            )
        )


# ============================================================
# 8-2 / 8-3
# ============================================================
#
# Les tables nationales 15-10 et 15-12 ne contiennent pas
# de ligne "Meknès" dans le diagnostic. On ne remplace donc PAS
# des données locales par les chiffres de Fès ou d'une autre ville.
# Conservation du modèle = choix sûr.
# ============================================================

print("[2/11] 8-2 — conservé (pas de donnée Meknès dans 15-10)")
print("[3/11] 8-3 — conservé (pas de donnée Meknès dans 15-12)")


# ============================================================
# 8-4 — COUR D'APPEL MEKNES
# ============================================================

print("[4/11] 8-4 Cour d'appel Meknès", flush=True)

table = verifier_table(document, 249, 5, 4)
b = bloc(blocs, 14)
ws = b["ws"]

r = trouver_ligne_dans_bloc(
    b,
    ["Meknès", "Meknes"]
)

remplacer_annee(table.cell(0, 0))
remplacer_annee(table.cell(0, 3))

# Source 15-14 :
# B pénal AI, C pénal AJ, D pénal AE,
# E civil AI, F civil AJ, G civil AE.
# Word :
# R3 enregistrées, R4 jugées, R5 instance.
valeurs = {
    2: (4, 7),  # enregistrées : D, G
    3: (3, 6),  # jugées       : C, F
    4: (2, 5),  # instance     : B, E
}

for wr, (cp, cc) in valeurs.items():
    remplacer_cellule(
        table.cell(wr, 1),
        formater(ws.cell(r, cp).value)
    )
    remplacer_cellule(
        table.cell(wr, 2),
        formater(ws.cell(r, cc).value)
    )


# ============================================================
# 8-5 — CONSERVER MODELE
# ============================================================

print("[5/11] 8-5 — conservé (aucune source Excel correspondante)")


# ============================================================
# 8-6 / 8-7 — TPI MEKNES
# ============================================================

classification = classer_blocs_15_15_a_20(blocs)

attendus = [
    ("civil", "enregistrees"),
    ("civil", "jugees"),
    ("civil", "instance"),
    ("penal", "enregistrees"),
    ("penal", "jugees"),
    ("penal", "instance"),
]

manquants = [
    x for x in attendus
    if x not in classification
]

if manquants:
    raise RuntimeError(
        "Impossible de classer toutes les tables 15-15 à 15-20 : "
        + str(manquants)
    )


def remplir_tpi(
    objet_word,
    lignes_word,
    nature,
    categories_mots
):
    table = document.tables[objet_word - 1]

    # Année dans les deux côtés du premier rang.
    remplacer_annee(table.cell(0, 0))
    remplacer_annee(
        table.cell(0, len(table.columns) - 1)
    )

    blocs_etat = {
        "enregistrees": classification[(nature, "enregistrees")],
        "jugees": classification[(nature, "jugees")],
        "instance": classification[(nature, "instance")],
    }

    # lignes Word : premier bloc R.C.A. Meknès + bloc Préfecture Meknès
    zones = [
        ("R.C.A. Meknès", lignes_word[0]),
        ("Meknès", lignes_word[1]),
    ]

    for libelle_excel, ligne_depart in zones:
        for col_word, etat in zip(
            (1, 2, 3),
            ("enregistrees", "jugees", "instance")
        ):
            b = blocs_etat[etat]
            ws = b["ws"]

            r = trouver_ligne_dans_bloc(
                b,
                [
                    libelle_excel,
                    libelle_excel.replace("Meknès", "Meknes")
                ]
            )

            # Détecter les colonnes à partir des en-têtes.
            cols = {}

            for cle, mots in categories_mots.items():
                cols[cle] = detecter_colonne(b, mots)

            # Total puis catégories dans l'ordre du Word.
            ordre = ["total"] + [
                k for k in categories_mots
                if k != "total"
            ]

            for offset, cle in enumerate(ordre):
                remplacer_cellule(
                    table.cell(
                        ligne_depart + offset,
                        col_word
                    ),
                    formater(
                        ws.cell(
                            r,
                            cols[cle]
                        ).value
                    )
                )


print("[6/11] 8-6 TPI Meknès — affaires civiles", flush=True)

remplir_tpi(
    254,
    # ligne de Total de chaque zone en indices 0-based :
    # R2 et R10
    (1, 9),
    "civil",
    {
        "total": ["total"],
        "travail": ["contentieux", "travail"],
        "immobilier": ["affaires", "immobil"],
        "famille": ["justice", "famille"],
        "referees": ["affaires", "refere"],
        "accidents": ["accidents", "travail"],
    }
)


print("[7/11] 8-7 TPI Meknès — affaires pénales", flush=True)

remplir_tpi(
    257,
    # R2 et R9
    (1, 8),
    "penal",
    {
        "total": ["total"],
        "delits": ["delit"],
        "mineurs": ["mineur"],
        "contraventions": ["contravention"],
        "accidents": ["accident", "circulation"],
    }
)


# ============================================================
# 8-8 — JURIDICTIONS DE PROXIMITE AU SEIN DU TPI
# ============================================================

print("[8/11] 8-8 Juridictions de proximité — TPI Meknès", flush=True)

table = verifier_table(document, 259, 6, 5)
b = bloc(blocs, 21)
ws = b["ws"]

remplacer_annee(table.cell(0, 0))
remplacer_annee(table.cell(0, 4))

for libelle, ligne_penal, ligne_civil in [
    ("R.C.A. Meknès", 1, 2),
    ("Meknès", 4, 5),
]:
    r = trouver_ligne_dans_bloc(
        b,
        [
            libelle,
            libelle.replace("Meknès", "Meknes")
        ]
    )

    # structure standard :
    # B/C/D = Pénal AI/AJ/AE
    # E/F/G = Civil AI/AJ/AE
    for col_word, col_excel in zip((1, 2, 3), (2, 3, 4)):
        remplacer_cellule(
            table.cell(ligne_penal, col_word),
            formater(ws.cell(r, col_excel).value)
        )

    for col_word, col_excel in zip((1, 2, 3), (5, 6, 7)):
        remplacer_cellule(
            table.cell(ligne_civil, col_word),
            formater(ws.cell(r, col_excel).value)
        )


# ============================================================
# 8-9 / 8-10 — CENTRES JUDICIAIRES
# ============================================================

def remplir_centre(objet_word, numero_source):
    table = verifier_table(document, objet_word, 6, 5)
    b = bloc(blocs, numero_source)
    ws = b["ws"]

    remplacer_annee(table.cell(0, 0))
    remplacer_annee(table.cell(0, 4))

    for libelle, lp, lc in [
        ("Ensemble", 1, 2),
        ("Meknès", 4, 5),
    ]:
        r = trouver_ligne_dans_bloc(
            b,
            [
                libelle,
                libelle.replace("Meknès", "Meknes"),
                "Total" if libelle == "Ensemble" else libelle,
            ]
        )

        for cw, ce in zip((1, 2, 3), (2, 3, 4)):
            remplacer_cellule(
                table.cell(lp, cw),
                formater(ws.cell(r, ce).value)
            )

        for cw, ce in zip((1, 2, 3), (5, 6, 7)):
            remplacer_cellule(
                table.cell(lc, cw),
                formater(ws.cell(r, ce).value)
            )


print("[9/11] 8-9 Centres judiciaires", flush=True)
remplir_centre(262, 22)

print("[10/11] 8-10 Juridictions de proximité — Centres judiciaires", flush=True)
remplir_centre(265, 23)


# ============================================================
# 8-11 — MINEURS PLACES
# ============================================================

print("[11/11] 8-11 Mineurs placés", flush=True)

table = verifier_table(document, 267, 4, 9)
b = bloc(blocs, 33)
ws = b["ws"]

r_total = trouver_ligne_total(b)
r_fem = trouver_ligne_dans_bloc(
    b,
    ["Féminin", "Feminin"]
)
r_meknes = trouver_ligne_dans_bloc(
    b,
    ["Meknès", "Meknes"]
)

# Lignes physiques du tableau Word, plus sûres avec les cellules fusionnées.
rows_physiques = table._tbl.tr_lst

# Header : cellules physiques C2..C6 = années.
header_cells = [
    _Cell(tc, table)
    for tc in rows_physiques[0].tc_lst
]

# Nation / Féminin / Meknès
nation_cells = [
    _Cell(tc, table)
    for tc in rows_physiques[1].tc_lst
]
fem_cells = [
    _Cell(tc, table)
    for tc in rows_physiques[2].tc_lst
]
meknes_cells = [
    _Cell(tc, table)
    for tc in rows_physiques[3].tc_lst
]

annees = ["2024", "2023", "2022", "2021", "2020"]

# D'après la structure diagnostiquée :
# header : [vide, y1, y2, y3, y4, y5, ...]
for i, annee in enumerate(annees, start=1):
    remplacer_cellule(
        header_cells[i],
        annee
    )

# Source Excel C:G = 2024..2020
for i, col_excel in enumerate(range(3, 8), start=1):
    remplacer_cellule(
        nation_cells[i],
        formater(
            ws.cell(r_total, col_excel).value
        )
    )

    remplacer_cellule(
        fem_cells[i],
        formater(
            ws.cell(r_fem, col_excel).value
        )
    )

# Ligne Meknès : première cellule = Meknès,
# deuxième = type établissement, puis 5 valeurs.
for i, col_excel in enumerate(range(3, 8), start=2):
    remplacer_cellule(
        meknes_cells[i],
        formater(
            ws.cell(r_meknes, col_excel).value
        )
    )


# ============================================================
# SAUVEGARDE
# ============================================================

wb.close()

print()
print("Sauvegarde du Word...", flush=True)

document.save(FICHIER_SORTIE)

print()
print("=" * 92)
print("CHAPITRE JUSTICE GENERE")
print("=" * 92)
print("8-1  : actualisé 2024")
print("8-2  : conservé — aucune donnée Meknès dans la source 15-10")
print("8-3  : conservé — aucune donnée Meknès dans la source 15-12")
print("8-4  : actualisé 2024")
print("8-5  : conservé — aucune source correspondante")
print("8-6  : actualisé à partir des blocs civil enregistrées/jugées/instance")
print("8-7  : actualisé à partir des blocs pénal enregistrées/jugées/instance")
print("8-8  : actualisé 2024")
print("8-9  : actualisé 2024")
print("8-10 : actualisé 2024")
print("8-11 : années 2024 à 2020 actualisées")
print()
print("Aucune source Word n'a été modifiée.")
print(f"Document créé : {FICHIER_SORTIE}")
print("=" * 92)
