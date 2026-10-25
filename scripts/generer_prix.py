from pathlib import Path
from datetime import datetime, date
from openpyxl import load_workbook
from docx import Document
import re
import unicodedata


# ============================================================
# CHEMINS
# ============================================================

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_MODELES = Path("modeles")
DOSSIER_SORTIE = Path("sorties")

FICHIER_PRIX = DOSSIER_SOURCE / "17.Indice des prix_AS  2025.xlsx"

# On continue à partir du Word où Transport est déjà actualisé.
FICHIER_ENTREE_TRANSPORT = DOSSIER_SORTIE / "Annuaire_HCP_Transport.docx"

FICHIER_SORTIE = DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix.docx"


# ============================================================
# OUTILS
# ============================================================

def normaliser(texte):
    texte = str(texte or "").strip().lower()
    texte = texte.replace("’", "'")

    texte = "".join(
        c
        for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )

    texte = re.sub(r"\s+", " ", texte)

    return texte


def feuille(classeur, nom):
    cible = normaliser(nom)

    for ws in classeur.worksheets:
        if normaliser(ws.title) == cible:
            return ws

    raise KeyError(
        f"Feuille introuvable : {nom}"
    )


def remplacer_texte_preserver_style(cellule, nouveau_texte):
    """
    Remplace le texte sans recréer la cellule Word.
    Les bordures, couleurs, dimensions et alignements restent inchangés.
    """
    nouveau_texte = str(nouveau_texte)

    if not cellule.paragraphs:
        cellule.text = nouveau_texte
        return

    paragraphe = cellule.paragraphs[0]

    if paragraphe.runs:
        paragraphe.runs[0].text = nouveau_texte

        for run in paragraphe.runs[1:]:
            run.text = ""
    else:
        paragraphe.add_run(nouveau_texte)

    for autre_paragraphe in cellule.paragraphs[1:]:
        for run in autre_paragraphe.runs:
            run.text = ""


def formater_decimal(valeur, decimales_max=4, forcer_une_decimale=False):
    if valeur in (None, ""):
        return ""

    if isinstance(valeur, bool):
        return str(valeur)

    try:
        nombre = float(valeur)
    except (TypeError, ValueError):
        return str(valeur).strip()

    if forcer_une_decimale:
        texte = f"{nombre:.1f}"
    else:
        texte = f"{nombre:.{decimales_max}f}".rstrip("0").rstrip(".")

    return texte.replace(".", ",")


def formater_date_jjmmaaaa(valeur):
    """
    2024-01-01 -> 01/01/2024
    """
    if isinstance(valeur, (datetime, date)):
        return valeur.strftime("%d/%m/%Y")

    texte = str(valeur or "").strip()

    # yyyy-mm-dd ...
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", texte)
    if m:
        return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"

    # yyyy/mm/dd
    m = re.match(r"^(\d{4})/(\d{2})/(\d{2})$", texte)
    if m:
        return f"{m.group(3)}/{m.group(2)}/{m.group(1)}"

    return texte


def formater_date_point(valeur):
    """
    Format du tableau électricité :
    2017/01/01 -> 1.1.2017
    """
    if isinstance(valeur, (datetime, date)):
        return f"{valeur.day}.{valeur.month}.{valeur.year}"

    texte = str(valeur or "").strip()

    m = re.match(r"^(\d{4})[/-](\d{2})[/-](\d{2})", texte)

    if m:
        return (
            f"{int(m.group(3))}."
            f"{int(m.group(2))}."
            f"{m.group(1)}"
        )

    return texte


def formater_date_aaaammjj(valeur):
    """
    Format du tableau eau :
    2014-08-01 -> 2014/08/01
    """
    if isinstance(valeur, (datetime, date)):
        return valeur.strftime("%Y/%m/%d")

    texte = str(valeur or "").strip()

    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", texte)

    if m:
        return f"{m.group(1)}/{m.group(2)}/{m.group(3)}"

    return texte


# ============================================================
# OUVRIR LES SOURCES
# ============================================================

if not FICHIER_PRIX.exists():
    raise FileNotFoundError(
        f"Source PRIX introuvable : {FICHIER_PRIX}"
    )

classeur = load_workbook(
    FICHIER_PRIX,
    data_only=True,
    read_only=False
)

ws_ipc = feuille(classeur, "mek")
ws_tabac = feuille(classeur, "t 23")
ws_electricite = feuille(classeur, "t 25")
ws_eau = feuille(classeur, "t26")


# ============================================================
# OUVRIR LE WORD DE TRAVAIL
# ============================================================

if FICHIER_ENTREE_TRANSPORT.exists():
    fichier_entree = FICHIER_ENTREE_TRANSPORT
else:
    fichiers_modele = [
        f
        for f in DOSSIER_MODELES.glob("*.docx")
        if not f.name.startswith("~$")
    ]

    if not fichiers_modele:
        raise FileNotFoundError(
            "Aucun modèle Word trouvé."
        )

    fichier_entree = fichiers_modele[0]


document = Document(fichier_entree)


# ============================================================
# 15-1 — INDICE DES PRIX A LA CONSOMMATION, MEKNES
# ============================================================
#
# Le tableau Word 15-1 est physiquement divisé en DEUX objets :
#   - objet Word 386 : première partie
#   - objet Word 387 : suite
#
# Cette structure n'avait pas été visible dans le premier mapping.
# ============================================================

table_ipc_1 = document.tables[386 - 1]
table_ipc_2 = document.tables[387 - 1]

if len(table_ipc_1.rows) != 28 or len(table_ipc_1.columns) != 5:
    raise RuntimeError(
        "Structure Word inattendue pour la première partie du tableau 15-1."
    )

if len(table_ipc_2.rows) != 34 or len(table_ipc_2.columns) != 5:
    raise RuntimeError(
        "Structure Word inattendue pour la suite du tableau 15-1."
    )


# Années 2024 / 2023 / 2022 dans les DEUX parties.
annees_ipc = [
    str(ws_ipc.cell(6, colonne).value)
    for colonne in (2, 3, 4)
]

for table_ipc in (table_ipc_1, table_ipc_2):
    for ligne_word in (0, 1):
        for colonne_word, annee in enumerate(
            annees_ipc,
            start=1
        ):
            remplacer_texte_preserver_style(
                table_ipc.cell(
                    ligne_word,
                    colonne_word
                ),
                annee
            )


# Première partie : Word R3:R27 <- Excel 9:33
for index, ligne_excel in enumerate(
    range(9, 34),
    start=2
):
    # On conserve les libellés Word FR/AR, on actualise les 3 indices.
    for colonne_word, colonne_excel in zip(
        (1, 2, 3),
        (2, 3, 4)
    ):
        remplacer_texte_preserver_style(
            table_ipc_1.cell(
                index,
                colonne_word
            ),
            formater_decimal(
                ws_ipc.cell(
                    ligne_excel,
                    colonne_excel
                ).value,
                decimales_max=1,
                forcer_une_decimale=True
            )
        )


# Source à la fin de la première partie.
remplacer_texte_preserver_style(
    table_ipc_1.cell(27, 0),
    str(ws_ipc["A75"].value or "")
)

remplacer_texte_preserver_style(
    table_ipc_1.cell(27, 4),
    str(ws_ipc["E75"].value or "")
)


# Deuxième partie : correspondance Word -> lignes Excel.
# Les lignes 43+44 et 46+47 sont des libellés coupés sur deux lignes dans Excel ;
# leurs valeurs sont portées par 44 et 47.
lignes_excel_ipc_suite = [
    34,  # Transport
    35,  # Achat de véhicules
    36,  # Dépenses d'utilisation des véhicules
    37,  # Services de transport
    38,  # Communications
    39,  # Services postaux
    40,  # Matériel de téléphonie...
    41,  # Services de téléphonie...
    42,  # Loisirs et culture
    44,  # Matériel audiovisuel... + de l'information
    45,  # Autres biens durables...
    47,  # Autres articles... + animaux de compagnie
    48,  # Services récréatifs...
    49,  # Journaux...
    50,  # Forfaits touristiques
    51,  # Enseignement
    52,  # Enseignement préélémentaire...
    53,  # Enseignement secondaire
    54,  # Postsecondaire non supérieur
    55,  # Enseignement supérieur
    56,  # Enseignement non défini
    57,  # Restaurants et hôtels
    58,  # Services de restauration
    59,  # Services d'hébergement
    60,  # Biens et services divers
    61,  # Soins corporels
    62,  # Effets personnels...
    63,  # Assurance
    64,  # Services financiers
    65,  # Autres services
    66,  # Indice général
    67,  # Variation
]

if len(lignes_excel_ipc_suite) != 32:
    raise RuntimeError(
        "La recette IPC suite doit contenir exactement 32 lignes."
    )

for index, ligne_excel in enumerate(
    lignes_excel_ipc_suite,
    start=2
):
    for colonne_word, colonne_excel in zip(
        (1, 2, 3),
        (2, 3, 4)
    ):
        remplacer_texte_preserver_style(
            table_ipc_2.cell(
                index,
                colonne_word
            ),
            formater_decimal(
                ws_ipc.cell(
                    ligne_excel,
                    colonne_excel
                ).value,
                decimales_max=1,
                forcer_une_decimale=True
            )
        )


# ============================================================
# 15-2 — PRIX REGLEMENTES DU TABAC
# ============================================================

table_tabac = document.tables[390 - 1]

if len(table_tabac.rows) != 35 or len(table_tabac.columns) != 7:
    raise RuntimeError(
        "Structure Word inattendue pour 15-2."
    )


# Entête : 5 dates actualisées.
for colonne_word, colonne_excel in zip(
    range(1, 6),
    range(2, 7)
):
    remplacer_texte_preserver_style(
        table_tabac.cell(
            0,
            colonne_word
        ),
        formater_date_jjmmaaaa(
            ws_tabac.cell(
                6,
                colonne_excel
            ).value
        )
    )


# Les 34 lignes qui structurent le tableau actuel.
lignes_tabac = [
    10, 11, 12,
    14, 15, 16, 17, 18, 19, 20,
    22, 23, 24, 25, 26, 27,
    29, 30, 31, 32,
    34, 35, 36, 37, 38, 39, 40, 41,
    43, 44, 45, 46, 47, 48
]

if len(lignes_tabac) != 34:
    raise RuntimeError(
        "La recette tabac doit contenir 34 lignes."
    )


for ligne_word, ligne_excel in enumerate(
    lignes_tabac,
    start=1
):
    # Français
    remplacer_texte_preserver_style(
        table_tabac.cell(
            ligne_word,
            0
        ),
        str(
            ws_tabac.cell(
                ligne_excel,
                1
            ).value or ""
        )
    )

    # Prix
    for colonne_word, colonne_excel in zip(
        range(1, 6),
        range(2, 7)
    ):
        valeur = ws_tabac.cell(
            ligne_excel,
            colonne_excel
        ).value

        remplacer_texte_preserver_style(
            table_tabac.cell(
                ligne_word,
                colonne_word
            ),
            (
                ""
                if valeur in (None, "")
                else formater_decimal(
                    valeur,
                    decimales_max=1,
                    forcer_une_decimale=True
                )
            )
        )

    # Arabe
    remplacer_texte_preserver_style(
        table_tabac.cell(
            ligne_word,
            6
        ),
        str(
            ws_tabac.cell(
                ligne_excel,
                7
            ).value or ""
        )
    )


# ============================================================
# 15-3 — TARIFS DE L'ELECTRICITE
# ============================================================

table_electricite = document.tables[393 - 1]

if len(table_electricite.rows) != 40 or len(table_electricite.columns) != 7:
    raise RuntimeError(
        "Structure Word inattendue pour 15-3."
    )


# Dates actuelles : 2017 / 2016 / 2015 / 2014 / 2009
dates_electricite = [
    formater_date_point(
        ws_electricite.cell(
            6,
            colonne
        ).value
    )
    for colonne in range(2, 7)
]

for ligne_word in (0, 1):
    for colonne_word, valeur in enumerate(
        dates_electricite,
        start=1
    ):
        remplacer_texte_preserver_style(
            table_electricite.cell(
                ligne_word,
                colonne_word
            ),
            valeur
        )


# Mapping des lignes Word vers les lignes Excel actuelles.
# Les anciennes tranches 101-200 sont conservées visuellement dans Word ;
# on prend la première des deux tranches actuelles (101-150), les tarifs
# étant identiques sur la coupure concernée.
mapping_electricite = {
    3: 10,   # Usage domestique
    4: 11,
    5: 12,   # 101-200 Word <- 101-150 source
    6: 14,
    7: 15,
    8: 16,

    9: 18,   # Eclairage privé
    10: 19,
    11: 20,  # 101-200 Word <- 101-150 source
    12: 22,
    13: 23,
    14: 24,

    15: 26,  # Eclairage patenté
    16: 27,
    17: 28,

    18: 30,  # administratif
    19: 32,  # public

    20: 34,  # Force motrice
    21: 35,
    22: 36,
    23: 37,

    24: 40,  # Moyenne tension
    25: 41,
    26: 42,
    27: 43,
    28: 44,
    29: 45,
    30: 46,

    31: 49,  # Haute / très haute tension
    32: 50,
    33: 51,
    34: 52,
    35: 53,
    36: 54,
    37: 55,
}


for ligne_word_1based, ligne_excel in mapping_electricite.items():
    ligne_word = ligne_word_1based - 1

    # Valeurs numériques uniquement : les libellés Word restent en place
    # pour conserver les regroupements historiques du modèle.
    for colonne_word, colonne_excel in zip(
        range(1, 6),
        range(2, 7)
    ):
        valeur = ws_electricite.cell(
            ligne_excel,
            colonne_excel
        ).value

        remplacer_texte_preserver_style(
            table_electricite.cell(
                ligne_word,
                colonne_word
            ),
            (
                ""
                if valeur in (None, "")
                else formater_decimal(
                    valeur,
                    decimales_max=4,
                    forcer_une_decimale=False
                )
            )
        )


# Source condensée et exacte.
remplacer_texte_preserver_style(
    table_electricite.cell(39, 0),
    str(ws_electricite["A66"].value or "")
)

remplacer_texte_preserver_style(
    table_electricite.cell(39, 6),
    str(ws_electricite["G66"].value or "")
)


# ============================================================
# 15-4 — PRIX DE L'EAU POTABLE : MEKNES
# ============================================================

table_eau = document.tables[396 - 1]

if len(table_eau.rows) != 8 or len(table_eau.columns) != 7:
    raise RuntimeError(
        "Structure Word inattendue pour 15-4."
    )


# Entête dates.
for colonne_word, colonne_excel in zip(
    range(1, 6),
    range(2, 7)
):
    remplacer_texte_preserver_style(
        table_eau.cell(
            0,
            colonne_word
        ),
        formater_date_aaaammjj(
            ws_eau.cell(
                79,
                colonne_excel
            ).value
        )
    )


# Meknès : lignes 101 à 107
# Word R2:R8 <- Excel 101:107.
for ligne_word, ligne_excel in enumerate(
    range(101, 108),
    start=1
):
    # Français
    remplacer_texte_preserver_style(
        table_eau.cell(
            ligne_word,
            0
        ),
        str(
            ws_eau.cell(
                ligne_excel,
                1
            ).value or ""
        )
    )

    # Valeurs
    for colonne_word, colonne_excel in zip(
        range(1, 6),
        range(2, 7)
    ):
        valeur = ws_eau.cell(
            ligne_excel,
            colonne_excel
        ).value

        remplacer_texte_preserver_style(
            table_eau.cell(
                ligne_word,
                colonne_word
            ),
            (
                ""
                if valeur in (None, "")
                else formater_decimal(
                    valeur,
                    decimales_max=3,
                    forcer_une_decimale=False
                )
            )
        )

    # Arabe
    remplacer_texte_preserver_style(
        table_eau.cell(
            ligne_word,
            6
        ),
        str(
            ws_eau.cell(
                ligne_excel,
                7
            ).value or ""
        )
    )


# ============================================================
# SAUVEGARDE
# ============================================================

classeur.close()

DOSSIER_SORTIE.mkdir(
    exist_ok=True
)

document.save(
    FICHIER_SORTIE
)


# ============================================================
# RESULTAT
# ============================================================

print("=" * 80)
print("CHAPITRE PRIX GENERE")
print("=" * 80)

print()
print("15-1 IPC Meknès")
print(
    "Années : "
    + " / ".join(annees_ipc)
)
print(
    "Deux parties Word actualisées : objets 386 + 387"
)

print()
print("15-2 Tabac")
print(
    "Dates : "
    + " | ".join(
        formater_date_jjmmaaaa(
            ws_tabac.cell(6, c).value
        )
        for c in range(2, 7)
    )
)

print()
print("15-3 Electricité")
print(
    "Dates : "
    + " | ".join(
        dates_electricite
    )
)

print()
print("15-4 Eau potable — Meknès")
print(
    "Données Meknès : lignes Excel 101 à 107"
)

print()
print(
    f"Document créé : {FICHIER_SORTIE}"
)

print()
print(
    f"Document d'entrée : {fichier_entree}"
)

print(
    "Le modèle original n'a pas été modifié."
)

print("=" * 80)
