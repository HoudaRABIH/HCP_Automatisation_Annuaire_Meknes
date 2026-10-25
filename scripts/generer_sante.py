from pathlib import Path
from openpyxl import load_workbook
from docx import Document
import re
import unicodedata


# ============================================================
# CHEMINS
# ============================================================

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_SORTIE = Path("sorties")

FICHIER_SANTE = DOSSIER_SOURCE / "12. Santé_AS 2025.xlsx"

# On repart du Word cumulé validé après Agriculture.
CANDIDATS_ENTREE = [
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_VALIDE.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_FINAL.docx",
    DOSSIER_SORTIE / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_OK.docx",
]

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme_Agriculture_Sante.docx"
)


# ============================================================
# OUTILS GENERAUX
# ============================================================

def normaliser(texte):
    texte = str(texte or "").strip().lower()
    texte = (
        texte
        .replace("’", "'")
        .replace("–", "-")
        .replace("—", "-")
        .replace("‐", "-")
    )

    texte = "".join(
        c
        for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )

    texte = re.sub(r"\s+", " ", texte)
    texte = re.sub(r"\s*-\s*", " - ", texte)

    return texte.strip()


def remplacer_texte_preserver_style(cellule, nouveau_texte):
    """
    Remplace uniquement le texte visible dans la cellule et conserve
    bordures, trames, tailles, alignements et le style du premier run.
    """
    nouveau_texte = str(nouveau_texte)

    if not cellule.paragraphs:
        cellule.text = nouveau_texte
        return

    p = cellule.paragraphs[0]

    if p.runs:
        p.runs[0].text = nouveau_texte

        for run in p.runs[1:]:
            run.text = ""
    else:
        p.add_run(nouveau_texte)

    for p2 in cellule.paragraphs[1:]:
        for run in p2.runs:
            run.text = ""


def remplacer_annee_preserver_texte(cellule, annee):
    """
    Remplace seulement l'année 20xx dans le texte déjà présent.
    Utile pour conserver l'ordre RTL arabe tel qu'il existe dans le modèle.
    """
    texte = cellule.text or ""

    nouveau = re.sub(
        r"20\d{2}",
        str(annee),
        texte,
        count=1
    )

    if nouveau != texte:
        remplacer_texte_preserver_style(
            cellule,
            nouveau
        )


def formater_nombre(valeur):
    if valeur is None or valeur == "":
        return "-"

    if isinstance(valeur, str):
        texte = valeur.strip()

        if texte in ("-", "–", "…", "..."):
            return texte

        try:
            nombre = float(
                texte
                .replace(" ", "")
                .replace(",", ".")
            )
        except ValueError:
            return texte
    else:
        try:
            nombre = float(valeur)
        except (TypeError, ValueError):
            return str(valeur)

    # Tous les indicateurs Santé concernés sont des effectifs entiers.
    if abs(nombre - round(nombre)) < 1e-9:
        return f"{int(round(nombre)):,}".replace(",", " ")

    texte = f"{nombre:.1f}".rstrip("0").rstrip(".")
    return texte.replace(".", ",")


def verifier_table(document, no_objet, lignes, colonnes):
    table = document.tables[
        no_objet - 1
    ]

    if (
        len(table.rows) != lignes
        or len(table.columns) != colonnes
    ):
        raise RuntimeError(
            f"Objet Word {no_objet}: structure "
            f"{len(table.rows)}x{len(table.columns)} "
            f"au lieu de {lignes}x{colonnes}."
        )

    return table


# ============================================================
# OUTILS EXCEL
# ============================================================

def trouver_ligne(ws, libelle):
    cible = normaliser(libelle)

    for r in range(
        1,
        ws.max_row + 1
    ):
        if normaliser(
            ws.cell(r, 1).value
        ) == cible:
            return r

    raise ValueError(
        f"{libelle!r} introuvable dans la feuille {ws.title!r}"
    )


def trouver_ligne_nation(ws):
    """
    Les feuilles Santé utilisent selon les cas 'Total' ou 'Ensemble'.
    On cherche depuis le bas pour récupérer le total national final.
    """
    cibles = {
        "total",
        "ensemble",
    }

    for r in range(
        ws.max_row,
        0,
        -1
    ):
        if normaliser(
            ws.cell(r, 1).value
        ) in cibles:
            return r

    raise ValueError(
        f"Total/Ensemble introuvable dans {ws.title!r}"
    )


def lignes_geo(ws):
    return {
        "nation": trouver_ligne_nation(ws),
        "region": trouver_ligne(ws, "Fès - Meknès"),
        "meknes": trouver_ligne(ws, "Meknès"),
    }


def annee_source(ws):
    """
    Repère la première mention 'Année 20xx' dans les 20 premières lignes.
    """
    for r in range(
        1,
        min(ws.max_row, 20) + 1
    ):
        for c in range(
            1,
            min(ws.max_column, 10) + 1
        ):
            valeur = ws.cell(r, c).value

            if valeur is None:
                continue

            m = re.search(
                r"\b(20\d{2})\b",
                str(valeur)
            )

            if m:
                return m.group(1)

    raise ValueError(
        f"Année source introuvable dans {ws.title!r}"
    )


def ecrire_geo(
    table,
    ws,
    word_rows,
    source_cols,
    word_cols=None
):
    """
    word_rows :
        {"nation": 2, "region": 3, "meknes": 4}
        indices Python (0-based)

    source_cols :
        [2,3,4] = colonnes Excel B,C,D

    word_cols :
        colonnes Word 0-based ; défaut 1..N
    """
    geo = lignes_geo(ws)

    if word_cols is None:
        word_cols = list(
            range(
                1,
                1 + len(source_cols)
            )
        )

    if len(word_cols) != len(source_cols):
        raise ValueError(
            "word_cols et source_cols doivent avoir la même longueur."
        )

    for cle_geo, ligne_word in word_rows.items():
        ligne_excel = geo[cle_geo]

        for col_word, col_excel in zip(
            word_cols,
            source_cols
        ):
            remplacer_texte_preserver_style(
                table.cell(
                    ligne_word,
                    col_word
                ),
                formater_nombre(
                    ws.cell(
                        ligne_excel,
                        col_excel
                    ).value
                )
            )


def actualiser_annee_table(table, ws):
    annee = annee_source(ws)

    remplacer_annee_preserver_texte(
        table.cell(0, 0),
        annee
    )

    remplacer_annee_preserver_texte(
        table.cell(
            0,
            len(table.columns) - 1
        ),
        annee
    )

    return annee


# ============================================================
# SOURCES ANNUAIRE 2023 -> 2025
# ============================================================

def changer_sources_annuaire_2023_vers_2025(document):
    """
    Règle validée avec l'utilisateur :
    on touche UNIQUEMENT à l'année des mentions
    'Annuaire Statistique du Maroc 2023'
    / équivalent arabe.
    Les autres sources (Ministères, HCP, RGPH...) restent intactes.
    """
    marqueurs = (
        "Annuaire Statistique du Maroc",
        "النشرة الإحصائية السنوية للمغرب",
    )

    def traiter_paragraphe(p):
        texte = p.text or ""

        if (
            "2023" not in texte
            or not any(
                marqueur in texte
                for marqueur in marqueurs
            )
        ):
            return

        for run in p.runs:
            if "2023" in run.text:
                run.text = run.text.replace(
                    "2023",
                    "2025"
                )

    cellules_vues = set()

    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                identifiant = id(
                    cell._tc
                )

                if identifiant in cellules_vues:
                    continue

                cellules_vues.add(
                    identifiant
                )

                for p in cell.paragraphs:
                    traiter_paragraphe(p)

    for p in document.paragraphs:
        traiter_paragraphe(p)

    for section in document.sections:
        for p in section.header.paragraphs:
            traiter_paragraphe(p)

        for p in section.footer.paragraphs:
            traiter_paragraphe(p)


# ============================================================
# OUVERTURE
# ============================================================

fichier_entree = None

for candidat in CANDIDATS_ENTREE:
    if candidat.exists():
        fichier_entree = candidat
        break

if fichier_entree is None:
    raise FileNotFoundError(
        "Aucun Word Agriculture VALIDE/FINAL/OK trouvé dans sorties."
    )

if not FICHIER_SANTE.exists():
    raise FileNotFoundError(
        f"Excel Santé introuvable : {FICHIER_SANTE}"
    )


print("=" * 88)
print("GENERATION DU CHAPITRE SANTE")
print("=" * 88)
print(f"Word d'entrée : {fichier_entree}")
print()

print("Ouverture du Word...", flush=True)

document = Document(
    fichier_entree
)

print("Ouverture de l'Excel Santé...", flush=True)

wb = load_workbook(
    FICHIER_SANTE,
    data_only=True,
    read_only=False
)


# ============================================================
# 2-1 — ETABLISSEMENTS DE SOINS PRIMAIRES PUBLICS
# Objet Word 54 — 8x7
# Source : 12-1 / feuille 1
# ============================================================

print("[1/23] 2-1", flush=True)

ws = wb["1"]

table = verifier_table(
    document,
    54,
    8,
    7
)

actualiser_annee_table(
    table,
    ws
)

geo = lignes_geo(ws)

# Word :
# R4 CSR-2, R5 CSR-1, R6 DR, R7 CSU-2, R8 CSU-1
# Colonnes géographiques Word :
# C3 Meknès, C4 Région, C5 Nation
categories = [
    (3, 5),  # CSR-2 <- Excel E
    (4, 4),  # CSR-1 <- Excel D
    (5, 6),  # DR    <- Excel F
    (6, 3),  # CSU-2 <- Excel C
    (7, 2),  # CSU-1 <- Excel B
]

geo_word_cols = {
    "meknes": 2,
    "region": 3,
    "nation": 4,
}

for ligne_word, col_excel in categories:
    for cle_geo, col_word in geo_word_cols.items():
        remplacer_texte_preserver_style(
            table.cell(
                ligne_word,
                col_word
            ),
            formater_nombre(
                ws.cell(
                    geo[cle_geo],
                    col_excel
                ).value
            )
        )


# ============================================================
# 2-2 — HOPITAUX PUBLICS ET NOMBRE DE LITS
# Objet 56 — composite 12-2 + 12-3
# ============================================================

print("[2/23] 2-2", flush=True)

ws_hop = wb["2"]
ws_lits = wb["3"]

table = verifier_table(
    document,
    56,
    5,
    4
)

actualiser_annee_table(
    table,
    ws_hop
)

geo_hop = lignes_geo(ws_hop)
geo_lits = lignes_geo(ws_lits)

word_rows = {
    "nation": 2,
    "region": 3,
    "meknes": 4,
}

for cle_geo, ligne_word in word_rows.items():
    remplacer_texte_preserver_style(
        table.cell(ligne_word, 1),
        formater_nombre(
            ws_hop.cell(
                geo_hop[cle_geo],
                2
            ).value
        )
    )

    remplacer_texte_preserver_style(
        table.cell(ligne_word, 2),
        formater_nombre(
            ws_lits.cell(
                geo_lits[cle_geo],
                2
            ).value
        )
    )


# ============================================================
# 2-3 — CLINIQUES PRIVEES
# ============================================================

print("[3/23] 2-3", flush=True)

ws = wb["10"]

table = verifier_table(
    document,
    59,
    5,
    3
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 2, "region": 3, "meknes": 4},
    [2]
)


# ============================================================
# 2-4 — MEDECINS PAR SECTEUR
# Word : Privé / Public
# Excel : Public / Privé
# ============================================================

print("[4/23] 2-4", flush=True)

ws = wb["5"]

table = verifier_table(
    document,
    61,
    5,
    4
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 2, "region": 3, "meknes": 4},
    [3, 2],       # Excel Privé, Public
    [1, 2]        # Word C2, C3
)


# ============================================================
# 2-5 — MEDECINS DU MINISTERE
# Word : Total / spécialistes / généralistes
# Excel : généralistes / spécialistes / Total
# ============================================================

print("[5/23] 2-5", flush=True)

ws = wb["6"]

table = verifier_table(
    document,
    64,
    5,
    5
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 2, "region": 3, "meknes": 4},
    [4, 3, 2]
)


# ============================================================
# 2-6 — MEDECINS DU MINISTERE DANS LE RESSP
# ============================================================

print("[6/23] 2-6", flush=True)

ws = wb["7"]

table = verifier_table(
    document,
    66,
    5,
    5
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 2, "region": 3, "meknes": 4},
    [4, 3, 2]
)


# ============================================================
# 2-7 — MEDECINS PRIVES
# ============================================================

print("[7/23] 2-7", flush=True)

ws = wb["8"]

table = verifier_table(
    document,
    69,
    5,
    5
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 2, "region": 3, "meknes": 4},
    [4, 3, 2]
)


# ============================================================
# 2-8 — DENTISTES PUBLICS + CABINETS PRIVES
# composite 12-9 + 12-11
# ============================================================

print("[8/23] 2-8", flush=True)

ws_public = wb["9"]
ws_prive = wb["11"]

table = verifier_table(
    document,
    71,
    5,
    4
)

actualiser_annee_table(
    table,
    ws_public
)

geo_public = lignes_geo(
    ws_public
)

geo_prive = lignes_geo(
    ws_prive
)

for cle_geo, ligne_word in {
    "nation": 2,
    "region": 3,
    "meknes": 4,
}.items():

    # C2 = cabinet privé
    remplacer_texte_preserver_style(
        table.cell(
            ligne_word,
            1
        ),
        formater_nombre(
            ws_prive.cell(
                geo_prive[cle_geo],
                2
            ).value
        )
    )

    # C3 = chirurgiens dentistes publics
    remplacer_texte_preserver_style(
        table.cell(
            ligne_word,
            2
        ),
        formater_nombre(
            ws_public.cell(
                geo_public[cle_geo],
                2
            ).value
        )
    )


# ============================================================
# 2-9 — PERSONNEL PARAMEDICAL PUBLIC
# ============================================================

print("[9/23] 2-9", flush=True)

ws = wb["14"]

table = verifier_table(
    document,
    73,
    5,
    6
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 2, "region": 3, "meknes": 4},
    [2, 3, 4, 5]
)


# ============================================================
# 2-10 — PERSONNEL PARAMEDICAL DANS LE RESSP
# ============================================================

print("[10/23] 2-10", flush=True)

ws = wb["15"]

table = verifier_table(
    document,
    75,
    5,
    6
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 2, "region": 3, "meknes": 4},
    [2, 3, 4, 5]
)


# ============================================================
# 2-11 — CONSULTATIONS CURATIVES
# ============================================================

print("[11/23] 2-11", flush=True)

ws = wb["16"]

table = verifier_table(
    document,
    78,
    6,
    5
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 3, "region": 4, "meknes": 5},
    [2, 3, 4]
)


# ============================================================
# 2-12 — CONSULTATIONS PRENATALES
# ============================================================

print("[12/23] 2-12", flush=True)

ws = wb["17"]

table = verifier_table(
    document,
    80,
    6,
    5
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 3, "region": 4, "meknes": 5},
    [2, 3, 4]
)


# ============================================================
# 2-13 — ACCOUCHEMENTS
# ============================================================

print("[13/23] 2-13", flush=True)

ws = wb["18"]

table = verifier_table(
    document,
    82,
    5,
    7
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 2, "region": 3, "meknes": 4},
    [2, 3, 4, 5, 6]
)


# ============================================================
# 2-14 — PLANIFICATION FAMILIALE : NOUVELLES ACCEPTANTES
# Le modèle Word ne reprend que Injections/DIU/Condom/Pilules.
# ============================================================

print("[14/23] 2-14", flush=True)

ws = wb["19"]

table = verifier_table(
    document,
    85,
    5,
    6
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 2, "region": 3, "meknes": 4},
    [2, 3, 4, 5]
)


# ============================================================
# 2-15 — PLANIFICATION FAMILIALE : ANCIENNES ACCEPTANTES
# ============================================================

print("[15/23] 2-15", flush=True)

ws = wb["20"]

table = verifier_table(
    document,
    87,
    4,
    6
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 1, "region": 2, "meknes": 3},
    [2, 3, 4, 5]
)


# ============================================================
# 2-16 — PNI
# ============================================================

print("[16/23] 2-16", flush=True)

ws = wb["21"]

table = verifier_table(
    document,
    90,
    5,
    8
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 2, "region": 3, "meknes": 4},
    [2, 3, 4, 5, 6, 7]
)


# ============================================================
# 2-17 — PNLMD
# ============================================================

print("[17/23] 2-17", flush=True)

ws = wb["22"]

table = verifier_table(
    document,
    92,
    7,
    7
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 4, "region": 5, "meknes": 6},
    [2, 3, 4, 5, 6]
)


# ============================================================
# 2-18 — MOUVEMENTS DES MALADES
# ============================================================

print("[18/23] 2-18", flush=True)

ws = wb["23"]

table = verifier_table(
    document,
    94,
    6,
    6
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 3, "region": 4, "meknes": 5},
    [2, 3, 4, 5]
)


# ============================================================
# 2-19
# ============================================================
#
# Il n'existe pas dans la liste de recettes à générer.
# On ne touche donc PAS au tableau correspondant.
# ============================================================


# ============================================================
# 2-20 — CONSULTATIONS MEDICALES DANS LES HOPITAUX PUBLICS
# ============================================================

print("[19/23] 2-20", flush=True)

ws = wb["25"]

table = verifier_table(
    document,
    98,
    6,
    5
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 3, "region": 4, "meknes": 5},
    [2, 3, 4]
)


# ============================================================
# 2-21 — ACTIVITES DES LABORATOIRES DES HOPITAUX
# ============================================================

print("[20/23] 2-21", flush=True)

ws = wb["26"]

table = verifier_table(
    document,
    100,
    6,
    9
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 3, "region": 4, "meknes": 5},
    [2, 3, 4, 5, 6, 7, 8]
)


# ============================================================
# 2-22 — MALADIES SOUS SURVEILLANCE
# ============================================================

print("[21/23] 2-22", flush=True)

ws = wb["27"]

table = verifier_table(
    document,
    103,
    5,
    10
)

actualiser_annee_table(
    table,
    ws
)

ecrire_geo(
    table,
    ws,
    {"nation": 2, "region": 3, "meknes": 4},
    [2, 3, 4, 5, 6, 7, 8, 9]
)


# ============================================================
# 2-23 — LABORATOIRES D'ANALYSES + OFFICINES
# composite 12-13 + 12-12
# ============================================================

print("[22/23] 2-23", flush=True)

ws_pharm = wb["12"]
ws_labo = wb["13"]

table = verifier_table(
    document,
    105,
    6,
    4
)

actualiser_annee_table(
    table,
    ws_labo
)

geo_pharm = lignes_geo(
    ws_pharm
)

geo_labo = lignes_geo(
    ws_labo
)

for cle_geo, ligne_word in {
    "nation": 2,
    "region": 3,
    "meknes": 4,
}.items():

    # C2 = laboratoires d'analyses
    remplacer_texte_preserver_style(
        table.cell(
            ligne_word,
            1
        ),
        formater_nombre(
            ws_labo.cell(
                geo_labo[cle_geo],
                2
            ).value
        )
    )

    # C3 = officines de pharmacie
    remplacer_texte_preserver_style(
        table.cell(
            ligne_word,
            2
        ),
        formater_nombre(
            ws_pharm.cell(
                geo_pharm[cle_geo],
                2
            ).value
        )
    )


# ============================================================
# 2-24 — HANDICAP
# ============================================================

print("[23/23] 2-24 — conservé tel quel", flush=True)

table_2_24 = verifier_table(
    document,
    108,
    37,
    5
)

# Aucune source correspondante dans les 23 Excel :
# le modèle est déjà en 2024, donc on ne modifie rien.


# ============================================================
# SOURCES
# ============================================================

changer_sources_annuaire_2023_vers_2025(
    document
)


# ============================================================
# SAUVEGARDE
# ============================================================

wb.close()

print()
print("Sauvegarde du Word...", flush=True)

document.save(
    FICHIER_SORTIE
)


print()
print("=" * 88)
print("CHAPITRE SANTE GENERE")
print("=" * 88)
print("2-1 à 2-18 : actualisés")
print("2-20 à 2-23 : actualisés")
print("2-24 : conservé tel quel (pas de source Excel correspondante)")
print()
print(
    "Les années 2024/2023 sont reprises directement de chaque feuille Excel."
)
print(
    "Aucune année 2024 n'est inventée pour les tableaux dont la source reste 2023."
)
print()
print(
    "Sources : uniquement les mentions "
    "'Annuaire Statistique du Maroc 2023' passent à 2025."
)
print(
    "Les autres sources du Word (Ministère, HCP, etc.) restent intactes."
)
print()
print(f"Document créé : {FICHIER_SORTIE}")
print("=" * 88)
