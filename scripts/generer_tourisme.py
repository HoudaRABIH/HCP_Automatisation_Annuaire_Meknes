from pathlib import Path
from openpyxl import load_workbook
from docx import Document
import re
import unicodedata


# ============================================================
# CHEMINS
# ============================================================

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_SORTIE = Path("sorties")

FICHIER_TOURISME = DOSSIER_SOURCE / "10. Tourisme_AS 2025.xlsx"

# On continue depuis le Word cumulé validé avec les sources 2025.
FICHIER_ENTREE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Sources2025_OK.docx"
)

FICHIER_SORTIE = (
    DOSSIER_SORTIE
    / "Annuaire_HCP_Transport_Prix_Environnement_Tourisme.docx"
)


# ============================================================
# OUTILS
# ============================================================

def normaliser(texte):
    texte = str(texte or "").strip().lower()
    texte = texte.replace("’", "'")

    texte = "".join(
        c
        for c in unicodedata.normalize("NFD", texte)
        if unicodedata.category(c) != "Mn"
    )

    texte = re.sub(r"\s+", " ", texte)
    return texte


def remplacer_texte_preserver_style(cellule, nouveau_texte):
    nouveau_texte = str(nouveau_texte)

    if not cellule.paragraphs:
        cellule.text = nouveau_texte
        return

    paragraphe = cellule.paragraphs[0]

    if paragraphe.runs:
        paragraphe.runs[0].text = nouveau_texte

        for run in paragraphe.runs[1:]:
            run.text = ""
    else:
        paragraphe.add_run(nouveau_texte)

    for autre in cellule.paragraphs[1:]:
        for run in autre.runs:
            run.text = ""


def formater_nombre(valeur, zero_tiret=False):
    if valeur in (None, ""):
        return "-"

    if isinstance(valeur, str):
        texte = valeur.strip()

        if texte in ("-", "–"):
            return "-"

        if texte.startswith("#"):
            return texte

        try:
            nombre = float(
                texte.replace(
                    " ",
                    ""
                ).replace(
                    ",",
                    "."
                )
            )
        except ValueError:
            return texte
    else:
        nombre = float(valeur)

    if zero_tiret and abs(nombre) < 1e-12:
        return "-"

    if abs(nombre - round(nombre)) < 1e-9:
        return f"{int(round(nombre)):,}".replace(",", " ")

    # Cas exceptionnel d'une valeur non entière.
    texte = f"{nombre:.1f}".rstrip("0").rstrip(".")
    return texte.replace(".", ",")


def trouver_ligne(ws, libelle, colonne=1):
    cible = normaliser(libelle)

    for r in range(1, ws.max_row + 1):
        if normaliser(ws.cell(r, colonne).value) == cible:
            return r

    raise ValueError(
        f"'{libelle}' introuvable dans {ws.title}"
    )


def trouver_dernier_total(ws):
    candidats = []

    for r in range(1, ws.max_row + 1):
        val = normaliser(
            ws.cell(r, 1).value
        )

        if val == "total":
            candidats.append(r)

    if not candidats:
        raise ValueError(
            f"Total introuvable dans {ws.title}"
        )

    return candidats[-1]


def valeur(ws, ligne, colonne):
    return ws.cell(
        ligne,
        colonne
    ).value


def remplacer_source_2023_par_2025_dans_tableaux(document, indices_objets):
    """
    Règle utilisateur :
    - on ne touche PAS aux sources RGPH, HCP, ministères, etc.
    - on remplace seulement 2023 -> 2025 dans une source contenant
      'Annuaire Statistique du Maroc' ou sa version arabe.
    - on conserve les runs pour préserver le formatage.
    """
    marqueurs = (
        "Annuaire Statistique du Maroc",
        "النشرة الإحصائية السنوية للمغرب",
    )

    cellules_vues = set()

    for no_objet in indices_objets:
        debut = max(
            0,
            no_objet - 1
        )

        fin = min(
            len(document.tables),
            no_objet + 2
        )

        for idx in range(
            debut,
            fin
        ):
            table = document.tables[idx]

            for ligne in table.rows:
                for cellule in ligne.cells:
                    identifiant = id(
                        cellule._tc
                    )

                    if identifiant in cellules_vues:
                        continue

                    cellules_vues.add(
                        identifiant
                    )

                    for paragraphe in cellule.paragraphs:
                        texte = paragraphe.text or ""

                        if (
                            "2023" not in texte
                            or not any(
                                marqueur in texte
                                for marqueur in marqueurs
                            )
                        ):
                            continue

                        for run in paragraphe.runs:
                            if "2023" in run.text:
                                run.text = run.text.replace(
                                    "2023",
                                    "2025"
                                )


# ============================================================
# OUVERTURE DES FICHIERS
# ============================================================

if not FICHIER_ENTREE.exists():
    raise FileNotFoundError(
        f"Word d'entrée introuvable : {FICHIER_ENTREE}"
    )

if not FICHIER_TOURISME.exists():
    raise FileNotFoundError(
        f"Excel Tourisme introuvable : {FICHIER_TOURISME}"
    )

wb = load_workbook(
    FICHIER_TOURISME,
    data_only=True,
    read_only=False
)

document = Document(
    FICHIER_ENTREE
)


# ============================================================
# 12-1 — CAPACITE DES ETABLISSEMENTS CLASSES
# Objet Word 316
# Excel 10-1
# ============================================================

ws = wb["10-1"]

table = document.tables[
    316 - 1
]

if (
    len(table.rows) != 5
    or len(table.columns) != 5
):
    raise RuntimeError(
        "Structure Word inattendue pour 12-1."
    )


ligne_region = trouver_ligne(
    ws,
    "Fès - Meknès"
)

ligne_meknes = trouver_ligne(
    ws,
    "Meknès"
)

ligne_nation = trouver_dernier_total(
    ws
)


# Date
date_fr = str(
    ws["A7"].value or ""
)

date_ar = str(
    ws["E7"].value or ""
)

remplacer_texte_preserver_style(
    table.cell(0, 0),
    date_fr
)

remplacer_texte_preserver_style(
    table.cell(0, 4),
    date_ar
)


# Lits / chambres / établissements
for ligne_word, ligne_excel in (
    (2, ligne_nation),
    (3, ligne_region),
    (4, ligne_meknes),
):
    for col_word, col_excel in zip(
        (1, 2, 3),
        (2, 3, 4)
    ):
        remplacer_texte_preserver_style(
            table.cell(
                ligne_word,
                col_word
            ),
            formater_nombre(
                valeur(
                    ws,
                    ligne_excel,
                    col_excel
                )
            )
        )


# ============================================================
# 12-2 — LES ETABLISSEMENTS CLASSES
# Objet Word 319
# Excel 10-2 + 10-2suite
# ============================================================

ws_a = wb["10-2"]
ws_b = wb["10-2suite"]

table = document.tables[
    319 - 1
]

if (
    len(table.rows) != 16
    or len(table.columns) != 5
):
    raise RuntimeError(
        "Structure Word inattendue pour 12-2."
    )


# Date
remplacer_texte_preserver_style(
    table.cell(0, 0),
    str(
        ws_a["A6"].value or ""
    )
)

remplacer_texte_preserver_style(
    table.cell(0, 4),
    str(
        ws_a["G6"].value or ""
    )
)


# Lignes géographiques
geo_a = {
    "nation": trouver_dernier_total(ws_a),
    "region": trouver_ligne(ws_a, "Fès - Meknès"),
    "meknes": trouver_ligne(ws_a, "Meknès"),
}

geo_b = {
    "nation": trouver_dernier_total(ws_b),
    "region": trouver_ligne(ws_b, "Fès - Meknès"),
    "meknes": trouver_ligne(ws_b, "Meknès"),
}


# Colonnes Word : Nation / Région / Meknès
word_geo = {
    "nation": 1,
    "region": 2,
    "meknes": 3,
}


# 5*,4*,3*,2*,1* : source 10-2 B:F
for i, col_excel in enumerate(
    range(2, 7),
    start=2
):
    for cle in (
        "nation",
        "region",
        "meknes"
    ):
        remplacer_texte_preserver_style(
            table.cell(
                i,
                word_geo[cle]
            ),
            formater_nombre(
                valeur(
                    ws_a,
                    geo_a[cle],
                    col_excel
                ),
                zero_tiret=True
            )
        )


# Reste des catégories depuis 10-2suite.
# Word :
# R8 VVT -> Hôtel-club (I)
# R9 Hôtel de luxe -> Luxe (J)
# R10 Résidences -> H
# R11 Pensions -> G
# R12 Maisons -> F
# R13 Campings -> E
# R14 Auberges -> D
# R15 Autres -> C
# R16 Total -> B
mapping_suite_12_2 = [
    (7, 9),   # VVT / Hôtel-club
    (8, 10),  # Luxe
    (9, 8),   # Résidences
    (10, 7),  # Pensions
    (11, 6),  # Maisons d'hôtes
    (12, 5),  # Campings
    (13, 4),  # Auberges
    (14, 3),  # Autres
    (15, 2),  # Total
]

for ligne_word, col_excel in mapping_suite_12_2:
    for cle in (
        "nation",
        "region",
        "meknes"
    ):
        remplacer_texte_preserver_style(
            table.cell(
                ligne_word,
                word_geo[cle]
            ),
            formater_nombre(
                valeur(
                    ws_b,
                    geo_b[cle],
                    col_excel
                ),
                zero_tiret=(
                    ligne_word != 15
                )
            )
        )


# ============================================================
# 12-3 — CAPACITE EN LITS
# Objet Word 322
# Excel 10-3 + 10-3(suite)
# ============================================================

ws_a = wb["10-3"]
ws_b = wb["10-3(suite)"]

table = document.tables[
    322 - 1
]

if (
    len(table.rows) != 16
    or len(table.columns) != 5
):
    raise RuntimeError(
        "Structure Word inattendue pour 12-3."
    )


remplacer_texte_preserver_style(
    table.cell(0, 0),
    str(
        ws_a["A6"].value or ""
    )
)

remplacer_texte_preserver_style(
    table.cell(0, 4),
    str(
        ws_a["G6"].value or ""
    )
)


geo_a = {
    "nation": trouver_dernier_total(ws_a),
    "region": trouver_ligne(ws_a, "Fès - Meknès"),
    "meknes": trouver_ligne(ws_a, "Meknès"),
}

geo_b = {
    "nation": trouver_dernier_total(ws_b),
    "region": trouver_ligne(ws_b, "Fès - Meknès"),
    "meknes": trouver_ligne(ws_b, "Meknès"),
}


# 5*,4*,3*,2*,1*
for i, col_excel in enumerate(
    range(2, 7),
    start=2
):
    for cle in (
        "nation",
        "region",
        "meknes"
    ):
        remplacer_texte_preserver_style(
            table.cell(
                i,
                word_geo[cle]
            ),
            formater_nombre(
                valeur(
                    ws_a,
                    geo_a[cle],
                    col_excel
                ),
                zero_tiret=True
            )
        )


# Word :
# R8 Hôtel de luxe -> J
# R9 Hôtel-club -> I
# R10 Résidences -> H
# R11 Pensions -> G
# R12 Maisons -> F
# R13 Campings -> E
# R14 Auberges -> D
# R15 Autres -> C
# R16 Total -> B
mapping_suite_12_3 = [
    (7, 10),  # Luxe
    (8, 9),   # Hôtel-club
    (9, 8),   # Résidences
    (10, 7),  # Pensions
    (11, 6),  # Maisons
    (12, 5),  # Campings
    (13, 4),  # Auberges
    (14, 3),  # Autres
    (15, 2),  # Total
]

for ligne_word, col_excel in mapping_suite_12_3:
    for cle in (
        "nation",
        "region",
        "meknes"
    ):
        remplacer_texte_preserver_style(
            table.cell(
                ligne_word,
                word_geo[cle]
            ),
            formater_nombre(
                valeur(
                    ws_b,
                    geo_b[cle],
                    col_excel
                ),
                zero_tiret=(
                    ligne_word != 15
                )
            )
        )


# ============================================================
# 12-4 — ARRIVEES PAR TYPE DE TOURISME
# Objet Word 324
# Excel 10-4
# ============================================================

ws = wb["10-4"]

table = document.tables[
    324 - 1
]

if (
    len(table.rows) != 6
    or len(table.columns) != 5
):
    raise RuntimeError(
        "Structure Word inattendue pour 12-4."
    )


# Année
annee = str(
    ws["A5"].value or ""
).replace(
    "Année ",
    ""
).strip()

remplacer_texte_preserver_style(
    table.cell(0, 0),
    f"Année {annee}"
)

remplacer_texte_preserver_style(
    table.cell(0, 4),
    f"سنة {annee}"
)


geo = {
    "nation": trouver_dernier_total(ws),
    "region": trouver_ligne(ws, "Fès - Meknès"),
    "meknes": trouver_ligne(ws, "Meknès"),
}

word_rows = {
    "nation": 3,
    "region": 4,
    "meknes": 5,
}


# Word : Total / national / international
# Excel : B Total / C international / D national
for cle in (
    "nation",
    "region",
    "meknes"
):
    ligne_excel = geo[cle]
    ligne_word = word_rows[cle]

    valeurs = (
        valeur(ws, ligne_excel, 2),
        valeur(ws, ligne_excel, 4),
        valeur(ws, ligne_excel, 3),
    )

    for col_word, val in zip(
        (1, 2, 3),
        valeurs
    ):
        remplacer_texte_preserver_style(
            table.cell(
                ligne_word,
                col_word
            ),
            formater_nombre(
                val
            )
        )


# ============================================================
# 12-6 — NUITEES TOURISTIQUES PAR CATEGORIE
# Objet Word 329
# Excel 10-8 + 10-8(suite)
# ============================================================

ws_a = wb["10-8"]
ws_b = wb["10-8(suite)"]

table = document.tables[
    329 - 1
]

if (
    len(table.rows) != 16
    or len(table.columns) != 5
):
    raise RuntimeError(
        "Structure Word inattendue pour 12-6."
    )


# Année
annee = str(
    ws_a["A7"].value or ""
).replace(
    "Année ",
    ""
).strip()

remplacer_texte_preserver_style(
    table.cell(0, 0),
    f"Année {annee}"
)

remplacer_texte_preserver_style(
    table.cell(0, 4),
    f"سنة {annee}"
)


geo_a = {
    "nation": trouver_dernier_total(ws_a),
    "region": trouver_ligne(ws_a, "Fès - Meknès"),
    "meknes": trouver_ligne(ws_a, "Meknès"),
}

geo_b = {
    "nation": trouver_dernier_total(ws_b),
    "region": trouver_ligne(ws_b, "Fès - Meknès"),
    "meknes": trouver_ligne(ws_b, "Meknès"),
}


# 5*,4*,3*,2*,1*
for i, col_excel in enumerate(
    range(2, 7),
    start=2
):
    for cle in (
        "nation",
        "region",
        "meknes"
    ):
        remplacer_texte_preserver_style(
            table.cell(
                i,
                word_geo[cle]
            ),
            formater_nombre(
                valeur(
                    ws_a,
                    geo_a[cle],
                    col_excel
                ),
                zero_tiret=True
            )
        )


# Hôtel-club national : la cellule source D94 contient #VALUE!.
# On ne laisse pas l'erreur Excel dans le Word.
# On dérive la valeur à partir du TOTAL officiel moins toutes les autres
# catégories officielles du même tableau.
def hotel_club_national_12_6():
    total_global = float(
        valeur(
            ws_b,
            geo_b["nation"],
            2
        )
    )

    cinq_categories = sum(
        float(
            valeur(
                ws_a,
                geo_a["nation"],
                c
            )
        )
        for c in range(
            2,
            7
        )
    )

    # C, E, F, G, H, I, J = toutes les catégories de la suite
    # sauf D (Hôtel-club).
    autres_categories = sum(
        float(
            valeur(
                ws_b,
                geo_b["nation"],
                c
            )
        )
        for c in (
            3,
            5,
            6,
            7,
            8,
            9,
            10
        )
    )

    return (
        total_global
        - cinq_categories
        - autres_categories
    )


hotel_club_nation = (
    hotel_club_national_12_6()
)


# Word :
# R8 Hôtel de luxe -> E
# R9 VVT -> Hôtel-club D
# R10 Résidences -> G
# R11 Pensions -> J
# R12 Maisons -> F
# R13 Campings -> I
# R14 Auberges -> H
# R15 Autres -> C
# R16 Total -> B
mapping_suite_12_6 = [
    (7, 5),   # Luxe
    (8, 4),   # Hôtel-club
    (9, 7),   # Résidences
    (10, 10), # Pensions
    (11, 6),  # Maisons
    (12, 9),  # Campings
    (13, 8),  # Auberges
    (14, 3),  # Autres
    (15, 2),  # Total
]

for ligne_word, col_excel in mapping_suite_12_6:
    for cle in (
        "nation",
        "region",
        "meknes"
    ):
        if (
            ligne_word == 8
            and cle == "nation"
        ):
            val = hotel_club_nation
        else:
            val = valeur(
                ws_b,
                geo_b[cle],
                col_excel
            )

        remplacer_texte_preserver_style(
            table.cell(
                ligne_word,
                word_geo[cle]
            ),
            formater_nombre(
                val,
                zero_tiret=(
                    ligne_word != 15
                )
            )
        )


# ============================================================
# SOURCES ANNUAIRE 2023 -> 2025
# ============================================================

remplacer_source_2023_par_2025_dans_tableaux(
    document,
    [
        316,
        319,
        322,
        324,
        329,
    ]
)


# ============================================================
# SAUVEGARDE
# ============================================================

wb.close()

document.save(
    FICHIER_SORTIE
)


# ============================================================
# RESULTAT
# ============================================================

print("=" * 84)
print("CHAPITRE TOURISME GENERE")
print("=" * 84)

print()
print("12-1 Capacité des établissements classés : 31-12-2024")
print("12-2 Etablissements classés : 31-12-2024")
print("12-3 Capacité en lits : 31-12-2024")
print("12-4 Arrivées par type de tourisme : 2024")
print("12-6 Nuitées touristiques : 2024")

print()
print(
    "12-6 Hôtel-club Nation : "
    + formater_nombre(
        hotel_club_nation
    )
    + " (calculé à partir des totaux du même tableau car Excel contient #VALUE!)"
)

print()
print(
    f"Document créé : {FICHIER_SORTIE}"
)

print(
    f"Document d'entrée : {FICHIER_ENTREE}"
)

print()
print(
    "Sources : seules les mentions 'Annuaire Statistique du Maroc 2023'"
)
print(
    "ont été passées à 2025. Les autres sources du Word n'ont pas été touchées."
)

print("=" * 84)
