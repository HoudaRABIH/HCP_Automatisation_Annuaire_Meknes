from __future__ import annotations

from pathlib import Path
import os
import queue
import runpy
import shutil
import re
import subprocess
import sys
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from docx import Document


def _force_utf8_stdio():
    """Force UTF-8 before GUI/backend/worker dispatch in frozen Windows EXE."""
    for stream in (getattr(sys, "stdout", None), getattr(sys, "stderr", None)):
        if stream is None:
            continue
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass


_force_utf8_stdio()

def _is_frozen():
    return bool(getattr(sys, "frozen", False) and hasattr(sys, "_MEIPASS"))


def _bundle_root():
    if _is_frozen():
        return Path(sys._MEIPASS)
    return Path(__file__).resolve().parent.parent


def _data_root():
    if _is_frozen():
        return Path.home() / "Documents" / "Annuaire_HCP_2025"
    return Path(".")


ROOT = _data_root()
BACKEND = ROOT / "scripts" / "generer_annuaire_tous_secteurs_v5.py"
FINAL = ROOT / "sorties" / "Annuaire_HCP_2025_FINAL.docx"


def _resource(relative_path):
    return _bundle_root() / relative_path


def bootstrap_runtime():
    for folder in ("sorties", "references", "modeles", "donnees_brutes", "scripts"):
        (ROOT / folder).mkdir(parents=True, exist_ok=True)

    if not _is_frozen():
        return

    assets = [
        (
            _resource(Path("assets") / "mapping_final_recettes.xlsx"),
            ROOT / "sorties" / "mapping_final_recettes.xlsx",
        ),
        (
            _resource(Path("assets") / "REFERENCE_EDUCATION_2025.docx"),
            ROOT / "references" / "REFERENCE_EDUCATION_2025.docx",
        ),
    ]

    for src, dst in assets:
        if not src.exists():
            raise FileNotFoundError(f"Ressource intégrée absente : {src}")
        if not dst.exists():
            shutil.copy2(src, dst)


def backend_available():
    if _is_frozen():
        return _resource(
            Path("scripts") / "generer_annuaire_tous_secteurs_exe.py"
        ).exists()
    return BACKEND.exists()


def hide_console_window():
    if os.name != "nt" or not _is_frozen():
        return

    try:
        import ctypes
        hwnd = ctypes.windll.kernel32.GetConsoleWindow()
        if hwnd:
            ctypes.windll.user32.ShowWindow(hwnd, 0)
    except Exception:
        pass


def dispatch_frozen_mode():
    if not _is_frozen() or len(sys.argv) < 2:
        return False

    mode = sys.argv[1]

    if mode == "--backend":
        backend_script = _resource(
            Path("scripts") / "generer_annuaire_tous_secteurs_exe.py"
        )
        sys.argv = [str(backend_script)] + sys.argv[2:]
        runpy.run_path(str(backend_script), run_name="__main__")
        return True

    if mode == "--worker":
        if len(sys.argv) < 3:
            raise SystemExit("Nom du worker manquant.")

        script_name = Path(sys.argv[2]).name
        worker_script = _resource(Path("scripts") / script_name)

        if not worker_script.exists():
            raise FileNotFoundError(
                f"Worker intégré introuvable : {script_name}"
            )

        sys.argv = [str(worker_script)]
        runpy.run_path(str(worker_script), run_name="__main__")
        return True

    return False


BG = "#F5F6F8"
WHITE = "#FFFFFF"
TEXT = "#1F2937"
MUTED = "#6B7280"
BORDER = "#D9DEE7"
PRIMARY = "#9E163E"
PRIMARY_DARK = "#7D102F"
SUCCESS = "#16803F"
ERROR = "#B42318"


def count_excel(folder: Path | None) -> int:
    if not folder or not folder.exists() or not folder.is_dir():
        return 0
    return len([
        p for p in folder.glob("*.xlsx")
        if not p.name.startswith("~$")
    ])


def validate_word_model(path: Path):
    """
    Validation stricte du modèle correspondant au mapping actuel.
    """
    if not path.exists() or path.suffix.lower() != ".docx":
        return False, "Fichier Word invalide"

    try:
        doc = Document(path)
    except Exception as exc:
        return False, f"Impossible d'ouvrir le Word : {exc}"

    if len(doc.tables) != 414:
        return False, f"Mauvais modèle : {len(doc.tables)} tableaux détectés, 414 attendus"

    t141 = doc.tables[379 - 1]
    t142 = doc.tables[381 - 1]

    if len(t141.rows) != 8 or len(t141.columns) != 6:
        return False, (
            "Mauvais modèle : le tableau 14-1 n'a pas la structure 8×6 attendue"
        )

    if len(t142.rows) != 6 or len(t142.columns) != 4:
        return False, (
            "Mauvais modèle : le tableau 14-2 n'a pas la structure 6×4 attendue"
        )

    text141 = " ".join(
        c.text for row in t141.rows for c in row.cells
    ).lower()

    text142 = " ".join(
        c.text for row in t142.rows for c in row.cells
    ).lower()

    for keyword in ("nation", "région", "meknès"):
        if keyword not in text141:
            return False, f"Tableau 14-1 incompatible : '{keyword}' absent"

    for keyword in ("nation", "région", "meknès"):
        if keyword not in text142:
            return False, f"Tableau 14-2 incompatible : '{keyword}' absent"

    return True, "Word : PRÊT • structure 414 tableaux validée"


def open_path(path: Path):
    if not path.exists():
        raise FileNotFoundError(str(path))

    if os.name == "nt":
        os.startfile(str(path.resolve()))
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(path.resolve())])
    else:
        subprocess.Popen(["xdg-open", str(path.resolve())])


class App(tk.Tk):
    def __init__(self):
        super().__init__()

        self.title("Générateur Annuaire Statistique de Meknès 2025 - HCP")
        self.geometry("860x540")
        self.minsize(780, 500)
        self.configure(bg=BG)

        self.word = tk.StringVar()
        self.excel = tk.StringVar()
        self.status = tk.StringVar(value="En attente • 0/24")
        self.q = queue.Queue()
        self.proc = None
        self.backend_output = []

        self.word_ok = False
        self.excel_ok = False

        self._styles()
        self._ui()
        self.after(100, self._poll)

    def _styles(self):
        s = ttk.Style(self)
        try:
            s.theme_use("clam")
        except Exception:
            pass

        s.configure(
            "Primary.TButton",
            font=("Segoe UI", 11, "bold"),
            padding=(22, 12),
            foreground="white",
            background=PRIMARY,
            borderwidth=0,
        )
        s.map(
            "Primary.TButton",
            background=[
                ("active", PRIMARY_DARK),
                ("disabled", "#C8AEB7"),
            ],
        )

    def _ui(self):
        tk.Label(
            self,
            text="Générateur Annuaire Statistique de Meknès 2025 - HCP",
            bg=BG,
            fg=TEXT,
            font=("Segoe UI", 22, "bold"),
        ).pack(anchor="w", padx=28, pady=(28, 4))

        tk.Label(
            self,
            text="Générez l’Annuaire Statistique de Meknès 2025 à partir du modèle Word et des sources Excel.",
            bg=BG,
            fg=MUTED,
            font=("Segoe UI", 10),
        ).pack(anchor="w", padx=28, pady=(0, 20))

        card = tk.Frame(
            self,
            bg=WHITE,
            highlightbackground=BORDER,
            highlightthickness=1,
        )
        card.pack(fill="x", padx=28)
        card.columnconfigure(1, weight=1)

        # Word
        tk.Label(
            card,
            text="1. Word initial",
            bg=WHITE,
            fg=TEXT,
            font=("Segoe UI", 11, "bold"),
        ).grid(
            row=0,
            column=0,
            columnspan=3,
            sticky="w",
            padx=20,
            pady=(20, 8),
        )

        ttk.Entry(
            card,
            textvariable=self.word,
        ).grid(
            row=1,
            column=0,
            columnspan=2,
            sticky="ew",
            padx=(20, 10),
        )

        ttk.Button(
            card,
            text="Choisir le Word",
            command=self._choose_word,
        ).grid(
            row=1,
            column=2,
            padx=(0, 20),
        )

        self.word_state = tk.Label(
            card,
            text="Word : NON SÉLECTIONNÉ",
            bg=WHITE,
            fg=ERROR,
            font=("Segoe UI", 9, "bold"),
        )
        self.word_state.grid(
            row=2,
            column=0,
            columnspan=3,
            sticky="w",
            padx=20,
            pady=(7, 18),
        )

        ttk.Separator(card).grid(
            row=3,
            column=0,
            columnspan=3,
            sticky="ew",
            padx=20,
        )

        # Excel
        tk.Label(
            card,
            text="2. Dossier des fichiers Excel",
            bg=WHITE,
            fg=TEXT,
            font=("Segoe UI", 11, "bold"),
        ).grid(
            row=4,
            column=0,
            columnspan=3,
            sticky="w",
            padx=20,
            pady=(18, 8),
        )

        ttk.Entry(
            card,
            textvariable=self.excel,
        ).grid(
            row=5,
            column=0,
            columnspan=2,
            sticky="ew",
            padx=(20, 10),
        )

        ttk.Button(
            card,
            text="Choisir le dossier",
            command=self._choose_excel,
        ).grid(
            row=5,
            column=2,
            padx=(0, 20),
        )

        self.excel_state = tk.Label(
            card,
            text="Excel : NON SÉLECTIONNÉ",
            bg=WHITE,
            fg=ERROR,
            font=("Segoe UI", 9, "bold"),
        )
        self.excel_state.grid(
            row=6,
            column=0,
            columnspan=3,
            sticky="w",
            padx=20,
            pady=(7, 20),
        )

        self.run = ttk.Button(
            self,
            text="GÉNÉRER L'ANNUAIRE",
            style="Primary.TButton",
            command=self._start,
            state="disabled",
        )
        self.run.pack(pady=(20, 14))

        tk.Label(
            self,
            textvariable=self.status,
            bg=BG,
            fg=TEXT,
            font=("Segoe UI", 9, "bold"),
        ).pack(anchor="w", padx=28)

        self.bar = ttk.Progressbar(
            self,
            maximum=100,
        )
        self.bar.pack(fill="x", padx=28, pady=(7, 12))

        self.open_btn = ttk.Button(
            self,
            text="Ouvrir le Word final",
            command=self._open_final,
            state="disabled",
        )
        self.open_btn.pack()

    def _choose_word(self):
        p = filedialog.askopenfilename(
            title="Choisir le Word initial",
            filetypes=[("Word", "*.docx")],
        )
        if not p:
            return

        self.word.set(p)
        path = Path(p)

        self.word_state.configure(
            text="Word : VÉRIFICATION...",
            fg=MUTED,
        )
        self.update_idletasks()

        ok, msg = validate_word_model(path)
        self.word_ok = ok

        self.word_state.configure(
            text=msg,
            fg=SUCCESS if ok else ERROR,
        )

        if not ok:
            messagebox.showerror(
                "Modèle Word incompatible",
                msg,
            )

        self._refresh_button()

    def _choose_excel(self):
        p = filedialog.askdirectory(
            title="Choisir le dossier des 23 fichiers Excel",
        )
        if not p:
            return

        self.excel.set(p)

        n = count_excel(Path(p))
        self.excel_ok = n == 23

        self.excel_state.configure(
            text=(
                "Excel : PRÊT • 23/23"
                if self.excel_ok
                else f"Excel : {n}/23"
            ),
            fg=SUCCESS if self.excel_ok else ERROR,
        )

        if not self.excel_ok:
            messagebox.showerror(
                "Dossier Excel incomplet",
                f"23 fichiers .xlsx attendus.\nDétectés : {n}",
            )

        self._refresh_button()

    def _refresh_button(self):
        ready = (
            self.word_ok
            and self.excel_ok
            and backend_available()
            and self.proc is None
        )

        self.run.configure(
            state="normal" if ready else "disabled"
        )

        self.open_btn.configure(
            state="normal" if FINAL.exists() else "disabled"
        )

    def _start(self):
        # Double contrôle juste avant génération.
        ok, msg = validate_word_model(Path(self.word.get()))
        if not ok:
            self.word_ok = False
            self.word_state.configure(text=msg, fg=ERROR)
            self._refresh_button()
            messagebox.showerror("Word incompatible", msg)
            return

        n = count_excel(Path(self.excel.get()))
        if n != 23:
            self.excel_ok = False
            self.excel_state.configure(
                text=f"Excel : {n}/23",
                fg=ERROR,
            )
            self._refresh_button()
            messagebox.showerror(
                "Dossier Excel incomplet",
                f"23 fichiers attendus.\nDétectés : {n}",
            )
            return

        if not backend_available():
            messagebox.showerror(
                "Backend introuvable",
                "Le backend intégré n'est pas disponible.",
            )
            return

        self.run.configure(state="disabled")
        self.open_btn.configure(state="disabled")
        self.backend_output = []

        self.status.set("Démarrage... • 0/24")
        self.bar["value"] = 1

        threading.Thread(
            target=self._worker,
            daemon=True,
        ).start()

    def _worker(self):
        try:
            env = os.environ.copy()
            env["PYTHONIOENCODING"] = "utf-8"
            env["PYTHONUTF8"] = "1"

            flags = (
                getattr(subprocess, "CREATE_NO_WINDOW", 0)
                if os.name == "nt"
                else 0
            )

            self.proc = subprocess.Popen(
                (
                    [
                        sys.executable,
                        "--backend",
                        "--word",
                        self.word.get(),
                        "--excel-dir",
                        self.excel.get(),
                    ]
                    if _is_frozen()
                    else
                    [
                        sys.executable,
                        "-u",
                        str(BACKEND),
                        "--word",
                        self.word.get(),
                        "--excel-dir",
                        self.excel.get(),
                    ]
                ),
                cwd=str(ROOT.resolve()),
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                encoding="utf-8",
                errors="replace",
                env=env,
                creationflags=flags,
            )

            assert self.proc.stdout is not None

            for line in self.proc.stdout:
                self.q.put(("line", line))

            self.q.put(("done", self.proc.wait()))

        except Exception as exc:
            self.q.put(("error", str(exc)))

        finally:
            self.proc = None

    def _poll(self):
        try:
            while True:
                kind, payload = self.q.get_nowait()

                if kind == "line":
                    self.backend_output.append(str(payload).rstrip())
                    if len(self.backend_output) > 80:
                        self.backend_output = self.backend_output[-80:]

                    # V4 affiche 23 vraies sources Excel + 1 validation finale.
                    # Les anciens marqueurs techniques [x/16] sont volontairement ignorés.
                    m = re.search(
                        r"@@PROGRESS24@@\s+(\d+)/24\s+\|\s+(.+)",
                        str(payload),
                    )

                    if m:
                        step = int(m.group(1))
                        name = m.group(2).strip()
                        pct = round(step / 24 * 100)

                        self.bar["value"] = pct
                        self.status.set(
                            f"Étape {step}/24 — {name}"
                        )

                elif kind == "done":
                    code = int(payload)

                    if code == 0 and FINAL.exists():
                        self.bar["value"] = 100
                        self.status.set("Terminé avec succès")
                        self.open_btn.configure(state="normal")

                        messagebox.showinfo(
                            "Succès",
                            "L'annuaire a été généré avec succès.",
                        )
                    else:
                        self.status.set("Échec")
                        details = "\n".join(self.backend_output[-12:]).strip()
                        messagebox.showerror(
                            "Échec",
                            "La génération a échoué."
                            + ("\n\nDerniers détails :\n" + details if details else ""),
                        )

                    self._refresh_button()

                elif kind == "error":
                    self.status.set("Erreur")
                    messagebox.showerror(
                        "Erreur",
                        str(payload),
                    )
                    self._refresh_button()

        except queue.Empty:
            pass

        self.after(100, self._poll)

    def _open_final(self):
        try:
            open_path(FINAL)
        except Exception as exc:
            messagebox.showerror(
                "Ouverture impossible",
                str(exc),
            )


if __name__ == "__main__":
    if dispatch_frozen_mode():
        raise SystemExit(0)

    bootstrap_runtime()
    hide_console_window()
    App().mainloop()
