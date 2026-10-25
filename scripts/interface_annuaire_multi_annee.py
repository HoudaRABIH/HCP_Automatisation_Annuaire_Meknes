from __future__ import annotations

from pathlib import Path
import os
import queue
import re
import subprocess
import sys
import threading
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from docx import Document

ROOT = Path('.')
BACKEND = ROOT / 'scripts' / 'generer_annuaire_multi_annee.py'

BG = '#F5F6F8'
WHITE = '#FFFFFF'
TEXT = '#1F2937'
MUTED = '#6B7280'
BORDER = '#D9DEE7'
PRIMARY = '#9E163E'
PRIMARY_DARK = '#7D102F'
SUCCESS = '#16803F'
ERROR = '#B42318'


def excel_files(folder: Path | None):
    if not folder or not folder.exists() or not folder.is_dir():
        return []
    return sorted([
        p for p in folder.glob('*.xlsx')
        if not p.name.startswith('~$')
    ])


def count_excel(folder: Path | None) -> int:
    return len(excel_files(folder))


def detect_edition_year(folder: Path | None):
    files = excel_files(folder)
    if len(files) != 23:
        return None, f'{len(files)}/23 fichiers Excel'

    prefixes = {}
    years_by_file = {}

    for p in files:
        m = re.match(r'\s*(\d+)', p.name)
        if not m:
            return None, f'Numéro 1..23 absent : {p.name}'
        n = int(m.group(1))
        if not 1 <= n <= 23:
            return None, f'Numéro de source invalide : {p.name}'
        if n in prefixes:
            return None, f'Deux fichiers portent le numéro {n}'
        prefixes[n] = p.name

        years = set(re.findall(r'20\d{2}', p.name))
        if not years:
            return None, f'Année absente du nom : {p.name}'
        if len(years) != 1:
            return None, f'Plusieurs années dans : {p.name}'
        years_by_file[p.name] = next(iter(years))

    missing = sorted(set(range(1, 24)) - set(prefixes))
    if missing:
        return None, 'Sources manquantes : ' + ', '.join(map(str, missing))

    editions = sorted(set(years_by_file.values()))
    if len(editions) != 1:
        return None, 'Années mélangées : ' + ', '.join(editions)

    return editions[0], None


def validate_word_model(path: Path):
    if not path.exists() or path.suffix.lower() != '.docx':
        return False, 'Fichier Word invalide'

    try:
        doc = Document(path)
    except Exception as exc:
        return False, f"Impossible d'ouvrir le Word : {exc}"

    if len(doc.tables) != 414:
        return False, f'Mauvais modèle : {len(doc.tables)} tableaux détectés, 414 attendus'

    t141 = doc.tables[378]
    t142 = doc.tables[380]

    if len(t141.rows) != 8 or len(t141.columns) != 6:
        return False, "Mauvais modèle : tableau 14-1 différent de 8×6"
    if len(t142.rows) != 6 or len(t142.columns) != 4:
        return False, "Mauvais modèle : tableau 14-2 différent de 6×4"

    text141 = ' '.join(c.text for row in t141.rows for c in row.cells).lower()
    text142 = ' '.join(c.text for row in t142.rows for c in row.cells).lower()
    for keyword in ('nation', 'région', 'meknès'):
        if keyword not in text141:
            return False, f"Tableau 14-1 incompatible : '{keyword}' absent"
        if keyword not in text142:
            return False, f"Tableau 14-2 incompatible : '{keyword}' absent"

    return True, 'Word : PRÊT • structure 414 tableaux validée'


def open_path(path: Path):
    if not path.exists():
        raise FileNotFoundError(str(path))
    if os.name == 'nt':
        os.startfile(str(path.resolve()))
    elif sys.platform == 'darwin':
        subprocess.Popen(['open', str(path.resolve())])
    else:
        subprocess.Popen(['xdg-open', str(path.resolve())])


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('Générateur Annuaire Statistique de Meknès - HCP')
        self.geometry('900x585')
        self.minsize(800, 530)
        self.configure(bg=BG)

        self.word = tk.StringVar()
        self.excel = tk.StringVar()
        self.status = tk.StringVar(value='En attente • 0/24')
        self.year_text = tk.StringVar(value='Année détectée : —')

        self.q = queue.Queue()
        self.proc = None
        self.backend_output = []
        self.word_ok = False
        self.excel_ok = False
        self.year = None

        self._styles()
        self._ui()
        self.after(100, self._poll)

    def _styles(self):
        s = ttk.Style(self)
        try:
            s.theme_use('clam')
        except Exception:
            pass
        s.configure(
            'Primary.TButton',
            font=('Segoe UI', 11, 'bold'),
            padding=(22, 12),
            foreground='white',
            background=PRIMARY,
            borderwidth=0,
        )
        s.map(
            'Primary.TButton',
            background=[('active', PRIMARY_DARK), ('disabled', '#C8AEB7')],
        )

    def _ui(self):
        tk.Label(
            self,
            text='Générateur Annuaire Statistique de Meknès - HCP',
            bg=BG,
            fg=TEXT,
            font=('Segoe UI', 22, 'bold'),
        ).pack(anchor='w', padx=28, pady=(28, 4))

        tk.Label(
            self,
            text="Sélectionnez le modèle Word et les 23 sources Excel. L'année est détectée automatiquement.",
            bg=BG,
            fg=MUTED,
            font=('Segoe UI', 10),
        ).pack(anchor='w', padx=28, pady=(0, 20))

        card = tk.Frame(self, bg=WHITE, highlightbackground=BORDER, highlightthickness=1)
        card.pack(fill='x', padx=28)
        card.columnconfigure(1, weight=1)

        tk.Label(card, text='1. Word initial', bg=WHITE, fg=TEXT,
                 font=('Segoe UI', 11, 'bold')).grid(
            row=0, column=0, columnspan=3, sticky='w', padx=20, pady=(20, 8)
        )
        ttk.Entry(card, textvariable=self.word).grid(
            row=1, column=0, columnspan=2, sticky='ew', padx=(20, 10)
        )
        ttk.Button(card, text='Choisir le Word', command=self._choose_word).grid(
            row=1, column=2, padx=(0, 20)
        )
        self.word_state = tk.Label(
            card, text='Word : NON SÉLECTIONNÉ', bg=WHITE, fg=ERROR,
            font=('Segoe UI', 9, 'bold')
        )
        self.word_state.grid(row=2, column=0, columnspan=3, sticky='w', padx=20, pady=(7, 18))

        ttk.Separator(card).grid(row=3, column=0, columnspan=3, sticky='ew', padx=20)

        tk.Label(card, text='2. Dossier des fichiers Excel', bg=WHITE, fg=TEXT,
                 font=('Segoe UI', 11, 'bold')).grid(
            row=4, column=0, columnspan=3, sticky='w', padx=20, pady=(18, 8)
        )
        ttk.Entry(card, textvariable=self.excel).grid(
            row=5, column=0, columnspan=2, sticky='ew', padx=(20, 10)
        )
        ttk.Button(card, text='Choisir le dossier', command=self._choose_excel).grid(
            row=5, column=2, padx=(0, 20)
        )
        self.excel_state = tk.Label(
            card, text='Excel : NON SÉLECTIONNÉ', bg=WHITE, fg=ERROR,
            font=('Segoe UI', 9, 'bold')
        )
        self.excel_state.grid(row=6, column=0, columnspan=3, sticky='w', padx=20, pady=(7, 3))

        self.year_state = tk.Label(
            card, textvariable=self.year_text, bg=WHITE, fg=MUTED,
            font=('Segoe UI', 9, 'bold')
        )
        self.year_state.grid(row=7, column=0, columnspan=3, sticky='w', padx=20, pady=(0, 20))

        self.run = ttk.Button(
            self, text="GÉNÉRER L'ANNUAIRE", style='Primary.TButton',
            command=self._start, state='disabled'
        )
        self.run.pack(pady=(20, 14))

        tk.Label(self, textvariable=self.status, bg=BG, fg=TEXT,
                 font=('Segoe UI', 9, 'bold')).pack(anchor='w', padx=28)
        self.bar = ttk.Progressbar(self, maximum=100)
        self.bar.pack(fill='x', padx=28, pady=(7, 12))

        self.open_btn = ttk.Button(
            self, text='Ouvrir le Word final', command=self._open_final, state='disabled'
        )
        self.open_btn.pack()

    def _final_path(self):
        if not self.year:
            return None
        return ROOT / 'sorties' / f'Annuaire_HCP_{self.year}_FINAL.docx'

    def _choose_word(self):
        p = filedialog.askopenfilename(
            title='Choisir le Word initial', filetypes=[('Word', '*.docx')]
        )
        if not p:
            return
        self.word.set(p)
        self.word_state.configure(text='Word : VÉRIFICATION...', fg=MUTED)
        self.update_idletasks()
        ok, msg = validate_word_model(Path(p))
        self.word_ok = ok
        self.word_state.configure(text=msg, fg=SUCCESS if ok else ERROR)
        if not ok:
            messagebox.showerror('Modèle Word incompatible', msg)
        self._refresh_button()

    def _choose_excel(self):
        p = filedialog.askdirectory(title='Choisir le dossier des 23 fichiers Excel')
        if not p:
            return
        self.excel.set(p)
        folder = Path(p)
        n = count_excel(folder)
        year, error = detect_edition_year(folder)

        self.year = year
        self.excel_ok = n == 23 and year is not None and error is None

        if self.excel_ok:
            self.excel_state.configure(text='Excel : PRÊT • 23/23', fg=SUCCESS)
            self.year_text.set(f'Année détectée : {year}')
            self.year_state.configure(fg=SUCCESS)
        else:
            self.excel_state.configure(text=f'Excel : {n}/23', fg=ERROR)
            self.year_text.set('Année détectée : —')
            self.year_state.configure(fg=ERROR)
            messagebox.showerror('Sources Excel incompatibles', error or f'{n}/23 fichiers détectés')

        self._refresh_button()

    def _refresh_button(self):
        ready = self.word_ok and self.excel_ok and BACKEND.exists() and self.proc is None
        self.run.configure(state='normal' if ready else 'disabled')
        final = self._final_path()
        self.open_btn.configure(state='normal' if final and final.exists() else 'disabled')

    def _start(self):
        ok, msg = validate_word_model(Path(self.word.get()))
        if not ok:
            self.word_ok = False
            self.word_state.configure(text=msg, fg=ERROR)
            self._refresh_button()
            messagebox.showerror('Word incompatible', msg)
            return

        year, error = detect_edition_year(Path(self.excel.get()))
        if error or not year:
            self.excel_ok = False
            self.year = None
            self._refresh_button()
            messagebox.showerror('Sources Excel incompatibles', error or 'Année non détectée')
            return

        self.year = year
        self.year_text.set(f'Année détectée : {year}')

        if not BACKEND.exists():
            messagebox.showerror('Backend introuvable', str(BACKEND))
            return

        self.run.configure(state='disabled')
        self.open_btn.configure(state='disabled')
        self.backend_output = []
        self.status.set(f'Démarrage édition {year}... • 0/24')
        self.bar['value'] = 1
        threading.Thread(target=self._worker, daemon=True).start()

    def _worker(self):
        try:
            env = os.environ.copy()
            env['PYTHONIOENCODING'] = 'utf-8'
            env['PYTHONUTF8'] = '1'
            flags = getattr(subprocess, 'CREATE_NO_WINDOW', 0) if os.name == 'nt' else 0

            self.proc = subprocess.Popen(
                [
                    sys.executable, '-u', str(BACKEND),
                    '--word', self.word.get(),
                    '--excel-dir', self.excel.get(),
                ],
                cwd=str(ROOT.resolve()),
                stdout=subprocess.PIPE,
                stderr=subprocess.STDOUT,
                text=True,
                encoding='utf-8',
                errors='replace',
                env=env,
                creationflags=flags,
            )

            assert self.proc.stdout is not None
            for line in self.proc.stdout:
                self.q.put(('line', line))
            self.q.put(('done', self.proc.wait()))
        except Exception as exc:
            self.q.put(('error', str(exc)))
        finally:
            self.proc = None

    def _poll(self):
        try:
            while True:
                kind, payload = self.q.get_nowait()
                if kind == 'line':
                    line = str(payload)
                    self.backend_output.append(line.rstrip())
                    if len(self.backend_output) > 100:
                        self.backend_output = self.backend_output[-100:]

                    ym = re.search(r'@@YEAR@@\s+(20\d{2})', line)
                    if ym:
                        self.year = ym.group(1)
                        self.year_text.set(f'Année détectée : {self.year}')
                        self.year_state.configure(fg=SUCCESS)

                    m = re.search(r'@@PROGRESS24@@\s+(\d+)/24\s+\|\s+(.+)', line)
                    if m:
                        step = int(m.group(1))
                        name = m.group(2).strip()
                        self.bar['value'] = round(step / 24 * 100)
                        self.status.set(f'Étape {step}/24 — {name}')

                elif kind == 'done':
                    code = int(payload)
                    final = self._final_path()
                    if code == 0 and final and final.exists():
                        self.bar['value'] = 100
                        self.status.set(f'Terminé avec succès — Annuaire {self.year}')
                        self.open_btn.configure(state='normal')
                        messagebox.showinfo(
                            'Succès',
                            f"L'Annuaire Statistique de Meknès {self.year} a été généré avec succès.\n\n{final}",
                        )
                    else:
                        self.status.set('Échec')
                        details = '\n'.join(self.backend_output[-14:]).strip()
                        messagebox.showerror(
                            'Échec',
                            'La génération a échoué.' + ('\n\nDerniers détails :\n' + details if details else ''),
                        )
                    self._refresh_button()

                elif kind == 'error':
                    self.status.set('Erreur')
                    messagebox.showerror('Erreur', str(payload))
                    self._refresh_button()

        except queue.Empty:
            pass
        self.after(100, self._poll)

    def _open_final(self):
        final = self._final_path()
        try:
            if not final:
                raise FileNotFoundError('Aucun fichier final associé à une année détectée.')
            open_path(final)
        except Exception as exc:
            messagebox.showerror('Ouverture impossible', str(exc))


if __name__ == '__main__':
    App().mainloop()
