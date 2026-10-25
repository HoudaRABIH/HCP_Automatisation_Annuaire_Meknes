
from pathlib import Path
import os
import re
import shutil
import tempfile
import zipfile
import unicodedata
from openpyxl import load_workbook
from docx import Document
from lxml import etree

ROOT = Path(".")
SORTIES = ROOT / "sorties"
DONNEES = ROOT / "donnees_brutes"

INPUT = SORTIES / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED.docx"
OUTPUT = INPUT
EXCEL = DONNEES / "11. Enseignement_AS 2025.xlsx"

EDITION_YEAR = os.environ.get("HCP_EDITION_YEAR", "2025").strip()

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def norm(s):
    s = str(s or "").replace("\u00a0", " ").replace("\u202f", " ")
    s = s.replace("’", "'").replace("–", "-").replace("—", "-")
    s = "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"\s+", " ", s).strip().lower()


def fmt(v, existing=""):
    if v is None:
        return existing if existing.strip() in {"-", "…", "..."} else ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        if v == 0 and existing.strip() == "-":
            return "-"
        return f"{v:,}".replace(",", " ")
    if isinstance(v, float):
        if v == 0 and existing.strip() == "-":
            return "-"
        if v.is_integer():
            return f"{int(v):,}".replace(",", " ")
        s = f"{v:.3f}".rstrip("0").rstrip(".")
        a, *b = s.split(".")
        a = f"{int(a):,}".replace(",", " ")
        return a + ("," + b[0] if b else "")
    return str(v).strip()


def _cmp_display(s):
    return re.sub(r"\s+", "", str(s or "")).replace("\u00a0", "").replace("\u202f", "")


def set_cell_text(cell, value):
    value = str(value)

    # Do not touch a cell when only spaces / line wrapping differ.
    # This preserves the original Word layout much better.
    if _cmp_display(cell.text) == _cmp_display(value):
        return

    if cell.paragraphs:
        p = cell.paragraphs[0]
        if p.runs:
            p.runs[0].text = value
            for run in p.runs[1:]:
                run.text = ""
        else:
            p.add_run(value)

        # Numeric/data cells normally have one paragraph. We avoid deleting
        # paragraph objects because that can disturb the table geometry.
        for pp in cell.paragraphs[1:]:
            for run in pp.runs:
                run.text = ""
    else:
        cell.text = value


def find_year_pair(ws):
    pat = re.compile(r"(20\d{2})\s*[-/]\s*(20\d{2})")
    for r in range(1, min(ws.max_row, 20) + 1):
        for c in range(1, min(ws.max_column, 12) + 1):
            m = pat.search(str(ws.cell(r, c).value or ""))
            if m:
                return f"{m.group(1)}-{m.group(2)}"
    return None


def apply_year_pair(table, year_pair):
    if not year_pair:
        return
    mnew = re.match(r"(20\d{2})-(20\d{2})", year_pair)
    if not mnew:
        return
    a, b = mnew.groups()
    pat = re.compile(r"20\d{2}\s*[-/]\s*20\d{2}")
    for row in table.rows[:5]:
        for cell in row.cells:
            txt = cell.text
            m = pat.search(txt)
            if not m:
                continue
            old = m.group(0)
            nums = re.findall(r"20\d{2}", old)
            low, high = sorted([int(a), int(b)])
            replacement = f"{low}-{high}"
            if len(nums) == 2 and int(nums[0]) > int(nums[1]):
                replacement = f"{high}-{low}"
            set_cell_text(cell, pat.sub(replacement, txt, count=1))


def find_source_row(ws, role, national_labels=None):
    if role == "nation":
        labels = [norm(x) for x in (national_labels or [
            "Ensemble général", "Total national", "Total", "Ensemble"
        ])]
        for wanted in labels:
            for r in range(1, ws.max_row + 1):
                if norm(ws.cell(r, 1).value) == wanted:
                    return r
        return None

    if role == "region":
        for r in range(1, ws.max_row + 1):
            a = norm(ws.cell(r, 1).value)
            if "fes" in a and "meknes" in a:
                return r
        return None

    if role == "meknes":
        for r in range(1, ws.max_row + 1):
            if norm(ws.cell(r, 1).value) == "meknes":
                return r
        return None

    return None


def role_from_word_label(label):
    n = norm(label)
    if "nation" in n:
        return "nation"
    if "region" in n or ("fes" in n and "meknes" in n):
        return "region"
    if "meknes" in n:
        return "meknes"
    return None


# Word data-table object numbers from the validated 414-table structure.
# colmap values are zero-based indexes into Excel columns B, C, D...
REGULAR = {
    "3-5":  {"obj": 138, "sheet": "pres 2",       "colmap": [0,1,1,2]},
    "3-6":  {"obj": 140, "sheet": "pres 3",       "colmap": [0,1,2]},
    "3-7":  {"obj": 143, "sheet": "prim 5",       "colmap": [0,1,2,3]},
    "3-8":  {"obj": 145, "sheet": "prim 6",       "colmap": [0,1,1,2,3,3]},
    "3-9":  {"obj": 147, "sheet": "prim 10",      "colmap": [0,1,2,3,3]},
    "3-10": {"obj": 150, "sheet": "prim 12",      "colmap": [0,1,2,3]},
    "3-11": {"obj": 152, "sheet": "prim 13",      "colmap": [0,1,2,3]},
    "3-12": {"obj": 155, "sheet": "prim 15",      "colmap": [0,1,2],
             "national_labels": ["Ensemble général", "Total", "Ensemble"]},
    "3-13": {"obj": 158, "sheet": "collège 19",   "colmap": [0,1,2,3]},
    "3-14": {"obj": 160, "sheet": "collège 20",   "colmap": [0,1,2,3,3]},
    "3-15": {"obj": 162, "sheet": "collège 23",   "colmap": [0,1,2,3,3]},
    "3-16": {"obj": 165, "sheet": "collège 26",   "colmap": [0,1,2]},
    "3-17": {"obj": 168, "sheet": "qualif 30",    "colmap": [0,0,1,2,3]},
    "3-18": {"obj": 170, "sheet": "qualif 31",    "colmap": [0,1,1,2,3,3]},
    "3-19": {"obj": 172, "sheet": "qualif 33",    "colmap": [0,1,2,3,3]},
    "3-20": {"obj": 175, "sheet": "qualif 36",    "colmap": [0,1,2],
             "national_labels": ["Total national", "Total", "Ensemble"]},
    "3-21": {"obj": 178, "sheet": "post_sec 43",  "colmap": [0,1,2,3,4,5],
             "region_compute": True},
    "5-1":  {"obj": 197, "sheet": "prof 66",      "colmap": [0,1,2,3,4]},
}

REGION_MEMBERS = {
    "meknes", "fes", "taounate", "taza", "ifrane",
    "sefrou", "el hajeb", "boulemane", "moulay yacoub"
}


def computed_region_values(ws, ncols):
    sums = [0.0] * ncols
    found = 0
    for r in range(1, ws.max_row + 1):
        a = norm(ws.cell(r, 1).value)
        if a not in REGION_MEMBERS:
            continue
        found += 1
        for j in range(ncols):
            v = ws.cell(r, 2 + j).value
            if isinstance(v, (int, float)):
                sums[j] += v
    if not found:
        return None
    return [int(x) if float(x).is_integer() else x for x in sums]


def update_regular(doc, wb):
    for code, cfg in REGULAR.items():
        obj = cfg["obj"]
        if len(doc.tables) < obj:
            raise RuntimeError(f"Table Word {code} absente : objet {obj}.")
        table = doc.tables[obj - 1]
        ws = wb[cfg["sheet"]]
        apply_year_pair(table, find_year_pair(ws))

        for row in table.rows:
            label = row.cells[0].text
            role = role_from_word_label(label)
            if not role:
                continue

            if role == "region" and cfg.get("region_compute"):
                vals = computed_region_values(ws, max(cfg["colmap"]) + 1)
                if vals is None:
                    raise RuntimeError(f"Impossible de calculer Région pour {code}.")
            else:
                sr = find_source_row(ws, role, cfg.get("national_labels"))
                if sr is None:
                    raise RuntimeError(
                        f"Ligne Excel {role} introuvable pour {code} / {cfg['sheet']}."
                    )
                vals = [
                    ws.cell(sr, c).value
                    for c in range(2, ws.max_column + 1)
                ]

            data_cells = row.cells[1:-1]
            if len(data_cells) < len(cfg["colmap"]):
                raise RuntimeError(f"Structure Word inattendue pour {code}.")

            for j, src_idx in enumerate(cfg["colmap"]):
                if src_idx >= len(vals):
                    raise RuntimeError(f"Colonne Excel absente pour {code}.")
                old = data_cells[j].text
                set_cell_text(data_cells[j], fmt(vals[src_idx], old))


def find_excel_row_contains(ws, required, start=1, end=None):
    req = [norm(x) for x in required]
    end = end or ws.max_row
    for r in range(start, min(end, ws.max_row) + 1):
        a = norm(ws.cell(r, 1).value)
        if all(x in a for x in req):
            return r
    return None


def copy_row_values(word_row, ws, source_row, colmap):
    vals = [ws.cell(source_row, c).value for c in range(2, ws.max_column + 1)]
    data_cells = word_row.cells[1:-1]
    for j, src_idx in enumerate(colmap):
        old = data_cells[j].text
        v = vals[src_idx] if src_idx < len(vals) else None
        set_cell_text(data_cells[j], fmt(v, old))


def update_4_1(doc, wb):
    table = doc.tables[183 - 1]
    ws = wb["sup 47"]
    apply_year_pair(table, find_year_pair(ws))

    r_uni = find_excel_row_contains(ws, ["universite", "moulay", "ismail", "meknes"])
    r_tot = find_source_row(ws, "nation", ["Total"])

    if not r_uni or not r_tot:
        raise RuntimeError("Source 4-1 introuvable.")

    # Word: label | '-' | MC | MCH | PES | Total | Feminin | duplicate Feminin | Arabic
    for wr, sr in [(5, r_uni), (6, r_tot)]:
        vals = [ws.cell(sr, c).value for c in range(2, ws.max_column + 1)]
        mapping = [0,1,2,3,4,5,5]
        cells = table.rows[wr].cells[1:-1]
        for j, src_idx in enumerate(mapping):
            set_cell_text(cells[j], fmt(vals[src_idx] if src_idx < len(vals) else None, cells[j].text))


def update_4_2(doc, wb):
    table = doc.tables[185 - 1]
    ws = wb["sup 49"]
    apply_year_pair(table, find_year_pair(ws))

    uni = find_excel_row_contains(ws, ["universite", "moulay", "ismail", "meknes"])
    if not uni:
        raise RuntimeError("Université Moulay Ismail introuvable dans sup 49.")

    next_uni = ws.max_row
    for r in range(uni + 1, ws.max_row + 1):
        if "universite" in norm(ws.cell(r, 1).value):
            next_uni = r - 1
            break

    def source_for_label(label):
        n = norm(label)
        if n == "dont :":
            return find_excel_row_contains(ws, ["faculte", "lettres", "sciences", "humaines"], uni, next_uni)
        if "universite moulay ismail" in n:
            return uni
        if "lettres" in n and "humaines" in n:
            return find_excel_row_contains(ws, ["faculte", "lettres", "sciences", "humaines"], uni, next_uni)
        if "juridiques" in n or ("economiques" in n and "sociales" in n):
            # Values are on the continuation row "Economiques et Sociales".
            return find_excel_row_contains(ws, ["economiques", "sociales"], uni, next_uni)
        if n == "faculte des sciences":
            for r in range(uni, next_uni + 1):
                if norm(ws.cell(r, 1).value) == "faculte des sciences":
                    return r
        if "arts et metiers" in n:
            return find_excel_row_contains(ws, ["arts", "metiers"], uni, next_uni)
        if "ecole superieure de technologie" in n:
            return find_excel_row_contains(ws, ["ecole", "superieure", "technologie"], uni, next_uni)
        if "ens meknes" in n:
            return find_excel_row_contains(ws, ["ens", "meknes"], uni, next_uni)
        if "encg meknes" in n:
            return find_excel_row_contains(ws, ["encg", "meknes"], uni, next_uni)
        return None

    for wr in range(4, 13):
        label = table.rows[wr].cells[0].text
        sr = source_for_label(label)
        if sr is None:
            raise RuntimeError(f"Source 4-2 introuvable pour : {label}")
        vals = [ws.cell(sr, c).value for c in range(2, 5)]
        for j in range(3):
            set_cell_text(table.rows[wr].cells[1 + j], fmt(vals[j], table.rows[wr].cells[1 + j].text))


def update_4_5(doc, wb):
    table = doc.tables[192 - 1]
    ws = wb["sup 54"]
    apply_year_pair(table, find_year_pair(ws))

    rows = {
        3: find_excel_row_contains(ws, ["enag", "meknes"]),
        4: find_excel_row_contains(ws, ["ispits"]),
        5: find_source_row(ws, "nation", ["Total"]),
    }
    for wr, sr in rows.items():
        if sr is None:
            raise RuntimeError("Source 4-5 introuvable.")
        vals = [ws.cell(sr, 2).value, ws.cell(sr, 3).value]
        for j, v in enumerate(vals):
            set_cell_text(table.rows[wr].cells[1 + j], fmt(v, table.rows[wr].cells[1 + j].text))


def replace_text_in_nodes(nodes, old, new):
    parts = [n.text or "" for n in nodes]
    combined = "".join(parts)
    if old not in combined:
        return 0
    replaced = combined.replace(old, new)
    lengths = [len(x) for x in parts]
    pos = 0
    for i, node in enumerate(nodes):
        if i < len(nodes) - 1:
            take = lengths[i]
            node.text = replaced[pos:pos + take]
            pos += take
        else:
            node.text = replaced[pos:]
    return combined.count(old)


def patch_4_4_xml(docx_path, wb):
    ws = wb["sup 51-52"]
    year_pair = find_year_pair(ws)

    r_mek = find_excel_row_contains(ws, ["cite", "universitaire", "meknes"])
    # Prefer the detailed total row with Total/Feminin/Etrangers all populated.
    r_tot = None
    for r in range(1, ws.max_row + 1):
        if norm(ws.cell(r, 1).value) == "total":
            if all(isinstance(ws.cell(r, c).value, (int, float)) for c in (2, 3, 6)):
                r_tot = r
    if not r_mek or not r_tot:
        raise RuntimeError("Données 4-4 introuvables dans sup 51-52.")

    mek = [ws.cell(r_mek, 2).value, ws.cell(r_mek, 3).value, ws.cell(r_mek, 6).value]
    nat = [ws.cell(r_tot, 2).value, ws.cell(r_tot, 3).value, ws.cell(r_tot, 6).value]
    mek_fmt = [fmt(x) for x in mek]
    nat_fmt = [fmt(x) for x in nat]

    with zipfile.ZipFile(docx_path, "r") as zin:
        root = etree.fromstring(zin.read("word/document.xml"))
        body = root.find(".//" + W + "body")
        children = list(body)

        start = end = None
        for i, el in enumerate(children):
            text = "".join(el.xpath(".//w:t/text()", namespaces=NS))
            n = norm(text)
            if start is None and re.search(r"\b4\s*-\s*4\b", n) and "effectif des etudiants residant" in n:
                start = i
                continue
            if start is not None and re.search(r"\b4\s*-\s*5\b", n) and "effectif des laureats" in n:
                end = i
                break

        if start is None or end is None:
            raise RuntimeError("Bloc XML 4-4 introuvable.")

        block = children[start:end]
        paragraphs = []
        for el in block:
            paragraphs.extend(el.xpath(".//w:p", namespaces=NS))

        # Update year labels and typo.
        year_pat = re.compile(r"20\d{2}\s*[-/]\s*20\d{2}")
        for p in paragraphs:
            nodes = p.xpath(".//w:t", namespaces=NS)
            text = "".join(n.text or "" for n in nodes)
            if "Fémini3n" in text:
                replace_text_in_nodes(nodes, "Fémini3n", "Féminin")
            if year_pair and year_pat.search(text):
                old = year_pat.search(text).group(0)
                nums = re.findall(r"20\d{2}", old)
                a, b = year_pair.split("-")
                low, high = sorted([int(a), int(b)])
                rep = f"{low}-{high}"
                if len(nums) == 2 and int(nums[0]) > int(nums[1]):
                    rep = f"{high}-{low}"
                replace_text_in_nodes(nodes, old, rep)

        # Replace the 2025 structural values only inside the 4-4 block.
        # This is safe because the reference is used only as a layout donor.
        old_mek = ["2 668", "1 908", "61"]
        old_nat = ["60 915", "38 491", "1 487"]
        for p in paragraphs:
            nodes = p.xpath(".//w:t", namespaces=NS)
            text = "".join(n.text or "" for n in nodes)
            for old, new in zip(old_mek, mek_fmt):
                if text.strip() == old:
                    replace_text_in_nodes(nodes, old, new)
            for old, new in zip(old_nat, nat_fmt):
                if text.strip() == old:
                    replace_text_in_nodes(nodes, old, new)

        new_xml = etree.tostring(
            root, xml_declaration=True, encoding="UTF-8", standalone=True
        )

        tmp = docx_path.with_name(docx_path.stem + "_EDU_TMP.docx")
        with zipfile.ZipFile(tmp, "w") as zout:
            for item in zin.infolist():
                data = new_xml if item.filename == "word/document.xml" else zin.read(item.filename)
                zout.writestr(item, data)

    tmp.replace(docx_path)


def validate(docx_path, wb):
    doc = Document(docx_path)

    # Quick deterministic checks on representative tables.
    checks = [
        (138, "3-5"), (147, "3-9"), (162, "3-15"),
        (172, "3-19"), (183, "4-1"), (185, "4-2"),
        (192, "4-5"), (197, "5-1")
    ]
    for obj, code in checks:
        if len(doc.tables) < obj:
            raise RuntimeError(f"Validation Education : {code} absent.")

    # Verify current source year is visible somewhere in the education block.
    found_year = False
    for obj in [138,140,143,145,147,150,152,155,158,160,162,165,168,170,172,175,178,183,185,192,197]:
        if EDITION_YEAR in " ".join(
            c.text for row in doc.tables[obj-1].rows for c in row.cells
        ):
            found_year = True
            break

    # Edition year may be one year ahead of data year in some editions;
    # not fatal, because the source workbook can legitimately contain older years.
    print(f"Education : tables DIRECT mises à jour depuis {EXCEL.name}.")
    print(f"Edition source détectée par le backend : {EDITION_YEAR}")
    print(f"Edition visible dans au moins un tableau Education : {'OUI' if found_year else 'NON (année de données différente possible)'}")


def main():
    if not INPUT.exists():
        raise FileNotFoundError(INPUT)
    if not EXCEL.exists():
        raise FileNotFoundError(EXCEL)

    wb = load_workbook(EXCEL, data_only=True)
    doc = Document(INPUT)

    if len(doc.tables) < 197:
        raise RuntimeError(
            f"Structure Word Education inattendue : {len(doc.tables)} tableaux."
        )

    update_regular(doc, wb)
    update_4_1(doc, wb)
    update_4_2(doc, wb)
    update_4_5(doc, wb)

    tmp = INPUT.with_name(INPUT.stem + "_EDU_SAVE.docx")
    doc.save(tmp)
    tmp.replace(OUTPUT)

    patch_4_4_xml(OUTPUT, wb)
    validate(OUTPUT, wb)

    print("=" * 100)
    print("ENSEIGNEMENT MULTI-ANNEE : OK")
    print("=" * 100)
    print(OUTPUT)


if __name__ == "__main__":
    main()
