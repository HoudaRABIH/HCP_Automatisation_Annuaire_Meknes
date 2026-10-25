
from pathlib import Path
import os
import re
import shutil
import zipfile
import unicodedata
from docx import Document
from lxml import etree

ROOT = Path(".")
SORTIES = ROOT / "sorties"
YEAR = os.environ.get("HCP_EDITION_YEAR", "2025").strip()
TARGET = SORTIES / f"Annuaire_HCP_{YEAR}_FINAL.docx"

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def norm(s):
    s = str(s or "").replace("\u00a0", " ").replace("\u202f", " ")
    s = s.replace("’", "'").replace("–", "-").replace("—", "-")
    s = "".join(
        c for c in unicodedata.normalize("NFD", s)
        if unicodedata.category(c) != "Mn"
    )
    return re.sub(r"\s+", " ", s).strip().lower()


def fmt(v, existing=""):
    if v is None:
        return existing
    if isinstance(v, int):
        if v == 0 and existing.strip() == "-":
            return "-"
        return f"{v:,}".replace(",", " ")
    if isinstance(v, float):
        if v == 0 and existing.strip() == "-":
            return "-"
        if v.is_integer():
            return f"{int(v):,}".replace(",", " ")
        s = f"{v:.3f}".rstrip("0").rstrip(".")
        left, *rest = s.split(".")
        left = f"{int(left):,}".replace(",", " ")
        return left + ("," + rest[0] if rest else "")
    return str(v).strip()


def parse_number(s):
    s = str(s or "").replace("\u00a0", " ").replace("\u202f", " ")
    s = s.strip().replace(" ", "").replace(",", ".")
    if s in {"", "-", "…", "..."}:
        return 0.0
    try:
        return float(s)
    except Exception:
        return None


def set_cell(cell, value):
    value = str(value)
    if re.sub(r"\s+", "", cell.text) == re.sub(r"\s+", "", value):
        return

    original = cell.text
    paragraphs = cell.paragraphs

    if not paragraphs:
        cell.text = value
        return

    # Clear all existing text to avoid leaving the old value in a second
    # paragraph. Preserve a leading blank paragraph when the template uses it.
    for p in paragraphs:
        for run in p.runs:
            run.text = ""

    target_idx = 1 if original.startswith("\n") and len(paragraphs) > 1 else 0
    p = paragraphs[target_idx]
    if p.runs:
        p.runs[0].text = value
    else:
        p.add_run(value)


def table_text(table):
    return "\n".join(
        " | ".join(c.text for c in row.cells)
        for row in table.rows
    )


def code_pattern(code):
    a, b = code.split("-")
    return re.compile(
        rf"(?<!\d){int(a)}\s*[-–—]\s*{int(b)}(?!\d)"
    )


_TABLE_TEXT_CACHE = {}


def _cached_table_texts(doc):
    key = id(doc)
    cached = _TABLE_TEXT_CACHE.get(key)
    if cached is not None:
        return cached

    cached = []
    for table in doc.tables:
        txt = table_text(table)
        cached.append((txt, norm(txt)))

    _TABLE_TEXT_CACHE[key] = cached
    return cached


def find_data_table(doc, code, keyword=None):
    pat = code_pattern(code)
    nkw = norm(keyword) if keyword else None
    candidates = []
    cached = _cached_table_texts(doc)

    for i, (txt, nt) in enumerate(cached):
        if "liste des tableaux" in nt or "لائحة الجداول" in txt:
            continue
        if not pat.search(txt):
            continue
        if nkw and nkw not in nt:
            continue
        candidates.append(i)

    if not candidates:
        raise RuntimeError(
            f"Titre Word {code} introuvable"
            + (f" ({keyword})" if keyword else "")
        )

    for i in reversed(candidates):
        for j in range(i + 1, min(i + 4, len(doc.tables))):
            if len(doc.tables[j].rows) >= 3:
                return doc.tables[j]

    raise RuntimeError(f"Table de données {code} introuvable.")


def find_row(table, label, start=0):
    nl = norm(label)
    for i in range(start, len(table.rows)):
        first = norm(table.rows[i].cells[0].text)
        if first == nl or nl in first or first in nl:
            return table.rows[i]
    return None


def find_row_exactish(table, label):
    nl = norm(label)
    best = None
    for row in table.rows:
        first = norm(row.cells[0].text)
        if first == nl:
            return row
        if nl in first or first in nl:
            best = best or row
    return best


def year_values(text):
    return [int(x) for x in re.findall(r"20\d{2}", str(text or ""))]


def years_from_table(table):
    """
    Years appearing in the HEADER only. Stop before Nation/Region/Meknes
    so a source note such as "Annuaire ... 2025" cannot be mistaken for
    the statistical year.
    """
    vals = []
    for row in table.rows:
        first = norm(row.cells[0].text) if row.cells else ""
        if first in {"nation", "region"} or "meknes" in first:
            break
        for cell in row.cells:
            vals.extend(year_values(cell.text))
    return sorted(set(vals))


def latest_year(table):
    years = years_from_table(table)
    return max(years) if years else None


def year_range(*tables):
    years = sorted(set(
        y for table in tables
        for y in [latest_year(table)]
        if y is not None
    ))
    if not years:
        return None
    if len(years) == 1:
        return str(years[0])
    return f"{years[0]}-{years[-1]}"


def first_year_pair(table):
    pat = re.compile(r"(20\d{2})\s*[-/]\s*(20\d{2})")
    for row in table.rows[:5]:
        for cell in row.cells:
            m = pat.search(cell.text)
            if m:
                a, b = map(int, m.groups())
                return f"{min(a,b)}-{max(a,b)}"
    ys = years_from_table(table)
    if ys:
        return str(max(ys))
    return None


def numeric_cells(row):
    return [c.text.strip() for c in row.cells[1:-1]]


def meknes_row(table):
    for row in table.rows:
        if "meknes" in norm(row.cells[0].text):
            return row
    raise RuntimeError("Ligne Meknès introuvable.")


def update_summary_tables(doc):
    if len(doc.tables) < 3:
        raise RuntimeError("Tables de synthèse introuvables.")

    s1 = doc.tables[1]  # Principal indicators page 1
    s2 = doc.tables[2]  # Principal indicators page 2

    # ------------------------------------------------------------------
    # ADMINISTRATION + POPULATION
    # ------------------------------------------------------------------
    t11 = find_data_table(doc, "1-1", "Division administrative")
    m11 = meknes_row(t11)
    years = years_from_table(t11)
    if years:
        set_cell(s1.rows[2].cells[2], str(max(years)))
    # Total communes. Urban/rural subdivision is intentionally preserved
    # because 1-2 is not semantically compatible with the local model.
    set_cell(s1.rows[3].cells[2], m11.cells[2].text.strip())

    t13 = find_data_table(doc, "1-3", "Population légale")
    years = years_from_table(t13)
    if years:
        set_cell(s1.rows[6].cells[2], str(max(years)))

    # first Meknès row = total; second = urban; third = rural
    mek_rows = [
        row for row in t13.rows
        if "meknes" in norm(row.cells[0].text)
    ]
    if len(mek_rows) >= 2:
        set_cell(s1.rows[7].cells[2], mek_rows[0].cells[1].text.strip())
        set_cell(s1.rows[8].cells[2], mek_rows[1].cells[1].text.strip())

    # ------------------------------------------------------------------
    # HEALTH
    # ------------------------------------------------------------------
    t25 = find_data_table(doc, "2-5", "Médecins du Ministère")
    t27 = find_data_table(doc, "2-7", "Médecins privés")
    t218 = find_data_table(doc, "2-18", "Mouvements des malades")
    t219 = find_data_table(doc, "2-19", "occupation des lits")

    yr = year_range(t25, t27, t218, t219)
    if yr:
        set_cell(s1.rows[9].cells[1], yr)
        set_cell(s1.rows[9].cells[2], yr)

    pub = meknes_row(t25)
    priv = meknes_row(t27)
    public_total = parse_number(pub.cells[1].text)
    private_total = parse_number(priv.cells[1].text)

    if public_total is not None and private_total is not None:
        set_cell(
            s1.rows[10].cells[2],
            fmt(public_total + private_total)
        )
    set_cell(s1.rows[11].cells[2], pub.cells[1].text.strip())

    m218 = meknes_row(t218)
    # 2-18: days | admissions | functional capacity | establishments
    set_cell(s1.rows[12].cells[2], m218.cells[3].text.strip())
    set_cell(s1.rows[13].cells[2], m218.cells[2].text.strip())
    set_cell(s1.rows[14].cells[2], m218.cells[1].text.strip())

    m219 = meknes_row(t219)
    # 2-19: interval | rotation | stay | occupancy
    set_cell(s1.rows[15].cells[2], m219.cells[4].text.strip())
    set_cell(s1.rows[16].cells[2], m219.cells[3].text.strip())
    set_cell(s1.rows[17].cells[2], m219.cells[2].text.strip())

    # ------------------------------------------------------------------
    # EDUCATION
    # ------------------------------------------------------------------
    t35  = find_data_table(doc, "3-5", "Educateurs")
    t36  = find_data_table(doc, "3-6", "préscolaire moderne")
    t38  = find_data_table(doc, "3-8", "Personnel enseignant")
    t39  = find_data_table(doc, "3-9", "Elèves")
    t314 = find_data_table(doc, "3-14", "Personnel enseignant")
    t315 = find_data_table(doc, "3-15", "Elèves")
    t318 = find_data_table(doc, "3-18", "Personnel enseignant")
    t319 = find_data_table(doc, "3-19", "Elèves")
    t41  = find_data_table(doc, "4-1", "Personnel enseignant")
    t42  = find_data_table(doc, "4-2", "Etudiants")

    edu_year = first_year_pair(t35)
    if edu_year:
        set_cell(s1.rows[18].cells[1], edu_year)
        set_cell(s1.rows[18].cells[2], edu_year)

    set_cell(s1.rows[19].cells[2], meknes_row(t35).cells[1].text.strip())
    set_cell(s1.rows[20].cells[2], meknes_row(t36).cells[1].text.strip())
    set_cell(s1.rows[21].cells[2], meknes_row(t39).cells[1].text.strip())
    set_cell(s1.rows[22].cells[2], meknes_row(t315).cells[1].text.strip())
    set_cell(s1.rows[23].cells[2], meknes_row(t319).cells[1].text.strip())

    uni_row = find_row_exactish(t42, "Université Moulay Ismail-Meknès")
    if uni_row:
        set_cell(s1.rows[24].cells[2], uni_row.cells[1].text.strip())

    set_cell(s1.rows[25].cells[2], meknes_row(t38).cells[1].text.strip())
    set_cell(s1.rows[26].cells[2], meknes_row(t314).cells[1].text.strip())
    set_cell(s1.rows[27].cells[2], meknes_row(t318).cells[1].text.strip())

    umi_teacher = None
    for row in t41.rows:
        if "moulay ismail" in norm(row.cells[0].text):
            umi_teacher = row
            break
    if umi_teacher:
        # total column in the 4-1 layout
        set_cell(s1.rows[28].cells[2], umi_teacher.cells[5].text.strip())

    # Employment / salary stays CONSERVER_MODELE until a deterministic
    # source mapping is validated for the annual employment workbook.

    # ------------------------------------------------------------------
    # AGRICULTURE
    # ------------------------------------------------------------------
    t91 = find_data_table(doc, "9-1", "céréales")
    t93 = find_data_table(doc, "9-3", "légumineuses")
    t97 = find_data_table(doc, "9-7", "cheptel")

    ag_year = first_year_pair(t91)
    if ag_year:
        set_cell(s2.rows[2].cells[2], ag_year)

    m91 = meknes_row(t91)
    # 9-1 layout: latest Production, latest Superficie, previous...
    set_cell(s2.rows[3].cells[2], m91.cells[2].text.strip())
    set_cell(s2.rows[4].cells[2], m91.cells[1].text.strip())

    m93 = meknes_row(t93)
    set_cell(s2.rows[5].cells[2], m93.cells[2].text.strip())
    set_cell(s2.rows[6].cells[2], m93.cells[1].text.strip())

    m97 = meknes_row(t97)
    # 9-7: Caprins | Ovins | Bovins
    set_cell(s2.rows[7].cells[2], m97.cells[3].text.strip())
    set_cell(s2.rows[8].cells[2], m97.cells[2].text.strip())
    set_cell(s2.rows[9].cells[2], m97.cells[1].text.strip())

    # ------------------------------------------------------------------
    # ENERGY + WATER
    # ------------------------------------------------------------------
    t111 = find_data_table(doc, "11-1", "Office National")
    t112 = find_data_table(doc, "11-2", "Ventes")
    en_year = year_range(t111, t112)
    if en_year:
        set_cell(s2.rows[10].cells[2], en_year)

    m112 = meknes_row(t112)
    set_cell(s2.rows[11].cells[2], m112.cells[1].text.strip())

    m111 = meknes_row(t111)
    # production is the last numeric column before Arabic
    set_cell(s2.rows[12].cells[2], m111.cells[-2].text.strip())

    # ------------------------------------------------------------------
    # TOURISM
    # ------------------------------------------------------------------
    t121 = find_data_table(doc, "12-1", "Capacité")
    t124 = find_data_table(doc, "12-4", "arrivées")
    t125 = find_data_table(doc, "12-5", "nuitées")

    tour_year = year_range(t121, t124, t125)
    if tour_year:
        set_cell(s2.rows[13].cells[2], tour_year)

    m121 = meknes_row(t121)
    # 12-1: Lits | Chambres | Etablissements
    set_cell(s2.rows[14].cells[2], m121.cells[3].text.strip())
    set_cell(s2.rows[15].cells[2], m121.cells[1].text.strip())

    m125 = meknes_row(t125)
    set_cell(s2.rows[16].cells[2], m125.cells[1].text.strip())

    m124 = meknes_row(t124)
    set_cell(s2.rows[17].cells[2], m124.cells[1].text.strip())

    # ------------------------------------------------------------------
    # HABITAT (preserved if detailed table itself is preserved)
    # ------------------------------------------------------------------
    t139 = find_data_table(doc, "13-9", "autorisations")
    hab_years = years_from_table(t139)
    if hab_years:
        set_cell(s2.rows[18].cells[2], str(max(hab_years)))

    def source_row_value(label):
        row = find_row_exactish(t139, label)
        if row:
            # Meknès is the third geography numeric column in 13-9
            return row.cells[3].text.strip()
        return None

    for sr, label in [
        (19, "Nombre d'autorisations"),
        (20, "Nombre de logements"),
        (21, "Valeur prévue"),
        (22, "Surface bâtie"),
        (23, "Surface des planchers"),
    ]:
        val = source_row_value(label)
        if val is not None:
            set_cell(s2.rows[sr].cells[2], val)

    # ------------------------------------------------------------------
    # TRANSPORT
    # ------------------------------------------------------------------
    t141 = find_data_table(doc, "14-1", "Immatriculation")
    t142 = find_data_table(doc, "14-2", "Permis")
    tr_years = years_from_table(t141) + years_from_table(t142)
    if tr_years:
        set_cell(s2.rows[24].cells[2], str(max(tr_years)))

    set_cell(s2.rows[25].cells[2], meknes_row(t141).cells[1].text.strip())
    set_cell(s2.rows[26].cells[2], meknes_row(t142).cells[1].text.strip())

    # ------------------------------------------------------------------
    # CONSUMER PRICE INDEX
    # ------------------------------------------------------------------
    t151 = find_data_table(doc, "15-1", "Indice des prix")

    # 15-1 is split over two consecutive Word tables in the model.
    price_tables = [t151]
    idx151 = None
    for i, table in enumerate(doc.tables):
        if table._tbl is t151._tbl:
            idx151 = i
            break
    if idx151 is not None and idx151 + 1 < len(doc.tables):
        nxt = doc.tables[idx151 + 1]
        if len(nxt.rows) >= 3 and years_from_table(nxt) == years_from_table(t151):
            price_tables.append(nxt)

    current_year = latest_year(t151)
    if current_year:
        set_cell(s2.rows[27].cells[2], str(current_year))

    price_labels = [
        (28, "Produits alimentaires et boissons non alcoolisées"),
        (29, "Boissons alcoolisées et tabac"),
        (30, "Articles d'habillement et chaussures"),
        (31, "Logement, eau, électricité, gaz et autres combustibles"),
        (32, "Meubles, articles de ménage et entretien courant du foyer"),
        (33, "Santé"),
        (34, "Transport"),
        (35, "Communications"),
        (36, "Loisirs et culture"),
        (37, "Enseignement"),
        (38, "Restaurants et hôtels"),
        (39, "Indice général"),
    ]

    def price_row(label):
        for table in price_tables:
            row = find_row_exactish(table, label)
            if row:
                return row
        return None

    for target_row, label in price_labels:
        src = price_row(label)
        if src:
            set_cell(s2.rows[target_row].cells[2], src.cells[1].text.strip())

    general = price_row("Indice général")
    if general and len(general.cells) >= 4:
        current = parse_number(general.cells[1].text)
        previous = parse_number(general.cells[2].text)
        if current is not None and previous not in (None, 0):
            variation = ((current - previous) / previous) * 100
            set_cell(
                s2.rows[40].cells[2],
                f"{variation:.1f}".replace(".", ",")
            )


def replace_across_nodes(nodes, old, new):
    parts = [n.text or "" for n in nodes]
    combined = "".join(parts)
    if old not in combined:
        return 0

    replaced = combined.replace(old, new)
    lengths = [len(x) for x in parts]
    pos = 0

    for i, node in enumerate(nodes):
        if i < len(nodes) - 1:
            take = lengths[i]
            node.text = replaced[pos:pos + take]
            pos += take
        else:
            node.text = replaced[pos:]

    return combined.count(old)


def _arabic_number(display):
    """
    Keep the same logical RTL convention used by the Word template:
    4 347 -> 347 4 ; 783,1 -> 1،783.
    """
    s = str(display).strip()
    if "," in s:
        left, right = s.split(",", 1)
        return f"{right}،{left.replace(' ', '')}"
    parts = s.split()
    if len(parts) > 1:
        return " ".join(reversed(parts))
    return s


def _replace_number_at(nodes, index, new_value):
    combined = "".join(n.text or "" for n in nodes)
    matches = list(re.finditer(
        r"(?<!\d)\d+(?:[ \u00a0\u202f،,]\d+)*(?!\d)",
        combined
    ))
    if index >= len(matches):
        return False

    m = matches[index]
    replaced = combined[:m.start()] + str(new_value) + combined[m.end():]

    lengths = [len(n.text or "") for n in nodes]
    pos = 0
    for i, node in enumerate(nodes):
        if i < len(nodes) - 1:
            take = lengths[i]
            node.text = replaced[pos:pos + take]
            pos += take
        else:
            node.text = replaced[pos:]
    return True


def _rewrite_combined_nodes(nodes, new_text):
    """Redistribute a paragraph's combined text while preserving its runs/shapes."""
    lengths = [len(n.text or "") for n in nodes]
    pos = 0
    for i, node in enumerate(nodes):
        if i < len(nodes) - 1:
            take = lengths[i]
            node.text = new_text[pos:pos + take]
            pos += take
        else:
            node.text = new_text[pos:]


def _rtl_number(value):
    """Visual storage used by the existing Arabic floating boxes."""
    value = str(value or "").strip()
    if not value:
        return value
    if "," in value:
        left, right = value.split(",", 1)
        return f"{right}،{left.replace(' ', '')}"
    parts = value.split()
    return " ".join(reversed(parts)) if len(parts) > 1 else value


def _paragraph_text(p):
    return "".join(p.xpath(".//w:t/text()", namespaces=NS))


def update_source_years_xml(root):
    """Update only publication-year tokens in Annuaire source lines."""
    year_re = re.compile(r"20\d{2}")
    ar_marker = norm("النشرة الإحصائية السنوية للمغرب")
    changed = 0

    for p in root.xpath(".//w:p", namespaces=NS):
        all_nodes = p.xpath(".//w:t", namespaces=NS)

        # Direct nodes can hold a source line in a paragraph that also embeds
        # text boxes. Never touch nested data values from that parent.
        direct_nodes = []
        for node in all_nodes:
            ancestors = node.xpath("ancestor::w:p[1]", namespaces=NS)
            if ancestors and ancestors[0] is p:
                direct_nodes.append(node)
        if direct_nodes:
            direct_text = "".join(n.text or "" for n in direct_nodes)
            nt = norm(direct_text)
            if "annuaire statistique du maroc" in nt or ar_marker in nt:
                newer = year_re.sub(YEAR, direct_text)
                if newer != direct_text:
                    _rewrite_combined_nodes(direct_nodes, newer)
                    changed += 1

        if p.xpath(".//w:p", namespaces=NS):
            continue
        nodes = p.xpath(".//w:t", namespaces=NS)
        if not nodes:
            continue
        text = "".join(n.text or "" for n in nodes)
        nt = norm(text)
        if "annuaire statistique du maroc" not in nt and ar_marker not in nt:
            continue
        new_text = year_re.sub(YEAR, text)
        if new_text != text:
            _rewrite_combined_nodes(nodes, new_text)
            changed += 1

    return changed

def patch_floating_boxes(docx_path):
    """
    Synchronise the free-floating summary boxes from detailed Word tables.
    No 2025 numeric value is hard-coded. No absolute paragraph number is used.
    """
    doc = Document(docx_path)

    t66 = find_data_table(doc, "6-6", "Installations sportives")
    t69 = find_data_table(doc, "6-9", "biblioth")
    t611 = find_data_table(doc, "6-11", "conservatoires")
    t111 = find_data_table(doc, "11-1", "Office National")
    t112 = find_data_table(doc, "11-2", "Ventes")
    t121 = find_data_table(doc, "12-1", "Capacité")

    m66 = meknes_row(t66)
    sport_total = 0.0
    for cell in m66.cells[1:-1]:
        v = parse_number(cell.text)
        if v is not None:
            sport_total += v
    sport_total = fmt(int(round(sport_total)))

    m69 = meknes_row(t69)
    libraries = m69.cells[4].text.strip()
    library_members = m69.cells[2].text.strip()

    m611 = meknes_row(t611)
    conservatoires = m611.cells[3].text.strip()
    conservatoire_students = m611.cells[2].text.strip()

    m111 = meknes_row(t111)
    water = m111.cells[-2].text.strip()

    m112 = meknes_row(t112)
    electricity = m112.cells[1].text.strip()

    m121 = meknes_row(t121)
    beds = m121.cells[1].text.strip()
    establishments = m121.cells[3].text.strip()

    with zipfile.ZipFile(docx_path, "r") as zin:
        root = etree.fromstring(zin.read("word/document.xml"))
        paragraphs = root.xpath(".//w:p", namespaces=NS)
        texts = [_paragraph_text(p) for p in paragraphs]

        # ---------------- CULTURE BOX ----------------
        culture_starts = [
            i for i, t in enumerate(texts)
            if norm(t) == norm("ACTIVITES CULTURELLES ET LOISIRS")
        ]
        # The summary box is the occurrence followed by short indicator paragraphs.
        culture_start = None
        for i in culture_starts:
            window = texts[i:min(i + 50, len(texts))]
            if any(re.fullmatch(r"\s*\d+\s+Installations sportives\s*", x, re.I) for x in window):
                culture_start = i
                break

        if culture_start is not None:
            for i in range(culture_start, min(culture_start + 55, len(paragraphs))):
                p = paragraphs[i]
                nodes = p.xpath(".//w:t", namespaces=NS)
                if not nodes:
                    continue
                text = "".join(n.text or "" for n in nodes).strip()
                new = text

                if re.fullmatch(r"\d+\s+Installations sportives", text, re.I):
                    new = f"{sport_total} Installations sportives"
                elif re.fullmatch(r"\d+\s+تجهيز رياضي", text):
                    new = f"{_rtl_number(sport_total)} تجهيز رياضي"
                elif re.fullmatch(r"\d+\s+biblioth[eé]ques", text, re.I):
                    new = f"{libraries} bibliothéques"
                elif re.fullmatch(r"[\d\s]+inscrits", text, re.I):
                    new = f"{library_members} inscrits"
                elif re.fullmatch(r"\d+\s+مكتبات", text):
                    new = f"{_rtl_number(libraries)} مكتبات"
                elif re.fullmatch(r"[\d\s]+مسجل", text):
                    new = f"{_rtl_number(library_members)} مسجل"
                elif re.fullmatch(r"\d+\s+conservatoire", text, re.I):
                    new = f"{conservatoires} conservatoire"
                elif re.fullmatch(r"[\d\s]+élèves inscrits", text, re.I):
                    new = f"{conservatoire_students} élèves inscrits"
                elif "معهد موسيقي" in text and "تلميذ مسجل" in text:
                    new = re.sub(
                        r"^\s*\d+\s+معهد موسيقي\s+[\d\s]+\s+تلميذ مسجل\s*$",
                        f"{_rtl_number(conservatoires)} معهد موسيقي "
                        f"{_rtl_number(conservatoire_students)} تلميذ مسجل",
                        text,
                    )

                if new != text:
                    _rewrite_combined_nodes(nodes, new)

        # ---------------- WATER / ELECTRICITY / TOURISM BOXES ----------------
        for p in paragraphs:
            if p.xpath(".//w:p", namespaces=NS):
                continue
            nodes = p.xpath(".//w:t", namespaces=NS)
            if not nodes:
                continue
            text = "".join(n.text or "" for n in nodes)
            nt = norm(text)
            new = text

            if "production de la prefecture de meknes en eau potable" in nt:
                # French and/or parent paragraph that can also contain Arabic.
                new = re.sub(
                    r"\d[\d\s]*(?:,\d+)?(?=\s+milliers de mètre cube)",
                    water,
                    new,
                    flags=re.I,
                )
                # Arabic duplicate inside the parent paragraph.
                new = re.sub(
                    r"[\d\s]+(?=\s+ألف متر مكعب\s*:\s*إنتاج عمالة مكناس)",
                    _rtl_number(water),
                    new,
                )
            elif "انتاج عمالة مكناس من الماء الصالح للشرب" in nt:
                new = re.sub(
                    r"[\d\s]+(?=\s+ألف متر مكعب\s*:\s*إنتاج عمالة مكناس)",
                    _rtl_number(water),
                    new,
                )

            if "ventes de la prefecture de meknes en electricite" in nt:
                new = re.sub(
                    r"\d[\d\s]*(?:[,.]\d+)?(?=\s+millions de Kwh)",
                    electricity,
                    new,
                    flags=re.I,
                )
                new = re.sub(
                    r"[\d\s،,.]+(?=\s+مليون كيلوواط في الساعة)",
                    _rtl_number(electricity),
                    new,
                )
            elif "مبيعات عمالة مكناس" in text:
                new = re.sub(
                    r"[\d\s،,.]+(?=\s+مليون كيلوواط في الساعة)",
                    _rtl_number(electricity),
                    new,
                )

            if "etablissements touristiques classes" in nt and "lits" in nt:
                new = re.sub(
                    r"\d+\s+établissements touristiques classés",
                    f"{establishments} établissements touristiques classés",
                    new,
                    flags=re.I,
                )
                new = re.sub(
                    r"(?<!\d)\d(?:[\d ]*\d)?(?=\s+lits)",
                    beds,
                    new,
                    flags=re.I,
                )
                new = re.sub(
                    r"\d+\s+مؤسسة سياحية مصنفة",
                    f"{_rtl_number(establishments)} مؤسسة سياحية مصنفة",
                    new,
                )
                new = re.sub(
                    r"(?<!\d)\d(?:[\d ]*\d)?(?=\s+سرير)",
                    _rtl_number(beds),
                    new,
                )
            elif "مؤسسة سياحية مصنفة" in text and "سرير" in text:
                new = re.sub(
                    r"\d+\s+مؤسسة سياحية مصنفة",
                    f"{_rtl_number(establishments)} مؤسسة سياحية مصنفة",
                    new,
                )
                new = re.sub(
                    r"(?<!\d)\d(?:[\d ]*\d)?(?=\s+سرير)",
                    _rtl_number(beds),
                    new,
                )

            # Historical typo in the Education structural donor.
            new = new.replace("Fémini3n", "Féminin")

            if new != text:
                _rewrite_combined_nodes(nodes, new)

        source_changes = update_source_years_xml(root)

        new_xml = etree.tostring(
            root,
            xml_declaration=True,
            encoding="UTF-8",
            standalone=True,
        )

        tmp = docx_path.with_name(docx_path.stem + "_SYNTH_TMP.docx")
        with zipfile.ZipFile(tmp, "w") as zout:
            for item in zin.infolist():
                data = new_xml if item.filename == "word/document.xml" else zin.read(item.filename)
                zout.writestr(item, data)

    tmp.replace(docx_path)
    print(f"Sources Annuaire mises à l'édition {YEAR} : {source_changes} paragraphes.")

def main():
    if not TARGET.exists():
        raise FileNotFoundError(TARGET)

    backup_dir = SORTIES / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"{TARGET.stem}_avant_synthese.docx"
    shutil.copy2(TARGET, backup)

    doc = Document(TARGET)
    update_summary_tables(doc)

    tmp = TARGET.with_name(TARGET.stem + "_TABLES_TMP.docx")
    doc.save(tmp)
    tmp.replace(TARGET)

    patch_floating_boxes(TARGET)

    print("=" * 100)
    print("SYNCHRONISATION SYNTHESE + ENCADRES MULTI-ANNEE : OK")
    print("=" * 100)
    print(f"Edition : {YEAR}")
    print(f"Word final : {TARGET}")


if __name__ == "__main__":
    main()
