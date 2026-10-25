from pathlib import Path
from openpyxl import load_workbook
from docx import Document
import re
import unicodedata


# ============================================================
# CHEMINS
# ============================================================

DOSSIER_SOURCE = Path("donnees_brutes")
DOSSIER_MODELES = Path("modeles")
DOSSIER_SORTIE = Path("sorties")

FICHIER_MAPPING = DOSSIER_SORTIE / "mapping_final_recettes.xlsx"
FICHIER_SORTIE = DOSSIER_SORTIE / "TEST_14-1_Annuaire_HCP.docx"

FICHIER_TRANSPORT = DOSSIER_SOURCE / "8. Transport_AS 2025.xlsx"

CODE_WORD = "14-1"


# ============================================================
# OUTILS
# ============================================================

def normaliser(texte):
    texte = str(texte or "").strip().lower()
    texte = texte.replace("’", "'")

    texte = "".join(
        caractere
        for caractere in unicodedata.normalize("NFD", texte)
        if unicodedata.category(caractere) != "Mn"
    )

    texte = re.sub(r"\s+", " ", texte)

    return texte


def feuille_par_nom_approximatif(classeur, nom):
    cible = normaliser(nom)

    for feuille in classeur.worksheets:
        if normaliser(feuille.title) == cible:
            return feuille

    raise KeyError(
        f"Feuille introuvable : {nom}"
    )


def convertir_nombre(valeur):
    """
    Convertit une valeur Excel en affichage HCP :
    16873 -> '16 873'
    '51618*' -> '51 618*'
    """
    if valeur is None:
        return ""

    if isinstance(valeur, bool):
        return str(valeur)

    if isinstance(valeur, int):
        return f"{valeur:,}".replace(",", " ")

    if isinstance(valeur, float):
        if valeur.is_integer():
            return f"{int(valeur):,}".replace(",", " ")

        texte = f"{valeur:.2f}".rstrip("0").rstrip(".")
        return texte.replace(".", ",")

    texte = str(valeur).strip()

    etoile = ""

    if texte.endswith("*"):
        etoile = "*"
        texte = texte[:-1].strip()

    texte_sans_espaces = (
        texte
        .replace(" ", "")
        .replace("\u00A0", "")
    )

    try:
        nombre = int(texte_sans_espaces)

        return (
            f"{nombre:,}".replace(",", " ")
            + etoile
        )

    except ValueError:
        return str(valeur).strip()


def remplacer_texte_preserver_style(cellule, nouveau_texte):
    """
    Remplace le texte d'une cellule sans recréer la cellule ni
    son paragraphe, afin de conserver au maximum le format Word.
    """
    if not cellule.paragraphs:
        cellule.text = nouveau_texte
        return

    paragraphe = cellule.paragraphs[0]

    if paragraphe.runs:
        paragraphe.runs[0].text = nouveau_texte

        for run in paragraphe.runs[1:]:
            run.text = ""

    else:
        paragraphe.add_run(
            nouveau_texte
        )

    # Vider les éventuels paragraphes supplémentaires
    for autre_paragraphe in cellule.paragraphs[1:]:
        for run in autre_paragraphe.runs:
            run.text = ""


def trouver_ligne_par_libelle(
    feuille,
    libelle,
    colonne=1,
    debut=1,
    fin=None
):
    cible = normaliser(
        libelle
    )

    if fin is None:
        fin = feuille.max_row

    for ligne in range(
        debut,
        fin + 1
    ):
        valeur = feuille.cell(
            ligne,
            colonne
        ).value

        if normaliser(
            valeur
        ) == cible:
            return ligne

    raise ValueError(
        f"Libellé '{libelle}' introuvable dans {feuille.title}"
    )


def trouver_ligne_total_national(
    feuille,
    debut=1,
    fin=None
):
    """
    Dans la table 8-8, cherche la ligne A='Total'
    avec des valeurs B:E.
    """
    if fin is None:
        fin = feuille.max_row

    candidats = []

    for ligne in range(
        debut,
        fin + 1
    ):
        valeur_a = normaliser(
            feuille.cell(
                ligne,
                1
            ).value
        )

        if valeur_a != "total":
            continue

        valeurs = [
            feuille.cell(
                ligne,
                colonne
            ).value
            for colonne in range(
                2,
                6
            )
        ]

        if all(
            valeur not in (
                None,
                ""
            )
            for valeur in valeurs
        ):
            candidats.append(
                ligne
            )

    if not candidats:
        raise ValueError(
            "Aucune ligne Total nationale trouvée dans la table 8-8."
        )

    # Dans le classeur actuel il s'agit de la ligne finale de total.
    return candidats[-1]


def lire_numero_table_word():
    wb = load_workbook(
        FICHIER_MAPPING,
        data_only=True
    )

    ws = wb[
        "Mapping_final"
    ]

    entetes = {
        cellule.value: cellule.column
        for cellule in ws[1]
    }

    numero_table = None

    for ligne in range(
        2,
        ws.max_row + 1
    ):
        code = str(
            ws.cell(
                ligne,
                entetes["N° Word"]
            ).value or ""
        ).strip()

        if code == CODE_WORD:
            numero_table = int(
                ws.cell(
                    ligne,
                    entetes["Table Word données"]
                ).value
            )

            break

    wb.close()

    if numero_table is None:
        raise RuntimeError(
            f"Tableau Word {CODE_WORD} introuvable dans le mapping."
        )

    return numero_table


# ============================================================
# 1) EXTRACTION EXCEL
# ============================================================

if not FICHIER_TRANSPORT.exists():
    raise FileNotFoundError(
        f"Fichier source introuvable : {FICHIER_TRANSPORT}"
    )

classeur = load_workbook(
    FICHIER_TRANSPORT,
    data_only=True,
    read_only=False
)

# 8-7 = détail par centre d'immatriculation :
# Région Fès-Meknès + Meknès
feuille_87 = feuille_par_nom_approximatif(
    classeur,
    "7(1)7(2)"
)

# 8-8 = total national
feuille_88 = feuille_par_nom_approximatif(
    classeur,
    "8"
)


# Année
annee = None

for ligne in range(
    1,
    15
):
    valeur = feuille_88.cell(
        ligne,
        1
    ).value

    correspondance = re.search(
        r"\b(20\d{2})\b",
        str(
            valeur or ""
        )
    )

    if correspondance:
        annee = correspondance.group(1)
        break

if annee is None:
    raise RuntimeError(
        "Année introuvable dans la table Excel 8-8."
    )


# Nation
ligne_nation = trouver_ligne_total_national(
    feuille_88,
    debut=3,
    fin=56
)

nation = [
    feuille_88.cell(
        ligne_nation,
        colonne
    ).value
    for colonne in range(
        2,
        6
    )
]


# Région Fès-Meknès
ligne_region = trouver_ligne_par_libelle(
    feuille_87,
    "Fès - Meknès",
    colonne=1,
    debut=3,
    fin=122
)

region = [
    feuille_87.cell(
        ligne_region,
        colonne
    ).value
    for colonne in range(
        2,
        6
    )
]


# Préfecture / centre Meknès
ligne_meknes = trouver_ligne_par_libelle(
    feuille_87,
    "Meknès",
    colonne=1,
    debut=3,
    fin=122
)

meknes = [
    feuille_87.cell(
        ligne_meknes,
        colonne
    ).value
    for colonne in range(
        2,
        6
    )
]


classeur.close()


# ============================================================
# 2) OUVRIR LE MODELE WORD
# ============================================================

fichiers_word = [
    fichier
    for fichier in DOSSIER_MODELES.glob(
        "*.docx"
    )
    if not fichier.name.startswith(
        "~$"
    )
]

if not fichiers_word:
    raise FileNotFoundError(
        "Aucun fichier .docx trouvé dans modeles."
    )

fichier_modele = fichiers_word[0]

document = Document(
    fichier_modele
)

numero_table_word = (
    lire_numero_table_word()
)

tableau = document.tables[
    numero_table_word - 1
]


# Sécurité : le tableau 14-1 doit avoir 8 x 6.
if (
    len(tableau.rows) != 8
    or len(tableau.columns) != 6
):
    raise RuntimeError(
        "Structure inattendue du tableau Word 14-1 : "
        f"{len(tableau.rows)} x {len(tableau.columns)}"
    )


# ============================================================
# 3) MISE A JOUR DU TABLEAU WORD
# ============================================================

# Année française / arabe
remplacer_texte_preserver_style(
    tableau.cell(
        0,
        0
    ),
    f"Année {annee}*"
)

remplacer_texte_preserver_style(
    tableau.cell(
        0,
        5
    ),
    f"سنة {annee}*"
)


# Les lignes Word :
# 6 = Nation
# 7 = Région
# 8 = Préfecture Meknès
#
# Les colonnes Word :
# C2 = Total général
# C3 = Véhicules utilitaires
# C4 = Voitures de tourisme
# C5 = Motocyclettes

donnees = [
    nation,
    region,
    meknes
]

for index_ligne, valeurs in enumerate(
    donnees,
    start=5
):
    for index_colonne, valeur in enumerate(
        valeurs,
        start=1
    ):
        remplacer_texte_preserver_style(
            tableau.cell(
                index_ligne,
                index_colonne
            ),
            convertir_nombre(
                valeur
            )
        )


# ============================================================
# 4) SAUVEGARDE
# ============================================================

DOSSIER_SORTIE.mkdir(
    exist_ok=True
)

document.save(
    FICHIER_SORTIE
)


# ============================================================
# 5) RESULTAT
# ============================================================

print("=" * 76)
print("TEST WORD 14-1 TERMINE")
print("=" * 76)

print(
    f"Année : {annee}"
)

print(
    "Nation : "
    + " | ".join(
        convertir_nombre(
            valeur
        )
        for valeur in nation
    )
)

print(
    "Région Fès-Meknès : "
    + " | ".join(
        convertir_nombre(
            valeur
        )
        for valeur in region
    )
)

print(
    "Meknès : "
    + " | ".join(
        convertir_nombre(
            valeur
        )
        for valeur in meknes
    )
)

print()

print(
    f"Document créé : {FICHIER_SORTIE}"
)

print()
print(
    "IMPORTANT : le fichier modeles.docx n'a pas été modifié."
)

print("=" * 76)
