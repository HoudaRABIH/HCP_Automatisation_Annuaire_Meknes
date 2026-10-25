from pathlib import Path
from openpyxl import load_workbook
from docx import Document
from zipfile import ZipFile
from lxml import etree
import re
import unicodedata

BRUT = Path("donnees_brutes")
SORTIES = Path("sorties")

DOCX = SORTIES / "Annuaire_HCP_FINAL_CLEAN_EDUCATION_CORRECTED_JUSTICE_CULTURE_HABITAT_FONCIER_POPULATION_INDUSTRIE_ENERGIE_EAU.docx"
OUT = SORTIES / "verification_derniers_points_final.txt"

def txt(v):
    return "" if v is None else str(v).replace("\n", " ").strip()

def norm(v):
    s = txt(v).lower()
    s = (s.replace("’","'").replace("–","-").replace("—","-")
           .replace("\u202f"," ").replace("\xa0"," "))
    s = "".join(c for c in unicodedata.normalize("NFD", s)
                if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s).strip()

def find_xlsx(*words):
    candidates = []
    nw = [norm(w) for w in words]
    for p in BRUT.glob("*.xlsx"):
        if p.name.startswith("~$"):
            continue
        n = norm(p.name)
        if all(w in n for w in nw):
            candidates.append(p)
    if len(candidates) != 1:
        raise RuntimeError(f"Excel ambigu/introuvable {words}: {[p.name for p in candidates]}")
    return candidates[0]

def row_text(ws, r, max_col=12):
    vals = []
    for c in range(1, min(ws.max_column, max_col) + 1):
        v = ws.cell(r,c).value
        if v not in (None, ""):
            vals.append(f"{ws.cell(r,c).coordinate}={txt(v)!r}")
    return " | ".join(vals)

def find_title_row(ws, regex, max_rows=100):
    pat = re.compile(regex, re.I)
    for r in range(1, min(ws.max_row, max_rows) + 1):
        for c in range(1, min(ws.max_column, 10) + 1):
            v = ws.cell(r,c).value
            if isinstance(v, str) and pat.search(v):
                return r
    return None

def find_rows_with(ws, labels):
    wanted = [norm(x) for x in labels]
    found = []
    for r in range(1, ws.max_row + 1):
        first = norm(ws.cell(r,1).value)
        if any(w == first or w in first for w in wanted):
            found.append(r)
    return found

report = []
def add(*args):
    report.extend(str(x) for x in args)

add("="*120)
add("VERIFICATION DES DERNIERS POINTS AVANT GEL FINAL")
add("="*120)
add(f"Word : {DOCX}")
add("")

# =====================================================================================
# 1. 13-12 : source réelle des immatriculations, surtout rural 2024
# =====================================================================================
print("[1/4] 13-12 Foncier / immatriculations...", flush=True)
add("#"*120)
add("1) 13-12 — IMMATRICULATIONS / RURAL 2024")
add("#"*120)

foncier = find_xlsx("foncier", "2025")
wb = load_workbook(foncier, data_only=True, read_only=True)

target_sheets = []
for ws in wb.worksheets:
    top = []
    for r in range(1, min(ws.max_row, 90)+1):
        line = row_text(ws, r)
        if line:
            top.append(line)
    joined = norm(" ".join(top[:40]))
    if "immatric" in joined:
        target_sheets.append(ws)

add(f"Feuilles contenant 'immatric' : {[ws.title for ws in target_sheets]}")

for ws in target_sheets:
    add(f"\n--- Feuille {ws.title} ---")
    for r in range(1, min(ws.max_row, 100)+1):
        line = row_text(ws, r)
        nl = norm(line)
        if (
            "2024" in line or "2023" in line or "2022" in line
            or "nation" in nl or "ensemble" in nl
            or "fes" in nl and "meknes" in nl
            or "meknes" in nl
            or "rural" in nl or "urbain" in nl
            or "immatric" in nl
        ):
            add(line)

wb.close()

# =====================================================================================
# 2. 6-7 cinéma : est-ce que Meknès existe dans la source 2024 ?
# =====================================================================================
print("[2/4] 6-7 Activité cinématographique...", flush=True)
add("")
add("#"*120)
add("2) 6-7 — ACTIVITE CINEMATOGRAPHIQUE : PRESENCE DE MEKNES")
add("#"*120)

culture = find_xlsx("activites", "culturelles", "2025")
wb = load_workbook(culture, data_only=True, read_only=True)

for ws in wb.worksheets:
    title_match = False
    for r in range(1, min(ws.max_row, 80)+1):
        for c in range(1, min(ws.max_column, 10)+1):
            v = ws.cell(r,c).value
            if isinstance(v, str) and re.match(r"^\s*16\s*-\s*11\b", v, re.I):
                title_match = True
                break
        if title_match:
            break
    if not title_match:
        continue

    add(f"Feuille source 16-11 : {ws.title}")
    mek_rows = find_rows_with(ws, ["Meknès","Meknes"])
    add(f"Lignes Meknès détectées : {mek_rows}")
    for r in range(max(1, r-3), min(ws.max_row, r+30)+1):
        line = row_text(ws, r)
        if line:
            add(line)
    for mr in mek_rows:
        add("MEKNES -> " + row_text(ws, mr))

wb.close()

# =====================================================================================
# 3. Justice 8-2 / 8-3 : vérifier absence de Meknès dans sources 2024
# =====================================================================================
print("[3/4] Justice 8-2 / 8-3...", flush=True)
add("")
add("#"*120)
add("3) JUSTICE 8-2 / 8-3 — PRESENCE OU ABSENCE DE MEKNES DANS LES SOURCES 2024")
add("#"*120)

justice = find_xlsx("justice", "2025")
wb = load_workbook(justice, data_only=True, read_only=True)

for source_code in ["15-10","15-12"]:
    a,b = source_code.split("-")
    pat = re.compile(rf"^\s*{a}\s*-\s*{b}\b", re.I)
    matched = []
    for ws in wb.worksheets:
        found = False
        for r in range(1, min(ws.max_row, 30)+1):
            for c in range(1, min(ws.max_column, 10)+1):
                v = ws.cell(r,c).value
                if isinstance(v, str) and pat.match(v):
                    found = True
                    break
            if found:
                break
        if found:
            matched.append(ws)

    add(f"\nSource {source_code} -> feuilles {[ws.title for ws in matched]}")
    for ws in matched:
        mek_rows = find_rows_with(ws, ["Meknès","Meknes"])
        add(f"  {ws.title}: lignes Meknès = {mek_rows}")
        if mek_rows:
            for mr in mek_rows:
                add("    " + row_text(ws, mr))

wb.close()

# =====================================================================================
# 4. Word 4-4 : vérifier présence des valeurs réelles dans XML (floating textbox)
# =====================================================================================
print("[4/4] 4-4 valeurs dans le Word/XML...", flush=True)
add("")
add("#"*120)
add("4) 4-4 — VALEURS CITE UNIVERSITAIRE DANS LE XML WORD")
add("#"*120)

expected = ["2 668", "1 908", "61", "60 915", "38 491", "1 487"]

with ZipFile(DOCX, "r") as zf:
    xml_texts = []
    for name in zf.namelist():
        if name.startswith("word/") and name.endswith(".xml"):
            try:
                root = etree.fromstring(zf.read(name))
            except Exception:
                continue
            xml_texts.extend(node.text for node in root.iter() if node.text)

whole = "\n".join(xml_texts)
compact = whole.replace("\u00a0"," ").replace("\u202f"," ")

for val in expected:
    present = val in compact
    add(f"{val}: {'PRESENT' if present else 'ABSENT'}")

OUT.write_text("\n".join(report), encoding="utf-8")

print()
print("="*100)
print("VERIFICATION TERMINEE")
print("="*100)
print(f"Rapport créé : {OUT}")
print("="*100)
